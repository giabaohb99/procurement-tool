# -*- coding: utf-8 -*-
"""duoc-CR-598 mục 3 — nạp danh mục hóa chất NĐ 24/2026 phụ lục I–IV (STT, tên khoa học, tên chất,
số CAS, công thức, ngưỡng phụ lục IV) vào `tab_customs_regulation`.

Nạp vào DB (local / prod — đọc bản JSON đã kèm trong repo):
    docker compose exec -T api python -m scripts.load_nd24_regulations
    docker compose exec -T api python -m scripts.load_nd24_regulations --dry-run   # chỉ đếm, không ghi

Đọc lại từ tệp Excel mới (khi phòng Thu mua sửa tệp) rồi ghi đè bản JSON trong repo:
    docker compose exec -T api python -m scripts.load_nd24_regulations --from-xlsx /app/tmp/tep.xlsx --write-json

Luật đồng bộ ở `app.modules.customs.nd24_regulation_loader.apply_rows`: cập nhật dòng đã có, thêm dòng
mới, NGỪNG DÙNG (không xóa) dòng phụ lục I–IV không còn trong tệp; TT 75 + TT 01 không đụng tới.
"""
import argparse
import json
import sys
from pathlib import Path

import app.core.all_models  # noqa: F401 — nạp đủ mapper
from app.core.database import SessionLocal
from app.modules.customs import nd24_regulation_loader as loader

DATA_FILE = Path(loader.__file__).parent / "data" / "nd24_2026_regulations.json"


def read_xlsx(path: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)   # KHÔNG data_only: lấy giá trị gốc của ô
    if loader.SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Tệp không có sheet «{loader.SHEET_NAME}»")
    return loader.parse_sheet(wb[loader.SHEET_NAME].iter_rows(max_col=7, values_only=True))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--from-xlsx", help="đọc từ tệp Excel thay vì bản JSON trong repo")
    ap.add_argument("--write-json", action="store_true", help="ghi kết quả đọc Excel ra bản JSON trong repo")
    ap.add_argument("--dry-run", action="store_true", help="chạy thử, không commit")
    args = ap.parse_args()

    if args.from_xlsx:
        items = read_xlsx(args.from_xlsx)
        if args.write_json:
            DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
            DATA_FILE.write_text(json.dumps(items, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
            print(f"Đã ghi {len(items)} dòng → {DATA_FILE}")
            return 0
    else:
        items = json.loads(DATA_FILE.read_text(encoding="utf-8"))

    db = SessionLocal()
    try:
        stats = loader.apply_rows(db, items)
        if args.dry_run:
            db.rollback()
        else:
            db.commit()
        print(("[CHẠY THỬ] " if args.dry_run else "") +
              f"{stats['total']} dòng: thêm {stats['inserted']} · cập nhật {stats['updated']} · "
              f"ngừng dùng {stats['deactivated']}")
    finally:
        db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
