"""Excel "Danh sách ký nhận" của báo cáo H1 (duoc-CR-688): MỖI PHÁP NHÂN MỘT SHEET.

Mỗi sheet: tiêu đề "DANH SÁCH NHẬN QUÀ <tên dịp>", STT + cột như `/rows`, ô "Ký nhận" để trống, dòng
Tổng tiền, khối ký người lập / người duyệt. Chuỗi do người nhập (tên dịp, họ tên, chức vụ) đi qua
`text_cell` nên không bị Excel hiểu thành công thức. Không có ngày sinh / tên con.
"""
from __future__ import annotations

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.modules.labor_register.export_common import (
    BOLD, BORDER, CENTER, HEAD_FILL, UNASSIGNED, empty_workbook, text_cell, unique_sheet_name, xlsx_download)

from .benefit_report_data import BenefitRow
from .benefit_report_inputs import Filters

WIDTHS = {"stt": 6, "employee_code": 12, "employee_name": 28, "department": 24, "position": 22,
          "children": 9, "amount": 15, "signature": 22}


def _title(label: str) -> str:
    return f"DANH SÁCH NHẬN QUÀ {label}".strip()


def _sheet(ws, company: str, rows: list[BenefitRow], flt: Filters, department_name: dict) -> None:
    with_kids = flt.child_max_age is not None
    heads = [("stt", "STT"), ("employee_code", "Mã NV"), ("employee_name", "Họ tên"),
             ("department", "Phòng ban"), ("position", "Chức vụ")]
    heads += [("children", "Số con")] if with_kids else []
    heads += [("amount", "Mức quà (đồng)"), ("signature", "Ký nhận")]
    n = len(heads)
    text_cell(ws, 1, 1, company, font=BOLD)
    text_cell(ws, 2, 1, _title(flt.occasion_label), font=Font(bold=True, size=14))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws.cell(row=2, column=1).alignment = CENTER
    for i, (key, label) in enumerate(heads, 1):
        text_cell(ws, 4, i, label, font=BOLD, fill=HEAD_FILL, alignment=CENTER, border=BORDER)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[key]
    r = 4
    for idx, p in enumerate(rows, 1):
        r += 1
        vals = [idx, p.employee_code, p.employee_name, department_name.get(p.department_id, ""), p.position]
        vals += [p.units] if with_kids else []
        vals += [p.amount, ""]
        for i, v in enumerate(vals, 1):
            c = text_cell(ws, r, i, v, border=BORDER)
            if heads[i - 1][0] == "amount":
                c.number_format = "#,##0"
    r += 1
    last = n - 1     # cột "Mức quà"
    text_cell(ws, r, 1, "Tổng tiền", font=BOLD, border=BORDER)
    for i in range(2, n + 1):
        ws.cell(row=r, column=i).border = BORDER
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last - 1)
    total = ws.cell(row=r, column=last, value=sum(p.amount for p in rows))
    total.font, total.number_format = BOLD, "#,##0"
    r += 2
    mid = max(2, n // 2)
    for col, who in ((2, "Người lập"), (mid + 1, "Người duyệt")):
        text_cell(ws, r, col, who, font=BOLD, alignment=CENTER)
        text_cell(ws, r + 1, col, "(Ký, ghi rõ họ tên)", alignment=CENTER)


def build_workbook(rows: list[BenefitRow], flt: Filters, company_name: dict, department_name: dict,
                   company_short: dict | None = None) -> Workbook:
    """`company_short`: {id: tên ngắn} — tên sheet ưu tiên tên ngắn, thiếu thì dùng tên đầy đủ."""
    if not rows:
        return empty_workbook("Không có ai thuộc diện nhận quà theo các điều kiện đã chọn.")
    by_co: dict[int, list[BenefitRow]] = {}
    for p in rows:
        by_co.setdefault(p.company_id, []).append(p)
    wb, used = Workbook(), set()
    first = True
    for cid in sorted(by_co, key=lambda c: (not c, company_name.get(c, ""))):
        name = company_name.get(cid) or UNASSIGNED
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = unique_sheet_name((company_short or {}).get(cid) or name, used)
        _sheet(ws, name, by_co[cid], flt, department_name)
    return wb


def benefit_export(rows: list[BenefitRow], flt: Filters, company_name: dict, department_name: dict,
                   company_short: dict | None = None) -> Response:
    return xlsx_download(build_workbook(rows, flt, company_name, department_name, company_short), "danh-sach-ky-nhan-phuc-loi")
