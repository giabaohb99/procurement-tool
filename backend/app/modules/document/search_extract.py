"""Trích CHỮ THUẦN từ nội dung soạn thảo (HTML) và tệp đính kèm.

Nguyên tắc cứng: một tệp hỏng / có mật khẩu / không hỗ trợ chỉ làm CHỈ MỤC CỦA
ĐÚNG TỆP ĐÓ rỗng, ghi log rồi thôi — không bao giờ được ném lỗi làm hỏng cả
lượt dựng chỉ mục của văn bản (`search_index_service.reindex`). Không OCR
(chốt 23/09/2026): PDF/ảnh scan không có lớp chữ → chuỗi rỗng, kết quả tìm vẫn
trúng qua siêu dữ liệu.
"""
import io
import logging
import re

log = logging.getLogger("app.document.search")

#  Tệp lớn hơn mức này KHÔNG trích — chặn cả CPU (parse PDF/docx nặng) lẫn bộ
#  nhớ container `api` (xem "Rủi ro" của phase 07).
MAX_FILE_BYTES = 20 * 1024 * 1024

#  Trần SỐ TRANG của một PDF (M11, rà soát 23/09/2026) — `MAX_FILE_BYTES` chặn
#  được DUNG LƯỢNG nhưng không chặn được SỐ TRANG: một PDF nhiều nghìn trang
#  gần như trắng vẫn lọt dưới 20MB mà `extract_text()` tuyến tính theo trang
#  vẫn ăn hết CPU của tiến trình `api` (đường đồng bộ — cả request tạo/ban
#  hành văn bản lẫn dự phòng của `queue_reindex` khi broker sập).
MAX_PDF_PAGES = 500

#  Đuôi tệp có hàm trích riêng — khớp `_extract_file_text`.
SUPPORTED_EXTENSIONS = ("docx", "xlsx", "pdf", "txt", "csv")


def strip_html(html: str) -> str:
    """Bỏ thẻ HTML, gộp khoảng trắng — dùng cho `content_html` của phiên bản.

    Không dùng thư viện phân tích HTML đầy đủ: nội dung do chính trình soạn
    thảo (tiptap) sinh ra, không phải HTML tùy ý từ ngoài — regex đủ và rẻ.
    """
    text = re.sub(r"<[^>]+>", " ", html or "")
    text = (text.replace("&nbsp;", " ").replace("&amp;", "&")
                .replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"'))
    return re.sub(r"\s+", " ", text).strip()


def extract_file_text(filename: str, data: bytes) -> str:
    """Trích chữ theo ĐUÔI TỆP — docx/xlsx/pdf/txt/csv. Đuôi lạ hoặc tệp quá
    20MB → chuỗi rỗng (không trích, không phải lỗi)."""
    if not data or len(data) > MAX_FILE_BYTES:
        return ""
    name = filename or ""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in SUPPORTED_EXTENSIONS:
        return ""
    try:
        if ext == "docx":
            return _extract_docx(data)
        if ext == "xlsx":
            return _extract_xlsx(data)
        if ext == "pdf":
            return _extract_pdf(data)
        return _extract_plain_text(data)
    except Exception:  # noqa: BLE001 — tệp hỏng không được làm hỏng cả lượt
        log.warning("Không trích được chữ từ tệp %r (đuôi .%s)", filename, ext, exc_info=True)
        return ""


def _extract_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text]
    #  Cả BẢNG — hợp đồng/quy chế phần lớn thông tin cụ thể nằm trong bảng.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    #  `read_only=True`: chỉ đọc giá trị, không tải toàn bộ workbook vào bộ
    #  nhớ dưới dạng đối tượng có thể sửa — rẻ hơn nhiều cho tệp lớn.
    #  `data_only=True`: đọc GIÁ TRỊ đã tính của công thức, không đọc công thức.
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is not None and str(value).strip():
                        parts.append(str(value))
        return "\n".join(parts)
    finally:
        wb.close()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        #  PDF "mã hóa" chỉ để hạn chế in ấn/sửa thường vẫn mở được bằng mật
        #  khẩu rỗng — thử một lần, không được thì bỏ qua (không phải lỗi).
        try:
            if reader.decrypt("") == 0:
                return ""
        except Exception:  # noqa: BLE001
            return ""
    parts = []
    for i, page in enumerate(reader.pages):
        if i >= MAX_PDF_PAGES:
            log.warning("PDF vượt trần %s trang — chỉ trích %s trang đầu, bỏ phần còn lại",
                       MAX_PDF_PAGES, MAX_PDF_PAGES)
            break
        #  PDF scan không có lớp chữ → `extract_text()` trả rỗng cho từng
        #  trang — đúng ý "không OCR", không phải lỗi trích.
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n".join(parts)


def _extract_plain_text(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")
