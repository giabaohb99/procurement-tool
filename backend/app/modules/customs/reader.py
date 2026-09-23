"""Đọc tệp kết xuất GTT02 thành dòng hàng đã chuẩn hóa — KHÔNG đụng DB (bao-CR-470).

Tách khỏi phần ghi để kiểm được bằng dữ liệu giả, không cần cơ sở dữ liệu. Mọi
luật đọc nằm ở đây; thiết kế ở `doc/erp/hai-quan/02-thiet-ke-ky-thuat.md` §4.

Ba cái bẫy đo được trên 5 tệp thật, cả ba im lặng nếu bỏ qua:

1. **Cột Ngày đăng ký bị đảo ngày/tháng ở 42% số dòng.** Nguồn là chữ
   `DD-MM-YYYY`; hễ ngày ≤ 12 thì Excel hiểu nhầm thành `MM-DD` rồi đổi thành ô
   ngày, ngày > 12 thì không hiểu được nên để nguyên chữ. Dấu vân tay: ô chữ chỉ
   chứa ngày 13–31, ô ngày suy ra "ngày" chỉ từ 1–12. `_detect_date_swap` soi
   đúng dấu vân tay đó trên CẢ TỆP rồi mới vá — nguồn mà sửa khuôn xuất (ô ngày
   đúng, có ngày > 12) thì tự thôi vá, không làm hỏng dữ liệu đúng.
2. **Cột Ngày hợp đồng là ô ngày THẬT, tuyệt đối KHÔNG vá** — 10.855 ô có ngày
   > 12. Vá nhầm là hỏng 39% số ngày hợp đồng đang đúng. Luật vá gắn vào đúng cột
   `reg_date`, không có vòng lặp "vá mọi cột ngày".
3. **Mã có dấu `'` đầu** (`'0500590269`, `'38089990`) — mẹo giữ số 0 đầu của Excel,
   bóc bỏ. Mã HS giữ dạng chữ vì có số 0 đầu (`01012100`).
"""
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from app.modules.import_tool.catalog_import import ImportValidationError

from .constants import COLUMNS, DECIMAL_KEYS, TEXT_LIMITS, TransportMode

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0"   # .xls đời cũ (BIFF)
_ZIP_MAGIC = b"PK"                  # .xlsx

EMPTY, TEXT, NUMBER, DATE = "empty", "text", "number", "date"


class CustomsFileError(ImportValidationError):
    """Tệp không dùng được (sai định dạng, thiếu cột, mã lạ) — từ chối CẢ LÔ.

    Kế thừa `ImportValidationError` để tác vụ nền của `import_tool` ghi đúng câu
    thân thiện vào lô, không kèm vết lỗi."""


@dataclass
class Cell:
    kind: str
    value: object = None


@dataclass
class ParseResult:
    rows: list[dict] = field(default_factory=list)
    logs: list[tuple[int, int, str]] = field(default_factory=list)   # (dòng, LogLevel, câu)
    skipped: int = 0
    date_swap: bool = False


def normalize_text(value: object) -> str:
    """Chữ thường, bỏ dấu, gộp khoảng trắng — dùng để khớp tiêu đề và chống trùng tên.

    Nhờ bỏ dấu nên «Tên nuớc xuất xứ» (lỗi chính tả của nguồn) và «Tên nước xuất
    xứ» ra cùng một chuỗi — nguồn có sửa chính tả thì tệp vẫn khớp.
    """
    s = unicodedata.normalize("NFD", str(value or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").lower()
    return re.sub(r"\s+", " ", s).strip()


def clean_party_name(value: object) -> str:
    """Sửa lỗi hoa-thường của nguồn: `CôNG TY TNHH BAYER VIệT NAM` → `CÔNG TY TNHH BAYER VIỆT NAM`.

    Kết xuất GTT02 viết hoa toàn tên nhưng lại để thường đúng các chữ có dấu. Tên mà phần
    lớn chữ cái đã hoa (≥ 70%) thì đưa về in hoa đều; tên viết kiểu thường ("Công Ty
    TNHH…") giữ nguyên — đừng biến mọi tên thành in hoa.
    """
    s = re.sub(r"\s+", " ", str(value or "")).strip()
    letters = [ch for ch in s if ch.isalpha()]
    if letters and sum(ch.isupper() for ch in letters) / len(letters) >= 0.7:
        return s.upper()
    return s


def hash_name(value: object) -> str:
    return hashlib.md5(normalize_text(value).encode("utf-8")).hexdigest()


# ── Đọc tệp thành lưới ô ────────────────────────────────────────────────────
def load_cells(raw: bytes, filename: str = "") -> list[list[Cell]]:
    """Tệp → lưới `Cell`. Hai định dạng đi chung một đầu ra để luật đọc chỉ viết một lần."""
    if raw[:4] == _OLE2_MAGIC:
        return _cells_from_xls(raw)
    if raw[:2] == _ZIP_MAGIC:
        return _cells_from_xlsx(raw)
    raise CustomsFileError(f"Tệp «{filename}» không phải Excel (.xls hoặc .xlsx).")


def _cells_from_xls(raw: bytes) -> list[list[Cell]]:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception as e:  # noqa: BLE001
        raise CustomsFileError(f"Không đọc được tệp .xls: {e}") from e
    return cells_from_xlrd_sheet(book.sheet_by_index(0), book.datemode)


def cells_from_xlrd_sheet(sheet, datemode: int) -> list[list[Cell]]:
    """Tách riêng để kiểm được bằng một sheet giả (bộ test không dựng nổi tệp .xls)."""
    import xlrd
    grid = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(Cell(DATE, xlrd.xldate_as_datetime(cell.value, datemode).date()))
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                row.append(Cell(NUMBER, cell.value))
            elif cell.ctype == xlrd.XL_CELL_TEXT and str(cell.value).strip():
                row.append(Cell(TEXT, str(cell.value)))
            else:
                row.append(Cell(EMPTY))
        grid.append(row)
    return grid


def _cells_from_xlsx(raw: bytes) -> list[list[Cell]]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        raise CustomsFileError(f"Không đọc được tệp .xlsx: {e}") from e
    grid = []
    for values in wb.worksheets[0].iter_rows(values_only=True):
        row = []
        for v in values:
            if isinstance(v, datetime):
                row.append(Cell(DATE, v.date()))
            elif isinstance(v, date):
                row.append(Cell(DATE, v))
            elif isinstance(v, bool) or v is None:
                row.append(Cell(EMPTY))
            elif isinstance(v, (int, float)):
                row.append(Cell(NUMBER, v))
            elif str(v).strip():
                row.append(Cell(TEXT, str(v)))
            else:
                row.append(Cell(EMPTY))
        grid.append(row)
    wb.close()
    return grid


# ── Tiêu đề ─────────────────────────────────────────────────────────────────
def map_header(header: list[Cell]) -> dict[str, int]:
    """Khớp 32 cột theo CHỮ đã chuẩn hóa, không theo vị trí. Thiếu cột nào → từ chối lô."""
    found = {normalize_text(c.value): i for i, c in enumerate(header) if c.kind != EMPTY}
    col_of, missing = {}, []
    for key, label in COLUMNS:
        idx = found.get(normalize_text(label))
        if idx is None:
            missing.append(label)
        else:
            col_of[key] = idx
    if missing:
        raise CustomsFileError("Tệp thiếu cột: " + ", ".join(f"«{m}»" for m in missing)
                               + ". Tệp phải là kết xuất tra cứu GTT02 đủ 32 cột.")
    return col_of


def check_headers(raw: bytes, filename: str = "") -> None:
    """Kiểm đồng bộ ngay lúc tải lên — sai là báo liền, không tạo lô."""
    grid = load_cells(raw, filename)
    if not grid:
        raise CustomsFileError(f"Tệp «{filename}» rỗng.")
    map_header(grid[0])


# ── Luật từng ô ─────────────────────────────────────────────────────────────
def _detect_date_swap(cells: list[Cell]) -> bool:
    """Tệp có đúng dấu vân tay "Excel đọc nhầm DD-MM thành MM-DD" không.

    Có: vừa có ô CHỮ, vừa có ô NGÀY, và MỌI ô ngày đều có "ngày" ≤ 12. Ô ngày thật
    của một kết xuất đúng khuôn sẽ có ngày > 12 → trả False → không vá gì.
    """
    dates = [c.value for c in cells if c.kind == DATE]
    texts = [c for c in cells if c.kind == TEXT]
    return bool(dates) and bool(texts) and all(d.day <= 12 for d in dates)


def _parse_text_date(text: str) -> date | None:
    m = re.fullmatch(r"\s*(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})\s*", text)
    if not m:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def _read_date(cell: Cell, swap: bool) -> tuple[date | None, bool]:
    """→ (ngày, đã vá?). `swap` chỉ bao giờ được bật cho cột Ngày đăng ký."""
    if cell.kind == DATE:
        d = cell.value
        if swap:
            return date(d.year, d.day, d.month), True     # ô ngày: ngày ↔ tháng đã bị đảo
        return d, False
    if cell.kind == TEXT:
        return _parse_text_date(cell.value), False
    if cell.kind == NUMBER:   # ô số trần mang giá trị ngày Excel (không định dạng ngày)
        d = (datetime(1899, 12, 30) + timedelta(days=int(cell.value))).date()
        return (date(d.year, d.day, d.month), True) if swap and d.day <= 12 else (d, False)
    return None, False


def _read_code(cell: Cell, width: int = 0) -> str:
    """Mã (mã số thuế, mã HS): bóc dấu `'`; ô số thì bù số 0 đầu đã mất."""
    if cell.kind == NUMBER:
        s = str(int(cell.value))
        return s.zfill(width) if width else s
    if cell.kind == TEXT:
        return str(cell.value).strip().lstrip("'").strip()
    return ""


def _read_decimal(cell: Cell) -> Decimal | None:
    if cell.kind == NUMBER:
        return Decimal(repr(float(cell.value)))
    if cell.kind == TEXT:
        try:
            return Decimal(str(cell.value).strip().replace(",", ""))
        except InvalidOperation:
            raise ValueError(f"không phải số: «{cell.value}»") from None
    return None


def _read_text(cell: Cell) -> str:
    if cell.kind == NUMBER:
        v = cell.value
        return str(int(v)) if float(v).is_integer() else str(v)
    return str(cell.value).strip() if cell.kind == TEXT else ""


def _read_transport(cell: Cell) -> int | None:
    """`2-Đường biển (container)` → 2. Mã lạ → từ chối CẢ LÔ, không lưu bừa thành 9:
    mã lạ nghĩa là GTT02 đã đổi khuôn, cần người xem."""
    text = _read_text(cell)
    if not text:
        return None
    m = re.match(r"\s*(\d+)", text)
    code = int(m.group(1)) if m else None
    if code not in {int(t) for t in TransportMode}:
        raise CustomsFileError(f"Phương tiện vận chuyển lạ: «{text}». Nguồn có thể đã đổi khuôn.")
    return code


# ── Đọc cả tệp ──────────────────────────────────────────────────────────────
def parse(raw: bytes, filename: str = "", today: date | None = None) -> ParseResult:
    from app.modules.import_tool.model import LogLevel

    grid = load_cells(raw, filename)
    if not grid:
        raise CustomsFileError(f"Tệp «{filename}» rỗng.")
    col = map_header(grid[0])
    body = [r for r in grid[1:] if any(c.kind != EMPTY for c in r)]
    width = max((len(r) for r in body), default=0)
    body = [r + [Cell(EMPTY)] * (width - len(r)) for r in body]

    res = ParseResult(date_swap=_detect_date_swap([r[col["reg_date"]] for r in body]))
    for offset, cells in enumerate(body):
        row_no = offset + 2                      # dòng thật trong tệp (dòng 1 là tiêu đề)
        get = lambda key: cells[col[key]]        # noqa: E731
        reg_date, fixed = _read_date(get("reg_date"), res.date_swap)
        if reg_date is None:
            res.skipped += 1
            res.logs.append((row_no, LogLevel.ERROR, "Không đọc được Ngày đăng ký — bỏ dòng"))
            continue
        out = {"source_row": row_no, "date_fixed": 1 if fixed else 0, "reg_date": reg_date,
               "contract_date": _read_date(get("contract_date"), False)[0],   # ⚠️ KHÔNG vá
               "importer_tax_code": _read_code(get("importer_tax_code"), 10),
               "hs_code": _read_code(get("hs_code"), 8),
               "transport_mode": _read_transport(get("transport_mode"))}
        line_no = get("line_no")
        out["line_no"] = int(float(line_no.value)) if line_no.kind == NUMBER else int(_read_text(line_no) or 0)
        for key in ("office_code", "importer_name", "partner_name", "product_name", "currency",
                    "unit_code", "origin_country", "contract_no", "incoterm", "import_country"):
            out[key] = _read_text(get(key))
        out["importer_name"] = clean_party_name(out["importer_name"])
        out["partner_name"] = clean_party_name(out["partner_name"])
        for key in DECIMAL_KEYS:
            try:
                out[key] = _read_decimal(get(key))
            except ValueError as e:
                out[key] = None
                res.logs.append((row_no, LogLevel.WARNING, f"Cột «{dict(COLUMNS)[key]}» {e} — để trống"))
        for key, limit in TEXT_LIMITS.items():
            if len(out[key]) > limit:
                res.logs.append((row_no, LogLevel.WARNING,
                                 f"Cột «{dict(COLUMNS)[key]}» dài hơn {limit} ký tự — đã cắt"))
                out[key] = out[key][:limit]
        res.rows.append(out)

    today = today or date.today()
    if any(r["reg_date"] > today for r in res.rows):
        res.logs.append((0, LogLevel.WARNING,
                         "Có Ngày đăng ký nằm trong TƯƠNG LAI — kiểm tra lại tệp: có thể nguồn đã "
                         "đổi cách xuất ngày và luật vá ngày không còn đúng"))
    return res
