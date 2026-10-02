"""Xuất Excel danh mục Thuốc BVTV — mục «Thuốc BVTV» của Tra cứu thị trường (02/10/2026).

Hai phạm vi:
  · `scope=all`  — CẢ danh mục (kể cả thuốc tự thêm, xem `pesticide_export_columns.export_id`),
    BỎ QUA bộ lọc: nạp lại («Nạp danh mục») là THAY TOÀN BỘ, nên chỉ tệp đầy đủ mới nạp lại an
    toàn — lọc bớt rồi nạp lại là xóa sạch phần không nằm trong tệp.
  · `scope=page` — ĐÚNG các dòng của MỘT TRANG màn tra cứu, dùng lại `pesticide_service.
    export_query` (cùng bộ lọc + sắp xếp với `list_pesticides`) — để xem/đối chiếu, không nhằm
    nạp lại (trang là một phần, nạp lại sẽ xóa phần còn lại của danh mục).

Dữ liệu lớn (`scope=all`): truy vấn + ghi THEO LÔ bằng `Workbook(write_only=True)` ra tệp TẠM
TRÊN ĐĨA — không giữ cả danh mục trong bộ nhớ (container `api` chỉ 2 GB, xem
`backend-input-limits.md`). Người gọi (controller) chịu trách nhiệm xóa tệp tạm sau khi gửi
(`BackgroundTask`).

Trần = trần của CHÍNH bộ đọc (`pesticide_reader.MAX_RECORDS` / `MAX_USE_ROWS`, đọc lại thuộc
tính mỗi lần — không chép hằng số) — tệp vượt trần thì `pesticide_reader` sẽ từ chối nạp, nên
chặn ngay lúc XUẤT cho người dùng biết sớm, rõ hơn là để họ nạp lên rồi mới thấy lỗi.

Chặn bấm dồn `scope=all` (M1) — xem `pesticide_export_lock` (tách riêng, tự kiểm được bằng db
giả, không cần ORM Session thật).

Tệp xuất ghi kèm sheet ẨN `_xuat` (`pesticide_export_marker`) nêu PHẠM VI đã xuất — bộ đọc dùng
nó để chọn chế độ nạp lại: `scope=all` → THAY TOÀN BỘ như cũ; `scope=page` → chỉ CẬP NHẬT đúng
thuốc có trong tệp (`pesticide_merge_service`, C1). Thuốc tự thêm VÀ thuốc nguồn trùng
`source_id` với nhau (dữ liệu nhiễm, hiếm) xuất với `id` ÂM — bộ đọc bỏ qua các dòng này, không
đụng/làm hỏng phần còn lại của tệp (C2, xem `pesticide_export_columns.export_id`). Ô chuỗi được
lọc ký tự điều khiển cấm của XML + ép chuỗi bắt đầu bằng =/+/-/@ thành DẠNG CHUỖI trước khi ghi
(H1/C3, `pesticide_export_safety`). Lỗi SAU khi đã ghi xong tệp (kiểm trần kích thước, ghi nhật
ký) tự xóa tệp tạm rồi ném lại (M2); tệp vượt trần `MAX_UPLOAD_BYTES` của bộ nạp thì 422 ngay,
không chờ người dùng nạp lên mới biết (M3).

Nhật ký xuất: ghi một dòng vào `tab_export_log` (Đ-13b, QĐ-I5 — "ghi nhận MỌI endpoint export")
NHƯNG không đăng ký vào `export_log.registry.EXPORT_ADAPTERS`/`run_export()`. Khung đó xuất MỘT
bảng phẳng bằng cột chung (`Col` + `xlsx_response`, trần 100.000 dòng, build cả workbook trong bộ
nhớ) — còn mục này xuất HAI sheet lồng nhau (thuốc + phạm vi sử dụng) với tên cột cố định khớp
đúng bộ đọc nạp, theo LÔ ra đĩa. Ép vào khung chung sẽ phá hỏng hợp đồng nạp lại hoặc mất hẳn lợi
ích về bộ nhớ — nên tách endpoint riêng, chỉ dùng chung bảng nhật ký để vẫn truy vết được ai xuất
gì. Không lưu lại file lên kho lưu trữ (`file_id=0`, không có "tải lại đúng file đã xuất" như
`/api/exports/{id}/file`): đọc hết tệp tạm vào bộ nhớ lần nữa chỉ để lưu lại là triệt tiêu đúng
lợi ích streaming-ra-đĩa ở trên; cần xuất lại thì gọi lại API, dữ liệu danh mục hiếm khi đổi dồn.
"""
import os
import tempfile
from datetime import datetime

from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.export_xlsx import VN_OFFSET   # L6 — hằng giờ VN DÙNG CHUNG, không khai lại số
from app.modules.export_log.model import ExportLog

from . import pesticide_export_lock as lock
from . import pesticide_export_marker as marker
from . import pesticide_export_safety as safety
from . import pesticide_reader as reader
from .model import CustomsPesticide, CustomsPesticideUse
from .pesticide_export_columns import LIST_HEADER, USE_HEADER, list_row, use_rows
from .pesticide_service import export_query

#  Cùng cỡ lô với `pesticide_service._CHUNK` (lượt nạp) — đo trên dữ liệu thật đủ nhanh, không
#  giữ quá nhiều đối tượng ORM một lúc.
_CHUNK = 1000


def _check_caps(pesticide_count: int, use_count: int) -> None:
    if pesticide_count > reader.MAX_RECORDS:
        raise HTTPException(422, (
            f"Danh mục có {pesticide_count:,} thuốc, vượt trần {reader.MAX_RECORDS:,} của bộ nạp "
            "— tệp xuất ra sẽ không nạp lại được bằng «Nạp danh mục». Hãy xuất theo trang/bộ lọc."
        ).replace(",", "."))
    if use_count > reader.MAX_USE_ROWS:
        raise HTTPException(422, (
            f"Danh mục có {use_count:,} dòng phạm vi sử dụng, vượt trần {reader.MAX_USE_ROWS:,} "
            "của bộ nạp — tệp xuất ra sẽ không nạp lại được bằng «Nạp danh mục». Hãy xuất theo "
            "trang/bộ lọc."
        ).replace(",", "."))


def _uses_by_pesticide(db: Session, ids: list[int]) -> dict[int, list[CustomsPesticideUse]]:
    if not ids:
        return {}
    rows = (db.query(CustomsPesticideUse).filter(CustomsPesticideUse.pesticide_id.in_(ids))
            .order_by(CustomsPesticideUse.pesticide_id, CustomsPesticideUse.sort_order,
                     CustomsPesticideUse.id).all())
    out: dict[int, list[CustomsPesticideUse]] = {}
    for u in rows:
        out.setdefault(u.pesticide_id, []).append(u)
    return out


def _iter_all(db: Session):
    """Toàn danh mục, THEO LÔ keyset bằng id — không `.all()` cả bảng cùng lúc."""
    last_id = 0
    while True:
        batch = (db.query(CustomsPesticide).filter(CustomsPesticide.id > last_id)
                .order_by(CustomsPesticide.id).limit(_CHUNK).all())
        if not batch:
            return
        uses = _uses_by_pesticide(db, [p.id for p in batch])
        for p in batch:
            yield p, uses.get(p.id, [])
        last_id = batch[-1].id
        #  Vài nghìn đối tượng ORM (thuốc + phạm vi) ngồi lại trong session tới cuối request là
        #  phí bộ nhớ — dữ liệu của lô này đã được ghi vào worksheet (write_only) ở trên rồi.
        db.expunge_all()


def _iter_page(db: Session, q: str, status: int | None, pest_group: str, sector: str,
              banned_only: bool, banned_regulation_id: int | None, offset: int, limit: int):
    """ĐÚNG các dòng MỘT TRANG — cùng bộ lọc + sắp xếp với `pesticide_service.list_pesticides`."""
    rows = export_query(db, q, status, pest_group, sector, banned_only,
                        banned_regulation_id).offset(offset).limit(limit).all()
    uses = _uses_by_pesticide(db, [p.id for p in rows])
    for p in rows:
        yield p, uses.get(p.id, [])


def _duplicate_source_ids(db: Session) -> frozenset[int]:
    """C2 — `source_id` DƯƠNG đang bị HAI thuốc nguồn (không thủ công) dùng chung trong DB hiện
    tại (dữ liệu nhiễm từ trước, vd `source_id = 0` từng được coi "luôn mới" ở
    `replace_catalog` khi nguồn thiếu mã). Truyền vào `export_id()` để các dòng này xuất với id
    ÂM thay vì `source_id` thật — tệp vẫn nạp lại được, phần trùng bị bỏ qua lúc nạp. Một câu
    gộp rẻ (có chỉ mục `source_id`), chạy MỘT LẦN cho cả lượt xuất, không theo từng dòng."""
    rows = (db.query(CustomsPesticide.source_id)
            .filter(CustomsPesticide.is_manual.is_(False), CustomsPesticide.source_id != 0)
            .group_by(CustomsPesticide.source_id)
            .having(func.count(CustomsPesticide.id) > 1).all())
    return frozenset(r[0] for r in rows)


def _write_file(pesticides_iter, dup_source_ids: frozenset[int], scope: str,
                page: int | None) -> tuple[str, int, int]:
    fd, path = tempfile.mkstemp(suffix=".xlsx", prefix="thuoc-bvtv-")
    os.close(fd)
    try:
        wb = Workbook(write_only=True)
        ws_list = wb.create_sheet(reader.LIST_SHEET)
        ws_list.append(LIST_HEADER)
        ws_use = wb.create_sheet(reader.USE_SHEET)
        ws_use.append(USE_HEADER)
        n_pesticides = n_uses = 0
        for p, uses in pesticides_iter:
            #  H1 + C3 — bỏ ký tự điều khiển cấm của XML, ép chuỗi bắt đầu bằng =/+/-/@ thành
            #  DẠNG CHUỖI (không phải công thức) — xem `pesticide_export_safety`.
            ws_list.append(safety.safe_row(ws_list, list_row(p, uses, dup_source_ids)))
            for row in use_rows(p, uses, dup_source_ids):
                ws_use.append(safety.safe_row(ws_use, row))
            n_pesticides += 1
            n_uses += len(uses)
        #  C1 — sheet ẩn ghi PHẠM VI đã xuất, cho bộ đọc chọn chế độ nạp lại (thay toàn bộ / chỉ
        #  cập nhật). Gọi SAU khi đã có ws_list/ws_use: workbook write_only không cho ẩn sheet
        #  DUY NHẤT của tệp.
        marker.write_marker(wb, scope, page)
        wb.save(path)
    except Exception:
        os.remove(path)
        raise
    return path, n_pesticides, n_uses


def _stamp() -> str:
    return (datetime.utcnow() + VN_OFFSET).strftime("%Y%m%d-%H%M")


def _log_summary(scope: str, page: int, filters: str) -> str:
    """L2 — nhãn ĐỌC ĐƯỢC cho `/system/exports` (trước đây ghi thẳng `scope=all`/tham số thô)."""
    base = "Thuốc BVTV — toàn bộ" if scope == "all" else f"Thuốc BVTV — trang {page}"
    return f"{base} (lọc: {filters})" if filters else base


def _record_log(db: Session, user_id: int, row_count: int, filename: str, file_size: int,
                summary: str) -> None:
    db.add(ExportLog(entity="customs_price", fmt="xlsx", row_count=row_count, filename=filename,
                     file_size=file_size, file_id=0, filter_summary=summary,
                     created_by=user_id, updated_by=user_id))
    db.commit()


def _check_output_size(file_size: int) -> None:
    """M3 — tệp vừa ghi vượt trần của chính bộ nạp (`reader.MAX_UPLOAD_BYTES`) thì bộ nạp sẽ từ
    chối ngay khi người dùng thử nạp lại; báo ngay lúc xuất cho rõ hơn là để họ tải lên rồi mới
    thấy lỗi."""
    if file_size > reader.MAX_UPLOAD_BYTES:
        got_mb = file_size / (1024 * 1024)
        cap_mb = reader.MAX_UPLOAD_BYTES / (1024 * 1024)
        raise HTTPException(422, (
            f"Tệp xuất ra {got_mb:.1f} MB, vượt trần {cap_mb:.0f} MB của bộ nạp — sẽ không nạp "
            "lại được bằng «Nạp danh mục». Hãy xuất theo trang/bộ lọc."))


def export_to_tempfile(db: Session, user, scope: str, q: str, status: int | None, pest_group: str,
                       sector: str, banned_only: bool, banned_regulation_id: int | None,
                       page: int, offset: int, limit: int) -> tuple[str, str]:
    """→ (đường dẫn tệp tạm, tên tệp để tải). Tệp tạm do NGƯỜI GỌI xóa (route gắn `BackgroundTask`
    sau khi đã trả response — xóa trước là hỏng lượt tải)."""
    dup_source_ids = _duplicate_source_ids(db)
    if scope == "all":
        with lock.export_lock(db, user.id):
            total = db.query(func.count(CustomsPesticide.id)).scalar() or 0
            use_total = db.query(func.count(CustomsPesticideUse.id)).scalar() or 0
            _check_caps(total, use_total)
            path, n_pesticides, n_uses = _write_file(_iter_all(db), dup_source_ids, scope, page)
        filename = f"thuoc-bvtv-toan-bo-{_stamp()}.xlsx"
        filters = ""
    else:
        rows = list(_iter_page(db, q, status, pest_group, sector, banned_only,
                               banned_regulation_id, offset, limit))
        _check_caps(len(rows), sum(len(u) for _, u in rows))
        path, n_pesticides, n_uses = _write_file(iter(rows), dup_source_ids, scope, page)
        filename = f"thuoc-bvtv-trang-{page}-{_stamp()}.xlsx"
        filters = (f"q={q!r} status={status} pest_group={pest_group!r} sector={sector!r} "
                  f"banned_only={banned_only} banned_regulation_id={banned_regulation_id}")

    #  M2 (review 02/10/2026) — lỗi ở ĐÂY xảy ra SAU KHI đã ghi xong tệp (`getsize`, kiểm trần,
    #  ghi nhật ký) — không dọn thì tệp tạm RÁC lại trên đĩa mãi (không ai gọi `BackgroundTask`
    #  xóa, vì hàm chưa trả về đường dẫn cho controller).
    try:
        file_size = os.path.getsize(path)
        _check_output_size(file_size)
        _record_log(db, user.id, n_pesticides, filename, file_size,
                   _log_summary(scope, page, filters))
    except Exception:
        os.remove(path)
        raise
    return path, filename
