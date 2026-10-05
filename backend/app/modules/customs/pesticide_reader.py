"""Đọc tệp danh mục thuốc BVTV — mục «Thuốc BVTV» của Tra cứu thị trường (29/09/2026).

Nhận ĐÚNG hai tệp mà bản cào danhmuc.thuocbvtv.com xuất ra (`plans/260928-1553-thuoc-bvtv-danh-muc/`):
  · `thuoc-bvtv.json`  — mảng lồng nhau, mỗi thuốc mang sẵn `pham_vi_su_dung`;
  · `thuoc-bvtv.xlsx`  — sheet «Danh sach thuoc» + sheet «Pham vi su dung» (dòng phẳng, nối bằng `id`).

Hai tệp đọc ra CÙNG một dạng (`PesticideRecord` = dict), để phần ghi không phải biết tệp gốc là gì.
Đọc hết và kiểm hết TRƯỚC khi ghi: tệp hỏng thì ném `ValueError`, danh mục đang có không bị đụng.
"""
import html
import io
import json
import re
from collections import defaultdict
from datetime import date, datetime

from .constants import PESTICIDE_STATUS_BY_LABEL, PesticideStatus

LIST_SHEET = "Danh sach thuoc"
USE_SHEET = "Pham vi su dung"
#  Cột THÊM của hệ thống ở cuối sheet «Danh sach thuoc» — bản cào gốc không có cột này, nên
#  tệp cào thật luôn đọc rỗng ở đây. `pesticide_export_columns` (02/10) ghi NGUYÊN VĂN cột
#  `resistance` đã lưu vào đây — bộ đọc ưu tiên cột này nếu có+không rỗng, bỏ qua bước tách
#  rồi ghép lại qua `_resistance()` (H2, review 02/10/2026): chữ tự do có thể chứa `;`/`:`
#  KHÔNG phải dấu phân tách, mà bước tách-ghép lại hiểu nhầm là dấu phân tách và mất/méo chữ.
RAW_RESISTANCE_COLUMN = "quan_ly_tinh_khang_raw"
#  Trần số thuốc một tệp — nguồn hiện 6 919; gấp năm lần vẫn là tệp đúng, quá nữa là tệp nhầm.
MAX_RECORDS = 40_000
MAX_USES_PER_RECORD = 500
#  Trần TỔNG dòng phạm vi (nguồn hiện 15 309) — đếm ngay lúc đọc, tệp nhầm dừng sớm chứ không
#  nạp hết vào bộ nhớ rồi mới kiểm (container api chỉ có 2 GB).
MAX_USE_ROWS = 200_000
#  Trần kích thước tệp NẠP (M3, review 02/10/2026) — dùng CHUNG giữa `pesticide_controller`
#  (chặn tệp người dùng tải lên) và `pesticide_export_service` (chặn xuất ra một tệp mà bộ
#  đọc chính mình sẽ từ chối), không khai lại số ở hai nơi. Tệp JSON của bản cào ~18 MB
#  (Excel ~3,4 MB); trần 30 MB đủ cho nguồn lớn dần mà JSON đã giải ra vẫn nằm gọn trong trần
#  bộ nhớ 2 GB của container api.
MAX_UPLOAD_BYTES = 30 * 1024 * 1024
#  Cột `Text` của MySQL chứa tối đa 64 KB; dài hơn là lỗi 500. Nguồn dài nhất hiện ~3,4 nghìn ký tự.
_TEXT_LIMIT = 20_000

#  Độ dài phải khớp `String(n)` ở `model.py` — MySQL từ chối chuỗi dài hơn (lỗi 500, không phải 422).
_LIMITS = {"trade_name": 255, "trade_key": 255, "active_ingredient": 500, "pest_group": 100,
           "registrant": 255, "sector": 100, "concentration": 100, "registration_no": 60,
           "toxicity": 500, "source_url": 255}
_USE_LIMITS = {"crop": 255, "pest": 255, "dosage": 255, "pre_harvest_interval": 255}


def trade_key(trade_name: str) -> str:
    """`Bipyrhone 20EC` → `BIPYRHONE` (phần trước hàm lượng). Ngắn hơn 5 ký tự thì bỏ —
    tên quá ngắn dò trong tên hàng hải quan sẽ khớp bừa (`ingredient.IngredientTagger`)."""
    base = re.split(r"\s+\d", html.unescape(trade_name or "").upper())[0].strip()
    return base if len(base) >= 5 else ""


def _text(value) -> str:
    #  M5 (review 02/10/2026) — `.split()`/`" ".join(...)` GỘP khoảng trắng + xuống dòng liên
    #  tiếp thành một dấu cách. Hành vi SẴN CÓ từ trước (không phải lỗi mới), giữ nguyên: câu
    #  tóm tắt/cách dùng nhiều dòng của nguồn không cần giữ xuống dòng để hiển thị đúng.
    if value is None:
        return ""
    return " ".join(html.unescape(str(value)).split())


def _date(value) -> date | None:
    """ISO `2028-09-25` (JSON) hoặc ô ngày của Excel. Năm ngoài 1950–2100 coi như rác."""
    if isinstance(value, datetime):
        value = value.date()
    if not isinstance(value, date):
        try:
            value = date.fromisoformat(_text(value)[:10])
        except ValueError:
            return None
    return value if 1950 <= value.year <= 2100 else None


def _status(label) -> int:
    return int(PESTICIDE_STATUS_BY_LABEL.get(_text(label).lower(), PesticideStatus.UNKNOWN))


def _resistance(items: list[tuple[str, list[str]]]) -> str:
    """[(hoạt chất, [mã, nhóm, phương thức])] → `Tên: FRAC 19 | nhóm | phương thức; …`.

    Bỏ phần rỗng, và bỏ hẳn hoạt chất không có gì ngoài tên — nguồn ghi «Chitosan:  |  | »
    cho hoạt chất chưa xếp nhóm kháng, bày ra chỉ là một dòng trống có dấu hai chấm."""
    out = []
    for name, parts in items:
        parts = [p for p in (_text(x) for x in parts) if p]
        if _text(name) and parts:
            out.append(f"{_text(name)}: {' | '.join(parts)}")
    return "; ".join(out)


def _record(raw: dict, toxicity: str, resistance: str, uses: list[dict]) -> dict:
    name = _text(raw.get("ten_thuoc"))
    rec = {
        "source_id": _int(raw.get("id")),
        "trade_name": name,
        "trade_key": trade_key(name),
        "active_ingredient": _text(raw.get("hoat_chat")),
        "pest_group": _text(raw.get("phan_nhom")),
        "registrant": _text(raw.get("cong_ty_dang_ky")),
        "sector": _text(raw.get("linh_vuc")),
        "status": _status(raw.get("tinh_trang")),
        "concentration": _text(raw.get("ham_luong")),
        "registration_no": _text(raw.get("so_dang_ky")),
        "registered_on": _date(raw.get("ngay_cap")),
        "expires_on": _date(raw.get("ngay_het_han")),
        "toxicity": toxicity,
        "resistance": resistance,
        "source_url": _http_url(raw.get("url")),
        #  Câu mô tả của trang nguồn — cả tệp JSON lẫn sheet Excel đều có cột này (duoc-CR-495).
        "summary": _text(raw.get("tom_tat_su_dung")),
    }
    for key, size in _LIMITS.items():
        rec[key] = rec[key][:size]
    rec["resistance"] = rec["resistance"][:_TEXT_LIMIT]
    rec["summary"] = rec["summary"][:_TEXT_LIMIT]
    rec["uses"] = uses[:MAX_USES_PER_RECORD]
    return rec


def _use(raw: dict) -> dict:
    use = {"crop": _text(raw.get("cay_trong")), "pest": _text(raw.get("dich_hai")),
           "dosage": _text(raw.get("lieu_luong")),
           "pre_harvest_interval": _text(raw.get("thoi_gian_cach_ly")),
           "usage": _text(raw.get("cach_dung"))}
    for key, size in _USE_LIMITS.items():
        use[key] = use[key][:size]
    use["usage"] = use["usage"][:_TEXT_LIMIT]
    return use


def _http_url(value) -> str:
    """Chỉ nhận `http(s)://` — giá trị này thành `<a href>` trên màn, không để lọt `javascript:`."""
    url = _text(value)
    return url if url.lower().startswith(("http://", "https://")) else ""


def _too_many(what: str, cap: int) -> ValueError:
    return ValueError(f"Tệp có hơn {cap:,} {what} — quá trần, có vẻ nhầm tệp".replace(",", "."))


def _int(value) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def read_json(content: bytes) -> list[dict]:
    try:
        data = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, ValueError) as exc:
        raise ValueError("Tệp JSON không đọc được — cần đúng tệp thuoc-bvtv.json của bản cào") from exc
    if not isinstance(data, list):
        raise ValueError("Tệp JSON phải là một mảng thuốc (thuoc-bvtv.json)")
    if len(data) > MAX_RECORDS:
        raise _too_many("thuốc", MAX_RECORDS)
    records = []
    for raw in data:
        if not isinstance(raw, dict):
            continue
        toxicity = "; ".join(
            f"{_text(t.get('he'))} {_text(t.get('nhom'))} ({_text(t.get('mo_ta'))})"
            for t in raw.get("nhom_doc") or [] if isinstance(t, dict) and _text(t.get("nhom")))
        qltk = raw.get("quan_ly_tinh_khang") or {}
        resistance = _resistance([
            (a.get("ten"), [a.get("ma"), a.get("nhom"), a.get("phuong_thuc")])
            for a in (qltk.get("hoat_chat") or []) if isinstance(a, dict)])
        uses = [_use(u) for u in raw.get("pham_vi_su_dung") or [] if isinstance(u, dict)]
        records.append(_record(raw, toxicity[:500], resistance, uses))
    return _validate(records)


def read_xlsx(content: bytes) -> list[dict]:
    records, _mode = _read_xlsx(content)
    return records


def _read_xlsx(content: bytes) -> tuple[list[dict], str]:
    from openpyxl import load_workbook

    from . import pesticide_export_marker as marker
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl ném đủ loại lỗi cho tệp hỏng
        raise ValueError("Tệp Excel không mở được") from exc
    try:
        if LIST_SHEET not in wb.sheetnames:
            raise ValueError(f"Tệp Excel thiếu sheet «{LIST_SHEET}» — cần đúng tệp thuoc-bvtv.xlsx")
        #  Thiếu sheet phạm vi thì TỪ CHỐI, đừng nạp: nạp là thay toàn bộ, nên tệp thiếu sheet
        #  lặng lẽ xóa sạch mười lăm nghìn dòng phạm vi đang có (review 29/09).
        if USE_SHEET not in wb.sheetnames:
            raise ValueError(f"Tệp Excel thiếu sheet «{USE_SHEET}» — cần đúng tệp thuoc-bvtv.xlsx")
        mode = marker.detect_mode(wb)
        uses: dict[int, list[tuple[int, dict]]] = defaultdict(list)
        for n, row in enumerate(_sheet_dicts(wb[USE_SHEET]), start=1):
            if n > MAX_USE_ROWS:
                raise _too_many("dòng phạm vi sử dụng", MAX_USE_ROWS)
            rid = _int(row.get("id"))
            if rid <= 0:
                continue   # C2 — phạm vi của thuốc thủ công / thuốc nguồn trùng source_id lúc xuất
            uses[rid].append((_int(row.get("stt_pham_vi")), _use(row)))
        records = []
        seen_ids: set[int] = set()
        for n, row in enumerate(_sheet_dicts(wb[LIST_SHEET]), start=1):
            if n > MAX_RECORDS:
                raise _too_many("thuốc", MAX_RECORDS)
            rid = _int(row.get("id"))
            #  C2 (review 02/10/2026) — id <= 0 = thuốc thủ công (`export_id` xuất `-p.id`) HOẶC
            #  thuốc nguồn trùng `source_id` với thuốc khác lúc xuất (xuất id ÂM để tránh đụng
            #  khóa nối hai sheet) — cả hai dòng BỊ BỎ lúc nạp, không đụng tới thuốc khác.
            if rid <= 0:
                continue
            if rid in seen_ids:
                raise ValueError(
                    f"Tệp có id {rid} trùng nhau ở sheet «{LIST_SHEET}» — kiểm lại tệp trước khi nạp")
            seen_ids.add(rid)
            #  Ô nhóm kháng: «Tên: mã | nhóm | phương thức; Tên 2: …» → tách lại rồi dọn phần rỗng.
            items = []
            for chunk in _text(row.get("quan_ly_tinh_khang")).split(";"):
                #  Tách ở «: » chứ không ở «:» — nguồn có tên hoạt chất dính cả URL (`https://…`, id 3670).
                name, _, rest = chunk.partition(": ")
                items.append((name, rest.split("|")))
            raw_resistance = _text(row.get(RAW_RESISTANCE_COLUMN))
            resistance = raw_resistance if raw_resistance else _resistance(items)
            own = [u for _, u in sorted(uses.get(rid, []), key=lambda x: x[0])]
            records.append(_record(row, _text(row.get("nhom_doc"))[:500], resistance, own))
        return _validate(records), mode
    finally:
        wb.close()


def _sheet_dicts(ws):
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows, [])]
    for values in rows:
        if values and any(v not in (None, "") for v in values):
            yield dict(zip(header, values))


def _validate(records: list[dict]) -> list[dict]:
    records = [r for r in records if r["trade_name"]]
    if not records:
        raise ValueError("Tệp không có thuốc nào (thiếu cột tên thuốc `ten_thuoc`?)")
    if len(records) > MAX_RECORDS:
        raise _too_many("thuốc", MAX_RECORDS)
    use_total = sum(len(r["uses"]) for r in records)
    if use_total > MAX_USE_ROWS:
        raise _too_many("dòng phạm vi sử dụng", MAX_USE_ROWS)
    #  Nạp là thay toàn bộ: tệp không có dòng phạm vi nào là tệp hỏng / nhầm, không phải
    #  «danh mục mới không có phạm vi» — nạp vào là xóa sạch phạm vi đang có.
    if not use_total:
        raise ValueError("Tệp không có dòng phạm vi sử dụng nào — kiểm lại tệp trước khi nạp")
    return records


def read_file(filename: str, content: bytes) -> list[dict]:
    records, _mode = read_file_with_mode(filename, content)
    return records


def read_file_with_mode(filename: str, content: bytes) -> tuple[list[dict], str]:
    """(bản ghi, chế độ nạp) — `"replace"` (thay toàn bộ, như trước giờ) hay `"merge"` (C1,
    chỉ cập nhật đúng thuốc có trong tệp — xem `pesticide_export_marker.detect_mode`).

    Tệp JSON LUÔN `"replace"`: bản cào gốc không xuất JSON kèm khái niệm phạm vi trang."""
    name = (filename or "").lower()
    if name.endswith(".json"):
        return read_json(content), "replace"
    if name.endswith(".xlsx"):
        return _read_xlsx(content)
    raise ValueError("Chỉ nhận tệp .json hoặc .xlsx của bản cào danh mục thuốc BVTV")
