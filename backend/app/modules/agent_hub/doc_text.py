"""Đọc chữ trong tệp văn phòng (ai-CR-105) — pdf, Word, Excel, văn bản thường — để bot đọc báo cáo người dùng gửi
riêng hoặc tệp gửi trong nhóm. Chỉ đọc chữ; ảnh đi đường ảnh, ghi âm / video đi đường biên bản họp.
"""
from __future__ import annotations

import io
import logging
import subprocess
import tempfile
import time
from pathlib import Path

log = logging.getLogger("app.agent_hub.doc_text")

MAX_CHARS = 60_000
#  ai-CR-115: 5 → 30 trang tính (báo cáo nhân sự đại ca gửi 08/10 có 11 trang, bot chỉ thấy 5).
XLSX_SHEETS = 30
XLSX_ROWS = 300
#  Tự tính công thức (pycel) chỉ với tệp ≤ 5 MB, tối đa 5.000 ô / 10 giây — tệp lớn đọc chế độ chỉ đọc như cũ.
FORMULA_FILE_MAX = 5 * 1024 * 1024
FORMULA_EVAL_MAX = 5000
FORMULA_EVAL_SEC = 10

READABLE_EXT = {".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".csv", ".md", ".doc"}
READABLE_MIME = {
    "application/pdf", "text/plain", "text/csv", "text/markdown", "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel.sheet.macroenabled.12",
}


class DocTextError(ValueError):
    """Không đọc được tệp — câu nói được cho người dùng."""


def readable(name: str, mime: str) -> bool:
    return (mime or "") in READABLE_MIME or Path(name or "").suffix.lower() in READABLE_EXT


def extract(name: str, mime: str, data: bytes) -> str:
    ext = Path(name or "").suffix.lower()
    if ext == ".pdf" or mime == "application/pdf" or data[:4] == b"%PDF":
        text = _pdf(data)
    elif ext == ".docx" or mime.endswith("wordprocessingml.document"):
        text = _docx(data)
    elif ext in (".xlsx", ".xlsm") or "spreadsheetml" in (mime or "") or "macroenabled" in (mime or ""):
        text = _xlsx(data)
    elif ext == ".doc" or mime == "application/msword":
        text = _doc(data)
    else:
        text = data.decode("utf-8", errors="replace")
    text = "\n".join(line.rstrip() for line in text.splitlines() if line.strip())
    if not text.strip():
        raise DocTextError("tệp không có chữ đọc được (có thể là bản scan — gửi dạng ảnh để em đọc bằng mắt)")
    return text[:MAX_CHARS] + ("\n[… tệp dài, em chỉ đọc phần đầu]" if len(text) > MAX_CHARS else "")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:200])
    except Exception as e:  # noqa: BLE001 — pdf hỏng / có mật khẩu
        raise DocTextError(f"không mở được PDF ({type(e).__name__})") from None


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise DocTextError(f"không mở được tệp Word ({type(e).__name__})") from None
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _cell_text(value) -> str:
    """Giá trị ô → chữ gọn: số thực làm tròn 4 chữ số lẻ, ngày giờ dạng ISO, rỗng thì rỗng."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{round(value, 4)}"
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ") if hasattr(value, "hour") and (value.hour or value.minute) else \
            value.isoformat()[:10]
    return str(value)


def _eval_formulas(data: bytes, cells: list[tuple[str, str]]) -> dict[tuple[str, str], object]:
    """Tự tính các ô CÔNG THỨC mà tệp không lưu sẵn kết quả (ai-CR-115).

    Tệp do phần mềm / script sinh ra (chưa từng mở bằng Excel rồi lưu) chỉ có công thức, không có giá trị tính sẵn —
    openpyxl đọc ra ô trống, bot tưởng tệp «chưa có số» (đại ca gửi báo cáo nhân sự 08/10: biên lợi nhuận, ROA, % hoàn
    thành… đều trống). Dùng pycel tính lại; ô nào tính hỏng thì bỏ, để chỗ gọi in công thức gốc.
    """
    if not cells:
        return {}
    try:
        from pycel import ExcelCompiler
    except ImportError:
        return {}
    out: dict[tuple[str, str], object] = {}
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
        fh.write(data)
        path = fh.name
    try:
        xc = ExcelCompiler(filename=path)
        deadline = time.monotonic() + FORMULA_EVAL_SEC
        for sheet, coord in cells[:FORMULA_EVAL_MAX]:
            if time.monotonic() > deadline:
                break
            try:
                value = xc.evaluate(f"{sheet}!{coord}")
            except Exception:  # noqa: BLE001 — hàm lạ / vòng tham chiếu: bỏ ô đó
                continue
            if value is not None and not (isinstance(value, str) and value.startswith("#")):
                out[(sheet, coord)] = value
    except Exception as e:  # noqa: BLE001 — pycel không dựng được sổ: in công thức gốc
        log.info("doc_text: không tự tính được công thức: %s", str(e)[:200])
    finally:
        Path(path).unlink(missing_ok=True)
    return out


def _xlsx(data: bytes) -> str:
    """Mọi trang tính (tối đa XLSX_SHEETS), mỗi dòng có SỐ DÒNG ở đầu để model lần theo công thức. Ô công thức chưa có kết
    quả lưu sẵn thì tự tính; tính không được thì in «=công thức»."""
    from openpyxl import load_workbook

    small = len(data) <= FORMULA_FILE_MAX
    try:
        values = load_workbook(io.BytesIO(data), read_only=not small, data_only=True)
        formulas = load_workbook(io.BytesIO(data), read_only=False, data_only=False) if small else None
    except Exception as e:  # noqa: BLE001
        raise DocTextError(f"không mở được tệp Excel ({type(e).__name__})") from None
    pending: dict[tuple[str, str], str] = {}
    if formulas is not None:
        for ws in formulas.worksheets[:XLSX_SHEETS]:
            vws = values[ws.title]
            for row in ws.iter_rows(max_row=XLSX_ROWS):
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and vws[c.coordinate].value is None:
                        pending[(ws.title, c.coordinate)] = c.value
    computed = _eval_formulas(data, list(pending))
    out = []
    if pending:
        out.append(f"(Tệp lưu {len(pending)} ô công thức chưa có kết quả tính sẵn; em đã tự tính {len(computed)} ô, ô "
                   "nào còn dạng «=…» là công thức chưa tính được.)")
    sheets = values.worksheets
    for ws in sheets[:XLSX_SHEETS]:
        out.append(f"## Trang tính: {ws.title}")
        for i, row in enumerate(ws.iter_rows()):
            if i >= XLSX_ROWS:
                out.append("[… còn nữa]")
                break
            cells = []
            for c in row:
                v = getattr(c, "value", None)
                key = (ws.title, getattr(c, "coordinate", ""))
                if v is None and key in pending:
                    v = computed.get(key, pending[key])
                cells.append(_cell_text(v))
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                out.append(f"{i + 1}: " + " | ".join(cells))
    if len(sheets) > XLSX_SHEETS:
        out.append(f"[… còn {len(sheets) - XLSX_SHEETS} trang tính nữa em chưa đọc]")
    return "\n".join(out)


def _doc(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as fh:
        fh.write(data)
        path = fh.name
    try:
        proc = subprocess.run(["antiword", path], capture_output=True, text=True, timeout=60)
    except (FileNotFoundError, subprocess.TimeoutExpired):
        raise DocTextError("máy chủ chưa đọc được tệp .doc đời cũ — lưu lại thành .docx giúp em") from None
    finally:
        Path(path).unlink(missing_ok=True)
    if proc.returncode != 0:
        raise DocTextError("không đọc được tệp .doc này")
    return proc.stdout
