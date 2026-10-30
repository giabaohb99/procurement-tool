"""Phần dùng chung của hai tệp Excel (Sổ quản lý lao động, mẫu 01/PLI) — duoc-CR-687.

Mỗi pháp nhân MỘT sheet. Người chưa gắn pháp nhân vào sheet "Chưa gắn pháp nhân". Chuỗi người nhập tay
(tên, địa chỉ…) luôn ghi ở dạng chuỗi để Excel không hiểu nhầm thành công thức.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from app.core.export_xlsx import _FORMULA_TRIGGER_CHARS  # cùng bộ ký tự mở công thức với mọi tệp xuất khác
from app.core.vn_time import vn_today
from app.modules.company.model import Company

from .people_data import PersonRow, load_people

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
UNASSIGNED = "Chưa gắn pháp nhân"


@dataclass(frozen=True)
class CompanyInfo:
    id: int
    name: str
    tax_code: str
    address: str
    short_name: str = ""      # tên ngắn: ưu tiên làm tên sheet (tên đầy đủ dễ bị cắt trùng ở 31 ký tự)

    @property
    def sheet_label(self) -> str:
        return self.short_name or self.name


def load_export_people(employees_q, contracts_q, period, company_id, department_id) -> list[PersonRow]:
    people = load_people(employees_q, contracts_q, period.date_from, period.date_to, vn_today(),
                         company_id, department_id)
    people.sort(key=lambda r: r.employee_code)
    return people


def split_by_company(db: Session, people: list[PersonRow]) -> list[tuple[CompanyInfo, list[PersonRow]]]:
    """Nhóm người theo pháp nhân, xếp theo tên pháp nhân; 2 truy vấn cố định (pháp nhân · phòng ban ở nơi khác)."""
    by_co: dict[int, list[PersonRow]] = {}
    for p in people:
        by_co.setdefault(p.company_id, []).append(p)
    info = {c.id: CompanyInfo(c.id, c.name or "", c.tax_code or "", c.address or "",
                                (c.short_name or "").strip())
            for c in db.query(Company.id, Company.name, Company.tax_code, Company.address, Company.short_name)
            .filter(Company.id.in_([k for k in by_co if k])).all()} if by_co else {}
    out = [(info.get(cid) or CompanyInfo(cid, UNASSIGNED, "", ""), rows) for cid, rows in by_co.items()]
    #  Người chưa gắn pháp nhân (hoặc pháp nhân không còn trong bảng) xếp CUỐI.
    return sorted(out, key=lambda t: (t[0].id not in info, t[0].name))


def unique_sheet_name(name: str, used: set[str]) -> str:
    """Tên sheet hợp lệ của Excel (≤ 31 ký tự, không có `[]:*?/\\`), không trùng sheet đã có."""
    #  Excel cấm tên sheet bắt đầu / kết thúc bằng dấu nháy đơn.
    base = (_BAD_SHEET_CHARS.sub(" ", name or "").strip().strip("'").strip() or "Sheet")[:28].rstrip("' ") or "Sheet"
    cand, n = base, 1
    while cand.lower() in used:
        n += 1
        cand = f"{base[:27 - len(str(n))]} {n}"
    used.add(cand.lower())
    return cand


def text_cell(ws, row: int, col: int, value, **style):
    """Ghi ô chữ; chuỗi mở đầu bằng = + - @ (hay tab / xuống dòng) bị ép về kiểu chuỗi."""
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] in _FORMULA_TRIGGER_CHARS:
        c.data_type = "s"
    for k, v in style.items():
        setattr(c, k, v)
    return c


def money_text(n: int) -> str:
    """1234567 -> "1.234.567" (kiểu viết số tiền của Việt Nam)."""
    return f"{int(n):,}".replace(",", ".")


def xlsx_download(wb: Workbook, filename: str) -> Response:
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'})


def empty_workbook(message: str) -> Workbook:
    """Sổ trống khi kỳ không có ai — Excel buộc có ít nhất một sheet."""
    wb = Workbook()
    wb.active.title = "Không có dữ liệu"
    wb.active["A1"] = message
    return wb


def period_label(period) -> str:
    return f"từ {period.date_from:%d/%m/%Y} đến {period.date_to:%d/%m/%Y}"
