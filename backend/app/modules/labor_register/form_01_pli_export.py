"""Tệp Excel BÁO CÁO TÌNH HÌNH SỬ DỤNG LAO ĐỘNG — Mẫu số 01/PLI (duoc-CR-687, khóa 20), mỗi pháp nhân một sheet.

Bố cục: tên pháp nhân + quốc hiệu · tiêu đề · «Kính gửi: Sở Nội vụ» · mục 1 thông tin chung · bảng 27 cột
(tiêu đề 3 tầng gộp ô, dòng đánh số 1…27) · dòng Tổng · khối ký · chú thích (2). Chi tiết cột ở
`form_01_pli_layout.py`; ô không có dữ liệu trong ERP để trống cho Nhân sự điền tay.
"""
from __future__ import annotations

from collections import Counter

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.modules.employee.profile_gaps_report_data import load_group_names

from . import export_common as xc
from . import form_01_pli_layout as lay
from .people_data import PersonRow, is_civil_contract
from .report_service import civil_excluded_note

LAST = get_column_letter(lay.NUM_COLS)
HEAD_ROW = 15                    # dòng đầu của tiêu đề 3 tầng (dòng 1–14 là khối đầu mẫu)


def _merged(ws, r1: int, c1: int, r2: int, c2: int, text, **style):
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = xc.text_cell(ws, r1, c1, text, **style)
    return cell


def _top_block(ws, co: xc.CompanyInfo, period, excluded: int = 0) -> None:
    _merged(ws, 1, 1, 1, 8, f"Tên doanh nghiệp: {co.name}", font=xc.BOLD)
    _merged(ws, 1, 16, 1, 27, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", font=xc.BOLD, alignment=xc.CENTER)
    _merged(ws, 2, 16, 2, 27, "Độc lập - Tự do - Hạnh phúc", font=Font(bold=True, underline="single"),
            alignment=xc.CENTER)
    _merged(ws, 3, 1, 3, 4, "Mẫu số 01/PLI", font=Font(italic=True))
    _merged(ws, 5, 1, 5, 27, "BÁO CÁO TÌNH HÌNH SỬ DỤNG LAO ĐỘNG", font=Font(bold=True, size=14), alignment=xc.CENTER)
    _merged(ws, 6, 1, 6, 27, f"(Kỳ báo cáo: {xc.period_label(period)})", font=Font(italic=True), alignment=xc.CENTER)
    _merged(ws, 7, 1, 7, 27, "Kính gửi: Sở Nội vụ ……………………………………", font=xc.BOLD, alignment=xc.CENTER)
    lines = ["1. Thông tin chung", f"- Tên doanh nghiệp, cơ quan, tổ chức: {co.name}", f"- Địa chỉ: {co.address}",
             f"- Mã số thuế / mã số doanh nghiệp: {co.tax_code}",
             "- Điện thoại: ……………  Fax: ……………  Email: ……………  Website: ……………",
             "- Ngành, nghề kinh doanh chính: ……………………………………", "2. Tình hình sử dụng lao động"]
    for i, text in enumerate(lines):
        _merged(ws, 8 + i, 1, 8 + i, 27, text, font=Font(bold=text[0] in "12"))
    if excluded:
        _merged(ws, 4, 1, 4, 27, civil_excluded_note(excluded), font=Font(italic=True, color="9C5700"))


def _header(ws) -> None:
    for t1, c1, t2, c2, text in lay.HEADER_CELLS:
        _merged(ws, HEAD_ROW + t1, c1, HEAD_ROW + t2, c2, text, font=xc.BOLD, fill=xc.HEAD_FILL, alignment=xc.CENTER)
    for r in range(HEAD_ROW, HEAD_ROW + 3):
        for c in range(1, lay.NUM_COLS + 1):
            ws.cell(row=r, column=c).border = xc.BORDER
            ws.cell(row=r, column=c).fill = xc.HEAD_FILL
    for c in range(1, lay.NUM_COLS + 1):
        xc.text_cell(ws, HEAD_ROW + 3, c, c, font=Font(italic=True), alignment=xc.CENTER, border=xc.BORDER)
        ws.column_dimensions[get_column_letter(c)].width = lay.COLUMN_WIDTHS.get(c, 12)
    ws.row_dimensions[HEAD_ROW + 1].height = 30
    ws.row_dimensions[HEAD_ROW + 2].height = 45


def _body(ws, rows: list[PersonRow], dept_name: dict[int, str]) -> int:
    """Ghi từng người và dòng Tổng; trả số dòng kế tiếp dòng Tổng."""
    first = HEAD_ROW + 4
    for n, r in enumerate(rows, start=1):
        for c in range(1, lay.NUM_COLS + 1):
            cell = ws.cell(row=first + n - 1, column=c)
            cell.border = xc.BORDER
            cell.alignment = xc.WRAP
        for c, v in lay.row_values(n, r, dept_name.get(r.department_id, "")).items():
            cell = xc.text_cell(ws, first + n - 1, c, v, border=xc.BORDER, alignment=xc.WRAP)
            if hasattr(v, "year"):
                cell.number_format = "dd/mm/yyyy"
            elif isinstance(v, int) and c in (12, 16):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="top")
    total_row = first + len(rows)
    _merged(ws, total_row, 1, total_row, 7, f"Tổng ({len(rows)} lao động)", font=xc.BOLD, alignment=xc.CENTER)
    for c in range(8, 12):
        ws.cell(row=total_row, column=c, value=sum(1 for r in rows if lay.JOB_CATEGORY_COLUMN[r.job_category] == c))
    for c in range(1, lay.NUM_COLS + 1):
        ws.cell(row=total_row, column=c).border = xc.BORDER
        ws.cell(row=total_row, column=c).font = xc.BOLD
        if 8 <= c <= 11:
            ws.cell(row=total_row, column=c).alignment = xc.CENTER
    return total_row + 1


def _footer(ws, start: int) -> None:
    _merged(ws, start + 1, 20, start + 1, 27, "Ngày ….. tháng ….. năm …..", font=Font(italic=True), alignment=xc.CENTER)
    _merged(ws, start + 2, 20, start + 2, 27, "ĐẠI DIỆN DOANH NGHIỆP, CƠ QUAN, TỔ CHỨC", font=xc.BOLD, alignment=xc.CENTER)
    _merged(ws, start + 3, 20, start + 3, 27, "(Chữ ký, dấu)", font=Font(italic=True), alignment=xc.CENTER)
    for i, text in enumerate((lay.NOTE_JOB_CATEGORY, lay.BLANK_NOTE)):
        _merged(ws, start + 8 + i, 1, start + 8 + i, 19, text, alignment=xc.WRAP)
        ws.row_dimensions[start + 8 + i].height = 45


def export_form_01_pli(db: Session, employees_q, contracts_q, period, company_id, department_id) -> Response:
    everyone = xc.load_export_people(employees_q, contracts_q, period, company_id, department_id)
    people = [p for p in everyone if not is_civil_contract(p)]
    excluded = Counter(p.company_id for p in everyone if is_civil_contract(p))
    if not people:
        wb = xc.empty_workbook(" ".join(["Không có nhân sự nào làm việc trong kỳ đã chọn.", civil_excluded_note(len(everyone))]).strip())
    else:
        dept_name = load_group_names(db, None, always=True)[1]
        wb, used = Workbook(), set()
        wb.remove(wb.active)
        for co, rows in xc.split_by_company(db, people):
            ws = wb.create_sheet(xc.unique_sheet_name(co.sheet_label, used))
            _top_block(ws, co, period, excluded[co.id])
            _header(ws)
            _footer(ws, _body(ws, rows, dept_name))
    return xc.xlsx_download(wb, f"bao-cao-su-dung-lao-dong-01-pli-{period.date_from}-{period.date_to}")
