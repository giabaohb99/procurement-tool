"""Tệp Excel SỔ QUẢN LÝ LAO ĐỘNG (duoc-CR-687, khóa 19) — mỗi pháp nhân một sheet, mỗi người một dòng.

Cột theo Điều 3 khoản 2 Nghị định 145/2020. Cột ERP chưa có dữ liệu (quốc tịch, bậc kỹ năng nghề, giờ làm
thêm, đào tạo, kỷ luật, tai nạn lao động) để TRỐNG cho Nhân sự điền tay. Người thiếu thông tin bắt buộc
được tô vàng ở STT + Họ tên và liệt kê ở cột cuối.
"""
from __future__ import annotations

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.modules.employee.constants import EDUCATION_LEVEL_LABELS, GENDER_LABELS

from . import export_common as xc
from .leave_days import load_leave_days
from .people_data import PersonRow, hire_fallback_note
from .report_service import contract_type_label

#  (tiêu đề, độ rộng) theo thứ tự cột của sổ; cột cuối là cột đánh dấu thiếu thông tin.
COLUMNS = [("STT", 6), ("Mã NV", 11), ("Họ và tên", 24), ("Giới tính", 9), ("Ngày sinh", 12), ("Quốc tịch", 11),
           ("Nơi cư trú", 34), ("Số CCCD", 15), ("Trình độ chuyên môn", 16), ("Bậc trình độ kỹ năng nghề", 14),
           ("Vị trí việc làm (chức vụ)", 22), ("Loại hợp đồng lao động", 18), ("Thời điểm bắt đầu làm việc", 14),
           ("Số sổ BHXH", 15), ("Tiền lương (theo HĐ)", 15), ("Nâng lương", 34),
           ("Số ngày nghỉ có lương trong kỳ", 12), ("Số ngày nghỉ không lương trong kỳ", 12),
           ("Số giờ làm thêm", 10), ("Học nghề, đào tạo, bồi dưỡng", 18),
           ("Kỷ luật lao động, trách nhiệm vật chất", 18), ("Tai nạn lao động, bệnh nghề nghiệp", 18),
           ("Thời điểm chấm dứt HĐLĐ", 14), ("Lý do chấm dứt HĐLĐ", 22), ("Thiếu thông tin bắt buộc", 26)]
HEAD_ROW = 7
DATE_COLS = (5, 13, 23)
NUMBER_COLS = (15, 17, 18)


def raise_history(r: PersonRow) -> str:
    """«dd/mm/yyyy: mức lương → …» — chỉ khi người có từ hai mức lương trở lên."""
    if len(r.salary_history) < 2:
        return ""
    return " → ".join(f"{d:%d/%m/%Y}: {xc.money_text(s)}" for d, s in r.salary_history)


def _person_values(n: int, r: PersonRow, leave: dict | None) -> list:
    #  `leave is None` = người xuất không có quyền xem đơn nghỉ: để TRỐNG chứ không ghi 0.
    paid, unpaid = (None, None) if leave is None else leave.get(r.employee_id, (0, 0))
    return [n, r.employee_code, r.employee_name, GENDER_LABELS.get(r.gender, ""), r.date_of_birth, "",
            r.address, r.id_number, EDUCATION_LEVEL_LABELS.get(r.education_level, ""), "", r.position,
            contract_type_label(r.contract_type), None if r.hire_missing else r.hire_date,
            r.social_insurance_no, r.base_salary or None, raise_history(r), paid, unpaid, "", "", "", "",
            r.contract_terminated or r.resign_date, r.terminate_reason, ", ".join(r.missing)]


NO_LEAVE_NOTE = ("Bạn chưa có quyền xem đơn nghỉ phép nên hai cột số ngày nghỉ để trống — "
                 "không phải người nghỉ 0 ngày.")


def _fill_sheet(ws, co: xc.CompanyInfo, rows: list[PersonRow], period, leave: dict | None) -> None:
    xc.text_cell(ws, 1, 1, "SỔ QUẢN LÝ LAO ĐỘNG", font=Font(bold=True, size=14))
    xc.text_cell(ws, 2, 1, f"Đơn vị sử dụng lao động: {co.name}")
    xc.text_cell(ws, 3, 1, f"Mã số thuế: {co.tax_code}")
    xc.text_cell(ws, 4, 1, f"Địa chỉ: {co.address}")
    xc.text_cell(ws, 5, 1, f"Kỳ: {xc.period_label(period)}")
    notes = [n for n in (hire_fallback_note(rows), NO_LEAVE_NOTE if leave is None else "") if n]
    xc.text_cell(ws, 6, 1, " ".join(notes), font=Font(italic=True, color="9C5700"))
    for i, (title, width) in enumerate(COLUMNS, start=1):
        xc.text_cell(ws, HEAD_ROW, i, title, font=xc.BOLD, fill=xc.HEAD_FILL, alignment=xc.CENTER, border=xc.BORDER)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[HEAD_ROW].height = 48
    for n, r in enumerate(rows, start=1):
        row = HEAD_ROW + n
        for col, v in enumerate(_person_values(n, r, leave), start=1):
            c = xc.text_cell(ws, row, col, v, border=xc.BORDER, alignment=xc.WRAP)
            if col in DATE_COLS:
                c.number_format = "dd/mm/yyyy"
            elif col in NUMBER_COLS:
                c.number_format = "#,##0.##"
                c.alignment = Alignment(horizontal="right", vertical="top")
        if r.missing:
            ws.cell(row=row, column=1).fill = xc.WARN_FILL
            ws.cell(row=row, column=3).fill = xc.WARN_FILL
    ws.freeze_panes = ws.cell(row=HEAD_ROW + 1, column=4)


def export_labor_book(db: Session, employees_q, contracts_q, period, company_id, department_id,
                      leave_q=None) -> Response:
    people = xc.load_export_people(employees_q, contracts_q, period, company_id, department_id)
    if not people:
        wb = xc.empty_workbook("Không có nhân sự nào làm việc trong kỳ đã chọn.")
    else:
        leave = (load_leave_days(db, leave_q, [p.employee_id for p in people], period.date_from, period.date_to)
                 if leave_q is not None else None)
        wb, used = Workbook(), set()
        wb.remove(wb.active)
        for co, rows in xc.split_by_company(db, people):
            _fill_sheet(wb.create_sheet(xc.unique_sheet_name(co.sheet_label, used)), co, rows, period, leave)
    return xc.xlsx_download(wb, f"so-quan-ly-lao-dong-{period.date_from}-{period.date_to}")
