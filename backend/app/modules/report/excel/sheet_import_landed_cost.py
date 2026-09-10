"""Hai sheet Excel của báo cáo giá vốn lô hàng nhập khẩu (bao-CR-347).

* **GIA VON THEO LO** xoay dọc so với mọi báo cáo còn lại: mỗi ĐƠN là một CỘT, mỗi chỉ tiêu
  là một DÒNG. Vì thế không tái dùng được `build_pivot_sheet` — bảng kia luôn là thực thể
  theo dòng × tháng theo cột.
* **GIA VON THEO DONG HANG** là bảng ngang bình thường: mỗi mã hàng một dòng, chi phí của
  lô đã chia về tới dòng đó.

Tên các loại chi phí do `data["cost_types"]` quyết định — chỉ loại có phát sinh mới có
dòng/cột, nên bố cục đổi theo lần lọc.
"""
from openpyxl.utils import get_column_letter

from . import styles as S

INT = S.FMT_INT
_FMT_RATE = "#,##0.####"     # tỷ giá và giá nguyên tệ giữ phần lẻ
_FMT_QTY = "#,##0.###"


def _currency(o: dict):
    return "Nhiều loại" if o.get("currency_mixed") else o.get("currency", "")


def _goods_amount(o: dict):
    """Tiền hàng nguyên tệ. Cột TỔNG của nhóm đơn nhiều đồng tiền để TRỐNG chứ không cộng
    yên với đô lại thành một con số không thuộc đồng tiền nào."""
    return "" if o.get("currency_mixed") else o.get("goods_amount", 0)


def _label_col(ws, row: int, no: str, label: str, *, bold=False):
    ws.cell(row=row, column=1, value=no).alignment = S.CENTER
    c = ws.cell(row=row, column=2, value=label)
    c.alignment = S.LEFT
    if bold:
        c.font = S.FONT_TOTAL
    for col in (1, 2):
        ws.cell(row=row, column=col).border = S.BORDER


def _write_row(ws, row: int, values: list, *, fmt=INT, numeric=True, total=False):
    """Ghi một dòng chỉ tiêu vào các cột đơn hàng (bắt đầu từ cột 3)."""
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=3 + i, value=v)
        S.style_body_cell(c, numeric=numeric, total=total)
        if numeric:
            c.number_format = fmt


def build_import_landed_cost_sheet(ws, data: dict) -> None:
    """Sheet theo LÔ HÀNG — mỗi đơn một cột, cột TỔNG đứng cuối."""
    cost_types = data.get("cost_types", [])
    orders = data.get("orders", [])
    totals = data.get("totals", {})
    # Cột TỔNG đứng cuối như một "đơn" ảo để dùng chung mọi hàm ghi dòng bên dưới.
    columns = [{"title": o["code"] or f"#{o['po_id']}", "row": o} for o in orders]
    if orders:
        columns.append({"title": "TỔNG", "row": totals, "is_total": True})

    ws.cell(row=1, column=1, value="BÁO CÁO GIÁ VỐN LÔ HÀNG NHẬP KHẨU").font = S.FONT_TITLE
    ws.cell(row=2, column=1,
            value="Số liệu tính tại thời điểm xuất file.").font = S.FONT_DESC

    head = 4
    ws.cell(row=head, column=1, value="STT")
    ws.cell(row=head, column=2, value="CHỈ TIÊU")
    for col in (1, 2):
        S.style_header_cell(ws.cell(row=head, column=col))
    for i, c in enumerate(columns):
        S.style_header_cell(ws.cell(row=head, column=3 + i, value=c["title"]))

    r = head + 1
    # (nhãn, lấy số của đơn, định dạng — None là ô chữ, cột TỔNG có bỏ trống không).
    # Cột TỔNG không có ngày / nhà cung cấp / ghi chú của riêng nó — để trống chứ đừng bịa:
    # gộp ghi chú của mười đơn lại thành một ô là không ai đọc được.
    info_rows = [("Ngày đặt hàng", lambda o: o.get("order_date", ""), None, True),
                 ("Nhà cung cấp", lambda o: o.get("supplier_name", ""), None, True),
                 ("Trạng thái", lambda o: o.get("status_label", ""), None, True),
                 ("Ngày hàng rời cảng (ETD)", lambda o: o.get("etd_date", ""), None, True),
                 ("Số tờ khai hải quan", lambda o: o.get("customs_decl_no", ""), None, True),
                 ("Ngày tờ khai", lambda o: o.get("customs_decl_date", ""), None, True),
                 ("Ghi chú", lambda o: o.get("note", ""), None, True),
                 ("Số lượng", lambda o: o.get("qty_total", 0), _FMT_QTY, False),
                 ("Khối lượng (kg)", lambda o: o.get("weight_total", 0), _FMT_QTY, False),
                 ("Đồng tiền", _currency, None, False),
                 ("Tiền hàng (nguyên tệ)", _goods_amount, _FMT_RATE, False),
                 ("Tỷ giá", lambda o: o.get("exchange_rate", 0), _FMT_RATE, True)]
    for label, getter, fmt, blank_total in info_rows:
        _label_col(ws, r, "", label)
        for i, c in enumerate(columns):
            val = "" if c.get("is_total") and blank_total else getter(c["row"])
            cell = ws.cell(row=r, column=3 + i, value=val)
            S.style_body_cell(cell, numeric=isinstance(val, (int, float)))
            if isinstance(val, (int, float)) and fmt:
                cell.number_format = fmt
        r += 1

    _label_col(ws, r, "1", "Tiền hàng (vnd)", bold=True)
    _write_row(ws, r, [c["row"].get("goods_base", 0) for c in columns], total=True)
    r += 1

    _label_col(ws, r, "2", "Chi phí nhập khẩu (vnd)", bold=True)
    _write_row(ws, r, [c["row"].get("cost_total", 0) for c in columns], total=True)
    r += 1
    for g in cost_types:
        _label_col(ws, r, g["no"], g["label"])
        _write_row(ws, r, [c["row"].get("by_type", {}).get(str(g["code"]), 0) for c in columns])
        r += 1

    _label_col(ws, r, "", "Tổng giá vốn (vnd)", bold=True)
    _write_row(ws, r, [c["row"].get("landed_total", 0) for c in columns], total=True)
    r += 1
    _label_col(ws, r, "", "Giá vốn/Kg (vnd)", bold=True)
    _write_row(ws, r, [c["row"].get("price_per_kg", 0) for c in columns], total=True)
    r += 2

    _signatures(ws, r, step=3)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 34
    for i in range(len(columns)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 18
    ws.freeze_panes = ws.cell(row=head + 1, column=3)


def build_import_landed_cost_item_sheet(ws, data: dict) -> None:
    """Sheet theo DÒNG HÀNG — chi phí của lô đã chia về từng mã hàng."""
    cost_types = data.get("cost_types", [])
    rows = data.get("items", [])
    totals = data.get("item_totals", {})

    ws.cell(row=1, column=1, value="GIÁ VỐN NHẬP KHẨU THEO DÒNG HÀNG").font = S.FONT_TITLE
    ws.cell(row=2, column=1,
            value="Chi phí của lô được chia về từng mã hàng theo cách chia khai trên "
                  "từng khoản chi.").font = S.FONT_DESC

    # Mỗi cột khai bốn thứ một chỗ: (tiêu đề, lấy số của DÒNG, định dạng — None là ô chữ,
    # lấy số của dòng TỔNG — None là để trống). Khai kèm nhau để dòng tổng không bao giờ
    # trượt cột so với thân bảng khi thêm bớt loại chi phí.
    def _by_type(code: str):
        return lambda r: r.get("by_type", {}).get(code, 0)

    cols = [("Mã đơn", lambda r: r.get("code", ""), None, lambda t: "TỔNG"),
            ("Ngày ETD", lambda r: r.get("etd_date", ""), None, None),
            ("Mã hàng", lambda r: r.get("product_code", ""), None, None),
            ("Tên hàng", lambda r: r.get("product_name", ""), None, None),
            ("ĐVT", lambda r: r.get("unit", ""), None, None),
            ("Số lượng", lambda r: r.get("qty_order", 0), _FMT_QTY, lambda t: t.get("qty_order", 0)),
            ("Khối lượng (kg)", lambda r: r.get("weight_kg", 0), _FMT_QTY, lambda t: t.get("weight_kg", 0)),
            ("Tiền hàng (vnd)", lambda r: r.get("goods_base", 0), INT, lambda t: t.get("goods_base", 0))]
    for g in cost_types:
        key = str(g["code"])
        cols.append((g["label"], _by_type(key), INT, _by_type(key)))
    cols += [("Tổng chi phí (vnd)", lambda r: r.get("cost_base", 0), INT, lambda t: t.get("cost_base", 0)),
             ("Tổng giá vốn (vnd)", lambda r: r.get("landed_base", 0), INT, lambda t: t.get("landed_base", 0)),
             # Dòng TỔNG không có giá vốn/ĐVT: mỗi dòng một đơn vị tính khác nhau, cộng cái
             # kg với cái cái rồi chia ra là số không đọc được.
             ("Giá vốn/ĐVT (vnd)", lambda r: r.get("price_per_unit", 0), INT, None),
             ("Giá vốn/Kg (vnd)", lambda r: r.get("price_per_kg", 0), INT, lambda t: t.get("price_per_kg", 0))]

    head = 4
    S.style_header_cell(ws.cell(row=head, column=1, value="STT"))
    for i, (title, _g, _f, _t) in enumerate(cols):
        S.style_header_cell(ws.cell(row=head, column=2 + i, value=title))

    r = head + 1
    for n, row in enumerate(rows, start=1):
        c = ws.cell(row=r, column=1, value=n)
        S.style_body_cell(c, numeric=True)
        c.number_format = INT
        for i, (_t, getter, fmt, _tot) in enumerate(cols):
            cell = ws.cell(row=r, column=2 + i, value=getter(row))
            S.style_body_cell(cell, numeric=fmt is not None)
            if fmt:
                cell.number_format = fmt
        r += 1

    if rows:
        S.style_body_cell(ws.cell(row=r, column=1, value=""), total=True)
        for i, (_t, _g, fmt, total_of) in enumerate(cols):
            val = total_of(totals) if total_of else ""
            cell = ws.cell(row=r, column=2 + i, value=val)
            S.style_body_cell(cell, numeric=isinstance(val, (int, float)), total=True)
            if isinstance(val, (int, float)) and fmt:
                cell.number_format = fmt
        r += 2

    _signatures(ws, r, step=4)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 34
    for i in range(len(cols) - 4):
        ws.column_dimensions[get_column_letter(6 + i)].width = 16
    ws.freeze_panes = ws.cell(row=head + 1, column=2)


def _signatures(ws, row: int, *, step: int) -> None:
    """Ba ô ký tay: người lập · kế toán · quản lý duyệt."""
    for i, name in enumerate(["Người lập biểu", "Kế toán", "Quản lý duyệt"]):
        c = ws.cell(row=row, column=2 + i * step, value=name)
        c.font = S.FONT_TOTAL
        c.alignment = S.CENTER
        ws.cell(row=row + 1, column=2 + i * step, value="(Ký, ghi rõ họ tên)").font = S.FONT_DESC
