"""duoc-CR-688 — báo cáo A3 «Cơ cấu quản lý» (khóa 23): `/api/management-span/summary`.

Canh: định nghĩa từng chỉ số (staff · managers · direct_reports · avg_span · unassigned · span_one ·
span_wide · departures_under), chiều `manager`, phạm vi dữ liệu, `department_id` chỉ thu hẹp, gác khóa
báo cáo + quyền xuất ở CẢ ba đường, 422 và số truy vấn CỐ ĐỊNH. Phòng B.hc của thế giới mẫu không có
ai nên dựng riêng ở đó để số liệu cô lập.
"""
import uuid
from datetime import timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event

from app.core.auth import get_current_user, get_perm_profile
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.report_period import parse_period
from app.core.subject_match import SUBJECT_EMPLOYEE
from app.core.vn_time import vn_today
from app.main import app
from app.modules.employee import management_report_service as svc
from app.modules.employee.model import Employee
from app.modules.report_access.model import ReportAccess

TODAY = vn_today()
SUM = "/api/management-span/summary"
QS = {"preset": "custom", "date_from": (TODAY - timedelta(days=20)).isoformat(),
      "date_to": TODAY.isoformat(), "compare": "none"}


@pytest.fixture
def client_as(db):
    def make_client(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make_client
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def _emp(db, world, code, manager=None, **kw) -> Employee:
    base = dict(code=code, full_name=f"Người {code}", company_id=world.co["B"],
                department_id=world.dept["B.hc"], is_active=True, hire_date=TODAY - timedelta(days=900),
                manager_id=manager.id if manager else 0)
    base.update(kw)
    e = Employee(**base)
    db.add(e)
    db.flush()
    return e


def _dataset(db, world, wide=11):
    """boss ─┬ M1 ─ e1, e2 (+ d1 nghỉ trong kỳ)   ┬ M2 ─ e3   └ X ─ w1..w{wide}; d3 nghỉ, không quản lý."""
    boss = _emp(db, world, "boss")
    m1, m2, x = (_emp(db, world, c, boss) for c in ("m1", "m2", "x"))
    for c in ("e1", "e2"):
        _emp(db, world, c, m1)
    _emp(db, world, "e3", m2)
    for i in range(wide):
        _emp(db, world, f"w{i}", x)
    _emp(db, world, "d1", m1, status="resigned", resign_date=TODAY - timedelta(days=5))
    _emp(db, world, "d2", m1, status="resigned", resign_date=TODAY - timedelta(days=200))   # ngoài kỳ
    _emp(db, world, "d3", None, status="resigned", resign_date=TODAY - timedelta(days=2))
    db.commit()
    return dict(boss=boss, m1=m1, m2=m2, x=x)


def _actor(db, world, client_as, *, scope="all", report=True, actions=("read", "export"), key="a1"):
    world.grant(key, "employee", scope, actions=actions)
    actor = world.actor(key)
    if report:
        db.add(ReportAccess(report_key=int(ReportKey.MANAGEMENT_SPAN), subject_kind=SUBJECT_EMPLOYEE,
                            subject_id=actor.employee.id, effect=1, reason="t", created_by=0, updated_by=0))
        db.commit()
    return actor, client_as(actor.user)


def _get(client, path=SUM, **extra):
    dept = extra.pop("department_id", "")
    return client.get(path, params={**QS, "department_id": dept, **extra})



def test_dinh_nghia_chi_so(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    resp = _get(client, department_id=world.dept["B.hc"])
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    cur = data["totals"]["current"]
    assert cur["staff"] == 18                       # boss + 3 quản lý + e1..e3 + 11 w
    assert cur["managers"] == 4 and cur["direct_reports"] == 17
    assert cur["avg_span"] == pytest.approx(4.25, abs=0.06)
    assert cur["unassigned_manager"] == 1           # boss
    assert cur["span_one"] == 1 and cur["span_wide"] == 1
    assert cur["departures_under"] == 2             # d1 + d3; d2 ngoài kỳ, người nghỉ không tính vào staff
    assert data["totals"]["compare"] is None and data["trend"] == []
    assert data["meta"]["drilldown_metrics"] == list(svc.DRILLDOWN_METRICS)
    assert [d["key"] for d in data["meta"]["dimensions"]][0] == "department"
    assert {d["key"] for d in data["meta"]["dimensions"]} == {"department", "company", "job_level", "manager"}


def test_mot_cap_duoi_va_tren_muoi_la_ranh_gioi(db, world, client_as):
    _dataset(db, world, wide=10)                    # đúng 10 KHÔNG phải «trên 10»
    _, client = _actor(db, world, client_as)
    cur = _get(client, department_id=world.dept["B.hc"]).json()["data"]["totals"]["current"]
    assert cur["span_wide"] == 0 and cur["span_one"] == 1


def test_chieu_quan_ly_moi_quan_ly_mot_dong(db, world, client_as):
    ppl = _dataset(db, world)
    _, client = _actor(db, world, client_as)
    data = _get(client, department_id=world.dept["B.hc"], group_by="manager").json()["data"]
    g = {x["key"]: x for x in data["groups"]}
    m1 = g[str(ppl["m1"].id)]["current"]
    assert g[str(ppl["m1"].id)]["label"] == "Người m1"
    assert m1["staff"] == 2 and m1["managers"] == 1 and m1["direct_reports"] == 2
    assert m1["departures_under"] == 1 and m1["avg_span"] == 2
    assert g[str(ppl["x"].id)]["current"]["span_wide"] == 1
    assert g[str(ppl["boss"].id)]["current"]["staff"] == 3
    none = g[""]["current"]                          # boss (chưa gán) + d3 (nghỉ, không quản lý)
    assert none["staff"] == 1 and none["unassigned_manager"] == 1 and none["departures_under"] == 1
    assert g[""]["label"] == "(Chưa gắn)"


def test_nhom_phong_ban_cong_ty_cap_bac(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    for dim in ("department", "company", "job_level"):
        assert _get(client, group_by=dim).status_code == 200
    data = _get(client, group_by="department").json()["data"]
    assert next(g for g in data["groups"] if g["key"] == str(world.dept["B.hc"]))["current"]["staff"] == 18
    assert _get(client, group_by="khong_co").status_code == 422


def test_department_id_chi_thu_hep(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    assert _get(client).json()["data"]["totals"]["current"]["staff"] == 18 + 7     # 7 người thế giới mẫu
    assert _get(client, department_id=world.dept["A.kt"]).json()["data"]["totals"]["current"]["staff"] == 3
    assert _get(client, department_id="all").status_code == 200
    assert _get(client, department_id="abc").status_code == 422


def test_pham_vi_du_lieu(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as, scope="dept")      # a1 chỉ thấy phòng A.kt
    cur = _get(client).json()["data"]["totals"]["current"]
    assert cur["staff"] == 3 and cur["managers"] == 0 and cur["departures_under"] == 0
    assert _get(client, department_id=world.dept["B.hc"]).json()["data"]["totals"]["current"]["staff"] == 0


def test_quan_ly_ngoai_pham_vi_khong_sinh_hang_quan_ly(db, world, client_as):
    ppl = _dataset(db, world)
    ppl["m1"].department_id = world.dept["B.kt"]               # m1 ra khỏi phòng đang lọc
    db.commit()
    _, client = _actor(db, world, client_as)
    cur = _get(client, department_id=world.dept["B.hc"]).json()["data"]["totals"]["current"]
    assert cur["managers"] == 3 and cur["staff"] == 17
    assert cur["unassigned_manager"] == 1                      # e1, e2 vẫn có quản lý (m1) dù ngoài bộ lọc


def test_quan_ly_da_nghi_khong_tinh_la_da_gan(db, world, client_as):
    ppl = _dataset(db, world)
    ex = _emp(db, world, "ex", ppl["boss"], status="resigned", is_active=False,
              resign_date=TODAY - timedelta(days=300))               # quản lý cũ, nghỉ trước kỳ
    orphan = _emp(db, world, "orphan", ex)
    _emp(db, world, "orphan2", ex)
    db.commit()
    _, client = _actor(db, world, client_as)
    dept = world.dept["B.hc"]
    data = _get(client, department_id=dept).json()["data"]
    cur = data["totals"]["current"]
    assert cur["manager_resigned"] == 2 and "manager_resigned" in data["meta"]["drilldown_metrics"]
    assert cur["managers"] == 4 and cur["direct_reports"] == 17      # quản lý nghỉ không có hàng «m», không cộng cấp dưới
    assert cur["unassigned_manager"] == 1                           # chỉ boss: người có quản lý đã nghỉ KHÔNG tính chưa gán
    body = _get(client, SUM + "/rows", department_id=dept, metric="manager_resigned").json()["data"]
    assert {r["code"]: r["manager"] for r in body["rows"]} == {"orphan": "Người ex", "orphan2": "Người ex"}
    g = {x["key"]: x for x in _get(client, department_id=dept, group_by="manager").json()["data"]["groups"]}
    assert str(ex.id) not in g and g[""]["current"]["staff"] == 3   # boss + 2 mồ côi nằm ở «(Chưa gắn)»
    assert orphan.id  # đã ghi


def test_nhan_quan_ly_ngoai_pham_vi_chi_co_ten(db, world, client_as):
    ppl = _dataset(db, world)
    ex = _emp(db, world, "ex", ppl["boss"], status="resigned", is_active=False,
              resign_date=TODAY - timedelta(days=300), department_id=world.dept["B.kt"])
    _emp(db, world, "d4", ex, status="resigned", resign_date=TODAY - timedelta(days=3))
    db.commit()
    _, client = _actor(db, world, client_as)
    body = _get(client, SUM + "/rows", department_id=world.dept["B.hc"], metric="departures_under").json()["data"]
    assert {r["code"]: r["manager"] for r in body["rows"]}["d4"] == "Người ex"     # chỉ tên, không «(đã nghỉ)»


def test_rows(db, world, client_as):
    ppl = _dataset(db, world)
    _, client = _actor(db, world, client_as)
    dept = world.dept["B.hc"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="managers").json()["data"]
    assert body["total"] == 4 and body["truncated"] is False
    assert [c["key"] for c in body["columns"]] == ["code", "name", "department", "position", "reports"]
    assert body["rows"][0]["code"] == "x" and body["rows"][0]["reports"] == 11     # nhiều cấp dưới nhất lên đầu
    assert body["rows"][0]["link"] == f"/hr/employees/{ppl['x'].id}"
    body = _get(client, SUM + "/rows", department_id=dept, metric="span_wide").json()["data"]
    assert [r["code"] for r in body["rows"]] == ["x"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="span_one").json()["data"]
    assert [r["code"] for r in body["rows"]] == ["m2"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="unassigned_manager").json()["data"]
    assert [r["code"] for r in body["rows"]] == ["boss"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="departures_under").json()["data"]
    assert {r["code"]: r["manager"] for r in body["rows"]} == {"d1": "Người m1", "d3": ""}
    assert "resign_date" in {c["key"] for c in body["columns"]} and "manager" in {c["key"] for c in body["columns"]}
    body = _get(client, SUM + "/rows", department_id=dept, metric="departures_under", group_by="manager",
                group_key=str(ppl["m1"].id)).json()["data"]
    assert [r["code"] for r in body["rows"]] == ["d1"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="unassigned_manager", group_by="manager",
                group_key="").json()["data"]
    assert [r["code"] for r in body["rows"]] == ["boss"]


@pytest.mark.parametrize("metric", ["staff", "direct_reports", "avg_span", "salary", ""])
def test_rows_chi_so_khong_co_danh_sach_la_422(db, world, client_as, metric):
    _, client = _actor(db, world, client_as)
    assert _get(client, SUM + "/rows", metric=metric).status_code == 422


def test_rows_nhom_gop_422(db, world, client_as):
    _, client = _actor(db, world, client_as)
    r = _get(client, SUM + "/rows", metric="managers", group_by="manager", group_key="__other__")
    assert r.status_code == 422


def test_gac_khoa_bao_cao_o_ca_ba_duong(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as, report=False)
    for path, extra in ((SUM, {}), (SUM + "/export", {}), (SUM + "/rows", {"metric": "managers"})):
        assert _get(client, path, **extra).status_code == 403, path


def test_xuat_can_quyen_export(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as, actions=("read",))
    assert _get(client, SUM + "/export").status_code == 403


def test_xuat_xlsx(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    resp = _get(client, SUM + "/export", department_id=world.dept["B.hc"], group_by="manager")
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]
    ws = load_workbook(BytesIO(resp.content)).active
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Nhân sự đang làm" in text and "Người m1" in text


def test_so_truy_van_co_dinh(db, world):
    world.grant("a1", "employee", scope="all", actions=("read",))
    actor = world.actor("a1")
    prof = get_perm_profile(db, actor.user)
    period = parse_period(QS)

    def add(n, start):
        boss = _emp(db, world, f"qb{start}")
        for i in range(n):
            _emp(db, world, f"q{start}-{i}", boss)
            _emp(db, world, f"qd{start}-{i}", boss, status="resigned", resign_date=TODAY - timedelta(days=1))
        db.commit()

    def count():
        stmts = []
        eng = db.get_bind()

        def on(conn, cursor, statement, *a):
            stmts.append(statement)
        event.listen(eng, "before_cursor_execute", on)
        try:
            svc.build_summary(db, actor.user, prof, period, "manager")
        finally:
            event.remove(eng, "before_cursor_execute", on)
        return len(stmts)

    add(5, 0)
    c5 = count()
    add(50, 1)
    assert c5 > 0 and c5 == count()
