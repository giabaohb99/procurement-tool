"""bao-CR-691 — bảng Chi tiết YC mua hàng (GET /api/reports/pr-lines) thêm cột phía phiếu / dòng / đơn mua hàng,
lọc nhiều trạng thái dòng, xuất Excel (GET /api/reports/pr-lines/export).

Canh:
  - dòng có 2 đơn: po_codes theo ngày tạo tăng, po_count, NCC đơn gần nhất, hẹn giao gần nhất, số lần nhận hàng;
    đơn đã hủy / xóa mềm / dòng đơn Hủy không tính;
  - dòng chưa có đơn: null / 0;
  - hai dòng cùng mã hàng trong một phiếu nhận CÙNG tập đơn (hạn chế đã ghi);
  - thiếu supplier.read thì suggested_supplier / latest_po_supplier là null;
  - line_status nhiều giá trị (danh sách hoặc chuỗi phẩy), một giá trị như cũ;
  - export trả xlsx đủ cột tiếng Việt; thiếu report.export → 403.
"""
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.modules.purchase_order.model import PODelivery, POItem, PurchaseOrder
from app.modules.purchase_request.model import PurchaseRequest, PurchaseRequestItem
from app.modules.report import service as report_service

USER = SimpleNamespace(id=0)


def _see_all(monkeypatch):
    monkeypatch.setattr(report_service, "report_dept_scope", lambda db, user: None)


def _pr(db, seed, code, **kw):
    pr = PurchaseRequest(code=code, company_id=seed.company_id, requester="Người YC", requester_id=seed.emp_req_id,
                         department="Phòng Test", status="approved", request_date="2026-10-01", purpose="Mua bao bì",
                         is_urgent=True, need_date="2026-10-15", head_of_dept="TBP cũ", suggested_supplier="NCC Đề Xuất",
                         created_by=seed.u_req_id, updated_by=seed.u_req_id, **kw)
    db.add(pr)
    db.flush()
    return pr


def _line(db, pr, code, line_status="no_po", **kw):
    it = PurchaseRequestItem(pr_id=pr.id, product_code=code, product_name=f"SP {code}", qty=10, unit="thùng", price=100,
                             amount=1000, required_date="2026-10-10", line_status=line_status, qty_ordered=kw.pop("qty_ordered", 0),
                             progress_note="đang hỏi giá", note="giao kho 2", **kw)
    db.add(it)
    db.flush()
    return it


def _po(db, pr_code, code, supplier, created_at, status="approved", deleted=False):
    po = PurchaseOrder(code=code, pr_code=pr_code, supplier_name=supplier, status=status, is_deleted=deleted, created_at=created_at)
    db.add(po)
    db.flush()
    return po


def _po_item(db, po, product_code, progress="ordered"):
    item = POItem(po_id=po.id, product_code=product_code, qty_order=5, progress_status=progress)
    db.add(item)
    db.flush()
    return item


@pytest.fixture
def data(db, seed):
    from datetime import datetime

    pr = _pr(db, seed, "PYC-691-A")
    two = _line(db, pr, "SP-2DON", line_status="ordered", qty_ordered=10)
    none = _line(db, pr, "SP-KHONG", line_status="no_po")
    po1 = _po(db, pr.code, "DMH-1", "NCC Cũ", datetime(2026, 10, 2))
    po2 = _po(db, pr.code, "DMH-2", "NCC Mới", datetime(2026, 10, 5))
    dead = _po(db, pr.code, "DMH-HUY", "NCC Hủy", datetime(2026, 10, 6), status="cancelled")
    gone = _po(db, pr.code, "DMH-XOA", "NCC Xóa", datetime(2026, 10, 7), deleted=True)
    i1 = _po_item(db, po1, "SP-2DON")
    i2 = _po_item(db, po2, "SP-2DON")
    _po_item(db, dead, "SP-2DON")
    _po_item(db, gone, "SP-2DON")
    po3 = _po(db, pr.code, "DMH-3", "NCC Dòng Hủy", datetime(2026, 10, 8))
    _po_item(db, po3, "SP-2DON", progress="cancelled")
    db.add_all([
        PODelivery(po_id=po1.id, po_item_id=i1.id, delivery_no=1, received_qty=3, promised_date="2026-10-09", received_date="2026-10-09"),
        PODelivery(po_id=po2.id, po_item_id=i2.id, delivery_no=1, received_qty=2, promised_date="2026-10-20", received_date="2026-10-21"),
        PODelivery(po_id=po2.id, po_item_id=i2.id, delivery_no=2, received_qty=0, promised_date="2026-10-25", received_date=""),
    ])
    db.commit()
    return SimpleNamespace(pr=pr, two=two, none=none)


def _row(out, line_id):
    return next(r for r in out["items"] if r["id"] == line_id)


def test_dong_hai_don_va_dong_chua_co_don(db, seed, data, monkeypatch):
    _see_all(monkeypatch)
    out = report_service.compute_pr_lines(db, USER, year="2026", show_supplier=True)
    two = _row(out, data.two.id)
    assert two["po_codes"] == "DMH-1, DMH-2", "đơn hủy / xóa mềm / dòng đơn Hủy không tính, theo ngày tạo tăng"
    assert (two["po_count"], two["latest_po_supplier"]) == (2, "NCC Mới")
    assert two["latest_delivery_date"] == "2026-10-25"
    assert two["delivery_count"] == 2, "chỉ đếm lần giao có nhận hàng"
    assert two["qty_ordered"] == 10.0 and two["unit"] == "thùng" and two["required_date"] == "2026-10-10"
    assert two["progress_note"] == "đang hỏi giá" and two["line_note"] == "giao kho 2"
    assert two["purpose"] == "Mua bao bì" and two["is_urgent"] is True and two["need_date"] == "2026-10-15"
    assert two["head_of_dept_name"] == "TBP cũ", "không có id nhân sự thì lùi về bản chụp tên"
    assert two["suggested_supplier"] == "NCC Đề Xuất"
    none = _row(out, data.none.id)
    assert (none["po_codes"], none["po_count"], none["latest_po_supplier"], none["latest_delivery_date"], none["delivery_count"]) == (None, 0, None, None, 0)


def test_hai_dong_cung_ma_hang_nhan_cung_tap_don(db, seed, data, monkeypatch):
    twin = _line(db, data.pr, "SP-2DON", line_status="ordered")
    db.commit()
    _see_all(monkeypatch)
    out = report_service.compute_pr_lines(db, USER, year="2026", show_supplier=True)
    assert _row(out, twin.id)["po_codes"] == _row(out, data.two.id)["po_codes"] == "DMH-1, DMH-2"


def test_thieu_supplier_read_an_ncc(db, seed, data, monkeypatch):
    _see_all(monkeypatch)
    out = report_service.compute_pr_lines(db, USER, year="2026", show_supplier=False)
    two = _row(out, data.two.id)
    assert two["suggested_supplier"] is None and two["latest_po_supplier"] is None
    assert two["po_codes"] == "DMH-1, DMH-2", "mã đơn vẫn hiện, chỉ ẩn NCC"


def test_loc_nhieu_trang_thai_dong(db, seed, data, monkeypatch):
    _see_all(monkeypatch)
    _line(db, data.pr, "SP-XONG", line_status="completed")
    db.commit()
    codes = lambda out: sorted(r["product_code"] for r in out["items"])  # noqa: E731
    assert codes(report_service.compute_pr_lines(db, USER, year="2026", line_status=["no_po", "completed"])) == ["SP-KHONG", "SP-XONG"]
    assert codes(report_service.compute_pr_lines(db, USER, year="2026", line_status="ordered,completed")) == ["SP-2DON", "SP-XONG"]
    assert codes(report_service.compute_pr_lines(db, USER, year="2026", line_status="ordered")) == ["SP-2DON"]
    assert codes(report_service.compute_pr_lines(db, USER, year="2026", line_status="chua_dat")) == ["SP-KHONG"]


@pytest.fixture
def client_as(db):
    from app.core.auth import get_current_user
    from app.core.database import get_db
    from app.main import app

    def make(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app)
    yield make
    app.dependency_overrides.clear()


def _user(db, code):
    from app.modules.user.model import User

    user = User(email=f"{code.lower()}@x.vn", password_hash="x", is_active=True)
    db.add(user)
    db.flush()
    return user


def test_xuat_excel_va_chan_khi_thieu_quyen_xuat(db, seed, data, cap_quyen, client_as, monkeypatch):
    _see_all(monkeypatch)
    reader = _user(db, "U691R")
    cap_quyen(reader.id, "report", read=True)
    db.commit()
    assert client_as(reader).get("/api/reports/pr-lines/export", params={"year": "2026"}).status_code == 403

    exporter = _user(db, "U691E")
    cap_quyen(exporter.id, "report", read=True, export=True)
    db.commit()
    resp = client_as(exporter).get("/api/reports/pr-lines/export", params={"year": "2026", "line_status": "ordered,no_po"})
    assert resp.status_code == 200, resp.text
    ws = load_workbook(BytesIO(resp.content)).active
    header = [c.value for c in ws[1]]
    assert header[:3] == ["Mã YCMH", "Ngày lập", "Pháp nhân"] and "Mã đơn mua hàng" in header and "NCC đơn gần nhất" in header
    assert ws.max_row == 3, "hai dòng theo bộ lọc nhiều trạng thái"
    supplier_col = header.index("NCC đơn gần nhất") + 1
    assert all(ws.cell(row=r, column=supplier_col).value in (None, "") for r in (2, 3)), "không supplier.read thì cột NCC trống"
