"""bao-CR-608 — Tra cứu thị trường: cột «ID» ghi đè, cột «Thao tác» xóa, sửa / xóa tay từng dòng.

```
tệp không có ID / Thao tác     -> nạp như cũ (bao-CR-541)
ID có thật                     -> GHI ĐÈ đủ cột, giữ id + batch_id cũ, row_hash tính lại
ID có số mà không có           -> THÊM MỚI như dòng thường + cảnh báo «ID n không có — thêm mới»
Thao tác = xóa, ID có thật     -> XÓA (dòng xóa chỉ cần ô ID — thiếu Ngày đăng ký vẫn chạy)
Thao tác = xóa, ID trống/không -> BỎ QUA + cảnh báo
chạy thử                       -> đếm đủ, KHÔNG ghi gì
hoàn tác lô                    -> xóa dòng thêm mới, dựng lại dòng xóa (đúng id), trả dòng ghi đè
lô cũ deleted_count > 0 không bản chụp -> vẫn chặn hoàn tác
sửa / xóa tay                  -> chụp vào tab_customs_line_change + nhật ký thao tác
```
"""
import io
import json
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from fastapi import HTTPException
from pydantic import ValidationError

from app.modules.audit.model import AuditLog
from app.modules.customs import dedupe, importer, line_edit_service, reader
from app.modules.customs import service as S
from app.modules.customs.constants import (COLUMNS, CustomsLineChangeAction, CustomsLineChangeSource)
from app.modules.customs.model import CustomsLine, CustomsLineChange, CustomsParty
from app.modules.customs.schema import CustomsLineUpdate
from app.modules.import_tool import service as import_service
from app.modules.import_tool.model import ImportLog, ImportMode, ImportRowStatus, ImportStatus, LogLevel

from test_hai_quan_luu_bo_loc_log_dong_cr496 import _batch, _row, _statuses
from test_hai_quan_luu_bo_loc_log_dong_cr496 import _xlsx as _plain_xlsx

USER = 7


def _xlsx(rows, action_header="Thao tác") -> bytes:
    """Tệp có cột «ID» đầu + cột thao tác cuối. Ô thiếu khóa thì để trống."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID"] + [label for _, label in COLUMNS] + [action_header])
    for r in rows:
        ws.append([r.get("id")] + [r.get(key) for key, _ in COLUMNS] + [r.get("action")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _lines(db):
    db.expire_all()
    return db.query(CustomsLine).order_by(CustomsLine.id).all()


def _seed(db, *rows):
    """Nạp sẵn các dòng bằng tệp thường → (lô, các dòng)."""
    b = _batch(db)
    importer.run(db, b, _plain_xlsx(list(rows)), apply=True)
    return b, _lines(db)


def _messages(db, bid):
    return {x.row_no: x.message for x in db.query(ImportLog).filter(ImportLog.batch_id == bid,
                                                                    ImportLog.row_status != 0)}


def _warnings(db, bid):
    return [x.message for x in db.query(ImportLog).filter(ImportLog.batch_id == bid,
                                                          ImportLog.level == LogLevel.WARNING)]


# ── Đọc tệp ─────────────────────────────────────────────────────────────────

def test_reader_reads_id_and_action_columns_and_aliases():
    raw = _xlsx([{**_row(), "id": 12}, {"id": "'34", "action": "XÓA"}, {"id": 5, "action": "Del"},
                 {**_row(line_no=2), "id": "abc", "action": "sửa"}], action_header="Hành động")
    res = reader.parse(raw, today=date(2026, 9, 23))
    assert res.control_columns == ["ref_id", "row_action"]
    assert [r.get("ref_id") for r in res.rows] == [12, None]
    assert [(d["source_row"], d["ref_id"]) for d in res.deletes] == [(3, 34), (4, 5)]
    warns = [m for _, level, m in res.logs if level == LogLevel.WARNING]
    assert any("«ID» «abc»" in m for m in warns) and any("«Thao tác» «sửa»" in m for m in warns)
    #  Tiêu đề «Action» cũng nhận
    assert reader.parse(_xlsx([{"id": 1, "action": "delete"}], action_header="Action"),
                        today=date(2026, 9, 23)).deletes[0]["ref_id"] == 1


def test_file_without_id_and_action_imports_as_before(db):
    b, lines = _seed(db, _row(), _row(line_no=2))
    assert len(lines) == 2 and b.created_count == 2 and b.updated_count == 0 and b.deleted_count == 0
    info = json.loads(b.sheet_info)
    assert info["updated_rows"] == 0 and info["deleted_rows"] == 0 and info["id_not_found"] == 0
    assert not db.query(CustomsLineChange).count()


# ── Ghi đè ──────────────────────────────────────────────────────────────────

def test_existing_id_overwrites_all_columns_keeps_id_and_batch(db):
    first, (target, other) = _seed(db, _row(), _row(line_no=2))
    old_hash = target.row_hash
    second = _batch(db)
    importer.run(db, second, _xlsx([{**_row(reg_date="20-02-2025", price_usd=99.5, product_name="MANCOZEB 80% WP",
                                            partner_name="SYNGENTA AG", importer_tax_code="'0100100100",
                                            importer_name="Công ty Mới", currency="USD"), "id": target.id}]),
                 apply=True)
    line = db.get(CustomsLine, target.id)
    assert line.batch_id == first.id, "dòng ghi đè giữ batch_id CŨ — hoàn tác lô mới không được xóa nó"
    assert line.reg_date == date(2025, 2, 20) and line.price_usd == Decimal("99.5")
    assert line.product_name == "MANCOZEB 80% WP" and line.currency == "USD"
    assert db.get(CustomsParty, line.partner_id).name == "SYNGENTA AG"
    assert db.get(CustomsParty, line.importer_id).tax_code == "0100100100"
    assert line.row_hash != old_hash
    assert line.row_hash == dedupe.compute_row_hash({k: getattr(line, k) for k in dedupe.HASH_FIELDS},
                                                    "0100100100", reader.hash_name("SYNGENTA AG"))
    assert len(_lines(db)) == 2 and db.get(CustomsLine, other.id).price_usd == other.price_usd
    assert second.updated_count == 1 and second.created_count == 0
    assert _statuses(db, second.id) == [(2, ImportRowStatus.UPDATED)]
    assert _messages(db, second.id)[2] == f"Ghi đè ID {target.id}"
    snap = db.query(CustomsLineChange).one()
    assert (snap.batch_id, snap.line_id, snap.action, snap.source) == (
        second.id, target.id, CustomsLineChangeAction.UPDATE, CustomsLineChangeSource.IMPORT)
    assert snap.snapshot["reg_date"] == "2026-01-13" and snap.snapshot["batch_id"] == first.id
    assert snap.snapshot["row_hash"] == old_hash and snap.snapshot["importer_id"] == target.importer_id


def test_overwrite_row_identical_to_stored_is_skipped_as_existing(db):
    _, (target,) = _seed(db, _row())
    again = _batch(db)
    importer.run(db, again, _xlsx([{**_row(), "id": target.id}]), apply=True)
    assert again.updated_count == 0 and _statuses(db, again.id) == [(2, ImportRowStatus.EXISTING)]
    assert "dữ liệu không đổi" in _messages(db, again.id)[2]
    assert not db.query(CustomsLineChange).count()


def test_overwrite_bypasses_dedupe_and_is_never_duplicate(db):
    """Hai dòng có ID khác nhau mang CÙNG nội dung: cả hai là ghi đè, không cái nào «trùng trong tệp»."""
    _, (a, b) = _seed(db, _row(), _row(line_no=2))
    batch = _batch(db)
    importer.run(db, batch, _xlsx([{**_row(line_no=9), "id": a.id}, {**_row(line_no=9), "id": b.id}]), apply=True)
    assert _statuses(db, batch.id) == [(2, ImportRowStatus.UPDATED), (3, ImportRowStatus.UPDATED)]
    assert batch.updated_count == 2


def test_unknown_id_is_inserted_as_new_with_warning(db):
    _seed(db, _row())
    b = _batch(db)
    importer.run(db, b, _xlsx([{**_row(line_no=5), "id": 999}, {**_row(), "id": 998}]), apply=True)
    assert b.created_count == 1 and b.updated_count == 0
    msgs = _messages(db, b.id)
    assert msgs[2] == "ID 999 không có — thêm mới"
    assert msgs[3].startswith("ID 998 không có — Đã có trong bảng giá"), "ID không có vẫn qua chống trùng"
    assert "ID 999 không có — thêm mới" in _warnings(db, b.id)
    assert json.loads(b.sheet_info)["id_not_found"] == 2
    assert 999 not in [x.id for x in _lines(db)], "thêm mới với id tự tăng, không lấy id trong tệp"


# ── Xóa ─────────────────────────────────────────────────────────────────────

def test_delete_row_needs_only_the_id(db):
    _, (target, keep) = _seed(db, _row(), _row(line_no=2))
    b = _batch(db)
    importer.run(db, b, _xlsx([{"id": target.id, "action": "Xóa"}]), apply=True)
    assert [x.id for x in _lines(db)] == [keep.id]
    assert b.deleted_count == 1 and b.skipped_count == 0
    assert _statuses(db, b.id) == [(2, ImportRowStatus.DELETED)]
    assert _messages(db, b.id)[2] == f"Xóa ID {target.id}"
    snap = db.query(CustomsLineChange).one()
    assert snap.action == CustomsLineChangeAction.DELETE and snap.snapshot["id"] == target.id


def test_delete_with_blank_or_unknown_id_is_skipped_with_warning(db):
    _seed(db, _row())
    b = _batch(db)
    importer.run(db, b, _xlsx([{"action": "delete"}, {"id": 4242, "action": "xoa"}]), apply=True)
    assert len(_lines(db)) == 1 and b.deleted_count == 0
    assert _statuses(db, b.id) == [(2, ImportRowStatus.IGNORED), (3, ImportRowStatus.IGNORED)]
    warns = _warnings(db, b.id)
    assert "Xóa: ô ID trống — bỏ qua" in warns and "Xóa: ID 4242 không có — bỏ qua" in warns
    assert b.skipped_count == 2


def test_unknown_action_word_warns_and_row_is_treated_as_plain(db):
    b = _batch(db)
    importer.run(db, b, _xlsx([{**_row(), "action": "sửa"}]), apply=True)
    assert b.created_count == 1 and any("«Thao tác» «sửa»" in m for m in _warnings(db, b.id))


def test_same_id_twice_first_row_wins(db):
    _, (target,) = _seed(db, _row())
    b = _batch(db)
    importer.run(db, b, _xlsx([{**_row(price_usd=1), "id": target.id}, {"id": target.id, "action": "xóa"}]),
                 apply=True)
    assert _statuses(db, b.id) == [(2, ImportRowStatus.UPDATED), (3, ImportRowStatus.IGNORED)]
    assert len(_lines(db)) == 1 and b.deleted_count == 0


# ── Chạy thử ────────────────────────────────────────────────────────────────

def test_dry_run_counts_everything_and_writes_nothing(db):
    _, (a, b_line, c) = _seed(db, _row(), _row(line_no=2), _row(line_no=3))
    before = [(x.id, x.price_usd, x.row_hash) for x in _lines(db)]
    dry = _batch(db, mode=ImportMode.DRY_RUN)
    importer.run(db, dry, _xlsx([{**_row(price_usd=5), "id": a.id}, {"id": b_line.id, "action": "xóa"},
                                 {**_row(line_no=8), "id": None}, {"action": "xóa"}]), apply=False)
    assert (dry.created_count, dry.updated_count, dry.deleted_count, dry.skipped_count) == (1, 1, 1, 1)
    assert [(x.id, x.price_usd, x.row_hash) for x in _lines(db)] == before
    assert not db.query(CustomsLineChange).count()
    info = json.loads(dry.sheet_info)
    assert (info["updated_rows"], info["deleted_rows"], info["ignored_rows"]) == (1, 1, 1)


# ── Hoàn tác ────────────────────────────────────────────────────────────────

def test_revert_restores_overwritten_and_deleted_rows_and_drops_new(db):
    first, (a, b_line, keep) = _seed(db, _row(), _row(line_no=2), _row(line_no=3))
    original = {x.id: (x.reg_date, x.price_usd, x.row_hash, x.batch_id, x.importer_id, x.partner_id, x.source_row)
                for x in _lines(db)}
    second = _batch(db)
    importer.run(db, second, _xlsx([{**_row(reg_date="01-03-2024", price_usd=77, partner_name="KHÁC"), "id": a.id},
                                    {"id": b_line.id, "action": "xóa"},
                                    {**_row(line_no=50)}]), apply=True)
    assert (second.created_count, second.updated_count, second.deleted_count) == (1, 1, 1)
    assert len(_lines(db)) == 3 and b_line.id not in [x.id for x in _lines(db)]

    res = import_service.revert_batch(db, second, user_id=1)
    assert res["ok"], res
    assert second.status == ImportStatus.REVERTED
    after = {x.id: (x.reg_date, x.price_usd, x.row_hash, x.batch_id, x.importer_id, x.partner_id, x.source_row)
             for x in _lines(db)}
    assert after == original, "đủ ba dòng cũ, ĐÚNG id, đúng giá trị; dòng thêm mới đã xóa"
    assert "trả 1 dòng ghi đè" in res["message"] and "dựng lại 1 dòng đã xóa" in res["message"]
    assert "xóa 1 dòng thêm mới" in res["message"]


def test_revert_warns_when_row_was_touched_later(db):
    _, (a,) = _seed(db, _row())
    second = _batch(db)
    importer.run(db, second, _xlsx([{**_row(price_usd=50), "id": a.id}]), apply=True)
    line_edit_service.update_line(db, a.id, CustomsLineUpdate(price_usd=Decimal("60")), USER)

    res = import_service.revert_batch(db, second, user_id=1)
    assert res["ok"] and res["warnings"]
    assert "sửa / xóa tay" in res["message"]
    assert db.get(CustomsLine, a.id).price_usd == Decimal("17.3897"), "vẫn trả về bản trước lô"


def test_old_batch_without_snapshots_still_blocked(db):
    old, _ = _seed(db, _row())
    old.deleted_count = 5
    db.commit()
    res = import_service.revert_batch(db, old, user_id=1)
    assert res["ok"] is False and "cách cũ" in res["message"]


def test_new_batch_with_deletes_and_snapshots_can_revert(db):
    _, (a,) = _seed(db, _row())
    b = _batch(db)
    importer.run(db, b, _xlsx([{"id": a.id, "action": "xóa"}]), apply=True)
    assert b.deleted_count == 1
    assert import_service.revert_batch(db, b, user_id=1)["ok"]
    assert [x.id for x in _lines(db)] == [a.id]


# ── Sửa / xóa tay ───────────────────────────────────────────────────────────

def test_manual_update_snapshots_audits_and_rehashes(db):
    first, (a,) = _seed(db, _row())
    old_hash = a.row_hash
    line = line_edit_service.update_line(db, a.id, CustomsLineUpdate(
        price_usd=Decimal("12.5"), currency="usd", partner_name="Bayer AG", importer_tax_code="0311111111",
        importer_name="Công ty B", active_ingredient="Atrazine tay", price_vnd_flat=Decimal("1000")), USER)
    assert line.price_usd == Decimal("12.5") and line.currency == "USD" and line.batch_id == first.id
    assert line.active_ingredient == "Atrazine tay" and line.active_ingredient_from_file is True
    assert line.price_vnd_flat == Decimal("1000")
    assert db.get(CustomsParty, line.partner_id).name == "Bayer AG"
    assert db.get(CustomsParty, line.importer_id).tax_code == "0311111111"
    assert line.row_hash and line.row_hash != old_hash
    snap = db.query(CustomsLineChange).one()
    assert (snap.batch_id, snap.action, snap.source, snap.created_by) == (
        0, CustomsLineChangeAction.UPDATE, CustomsLineChangeSource.MANUAL, USER)
    assert snap.snapshot["price_usd"] == "17.3897"
    log = db.query(AuditLog).filter(AuditLog.entity == "customs_price").one()
    assert log.action == "update" and log.entity_id == a.id and "Đơn giá khai báo(USD)" in log.message
    #  Xóa trắng ô hoạt chất → trả về cho hệ thống suy ra
    line = line_edit_service.update_line(db, a.id, CustomsLineUpdate(active_ingredient=""), USER)
    assert line.active_ingredient_from_file is False


def test_manual_update_requires_tax_code_for_importer_name(db):
    _, (a,) = _seed(db, _row())
    with pytest.raises(HTTPException) as e:
        line_edit_service.update_line(db, a.id, CustomsLineUpdate(importer_tax_code="", importer_name="X"), USER)
    assert e.value.status_code == 400


def test_manual_delete_snapshots_and_audits(db):
    _, (a, b_line) = _seed(db, _row(), _row(line_no=2))
    line_edit_service.delete_line(db, a.id, USER)
    assert [x.id for x in _lines(db)] == [b_line.id]
    snap = db.query(CustomsLineChange).one()
    assert (snap.action, snap.source, snap.line_id) == (CustomsLineChangeAction.DELETE,
                                                       CustomsLineChangeSource.MANUAL, a.id)
    assert db.query(AuditLog).filter(AuditLog.entity == "customs_price", AuditLog.action == "delete").count() == 1
    with pytest.raises(HTTPException) as e:
        line_edit_service.delete_line(db, a.id, USER)
    assert e.value.status_code == 404


@pytest.mark.parametrize("field,limit", [
    ("office_code", 10), ("importer_tax_code", 14), ("importer_name", 255), ("partner_name", 255),
    ("hs_code", 8), ("product_name", 255), ("currency", 3), ("unit_code", 4), ("origin_country", 2),
    ("contract_no", 40), ("incoterm", 3), ("import_country", 2), ("active_ingredient", 255),
    ("formulation", 40),
])
def test_update_schema_caps_strings(field, limit):
    """SQLite không ép VARCHAR — chặn phải nằm ở schema (duoc-CR-316)."""
    CustomsLineUpdate.model_validate({field: "x" * limit})
    with pytest.raises(ValidationError):
        CustomsLineUpdate.model_validate({field: "x" * (limit + 1)})


def test_update_schema_caps_numbers_dates_and_codes():
    for bad in ({"reg_date": "1990-01-01"}, {"reg_date": None}, {"contract_date": "2200-01-01"},
                {"price_usd": "1e13"}, {"rate_import": 10000}, {"line_no": 40000}, {"transport_mode": 7},
                {"quantity": -1}, {"unknown": 1}):
        with pytest.raises(ValidationError):
            CustomsLineUpdate.model_validate(bad)
    assert CustomsLineUpdate.model_validate({"transport_mode": 2, "price_vnd_flat": None}).transport_mode == 2


# ── Excel xuất ──────────────────────────────────────────────────────────────

def test_export_has_id_first_and_round_trips_to_overwrite(db):
    _, (a,) = _seed(db, _row())
    ws = openpyxl.load_workbook(io.BytesIO(S.export_lines_xlsx(db, {}))).active
    assert ws.cell(1, 1).value == "ID" and ws.cell(2, 1).value == a.id
    assert ws.cell(1, ws.max_column).value == "Thao tác"
    #  Sửa giá trong tệp xuất rồi nạp lại → ghi đè đúng dòng đó, không thêm dòng.
    price_col = 1 + [k for k, _ in COLUMNS].index("price_usd") + 1
    ws.cell(2, price_col).value = 21
    buf = io.BytesIO()
    ws.parent.save(buf)
    b = _batch(db)
    importer.run(db, b, buf.getvalue(), apply=True)
    assert (b.created_count, b.updated_count) == (0, 1)
    line = db.get(CustomsLine, a.id)
    assert line.price_usd == Decimal("21")
    #  Ô suy ra / tính sẵn trong tệp xuất KHÔNG bị đóng băng thành «do người nhập»
    assert line.active_ingredient_from_file is False and line.price_vnd_flat is None


def test_line_edit_and_delete_routes_are_wired():
    """Hai đường mới nằm trong app thật (đăng ký router) — sửa = PATCH, xóa = DELETE."""
    import app.main as main
    wired = {(m, r.path) for r in main.app.routes for m in getattr(r, "methods", set())}
    assert ("PATCH", "/api/customs/lines/{line_id}") in wired
    assert ("DELETE", "/api/customs/lines/{line_id}") in wired
