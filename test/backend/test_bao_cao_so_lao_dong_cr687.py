"""duoc-CR-687 — Sổ quản lý lao động (khóa 19): `/api/labor-register/summary` (+ export, rows).

Canh: tập người theo kỳ, chọn HĐ áp cho từng người, chỉ số, phạm vi + bộ lọc chỉ thu hẹp, 403 thiếu quyền
(nhạy cảm / HĐLĐ / chưa giao báo cáo) ở CẢ ba đường, danh sách không lộ giá trị nhạy cảm, tệp Excel và số
truy vấn CỐ ĐỊNH. Kỳ cố định 2025 nên `as_of` = 31/12/2025, không phụ thuộc ngày chạy bài. Nhân sự của
thế giới mẫu được tạo hôm nay (ngoài kỳ 2025) nên không lẫn vào số liệu.
"""
import uuid
from datetime import date, datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event

from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.report_period import parse_period
from app.core.scoping import apply_scope
from app.core.subject_match import SUBJECT_EMPLOYEE
from app.main import app
from app.modules.employee.model import Employee
from app.modules.labor_contract.model import LaborContract
from app.modules.labor_register import report_service as svc
from app.modules.labor_register.labor_book_export import export_labor_book
from app.modules.labor_register.people_data import load_people
from app.modules.leave.request_model import LeaveRequest
from app.modules.report_access.model import ReportAccess

SIGNED, TERMINATED = 2, 4
FIXED, INDEFINITE = 2, 3
RESIGNED = "resigned"
SUM = "/api/labor-register/summary"
QS = {"preset": "custom", "date_from": "2025-01-01", "date_to": "2025-12-31", "compare": "none"}
_seq = {"n": 0}


@pytest.fixture
def client_as(db):
    def make_client(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make_client
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def _emp(db, world, tag, dept="B.hc", co="B", full=True, **kw) -> Employee:
    base = dict(code=f"L{tag}", full_name=f"Người {tag}", company_id=world.co[co], department_id=world.dept[dept],
                is_active=True, hire_date=date(2020, 1, 1))
    if full:
        base.update(date_of_birth=date(1990, 5, 6), id_number="079090123456", permanent_address="1 Lê Lợi, Q1",
                    education_level=4, social_insurance_no="7900000001", position="Chuyên viên")
    emp = Employee(**{**base, **kw})
    db.add(emp)
    db.flush()
    return emp


def _hd(db, world, emp, *, ctype=FIXED, status=SIGNED, start=date(2024, 1, 1), end=None, salary=10_000_000,
        co="B", dept="B.hc", **kw) -> LaborContract:
    _seq["n"] += 1
    c = LaborContract(code=f"HDS{_seq['n']:05d}", employee_id=emp.id, company_id=world.co[co],
                      department_id=world.dept[dept], contract_type=ctype, status=status, start_date=start,
                      end_date=end, base_salary=salary, **kw)
    db.add(c)
    db.flush()
    return c


def _dataset(db, world):
    a = _emp(db, world, "stay"); _hd(db, world, a, ctype=INDEFINITE, start=date(2020, 1, 1))
    b = _emp(db, world, "left", status=RESIGNED, resign_date=date(2025, 6, 30), date_of_birth=None)
    _hd(db, world, b, end=date(2025, 6, 30), status=TERMINATED, terminated_date=date(2025, 6, 30),
        terminate_reason="Hết hạn HĐ")
    _emp(db, world, "gone", status=RESIGNED, resign_date=date(2024, 12, 31))         # nghỉ trước kỳ
    _emp(db, world, "late", hire_date=date(2026, 2, 1))                               # vào sau kỳ
    _emp(db, world, "new", hire_date=date(2025, 3, 1))                                # vào trong kỳ, chưa HĐ
    _emp(db, world, "nohire", hire_date=None, created_at=datetime(2025, 5, 5))        # lấy ngày tạo hồ sơ
    _emp(db, world, "nodate", status=RESIGNED)                                        # nghỉ nhưng không ghi ngày
    db.commit()


def _actor(db, world, client_as, key="a1", scope="all", sensitive=True, contract=True, report=True):
    world.grant(key, "employee", scope, actions=("read", "export"))
    if sensitive:
        world.grant(key, "employee_sensitive", "all", actions=("read",))
    if contract:
        world.grant(key, "labor_contract", "all", actions=("read", "export"))
    actor = world.actor(key)
    if report:
        db.add(ReportAccess(report_key=int(ReportKey.LABOR_BOOK), subject_kind=SUBJECT_EMPLOYEE,
                            subject_id=actor.employee.id, effect=1, reason="t", created_by=0, updated_by=0))
        db.commit()
    return actor, client_as(actor.user)


def _get(client, path=SUM, **extra):
    return client.get(path, params={**QS, "department_id": extra.pop("department_id", ""), **extra})


def _people(db, **kw):
    return load_people(db.query(Employee), db.query(LaborContract), date(2025, 1, 1), date(2025, 12, 31),
                       date(2026, 10, 27), **kw)


def test_tap_nguoi_theo_ky(db, world):
    _dataset(db, world)
    codes = {p.employee_code for p in _people(db, department_id=world.dept["B.hc"])}
    assert codes == {"Lstay", "Lleft", "Lnew", "Lnohire"}      # gone/late/nodate bị loại
    nohire = next(p for p in _people(db, department_id=world.dept["B.hc"]) if p.employee_code == "Lnohire")
    assert nohire.hire_missing and nohire.hire_date == date(2025, 5, 5)


def test_chon_hop_dong_ap_cho_moi_nguoi(db, world):
    e1 = _emp(db, world, "two")
    _hd(db, world, e1, start=date(2023, 1, 1), end=date(2023, 12, 31), salary=8_000_000)
    _hd(db, world, e1, ctype=INDEFINITE, start=date(2024, 1, 1), salary=10_000_000)
    e2 = _emp(db, world, "expired")
    _hd(db, world, e2, start=date(2022, 1, 1), end=date(2022, 12, 31), salary=7_000_000)
    e3 = _emp(db, world, "future")
    _hd(db, world, e3, start=date(2026, 3, 1), end=date(2027, 3, 1))                   # chưa bắt đầu tại cuối kỳ
    db.commit()
    got = {p.employee_code: p for p in _people(db, department_id=world.dept["B.hc"])}
    assert got["Ltwo"].contract_type == INDEFINITE and got["Ltwo"].base_salary == 10_000_000
    assert [s for _d, s in got["Ltwo"].salary_history] == [8_000_000, 10_000_000]
    assert got["Lexpired"].contract_type == FIXED and got["Lexpired"].contract_end == date(2022, 12, 31)
    assert got["Lfuture"].contract_type == 0


def test_chi_so_va_nhom(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    resp = _get(client, department_id=world.dept["B.hc"])
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    assert data["totals"]["current"] == {"people": 4, "working": 3, "left_in_period": 1, "missing_required": 3}
    assert data["trend"] == [] and data["meta"]["drilldown_metrics"] == list(svc.BOOK.metrics)
    by_cat = _get(client, department_id=world.dept["B.hc"], group_by="job_category").json()["data"]["groups"]
    assert {g["label"]: g["current"]["people"] for g in by_cat}["Chuyên môn kỹ thuật bậc cao"] == 4


def test_rows_khop_so_va_khong_lo_gia_tri_nhay_cam(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    body = _get(client, SUM + "/rows", department_id=world.dept["B.hc"], metric="left_in_period").json()["data"]
    assert body["total"] == 1 and body["rows"][0]["employee_code"] == "Lleft"
    assert body["rows"][0]["missing"] == "Ngày sinh" and body["rows"][0]["link"].startswith("/hr/employees/")
    text = str(body)
    assert "079090123456" not in text and "7900000001" not in text and "1 Lê Lợi" not in text
    assert "10000000" not in text
    miss = _get(client, SUM + "/rows", department_id=world.dept["B.hc"], metric="missing_required").json()["data"]
    assert {r["employee_code"] for r in miss["rows"]} == {"Lleft", "Lnew", "Lnohire"}


@pytest.mark.parametrize("extra", [{"metric": "salary"}, {},
                                   {"metric": "people", "group_by": "department", "group_key": "__other__"}])
def test_rows_422(db, world, client_as, extra):
    _, client = _actor(db, world, client_as)
    assert _get(client, SUM + "/rows", **extra).status_code == 422


def test_pham_vi_va_bo_loc_chi_thu_hep(db, world, client_as):
    _emp(db, world, "pa", dept="A.kt", co="A")
    _emp(db, world, "pb")
    db.commit()
    _, client = _actor(db, world, client_as, scope="dept")
    people = _get(client).json()["data"]["totals"]["current"]["people"]
    assert people == 1                                               # chỉ phòng của a1 (A.kt)
    assert _get(client, department_id=world.dept["B.hc"]).json()["data"]["totals"]["current"]["people"] == 0
    assert _get(client, department_id="abc").status_code == 422


@pytest.mark.parametrize("suffix", ["", "/export", "/rows"])
@pytest.mark.parametrize("case,message", [("sensitive", "quyền xem thông tin nhạy cảm của nhân sự"), ("contract", "quyền xem hợp đồng lao động")])
def test_thieu_quyen_nhay_cam_hoac_hdld_403(db, world, client_as, suffix, case, message):
    _, client = _actor(db, world, client_as, sensitive=case != "sensitive", contract=case != "contract")
    resp = _get(client, SUM + suffix, metric="people")
    assert resp.status_code == 403 and message in resp.text


@pytest.mark.parametrize("suffix", ["", "/export", "/rows"])
def test_chua_giao_bao_cao_403(db, world, client_as, suffix):
    _, client = _actor(db, world, client_as, report=False)
    assert _get(client, SUM + suffix, metric="people").status_code == 403


def test_xuat_excel_moi_phap_nhan_mot_sheet(db, world, client_as):
    _dataset(db, world)
    _emp(db, world, "other-co", dept="A.kt", co="A", hire_date=date(2021, 1, 1))
    e = _emp(db, world, "raise"); _hd(db, world, e, start=date(2023, 1, 1), salary=8_000_000)
    _hd(db, world, e, start=date(2024, 6, 1), salary=9_500_000)
    db.commit()
    _, client = _actor(db, world, client_as)
    resp = _get(client, SUM + "/export")
    assert resp.status_code == 200 and "so-quan-ly-lao-dong-2025-01-01-2025-12-31" in resp.headers["content-disposition"]
    wb = load_workbook(BytesIO(resp.content))
    names = {c.name for c in db.query(__import__("app.modules.company.model", fromlist=["Company"]).Company).all()}
    assert len(wb.sheetnames) == 2 and set(wb.sheetnames) <= {n[:28] for n in names}
    ws = next(w for w in wb if any(str(c.value) == "Lstay" for row in w.iter_rows() for c in row))
    assert ws["A1"].value == "SỔ QUẢN LÝ LAO ĐỘNG" and "Mã số thuế" in ws["A3"].value
    heads = [c.value for c in ws[7]]
    assert heads[0] == "STT" and heads[7] == "Số CCCD" and len(heads) == 25
    assert heads[16].endswith("trong kỳ") and heads[17].endswith("trong kỳ") and heads[6] == "Nơi cư trú"
    rows = {r[1].value: [c.value for c in r] for r in ws.iter_rows(min_row=8)}
    stay, left = rows["Lstay"], rows["Lleft"]
    assert stay[7] == "079090123456" and stay[5] is None and stay[9] is None and stay[18] is None  # cột trống
    assert stay[14] == 10_000_000 and stay[15] in (None, "")
    assert rows["Lraise"][15] == "01/01/2023: 8.000.000 → 01/06/2024: 9.500.000"
    assert left[22].date() == date(2025, 6, 30) and left[23] == "Hết hạn HĐ" and left[24] == "Ngày sinh"
    assert ws.freeze_panes == "D8"


def test_so_ngay_nghi_co_luong_khong_luong(db, world, client_as):
    from app.modules.leave.constants import LR_APPROVED
    from app.modules.leave.catalog_model import LeaveType
    from app.modules.leave.request_model import LeaveRequest, LeaveRequestLine
    emp = _emp(db, world, "leave")
    paid = LeaveType(code="PN", name="Phép năm", is_paid=True)
    unpaid = LeaveType(code="KL", name="Không lương", is_paid=False)
    db.add_all([paid, unpaid])
    db.flush()
    req = LeaveRequest(code="NP1", employee_id=emp.id, company_id=emp.company_id, department_id=emp.department_id,
                       leave_type_id=paid.id, from_date=date(2025, 3, 3), to_date=date(2025, 3, 5), status=LR_APPROVED)
    db.add(req)
    db.flush()
    db.add_all([LeaveRequestLine(request_id=req.id, leave_type_id=paid.id, days=2),
                LeaveRequestLine(request_id=req.id, leave_type_id=unpaid.id, days=1)])
    #  Đơn đã hủy / chưa duyệt không tính; đơn đã xóa mềm cũng không.
    for code, status, deleted in (("NP2", 1, False), ("NP3", LR_APPROVED, True)):
        extra = LeaveRequest(code=code, employee_id=emp.id, company_id=emp.company_id, leave_type_id=paid.id,
                             from_date=date(2025, 4, 1), to_date=date(2025, 4, 2), status=status, is_deleted=deleted)
        db.add(extra)
        db.flush()
        db.add(LeaveRequestLine(request_id=extra.id, leave_type_id=paid.id, days=2))
    db.commit()
    from app.modules.labor_register.leave_days import load_leave_days
    assert load_leave_days(db, db.query(LeaveRequest), [emp.id], date(2025, 1, 1), date(2025, 12, 31)) == {emp.id: (2.0, 1.0)}
    assert load_leave_days(db, db.query(LeaveRequest), [emp.id], date(2024, 1, 1), date(2024, 12, 31)) == {}
    #  Đơn vắt qua biên năm: chỉ phần ngày rơi vào kỳ được tính (không phải cả 3 ngày).
    cross = LeaveRequest(code="NP4", employee_id=emp.id, company_id=emp.company_id, department_id=emp.department_id,
                         leave_type_id=paid.id, from_date=date(2025, 12, 31), to_date=date(2026, 1, 2),
                         status=LR_APPROVED)
    db.add(cross)
    db.flush()
    db.add(LeaveRequestLine(request_id=cross.id, leave_type_id=paid.id, days=3))
    db.commit()
    got_paid, got_unpaid = load_leave_days(db, db.query(LeaveRequest), [emp.id], date(2025, 1, 1), date(2025, 12, 31))[emp.id]
    assert 2.0 < got_paid < 5.0 and got_unpaid == 1.0


def test_so_truy_van_khong_tang_theo_so_nguoi(db, world):
    world.grant("a1", "employee", "all", actions=("read", "export"))
    world.grant("a1", "labor_contract", "all", actions=("read",))
    world.grant("a1", "leave_request", "all", actions=("read",))
    actor = world.actor("a1")
    prof = actor.profile()
    period = parse_period(QS)

    def run():
        eq = apply_scope(db.query(Employee), Employee, "employee", actor.user, prof)
        cq = apply_scope(db.query(LaborContract), LaborContract, "labor_contract", actor.user, prof)
        svc.build_summary(db, svc.BOOK, eq, cq, period, "department")
        svc.build_rows(db, svc.BOOK, eq, cq, period, "people", None, None)
        lq = apply_scope(db.query(LeaveRequest), LeaveRequest, "leave_request", actor.user, prof)
        export_labor_book(db, eq, cq, period, None, None, lq)

    def count() -> int:
        stmts: list[str] = []
        listener = lambda conn, cur, stmt, *a: stmts.append(stmt)  # noqa: E731
        engine = db.get_bind()
        event.listen(engine, "before_cursor_execute", listener)
        try:
            run()
        finally:
            event.remove(engine, "before_cursor_execute", listener)
        return len(stmts)

    def add(n, start):
        for i in range(n):
            e = _emp(db, world, f"q{start + i}")
            _hd(db, world, e, end=date(2026, 1, 1))
        db.commit()

    add(5, 0)
    c5 = count()
    add(45, 5)
    c50 = count()
    assert c5 > 0 and c5 == c50, (c5, c50)


def _approved_leave(db, emp, code="NPX"):
    from app.modules.leave.catalog_model import LeaveType
    from app.modules.leave.constants import LR_APPROVED
    from app.modules.leave.request_model import LeaveRequestLine
    paid = LeaveType(code=f"P{code}", name="Phép năm", is_paid=True)
    db.add(paid)
    db.flush()
    req = LeaveRequest(code=code, employee_id=emp.id, company_id=emp.company_id, department_id=emp.department_id,
                       leave_type_id=paid.id, from_date=date(2025, 3, 3), to_date=date(2025, 3, 4), status=LR_APPROVED)
    db.add(req)
    db.flush()
    db.add(LeaveRequestLine(request_id=req.id, leave_type_id=paid.id, days=2))


def _book_sheet(resp):
    wb = load_workbook(BytesIO(resp.content))
    return wb[wb.sheetnames[0]]


def test_xuat_tep_theo_pham_vi_export_khong_phai_read(db, world, client_as):
    me = world.actor("a1")
    me.employee.hire_date = date(2020, 1, 1)
    _emp(db, world, "other")
    db.commit()
    world.grant("a1", "employee", "all", actions=("read",))
    world.grant("a1", "employee", "own", actions=("export",))
    world.grant("a1", "employee_sensitive", "all", actions=("read",))
    world.grant("a1", "labor_contract", "all", actions=("read", "export"))
    db.add(ReportAccess(report_key=int(ReportKey.LABOR_BOOK), subject_kind=SUBJECT_EMPLOYEE,
                        subject_id=me.employee.id, effect=1, reason="t", created_by=0, updated_by=0))
    db.commit()
    client = client_as(me.user)
    assert _get(client).json()["data"]["totals"]["current"]["people"] >= 2           # xem: cả công ty
    ws = _book_sheet(_get(client, SUM + "/export"))
    codes = [r[1].value for r in ws.iter_rows(min_row=8) if r[1].value]
    assert codes == [me.employee.code]                                                # xuất: chỉ phần của mình


def test_xuat_tep_thieu_quyen_xuat_hdld_403(db, world, client_as):
    world.grant("a1", "employee", "all", actions=("read", "export"))
    world.grant("a1", "employee_sensitive", "all", actions=("read",))
    world.grant("a1", "labor_contract", "all", actions=("read",))
    actor = world.actor("a1")
    db.add(ReportAccess(report_key=int(ReportKey.LABOR_BOOK), subject_kind=SUBJECT_EMPLOYEE,
                        subject_id=actor.employee.id, effect=1, reason="t", created_by=0, updated_by=0))
    db.commit()
    client = client_as(actor.user)
    resp = _get(client, SUM + "/export")
    assert resp.status_code == 403 and "quyền xuất hợp đồng lao động" in resp.text
    assert _get(client).status_code == 200                                            # xem vẫn được


def test_cot_nghi_phep_theo_pham_vi_don_nghi(db, world, client_as):
    emp = _emp(db, world, "lv")
    _approved_leave(db, emp)
    db.commit()
    _, client = _actor(db, world, client_as)       # có leave_request? chưa -> cột nghỉ trống + ghi chú
    ws = _book_sheet(_get(client, SUM + "/export"))
    row = next(r for r in ws.iter_rows(min_row=8) if r[1].value == "Llv")
    assert row[16].value is None and row[17].value is None and "chưa có quyền xem đơn nghỉ phép" in ws["A6"].value
    world.grant("a1", "leave_request", "all", actions=("read",))
    from app.core.auth import _PERM_CACHE
    _PERM_CACHE.clear()
    ws = _book_sheet(_get(client, SUM + "/export"))
    row = next(r for r in ws.iter_rows(min_row=8) if r[1].value == "Llv")
    assert (row[16].value, row[17].value) == (2, 0) and "đơn nghỉ phép" not in (ws["A6"].value or "")


def test_ngay_tao_ho_so_utc_so_voi_cuoi_ky_gio_vn(db, world):
    #  created_at là UTC: 16:30 UTC 31/12 = 23:30 VN 31/12 (trong kỳ); 17:30 UTC = 00:30 VN 01/01/2026 (ngoài kỳ).
    _emp(db, world, "in", hire_date=None, created_at=datetime(2025, 12, 31, 16, 30))
    _emp(db, world, "out", hire_date=None, created_at=datetime(2025, 12, 31, 17, 30))
    db.commit()
    got = {p.employee_code: p for p in _people(db, department_id=world.dept["B.hc"])}
    assert set(got) == {"Lin"} and got["Lin"].hire_date == date(2025, 12, 31) and got["Lin"].hire_missing


def test_nghi_viec_sau_hom_nay_van_dang_lam(db, world):
    _emp(db, world, "soon", status=RESIGNED, resign_date=date(2025, 9, 1))
    _emp(db, world, "done", status=RESIGNED, resign_date=date(2025, 6, 1))
    db.commit()
    got = {p.employee_code: p for p in load_people(db.query(Employee), db.query(LaborContract), date(2025, 1, 1),
                                                   date(2025, 12, 31), date(2025, 6, 15),
                                                   department_id=world.dept["B.hc"])}
    assert got["Lsoon"].working_at_end and not got["Lsoon"].left_in_period
    assert not got["Ldone"].working_at_end and got["Ldone"].left_in_period


def test_ghi_chu_ngay_vao_lam_lay_tu_ngay_tao_ho_so(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    notes = _get(client, department_id=world.dept["B.hc"]).json()["data"]["notes"]
    assert any(n.startswith("Có 1 người chưa ghi ngày vào làm") for n in notes)
