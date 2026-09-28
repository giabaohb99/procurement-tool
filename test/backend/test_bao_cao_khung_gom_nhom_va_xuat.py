"""P01 khung báo cáo — `report_aggregate.py` + `report_export.py`.

Dữ liệu giả bằng `SimpleNamespace` (không đụng model nghiệp vụ nào — khung này KHÔNG thuộc
về một bảng cụ thể). Chiều `pic` CỐ Ý nhiều giá trị (một dòng có 2 người phụ trách) để chứng
minh luật cốt lõi: `totals` KHÔNG bằng tổng các `groups`.
"""
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.core.report_aggregate import (DerivedSpec, DimensionSpec, MetricSpec, ReportSpec,
                                        aggregate, build_report, compute_metrics)
from app.core.report_export import report_xlsx
from app.core.report_period import bucket_axis, bucket_of, parse_period


def _row(d, dept, pics, amount):
    return SimpleNamespace(date=d, dept=dept, pics=pics, amount=amount)


#  3 dòng tháng 9: dòng 1 có 2 PIC (A,B), dòng 3 KHÔNG có PIC nào (rơi vào "Chưa gắn").
ROWS_SEP = [
    _row("2026-09-05", "Sale", ["A", "B"], 100),
    _row("2026-09-10", "Sale", ["A"], 50),
    _row("2026-09-15", "Mkt", [], 30),
]
ROWS_AUG = [_row("2026-08-05", "Sale", ["A"], 40)]  # kỳ so sánh


def _spec() -> ReportSpec:
    return ReportSpec(
        date_of=lambda r: date.fromisoformat(r.date),
        metrics=[
            MetricSpec("count", "Số dòng", kind="int", value_of=lambda r: 1),
            MetricSpec("amount", "Giá trị", kind="money", value_of=lambda r: r.amount),
            MetricSpec("unique_dept", "Số phòng khác nhau", kind="int",
                       distinct_of=lambda r: r.dept or None),
        ],
        derived=[
            DerivedSpec("avg_amount", "Giá trị TB", num="amount", den="count",
                        kind="money", scale=1, ndigits=0),
            DerivedSpec("pct_full", "Tỷ lệ đủ", num="count", den="count", kind="percent"),
        ],
        dimensions={"pic": DimensionSpec(
            "pic", "Người phụ trách",
            key_of=lambda r: [(p, p) for p in r.pics] or [("", "nhãn sai — phải bị ghi đè")])},
    )


def _fetch_all(d_from, d_to):
    return [r for r in ROWS_SEP + ROWS_AUG if d_from <= date.fromisoformat(r.date) <= d_to]


# ── aggregate(): tổng KHÔNG bằng tổng các nhóm khi chiều nhiều giá trị ───────────
def test_tong_tinh_tu_hang_duy_nhat_khong_bang_tong_cac_nhom():
    data = aggregate(ROWS_SEP, _spec(), date(2026, 9, 1), date(2026, 9, 30), "day", "pic")
    assert data["totals"]["count"] == 3          # 3 DÒNG thật, không phải 3 PIC
    by_pic = {g["key"]: g["values"]["count"] for g in data["groups"]}
    assert by_pic == {"A": 2, "B": 1, "": 1}
    assert sum(by_pic.values()) == 4              # 4 != 3 — dòng 1 góp mặt ở cả A lẫn B


def test_nhan_chua_gan_cho_khoa_rong_bat_ke_nhan_goc():
    data = aggregate(ROWS_SEP, _spec(), date(2026, 9, 1), date(2026, 9, 30), "day", "pic")
    empty = next(g for g in data["groups"] if g["key"] == "")
    assert empty["label"] == "(Chưa gắn)"          # ghi đè nhãn sai mà key_of cố tình trả về


def test_derived_tinh_lai_o_moi_cap_tu_metric_da_cong():
    data = aggregate(ROWS_SEP, _spec(), date(2026, 9, 1), date(2026, 9, 30), "day", "pic")
    assert data["totals"]["avg_amount"] == 60.0    # 180/3
    assert data["totals"]["pct_full"] == 100.0
    group_a = next(g for g in data["groups"] if g["key"] == "A")
    assert group_a["values"]["avg_amount"] == 75.0  # (100+50)/2 — KHÔNG kế thừa từ tổng


def test_distinct_dem_gia_tri_khac_nhau_bo_qua_falsy():
    vals = compute_metrics(ROWS_SEP, [MetricSpec("d", "d", kind="int",
                                                   distinct_of=lambda r: r.dept or None)])
    assert vals["d"] == 2  # "Sale","Mkt" — dept rỗng (nếu có) bị loại


def test_group_by_khong_ton_tai_nem_422():
    with pytest.raises(HTTPException) as exc:
        aggregate(ROWS_SEP, _spec(), date(2026, 9, 1), date(2026, 9, 30), "day", "khong-co")
    assert exc.value.status_code == 422


def test_metric_spec_bat_buoc_dung_1_trong_value_of_distinct_of():
    with pytest.raises(ValueError):
        MetricSpec("x", "X")
    with pytest.raises(ValueError):
        MetricSpec("x", "X", value_of=lambda r: 1, distinct_of=lambda r: r)


# ── build_report(): hợp đồng JSON + số lần gọi fetch ─────────────────────────────
def test_compare_none_ra_null_va_fetch_goi_dung_1_lan():
    calls = []

    def fetch(d_from, d_to):
        calls.append((d_from, d_to))
        return _fetch_all(d_from, d_to)

    period = parse_period({"preset": "this_month", "compare": "none"}, today=date(2026, 9, 30))
    report = build_report(fetch, _spec(), period, group_by="pic")
    assert len(calls) == 1
    assert report["period"]["compare"] == "none"
    assert report["totals"]["compare"] is None
    assert all(t["compare"] is None for t in report["trend"])


def test_compare_previous_fetch_goi_2_lan_va_gop_theo_khoa():
    calls = []

    def fetch(d_from, d_to):
        calls.append((d_from, d_to))
        return _fetch_all(d_from, d_to)

    period = parse_period({"preset": "custom", "date_from": "2026-09-01", "date_to": "2026-09-30",
                            "compare": "previous"})
    report = build_report(fetch, _spec(), period, group_by="pic")
    assert len(calls) == 2
    assert report["totals"]["current"]["count"] == 3
    assert report["totals"]["compare"]["count"] == 1     # đúng 1 dòng tháng 8
    group_a = next(g for g in report["groups"] if g["key"] == "A")
    assert group_a["current"]["count"] == 2 and group_a["compare"]["count"] == 1
    group_b = next(g for g in report["groups"] if g["key"] == "B")
    assert group_b["compare"] is None                    # B không xuất hiện ở kỳ trước


# ── report_xlsx(): dòng Tổng đầu tiên + cột so sánh/±% ───────────────────────────
def test_xlsx_co_dong_tong_dau_va_cot_so_sanh():
    period = parse_period({"preset": "custom", "date_from": "2026-09-01", "date_to": "2026-09-30",
                            "compare": "previous"})
    report = build_report(_fetch_all, _spec(), period, group_by="pic")
    resp = report_xlsx("test-khung-bao-cao", report)

    wb = load_workbook(BytesIO(resp.body))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert any(h and "kỳ trước" in h for h in headers)
    assert any(h and "±%" in h for h in headers)
    assert ws.cell(row=2, column=1).value == "Tổng"       # dòng Tổng ngay sau tiêu đề
    assert ws.max_row == 2 + len(report["groups"])         # tiêu đề + Tổng + mỗi nhóm

    #  Chỉ số kind="percent" (P03) xuất Ô SỐ (cộng/lọc/pivot được), định dạng Excel gắn
    #  chữ "%" chứ không còn xuất chuỗi "xx.x%" — xem `export_xlsx._FORMATS["percent"]`.
    pct_col = headers.index(next(h for h in headers if h and h.startswith("Tỷ lệ đủ")
                                  and "kỳ trước" not in h and "±%" not in h)) + 1
    cell = ws.cell(row=2, column=pct_col)
    assert cell.value == 100.0
    assert cell.number_format == '0.0"%"'


def test_xlsx_khong_co_cot_so_sanh_khi_compare_none():
    period = parse_period({"preset": "this_month", "compare": "none"}, today=date(2026, 9, 30))
    report = build_report(_fetch_all, _spec(), period, group_by="pic")
    resp = report_xlsx("test-khung-bao-cao", report)

    headers = [c.value for c in load_workbook(BytesIO(resp.body)).active[1]]
    assert not any(h and "kỳ trước" in h for h in headers)
    assert not any(h and "±%" in h for h in headers)


# ── rank_by: khối "Top" xếp theo CHỈ SỐ, không theo số hàng ──────────────────────
def test_breakdown_xep_theo_rank_by_khong_theo_so_hang():
    """Sale có 2 dòng (150đ), Mkt có 1 dòng (300đ): Top theo `amount` phải đưa Mkt lên đầu —
    đếm số hàng thì Sale thắng, đúng lỗi khiến "Top NCC" ra NCC nhiều dòng thay vì chi nhiều."""
    rows = [_row("2026-09-05", "Sale", [], 100), _row("2026-09-06", "Sale", [], 50),
            _row("2026-09-07", "Mkt", [], 300)]
    dept = DimensionSpec("dept", "Phòng", key_of=lambda r: [(r.dept, r.dept)])
    spec = ReportSpec(date_of=lambda r: date.fromisoformat(r.date),
                      metrics=[MetricSpec("count", "Số dòng", value_of=lambda r: 1),
                               MetricSpec("amount", "Giá trị", kind="money", value_of=lambda r: r.amount)],
                      dimensions={"dept": dept}, breakdowns={"dept": dept}, rank_by="amount")
    data = aggregate(rows, spec, date(2026, 9, 1), date(2026, 9, 30), "day", "dept")
    assert [(i["key"], i["value"]) for i in data["breakdowns"]["dept"]] == [("Mkt", 300), ("Sale", 150)]
    assert [g["key"] for g in data["groups"]] == ["Mkt", "Sale"]


def test_breakdown_bo_nhom_bang_0_va_cat_ngon_10():
    rows = [_row("2026-09-05", f"P{i:02d}", [], i) for i in range(15)]  # P00 có amount 0
    dept = DimensionSpec("dept", "Phòng", key_of=lambda r: [(r.dept, r.dept)])
    spec = ReportSpec(date_of=lambda r: date.fromisoformat(r.date),
                      metrics=[MetricSpec("amount", "Giá trị", kind="money", value_of=lambda r: r.amount)],
                      breakdowns={"dept": dept})
    items = aggregate(rows, spec, date(2026, 9, 1), date(2026, 9, 30), "day", None)["breakdowns"]["dept"]
    assert len(items) == 10 and items[0]["key"] == "P14" and all(i["value"] for i in items)


def test_meta_khai_rank_by_mac_dinh_la_chi_so_dau():
    period = parse_period({"preset": "custom", "date_from": "2026-09-01", "date_to": "2026-09-30",
                           "compare": "none"})
    assert build_report(_fetch_all, _spec(), period)["meta"]["rank_by"] == "count"


# ── H1: trend kỳ so sánh dời lên trục/độ hạt của kỳ NÀY (không tự tính độ hạt riêng) ────────
def test_trend_ky_so_sanh_gop_theo_truc_hien_tai_khong_theo_do_hat_rieng():
    """`this_quarter` ngày 01/08/2026: kỳ này 07-01..08-01 (32 ngày -> tuần), kỳ so sánh (lịch
    lùi 3 tháng) 04-01..05-01 (31 ngày -> lẽ ra tự tính ra 'ngày'). Trước fix, hai `trend` khác
    độ hạt bị ghép theo CHỈ SỐ mốc (tuần 0 ↔ ngày 0...) — sai âm thầm. Sau fix: dòng của kỳ so
    sánh được DỜI đúng số ngày `date_from - compare_from` rồi gộp lên TRỤC TUẦN của kỳ này."""
    rows = [_row("2026-07-05", "Sale", [], 10),     # kỳ này (07-01..08-01)
            _row("2026-04-05", "Sale", [], 5)]       # kỳ so sánh (04-01..05-01)

    def fetch(d_from, d_to):
        return [r for r in rows if d_from <= date.fromisoformat(r.date) <= d_to]

    period = parse_period({"preset": "this_quarter", "compare": "previous"}, today=date(2026, 8, 1))
    assert period.granularity == "week"
    assert (period.compare_from, period.compare_to) == (date(2026, 4, 1), date(2026, 5, 1))

    report = build_report(fetch, _spec(), period)
    axis = bucket_axis(period.date_from, period.date_to, "week")
    assert len(report["trend"]) == len(axis)          # LUÔN cùng độ dài — cùng một trục

    target_key = bucket_of(date(2026, 7, 5), "week")  # mốc của dòng "07-05" (kỳ này)
    entry = next(t for t in report["trend"] if t["key"] == target_key)
    assert entry["current"]["count"] == 1             # dòng kỳ này đúng mốc của nó
    assert entry["compare"]["count"] == 1             # dòng "04-05" DỜI TỚI ĐÚNG mốc này (91 ngày)
    other_cmp_total = sum((t["compare"] or {}).get("count", 0)
                          for t in report["trend"] if t["key"] != target_key)
    assert other_cmp_total == 0                       # không rơi lạc sang mốc khác


# ── H2: chỉ số dẫn xuất chia cho 0 phải ra None ("chưa có dữ liệu"), không phải 0 ───────────
def test_derived_ve_null_khi_mau_so_bang_0_khong_phai_0():
    """Kỳ trống dữ liệu (vd 'Hôm nay' chưa có phát sinh): mọi metric = 0, `avg_amount`/
    `pct_full` (mẫu số = count = 0) phải là `None` — 0% cũ từng đọc nhầm thành '-100% so với
    kỳ trước' trên FE (report-kpi-row.tsx) và ghi đè giá trị giả trong Excel."""
    period = parse_period({"preset": "today", "compare": "none"}, today=date(2026, 9, 28))
    report = build_report(lambda d_from, d_to: [], _spec(), period, group_by="pic")
    assert report["totals"]["current"]["avg_amount"] is None
    assert report["totals"]["current"]["pct_full"] is None
    assert all(t["current"]["pct_full"] is None for t in report["trend"])
    assert report["groups"] == []                     # không có hàng nào -> không có nhóm nào


# ── L1: `groups` là HỢP khóa hai kỳ — nhóm hết việc ở kỳ này không được biến mất ────────────
def test_groups_hop_khoa_chi_co_o_ky_so_sanh_va_ve_0_khong_phai_bien_mat():
    rows_aug_extra = ROWS_AUG + [_row("2026-08-06", "Sale", ["Z"], 999)]  # "Z" chỉ có ở kỳ trước

    def fetch(d_from, d_to):
        pool = ROWS_SEP + rows_aug_extra
        return [r for r in pool if d_from <= date.fromisoformat(r.date) <= d_to]

    period = parse_period({"preset": "custom", "date_from": "2026-09-01", "date_to": "2026-09-30",
                           "compare": "previous"})
    report = build_report(fetch, _spec(), period, group_by="pic")
    z = next(g for g in report["groups"] if g["key"] == "Z")
    assert z["current"]["count"] == 0
    assert z["current"]["avg_amount"] is None          # số đo DẪN XUẤT về None, không phải 0
    assert z["compare"]["count"] == 1 and z["compare"]["amount"] == 999
    #  Vẫn sắp theo chỉ số XẾP HẠNG của kỳ NÀY, giảm dần — "Z" (0 ở kỳ này) rơi xuống cuối.
    assert report["groups"][-1]["key"] == "Z"


# ── M5: nhãn nhóm không bị Excel hiểu nhầm thành công thức ──────────────────────────────────
def test_xlsx_nhan_nhom_bat_dau_bang_dau_bang_khong_thanh_cong_thuc():
    """Tên NCC/phòng ban/nhóm hàng đều do người dùng gõ — một nhãn bắt đầu bằng '=' (vd NCC đặt
    tên '=1+1') không được để openpyxl/Excel hiểu là công thức và CHẠY nó (CSV/Excel injection)."""
    rows = [_row("2026-09-05", "=1+1", [], 100), _row("2026-09-06", "+HYPERLINK(\"x\")", [], 50)]
    dept = DimensionSpec("dept", "Phòng", key_of=lambda r: [(r.dept, r.dept)])
    spec = ReportSpec(date_of=lambda r: date.fromisoformat(r.date),
                      metrics=[MetricSpec("count", "Số dòng", value_of=lambda r: 1)],
                      dimensions={"dept": dept})
    period = parse_period({"preset": "custom", "date_from": "2026-09-01", "date_to": "2026-09-30",
                           "compare": "none"})
    report = build_report(lambda d_from, d_to: rows, spec, period, group_by="dept")
    resp = report_xlsx("test-cong-thuc", report)

    wb = load_workbook(BytesIO(resp.body))
    ws = wb.active
    dim_col = [c.value for c in ws[1]].index("Phòng") + 1
    labels = {ws.cell(row=r, column=dim_col).value for r in range(2, ws.max_row + 1)}
    assert {"=1+1", '+HYPERLINK("x")'} <= labels       # NGUYÊN VĂN, không bị nuốt/đổi
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=dim_col)
        if c.value in ("=1+1", '+HYPERLINK("x")'):
            assert c.data_type == "s"                  # ép CHUỖI — Excel không chạy như công thức
