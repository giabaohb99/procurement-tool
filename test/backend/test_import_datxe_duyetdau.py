"""Nhập hàng loạt danh mục Đặt xe (Xe · Tài xế) + Duyệt dấu (Loại con dấu).

Chạy thẳng `catalog_import.run` + `service.revert_batch` trên SQLite in-memory,
phủ: apply tạo mới, dedupe cập nhật (Xe theo biển số, Tài xế theo SĐT, Loại con dấu
theo tên), dry-run không ghi, revert xoá bản ghi mới.
"""
from app.modules.company.model import Company
from app.modules.import_tool import catalog_import, doc_import, service
from app.modules.import_tool.model import (ImportBatch, ImportChange, ImportMode,
                                           ImportModule, ImportStatus)
from app.modules.seal_request.model import SealRequest, SealType
from app.modules.seal_request.service import get_company_ids
from app.modules.vehicle_booking.model import Driver, Vehicle, VehicleBooking


def _wb(module: int, rows_by_attr):
    import openpyxl
    adapter = catalog_import.ADAPTERS[module]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = adapter["sheet"]
    for i, f in enumerate(adapter["fields"], start=1):
        ws.cell(row=1, column=i, value=f["header"])
    for ri, row in enumerate(rows_by_attr, start=2):
        for i, f in enumerate(adapter["fields"], start=1):
            if f["attr"] in row:
                ws.cell(row=ri, column=i, value=row[f["attr"]])
    return wb


def _batch(db, module, mode=ImportMode.APPLY):
    b = ImportBatch(module=module, mode=mode, filename="t.xlsx", file_id=0,
                    status=ImportStatus.QUEUED, created_by=1, updated_by=1)
    db.add(b); db.commit(); db.refresh(b)
    return b


def test_import_vehicle_creates_and_dedupes_by_plate(db):
    wb = _wb(ImportModule.VEHICLE, [
        {"license_plate": "51A-111.11", "model": "Vios", "type": "Xe con", "capacity": 4},
        {"license_plate": "51C-222.22", "model": "Transit", "type": "Xe tải", "capacity": 1.5},
    ])
    b = _batch(db, ImportModule.VEHICLE)
    catalog_import.run(db, b, wb, apply=True)
    assert b.status == ImportStatus.DONE and b.created_count == 2
    assert db.query(Vehicle).filter(Vehicle.license_plate == "51A-111.11").one().model == "Vios"

    # Cùng biển số → CẬP NHẬT, không tạo trùng.
    wb2 = _wb(ImportModule.VEHICLE, [{"license_plate": "51A-111.11", "model": "Vios 2024"}])
    b2 = _batch(db, ImportModule.VEHICLE)
    catalog_import.run(db, b2, wb2, apply=True)
    assert b2.updated_count == 1
    assert db.query(Vehicle).filter(Vehicle.license_plate == "51A-111.11").count() == 1
    assert db.query(Vehicle).filter(Vehicle.license_plate == "51A-111.11").one().model == "Vios 2024"


def test_import_driver_dedupes_by_phone(db):
    wb = _wb(ImportModule.DRIVER, [{"phone": "0909000001", "name": "Nguyễn Văn A", "license_class": "B2"}])
    b = _batch(db, ImportModule.DRIVER)
    catalog_import.run(db, b, wb, apply=True)
    assert b.created_count == 1
    wb2 = _wb(ImportModule.DRIVER, [{"phone": "0909000001", "name": "Nguyễn Văn A", "license_class": "C"}])
    b2 = _batch(db, ImportModule.DRIVER)
    catalog_import.run(db, b2, wb2, apply=True)
    assert b2.updated_count == 1
    assert db.query(Driver).filter(Driver.phone == "0909000001").one().license_class == "C"


def test_import_seal_type_dry_run_keeps_no_row(db):
    wb = _wb(ImportModule.SEAL_TYPE, [{"name": "Dấu tròn công ty", "description": "Dấu pháp nhân"}])
    b = _batch(db, ImportModule.SEAL_TYPE, ImportMode.DRY_RUN)
    catalog_import.run(db, b, wb, apply=False)
    assert b.status == ImportStatus.DONE and b.created_count == 1
    assert db.query(SealType).filter(SealType.name == "Dấu tròn công ty").count() == 0  # thử → không ghi


def test_revert_vehicle_deletes_new_rows(db):
    wb = _wb(ImportModule.VEHICLE, [{"license_plate": "51A-999.99", "model": "X"}])
    b = _batch(db, ImportModule.VEHICLE)
    catalog_import.run(db, b, wb, apply=True)
    assert db.query(Vehicle).filter(Vehicle.license_plate == "51A-999.99").count() == 1
    assert db.query(ImportChange).filter(ImportChange.batch_id == b.id).count() == 1

    service.revert_batch(db, b, user_id=1)
    assert db.query(Vehicle).filter(Vehicle.license_plate == "51A-999.99").count() == 0
    assert b.status == ImportStatus.REVERTED


# ── Chứng từ header-only (Yêu cầu đặt xe / đóng dấu) qua doc_import ────────────
def _doc_wb(module, rows):
    import openpyxl
    adapter = doc_import.DOC_ADAPTERS[module]
    fields = [adapter["code"], *adapter["header_fields"], *adapter.get("line_fields", [])]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = adapter["sheet"]
    for i, f in enumerate(fields, start=1):
        ws.cell(row=1, column=i, value=f["header"])
    for ri, row in enumerate(rows, start=2):
        for i, f in enumerate(fields, start=1):
            if f["attr"] in row:
                ws.cell(row=ri, column=i, value=row[f["attr"]])
    return wb


def test_import_vehicle_booking_document(db):
    wb = _doc_wb(ImportModule.VEHICLE_BOOKING,
                 [{"code": "DX900", "purpose": "Đi họp khách", "request_type": 1, "start_location": "VP"}])
    b = _batch(db, ImportModule.VEHICLE_BOOKING)
    doc_import.run(db, b, wb, apply=True)
    assert b.created_count == 1
    bk = db.query(VehicleBooking).filter(VehicleBooking.code == "DX900").one()
    assert bk.purpose == "Đi họp khách"
    assert bk.status == 5   # mặc định Hoàn thành (dữ liệu lịch sử)


def test_import_seal_request_document_sets_company_link(db):
    co = Company(code="CT1", name="Công ty 1")
    db.add(co)
    db.flush()
    wb = _doc_wb(ImportModule.SEAL_REQUEST,
                 [{"code": "DD900", "purpose": "Đóng dấu HĐ", "company_id": "CT1"}])
    b = _batch(db, ImportModule.SEAL_REQUEST)
    doc_import.run(db, b, wb, apply=True)
    req = db.query(SealRequest).filter(SealRequest.code == "DD900").one()
    assert req.company_id == co.id
    #  post_apply đã ghi BẢNG NỐI công ty → Văn thư/Giám đốc lọc đúng phạm vi.
    assert get_company_ids(db, req.id) == [co.id]


def test_import_document_dedupe_skips_existing_code(db):
    db.add(VehicleBooking(code="DX901", purpose="giữ nguyên", status=5))
    db.flush()
    wb = _doc_wb(ImportModule.VEHICLE_BOOKING, [{"code": "DX901", "purpose": "đè mới"}])
    b = _batch(db, ImportModule.VEHICLE_BOOKING)
    doc_import.run(db, b, wb, apply=True)
    assert b.skipped_count >= 1   # mã đã tồn tại → bỏ qua cả phiếu, KHÔNG đè
    assert db.query(VehicleBooking).filter(VehicleBooking.code == "DX901").one().purpose == "giữ nguyên"


def test_revert_document_deletes_new(db):
    wb = _doc_wb(ImportModule.VEHICLE_BOOKING, [{"code": "DX902", "purpose": "z"}])
    b = _batch(db, ImportModule.VEHICLE_BOOKING)
    doc_import.run(db, b, wb, apply=True)
    assert db.query(VehicleBooking).filter(VehicleBooking.code == "DX902").count() == 1
    service.revert_batch(db, b, user_id=1)
    assert db.query(VehicleBooking).filter(VehicleBooking.code == "DX902").count() == 0
    assert b.status == ImportStatus.REVERTED


# ── Cột theo NHÃN (khớp bản xuất hệ cũ): mã/tên, trạng thái, suy diễn, tệp ─────
def _named_wb(sheet: str, headers: list[str], rows: list[dict]):
    """Workbook với TIÊU ĐỀ tuỳ ý (mô phỏng file xuất hệ cũ), giá trị đọc theo header."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    for ri, row in enumerate(rows, start=2):
        for i, h in enumerate(headers, start=1):
            if h in row:
                ws.cell(row=ri, column=i, value=row[h])
    return wb


def test_precheck_headers_accepts_header_only_docs(db):
    #  Regression: precheck bước upload từng dùng adapter["line_fields"] -> KeyError với
    #  chứng từ header-only (Đặt xe / Duyệt dấu) => UI báo "Không đọc được file".
    from app.modules.import_tool.tasks import precheck_headers
    for module in (ImportModule.VEHICLE_BOOKING, ImportModule.SEAL_REQUEST):
        wb = _named_wb(doc_import.DOC_ADAPTERS[module]["sheet"], ["Mã yêu cầu", "Tiêu đề"],
                       [{"Mã yêu cầu": "X1", "Tiêu đề": "t"}])
        precheck_headers(module, wb)  # không được ném lỗi


def test_import_booking_resolves_company_by_name_and_status_by_label(db):
    db.add(Company(code="IDA", name="CÔNG TY TNHH XUẤT NHẬP KHẨU IDA GLOBAL", tax_code="0314562909"))
    db.flush()
    wb = _named_wb("Yeu cau dat xe",
                   ["Mã yêu cầu", "Tiêu đề", "Loại yêu cầu", "Trạng thái chung", "Công ty",
                    "Lộ trình", "Khứ hồi"],
                   [{"Mã yêu cầu": "DX010", "Tiêu đề": "Đi thăm khách",
                     "Loại yêu cầu": "Đi công tác", "Trạng thái chung": "Hoàn thành",
                     # Công ty ghi kiểu "TÊN - MST" như bản xuất cũ → vẫn khớp.
                     "Công ty": "CÔNG TY TNHH XUẤT NHẬP KHẨU IDA GLOBAL - 0314562909",
                     "Lộ trình": "Văn phòng -> An Giang", "Khứ hồi": "Có"}])
    b = _batch(db, ImportModule.VEHICLE_BOOKING)
    doc_import.run(db, b, wb, apply=True)
    bk = db.query(VehicleBooking).filter(VehicleBooking.code == "DX010").one()
    assert bk.company_id == db.query(Company).filter(Company.code == "IDA").one().id
    assert bk.status == 5          # "Hoàn thành" (nhãn) → BK_COMPLETED
    assert bk.request_type == 1    # "Đi công tác" → công tác
    assert bk.is_self_drive is False
    assert bk.start_location == "Văn phòng" and bk.end_location == "An Giang"
    assert bk.is_round_trip is True


def test_import_delivery_type_and_self_drive_and_vehicle_driver(db):
    db.add(Vehicle(license_plate="51D-629.48", model="MAZDA BT50"))
    db.add(Driver(name="Lưu Nhựt Minh", phone="0937187336"))
    db.flush()
    wb = _named_wb("Yeu cau dat xe",
                   ["Mã yêu cầu", "Loại yêu cầu", "Tên hàng hóa", "Biển số xe", "SĐT tài xế",
                    "Trạng thái tài xế", "Thời gian lấy hàng"],
                   [{"Mã yêu cầu": "DX011", "Loại yêu cầu": "Giao hàng tự lái",
                     "Tên hàng hóa": "N2 MgZn 70", "Biển số xe": "51D-629.48",
                     "SĐT tài xế": "0937187336", "Trạng thái tài xế": "Hoàn thành",
                     "Thời gian lấy hàng": "08:30:00 21/7/2026"}])
    b = _batch(db, ImportModule.VEHICLE_BOOKING)
    doc_import.run(db, b, wb, apply=True)
    bk = db.query(VehicleBooking).filter(VehicleBooking.code == "DX011").one()
    assert bk.request_type == 2 and bk.is_self_drive is True      # "Giao hàng tự lái"
    assert bk.goods_name == "N2 MgZn 70"
    assert bk.assigned_vehicle_id == db.query(Vehicle).one().id   # theo biển số
    assert bk.assigned_driver_id == db.query(Driver).one().id     # theo SĐT
    assert bk.driver_status == 4                                  # "Hoàn thành" → DRV_COMPLETED
    assert bk.start_time == "2026-07-21T08:30:00"                 # "Thời gian lấy hàng" → ISO


def test_import_seal_creates_attachment_placeholder_and_revert_cleans_it(db):
    from app.modules.attachment.model import FileLink, StoredFile
    db.add(Company(code="ABA", name="CÔNG TY TNHH HÓA CHẤT NÔNG NGHIỆP ABA", tax_code="1801818328"))
    db.flush()
    wb = _named_wb("Yeu cau dong dau",
                   ["Mã yêu cầu", "Tiêu đề", "Chi tiết loại", "Trạng thái chung", "Công ty",
                    "Tệp đính kèm"],
                   [{"Mã yêu cầu": "DD010", "Tiêu đề": "Duyệt dấu hợp đồng",
                     "Chi tiết loại": "Phê duyệt dấu", "Trạng thái chung": "Hoàn thành",
                     "Công ty": "CÔNG TY TNHH HÓA CHẤT NÔNG NGHIỆP ABA", "Tệp đính kèm": "AV-ABA.pdf"}])
    b = _batch(db, ImportModule.SEAL_REQUEST)
    doc_import.run(db, b, wb, apply=True)
    req = db.query(SealRequest).filter(SealRequest.code == "DD010").one()
    assert req.title == "Duyệt dấu hợp đồng" and req.status == 4
    assert get_company_ids(db, req.id) == [db.query(Company).filter(Company.code == "ABA").one().id]
    #  Tệp đính kèm giữ chỗ (tên tệp còn, chưa có nội dung) + liên kết đúng phiếu.
    link = db.query(FileLink).filter(FileLink.entity == "seal_request",
                                     FileLink.entity_id == req.id).one()
    assert db.get(StoredFile, link.file_id).filename == "AV-ABA.pdf"

    #  Revert xoá cả phiếu, bảng nối công ty lẫn tệp giữ chỗ.
    service.revert_batch(db, b, user_id=1)
    assert db.query(SealRequest).filter(SealRequest.code == "DD010").count() == 0
    assert db.query(FileLink).filter(FileLink.entity == "seal_request",
                                     FileLink.entity_id == req.id).count() == 0
