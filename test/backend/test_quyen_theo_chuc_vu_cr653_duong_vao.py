"""duoc-CR-653 — mọi ĐƯỜNG đổi chức vụ đều kích hoạt đối chiếu quyền: nhập Excel (có/không khớp danh mục), hoàn tác lô,
áp quá trình công tác, và lõi sửa hồ sơ trong savepoint (dry-run)."""
from datetime import date

import openpyxl
import pytest
from sqlalchemy import event

from app.core import auth
from cr653_tro_giup import mk_admin
from app.modules.audit.model import AuditLog
from app.core.hr_work_history_codes import WorkEventType
from app.modules.employee import position_role_service as svc
from app.modules.employee import service as emp_service
from app.modules.employee import work_history_apply_service as applysvc
from app.modules.employee.model import Employee
from app.modules.employee.position_model import JobPosition, JobPositionRole
from app.modules.employee.schema import EmployeeUpdate
from app.modules.employee.work_history_model import EmployeeWorkHistory
from app.modules.import_tool import catalog_import, service
from app.modules.import_tool.model import ImportBatch, ImportMode, ImportModule, ImportStatus
from app.modules.role.model import Role
from app.modules.user.model import User, UserRole


def mk_role(db, code):
    r = Role(code=code, name=f"Vai trò {code}")
    db.add(r)
    db.flush()
    return r


def mk_pos(db, name, roles=()):
    p = JobPosition(code=f"cv-{name}".lower(), name=name, is_active=True)
    db.add(p)
    db.flush()
    for r in roles:
        db.add(JobPositionRole(job_position_id=p.id, role_id=r.id))
    db.commit()
    return p


def rows_of(db, uid):
    db.expire_all()
    return {r.role_id: r.granted_by_position_id for r in db.query(UserRole).filter(UserRole.user_id == uid)}


def savepoint_that(db):
    db.connection().connection.dbapi_connection.isolation_level = None
    event.listen(db.get_bind(), "begin", lambda conn: conn.exec_driver_sql("BEGIN"))


def run_import(db, rows, uid, apply=True):
    headers = []
    for r in rows:
        headers += [k for k in r if k not in headers]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nhan su"
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    for ri, r in enumerate(rows, 2):
        for i, h in enumerate(headers, 1):
            if h in r:
                ws.cell(row=ri, column=i, value=r[h])
    b = ImportBatch(module=ImportModule.EMPLOYEE, mode=ImportMode.APPLY if apply else ImportMode.DRY_RUN,
                    filename="t.xlsx", file_id=0, status=ImportStatus.QUEUED, created_by=uid, updated_by=uid)
    db.add(b)
    db.commit()
    db.refresh(b)
    catalog_import.run(db, b, wb, apply=apply)
    db.refresh(b)
    return b


@pytest.fixture
def hr(db, seed, cap_quyen):
    savepoint_that(db)
    uid = seed.u_req_id
    cap_quyen(uid, "employee", scope="all", read=True, create=True, write=True, delete=True)
    cap_quyen(uid, "employee_sensitive", scope="all", read=True, write=True)
    return uid


def _user(db, seed):
    return db.query(User).filter(User.employee_id == seed.emp_nstm_id).one()


def test_nhap_excel_chuc_vu_khop_danh_muc_doi_goi(db, seed, hr):
    ra, rb = mk_role(db, "ZA"), mk_role(db, "ZB")
    px, py = mk_pos(db, "Chuyên viên", [ra]), mk_pos(db, "Trưởng nhóm", [rb])
    u = _user(db, seed)
    emp_service.update_employee(db, seed.emp_nstm_id, EmployeeUpdate(position_id=px.id), 1)
    assert rows_of(db, u.id) == {ra.id: px.id}
    b = run_import(db, [{"Mã NV *": "DEMONV", "Chức vụ (mã/tên)": "Trưởng nhóm"}], hr)
    assert b.error_count == 0
    assert rows_of(db, u.id) == {rb.id: py.id}


def test_nhap_excel_chuc_vu_chu_tu_do_go_dong_tu_cap(db, seed, hr):
    ra = mk_role(db, "ZA")
    px = mk_pos(db, "Chuyên viên", [ra])
    u = _user(db, seed)
    emp_service.update_employee(db, seed.emp_nstm_id, EmployeeUpdate(position_id=px.id), 1)
    run_import(db, [{"Mã NV *": "DEMONV", "Chức vụ (mã/tên)": "Chức danh lạ hoắc"}], hr)
    assert db.get(Employee, seed.emp_nstm_id).position_id == 0
    assert ra.id not in rows_of(db, u.id)


def test_nhap_excel_chay_thu_khong_ghi_quyen(db, seed, hr):
    ra = mk_role(db, "ZA")
    px = mk_pos(db, "Chuyên viên", [ra])
    u = _user(db, seed)
    run_import(db, [{"Mã NV *": "DEMONV", "Chức vụ (mã/tên)": "Chuyên viên"}], hr, apply=False)
    assert rows_of(db, u.id) == {}
    assert db.query(AuditLog).filter(AuditLog.action.in_(
        ["auto_grant_position", "auto_revoke_position"])).count() == 0


def test_hoan_tac_lo_nhap_thu_hoi_theo_chuc_vu_cu(db, seed, hr):
    ra, rb = mk_role(db, "ZA"), mk_role(db, "ZB")
    px, py = mk_pos(db, "Chuyên viên", [ra]), mk_pos(db, "Trưởng nhóm", [rb])
    u = _user(db, seed)
    emp_service.update_employee(db, seed.emp_nstm_id, EmployeeUpdate(position_id=px.id), 1)
    b = run_import(db, [{"Mã NV *": "DEMONV", "Chức vụ (mã/tên)": "Trưởng nhóm"}], hr)
    assert rows_of(db, u.id) == {rb.id: py.id}
    res = service.revert_batch(db, b, hr)
    assert res["ok"], res
    assert rows_of(db, u.id) == {ra.id: px.id}


def test_ap_qua_trinh_cong_tac_doi_chuc_vu_doi_goi(world):
    db = world.db
    ra = mk_role(db, "ZA")
    px = mk_pos(db, "Nhân viên KD", [ra])
    eid = world.emp["a3"]
    user = User(email="a3", employee_id=eid, password_hash="x", is_active=True)
    db.add(user)
    db.commit()
    a1 = world.grant("a1", "employee", scope="company", actions=("read", "write"))
    row = EmployeeWorkHistory(employee_id=eid, event_type=int(WorkEventType.APPOINT), from_date=date.today(),
                              position_id=px.id, position_label=px.name, created_by=1, updated_by=1)
    db.add(row)
    db.commit()
    applysvc.apply(db, row, a1.user, auth.get_perm_profile(db, a1.user))
    db.commit()
    assert rows_of(db, user.id) == {ra.id: px.id}


def test_xoa_cache_chay_sau_commit_khong_chi_luc_flush(db, seed):
    ra = mk_role(db, "ZA")
    px = mk_pos(db, "Chuyên viên", [ra])
    u = _user(db, seed)
    emp = db.get(Employee, seed.emp_nstm_id)
    emp.position_id = px.id
    svc.sync_employee_roles(db, emp, 1)                       # chỉ flush
    auth._PERM_CACHE[u.id] = ({"stale": True}, 10 ** 12)      # ai đó dựng lại cache giữa chừng
    db.commit()
    assert u.id not in auth._PERM_CACHE
