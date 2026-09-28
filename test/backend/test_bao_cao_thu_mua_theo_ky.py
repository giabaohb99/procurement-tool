"""P03 — Báo cáo mua hàng kiểu Haravan: `/api/reports/procurement/summary` MỚI + chế độ KÉP
(`preset` mới / `year` cũ) của bốn `/summary` Thu mua đã có. Canh:

  - kỳ so sánh (`compare=previous`) tính đúng cho HAI nguồn có HAI cột ngày khác nhau
    (`Payable.incur_date` cho chi phí, `PurchaseOrder.order_date` cho giá trị đặt/số ĐMH)
    gộp trong CÙNG một báo cáo (Q3.2);
  - công nợ còn lại/quá hạn (snapshot) loại đúng khoản nợ gắn ĐMH KHÔNG THẬT;
  - chiều NCC/NSPT vắng mặt khỏi `meta.dimensions` khi thiếu `purchase_order.read`, ép
    `group_by=supplier|nspt` → 403 (bao-CR-437 — lỗ cũ của `/api/reports/procurement`
    không hề chặn NCC ở backend, route MỚI này phải vá);
  - `year` cũ (bảng gốc) vẫn chạy nguyên hình dạng cũ khi không có `preset`;
  - mỗi `/summary/export` đòi ĐÚNG quyền export của bảng nguồn nó tách ra từ đó;
  - tổng số khớp dữ liệu dựng tay (không phải chỉ "gọi được").
"""
import json
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from starlette.datastructures import QueryParams

from app.core.auth import require
from app.core.report_period import parse_period
from app.modules.payable.model import Payable
from app.modules.purchase_order.model import POItem, PurchaseOrder
from app.modules.report import controller as report_ctrl
from app.modules.report import service as report_service
from app.modules.report.procurement_summary_service import compute_procurement_summary
from app.modules.purchase_progress import controller as pp_ctrl
from app.modules.survey import controller as report_ctrl_survey
from app.modules.survey_progress import controller as sp_ctrl


def make_request(**params):
    """`Request` giả — các route ở đây chỉ đụng `request.query_params`."""
    return SimpleNamespace(query_params=QueryParams([(k, str(v)) for k, v in params.items()]))


def unwrap(resp):
    return json.loads(resp.body)["data"]


def _see_all(monkeypatch):
    """`user=SimpleNamespace(id=0)` không có vai trò nào trong DB seed -> `report_dept_scope`
    thật trả `set()` (KHÔNG thấy phòng ban nào, đúng luật M4 mới) — test không canh scope thì
    tắt nó đi, giống `_see_all` của `test_bao_cao_dong_ycmh_tong_hop.py`."""
    monkeypatch.setattr(report_service, "report_dept_scope", lambda db, user: None)


def _po(db, company_id, order_date, status="approved", supplier="NCC A", nspt="NV A", dept="Kho"):
    po = PurchaseOrder(code=f"PO-{order_date}-{supplier}-{status}", company_id=company_id,
                       order_date=order_date, status=status, supplier_code=supplier,
                       supplier_name=supplier, nspt=nspt, department=dept)
    db.add(po)
    db.flush()
    return po


def _item(db, po, qty=1, price=100, vat=0, item_group="Nhãn"):
    it = POItem(po_id=po.id, product_code="SP1", item_group=item_group, qty_order=qty,
               price=price, vat=vat, exchange_rate=1)
    db.add(it)
    db.flush()
    return it


def _payable(db, company_id, incur_date, total, supplier="NCC A", status="unpaid",
            due_date="", po_id=0):
    p = Payable(company_id=company_id, incur_date=incur_date, total=total, remaining=total,
               status=status, due_date=due_date, supplier_code=supplier, supplier_name=supplier,
               po_id=po_id)
    db.add(p)
    db.flush()
    return p


class TestBaoCaoMuaHangTongHop:
    def test_ky_so_sanh_khop_hai_nguon_ngay_khac_nhau(self, db, seed, monkeypatch):
        _see_all(monkeypatch)
        cid = seed.company_id
        po_cur = _po(db, cid, "2026-09-10")
        _item(db, po_cur, qty=10, price=100)          # order_value kỳ này = 1000
        po_cmp = _po(db, cid, "2026-08-15")
        _item(db, po_cmp, qty=5, price=100)            # order_value kỳ trước = 500
        _payable(db, cid, "2026-09-05", 300)           # chi phí kỳ này
        _payable(db, cid, "2026-08-10", 150)           # chi phí kỳ trước
        db.commit()
        user = SimpleNamespace(id=0)

        period = parse_period({"preset": "custom", "date_from": "2026-09-01",
                               "date_to": "2026-09-30", "compare": "previous"})
        data = compute_procurement_summary(db, user, period, cid, None, show_ncc=True)
        cur, cmp = data["totals"]["current"], data["totals"]["compare"]
        assert cur["order_value"] == 1000.0 and cur["spend"] == 300.0 and cur["po_count"] == 1
        assert cmp["order_value"] == 500.0 and cmp["spend"] == 150.0 and cmp["po_count"] == 1

    def test_cong_no_con_lai_qua_han_loai_dmh_khong_that(self, db, seed, monkeypatch):
        _see_all(monkeypatch)
        cid = seed.company_id
        _payable(db, cid, "2026-09-01", 500, status="unpaid", due_date="2026-08-01")  # quá hạn
        _payable(db, cid, "2026-09-02", 200, status="unpaid", due_date="2026-12-01")  # còn hạn
        _payable(db, cid, "2026-09-03", 100, status="paid", due_date="2026-08-01")    # đã trả
        po_draft = _po(db, cid, "2026-09-01", status="draft")                          # ĐMH nháp
        _payable(db, cid, "2026-09-04", 999, po_id=po_draft.id)                        # -> phải loại
        db.commit()
        user = SimpleNamespace(id=0)

        period = parse_period({"preset": "custom", "date_from": "2026-09-01",
                               "date_to": "2026-09-30", "compare": "none"})
        data = compute_procurement_summary(db, user, period, cid, None, show_ncc=True)
        cur = data["totals"]["current"]
        assert cur["debt_remaining"] == 700.0      # 500 + 200, KHÔNG cộng khoản của ĐMH nháp
        assert cur["debt_overdue"] == 500.0
        assert cur["spend"] == 800.0               # 3 khoản hợp lệ, KHÔNG cộng khoản ĐMH nháp

    def test_group_by_ncc_bi_chan_khi_thieu_quyen_xem_dmh(self, db, seed):
        period = parse_period({"preset": "this_month"})
        user = SimpleNamespace(id=0)
        with pytest.raises(HTTPException) as exc:
            compute_procurement_summary(db, user, period, None, "supplier", show_ncc=False)
        assert exc.value.status_code == 403
        with pytest.raises(HTTPException) as exc2:
            compute_procurement_summary(db, user, period, None, "nspt", show_ncc=False)
        assert exc2.value.status_code == 403
        # Không ép group_by NCC/NSPT thì vẫn chạy được — chỉ thiếu 2 chiều trong meta.
        data = compute_procurement_summary(db, user, period, None, None, show_ncc=False)
        dim_keys = {d["key"] for d in data["meta"]["dimensions"]}
        assert "supplier" not in dim_keys and "nspt" not in dim_keys
        assert "breakdowns" in data and "supplier" not in data["breakdowns"]

    def test_dimension_ncc_hien_va_dung_so_khi_co_quyen(self, db, seed, monkeypatch):
        _see_all(monkeypatch)
        cid = seed.company_id
        po = _po(db, cid, "2026-09-10", supplier="NCC X", nspt="NV Y")
        _item(db, po, qty=2, price=100)
        db.commit()
        user = SimpleNamespace(id=0)
        period = parse_period({"preset": "custom", "date_from": "2026-09-01",
                               "date_to": "2026-09-30", "compare": "none"})
        data = compute_procurement_summary(db, user, period, cid, "supplier", show_ncc=True)
        dim_keys = {d["key"] for d in data["meta"]["dimensions"]}
        assert {"supplier", "nspt"} <= dim_keys
        group = next(g for g in data["groups"] if g["key"] == "NCC X")
        assert group["current"]["order_value"] == 200.0

    def test_export_excel_ra_dung_workbook_khop_so_tren_man(self, db, seed, monkeypatch):
        from io import BytesIO

        from openpyxl import load_workbook

        from app.modules.report.summary_controller import procurement_summary_export

        _see_all(monkeypatch)
        cid = seed.company_id
        po = _po(db, cid, "2026-09-10")
        _item(db, po, qty=1, price=1000)
        db.commit()
        user = SimpleNamespace(id=0)
        resp = procurement_summary_export(
            make_request(preset="custom", date_from="2026-09-01", date_to="2026-09-30",
                        compare="none", company_id=cid), db=db, user=user)
        wb = load_workbook(BytesIO(resp.body))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert any(h == "Giá trị đặt hàng" for h in headers)
        col = headers.index("Giá trị đặt hàng") + 1
        assert ws.cell(row=2, column=col).value == 1000.0   # dòng "Tổng" khớp số trên màn

    def test_scope_phong_ban_ap_dung_nhu_bang_matrix(self, db, seed, monkeypatch):
        """M4: `/procurement/summary` áp CÙNG `report_dept_scope` với bảng nguồn `/matrix` —
        phòng ban YÊU CẦU chỉ thấy chi phí/giá trị đặt của phòng mình, không phải toàn công ty."""
        cid = seed.company_id
        po_mine = _po(db, cid, "2026-09-10", dept="Phòng Test")
        _item(db, po_mine, qty=1, price=100)
        po_other = _po(db, cid, "2026-09-11", dept="Phòng Khác")
        _item(db, po_other, qty=1, price=900)
        db.commit()
        monkeypatch.setattr(report_service, "report_dept_scope", lambda db, user: {"Phòng Test"})
        user = SimpleNamespace(id=0)

        period = parse_period({"preset": "custom", "date_from": "2026-09-01",
                               "date_to": "2026-09-30", "compare": "none"})
        data = compute_procurement_summary(db, user, period, cid, None, show_ncc=True)
        assert data["totals"]["current"]["order_value"] == 100.0   # KHÔNG cộng đơn "Phòng Khác"

    def test_chi_phi_gan_duoc_phong_ban_va_chia_theo_nhom_hang(self, db, seed, monkeypatch):
        """Finding #3: `Payable` không tự có phòng ban/nhóm hàng — tra qua `department_id`
        (bảng Phòng ban) và qua ĐMH liên quan (`po_id`); ĐMH có NHIỀU nhóm hàng thì chia chi
        phí theo TỶ TRỌNG giá trị đặt của từng nhóm, không còn gộp hết vào '(Chưa gắn)'."""
        from app.modules.department.model import Department

        _see_all(monkeypatch)
        cid = seed.company_id
        dept = Department(code="PT-F3", name="Phòng Test F3", company_id=cid)
        db.add(dept)
        db.flush()
        po = _po(db, cid, "2026-09-10")
        _item(db, po, qty=1, price=300, item_group="Nhãn")     # order_value 300
        _item(db, po, qty=1, price=100, item_group="Thùng")    # order_value 100
        pay = Payable(company_id=cid, incur_date="2026-09-12", total=400, remaining=400,
                     status="unpaid", supplier_code="NCC A", supplier_name="NCC A",
                     po_id=po.id, department_id=dept.id)
        db.add(pay)
        db.commit()
        user = SimpleNamespace(id=0)

        period = parse_period({"preset": "custom", "date_from": "2026-09-01",
                               "date_to": "2026-09-30", "compare": "none"})
        by_dept = compute_procurement_summary(db, user, period, cid, "department", show_ncc=True)
        dept_group = next(g for g in by_dept["groups"] if g["key"] == "Phòng Test F3")
        assert dept_group["current"]["spend"] == 400.0   # tra qua bảng Phòng ban, không rơi "(Chưa gắn)"

        by_group = compute_procurement_summary(db, user, period, cid, "item_group", show_ncc=True)
        spend_by_group = {g["key"]: g["current"]["spend"] for g in by_group["groups"]}
        assert spend_by_group["Nhãn"] == 300.0 and spend_by_group["Thùng"] == 100.0  # tỷ trọng 300:100

    def test_cong_no_la_snapshot_vang_mat_o_nhom_khong_phai_0(self, db, seed, monkeypatch):
        """M1: `debt_remaining`/`debt_overdue` là chỉ số THỜI ĐIỂM — chỉ đúng ở Tổng. Nhóm phải
        THIẾU HẲN khóa (không phải 0 giả), và `meta` phải khai đúng cờ `snapshot`/`helper`."""
        _see_all(monkeypatch)
        cid = seed.company_id
        po = _po(db, cid, "2026-09-10", dept="Kho")
        _item(db, po, qty=1, price=100)
        _payable(db, cid, "2026-09-05", 500, status="unpaid", due_date="2026-08-01")
        db.commit()
        user = SimpleNamespace(id=0)
        period = parse_period({"preset": "custom", "date_from": "2026-09-01",
                               "date_to": "2026-09-30", "compare": "none"})
        data = compute_procurement_summary(db, user, period, cid, "department", show_ncc=True)
        assert "debt_remaining" in data["totals"]["current"]        # Tổng CÓ (snapshot ghi đè)
        for g in data["groups"]:
            assert "debt_remaining" not in g["current"]              # nhóm THIẾU khóa, không phải 0
        meta_by_key = {m["key"]: m for m in data["meta"]["metrics"]}
        assert meta_by_key["debt_remaining"]["snapshot"] is True
        assert meta_by_key["deliveries_done"]["helper"] is True
        assert meta_by_key["spend"]["helper"] is False and meta_by_key["spend"]["snapshot"] is False

    def test_company_id_khong_phai_so_nem_422(self, db, seed, monkeypatch):
        """L3: `company_id=abc` (query string sai kiểu) phải báo 422, không phải `ValueError`
        trần lộ ra 500 như `int("abc")` trước đây."""
        _see_all(monkeypatch)
        period = parse_period({"preset": "this_month"})
        user = SimpleNamespace(id=0)
        with pytest.raises(HTTPException) as exc:
            compute_procurement_summary(db, user, period, "abc", None, show_ncc=True)
        assert exc.value.status_code == 422


class TestPrLinesGhiChuKhopHanhViThat:
    def test_ghi_chu_khop_hanh_vi_that_dong_huy_ve_0_o_moi_cho(self, db, seed, monkeypatch):
        """M2: dòng đã hủy BỎ khỏi mọi chỉ số, kể cả khi "Xem theo" Tiến độ dòng — nhóm 'Đã
        hủy' vì vậy hiện 0 dòng; ghi chú phải khớp đúng hành vi này (ghi chú CŨ nói ngược lại)."""
        from test_bao_cao_dong_ycmh_ticket23 import _make_item, _make_pr

        from app.modules.report.pr_lines_period_service import compute_pr_lines_summary_period

        monkeypatch.setattr(report_service, "report_dept_scope", lambda db, user: None)
        pr = _make_pr(db, seed, "PYC-M2-01", "2026-09-05")
        _make_item(db, pr, "X1", line_status="cancelled")
        db.commit()
        user = SimpleNamespace(id=0)
        period = parse_period({"preset": "custom", "date_from": "2026-09-01",
                               "date_to": "2026-09-30", "compare": "none"})
        data = compute_pr_lines_summary_period(db, user, period, group_by="line_status")
        assert data["totals"]["current"]["lines"] == 0             # dòng hủy không cộng vào tổng
        cancelled = next(g for g in data["groups"] if g["key"] == "cancelled")
        assert cancelled["current"]["lines"] == 0                  # nhóm hủy cũng 0, đúng ghi chú mới
        assert any("MỌI chỉ số" in n for n in data["notes"])
        assert not any("vẫn đếm đủ" in n for n in data["notes"])   # câu SAI cũ đã bỏ

    def test_company_id_khong_phai_so_nem_422(self, db, seed, monkeypatch):
        """L3: `_pr_lines_base_query` (dùng chung bởi bản `year` cũ lẫn bản `preset` mới)."""
        from app.modules.report.pr_lines_period_service import compute_pr_lines_summary_period

        monkeypatch.setattr(report_service, "report_dept_scope", lambda db, user: None)
        user = SimpleNamespace(id=0)
        period = parse_period({"preset": "this_month"})
        with pytest.raises(HTTPException) as exc:
            compute_pr_lines_summary_period(db, user, period, company_id="abc")
        assert exc.value.status_code == 422


class TestTienDoMuaHangNgayNhanNgayDat:
    def test_lan_giao_theo_ngay_nhan_dong_theo_ngay_dat(self, db, seed, cap_quyen):
        """M3: chỉ số GIAO (deliveries/late/on_time) lọc theo NGÀY NHẬN của lần giao; số DÒNG
        lọc theo NGÀY ĐẶT của đơn — một ĐMH đặt tháng 8, nhận tháng 9 phải tách đúng hai kỳ."""
        from app.modules.purchase_order.model import PODelivery

        cid = seed.company_id
        cap_quyen(seed.u_nstm_id, "purchase_order", scope="all", read=True)
        user = SimpleNamespace(id=seed.u_nstm_id)
        po = _po(db, cid, "2026-08-25")            # ĐẶT tháng 8
        it = _item(db, po, qty=10, price=100)
        db.add(PODelivery(po_id=po.id, po_item_id=it.id, received_qty=10,
                          received_date="2026-09-05"))   # NHẬN tháng 9
        db.commit()

        sep = unwrap(pp_ctrl.progress_summary(
            make_request(company_id=cid, preset="custom", date_from="2026-09-01",
                        date_to="2026-09-30", compare="none"), year="", db=db, user=user))
        assert sep["totals"]["current"]["deliveries"] == 1     # nhận trong tháng 9 -> tính ở đây
        assert sep["totals"]["current"]["lines"] == 0           # đặt ở tháng 8 -> KHÔNG tính ở tháng 9

        aug = unwrap(pp_ctrl.progress_summary(
            make_request(company_id=cid, preset="custom", date_from="2026-08-01",
                        date_to="2026-08-31", compare="none"), year="", db=db, user=user))
        assert aug["totals"]["current"]["lines"] == 1            # đặt ở tháng 8 -> tính ở tháng 8
        assert aug["totals"]["current"]["deliveries"] == 0       # nhận ở tháng 9 -> KHÔNG tính ở tháng 8


class TestCheDoKepCuaBonBaoCaoCu:
    """`preset` vắng mặt -> hành vi `year` CŨ (bảng gốc không đổi); có `preset` -> hợp đồng MỚI."""

    def test_pr_lines_summary_khong_preset_la_hinh_dang_cu(self, db, seed, monkeypatch):
        from app.modules.purchase_request.model import PurchaseRequest, PurchaseRequestItem
        from app.modules.report import service as report_service

        monkeypatch.setattr(report_service, "report_dept_scope", lambda db, user: None)
        pr = PurchaseRequest(code="PYC-OLD-01", company_id=seed.company_id, department="Phòng Test",
                             request_date="2026-03-05", status="submitted", is_deleted=False)
        db.add(pr)
        db.flush()
        db.add(PurchaseRequestItem(pr_id=pr.id, product_code="X1", product_name="Sản phẩm X1",
                                   item_group="Nhãn", amount=100, line_status="no_po"))
        db.commit()
        user = SimpleNamespace(id=0)

        old = unwrap(report_ctrl.pr_lines_summary(make_request(year="2026"), db=db, user=user))
        assert set(old.keys()) == {"total", "by_line_status", "by_month", "by_department",
                                   "by_item_group", "by_assignee"}
        assert old["total"]["lines"] == 1

        new = unwrap(report_ctrl.pr_lines_summary(
            make_request(preset="custom", date_from="2026-03-01", date_to="2026-03-31",
                        compare="none"), db=db, user=user))
        assert set(new.keys()) >= {"period", "meta", "totals", "trend", "breakdowns", "notes"}
        assert new["totals"]["current"]["lines"] == 1

    def test_purchase_progress_summary_khong_preset_la_hinh_dang_cu(self, db, seed, cap_quyen):
        cid = seed.company_id
        po = _po(db, cid, "2026-09-10")
        _item(db, po, qty=10, price=100)
        db.commit()
        cap_quyen(seed.u_nstm_id, "purchase_order", scope="all", read=True)
        user = SimpleNamespace(id=seed.u_nstm_id)

        old = unwrap(pp_ctrl.progress_summary(make_request(company_id=cid), year="2026",
                                              db=db, user=user))
        assert "show_supplier" in old and "by_progress_status" in old

        new = unwrap(pp_ctrl.progress_summary(
            make_request(company_id=cid, preset="custom", date_from="2026-09-01",
                        date_to="2026-09-30", compare="none"), year="", db=db, user=user))
        assert new["totals"]["current"]["lines"] == 1     # khớp 1 dòng ĐMH vừa dựng
        assert set(new["meta"].keys()) == {"metrics", "dimensions", "group_by", "rank_by"}

    def test_purchase_progress_group_by_supplier_bi_chan_qua_route(self, db, seed):
        """Cùng luật NCC ở tầng route — không đợi tới `_can_see_ncc`, `supplier.read` gác
        đủ vì đây là `_show_supplier`, không phải `_can_see_ncc`."""
        user = SimpleNamespace(id=0, employee_id=0, email="x@dego.vn")
        with pytest.raises(HTTPException) as exc:
            pp_ctrl.progress_summary(
                make_request(preset="this_month", group_by="supplier"), year="", db=db, user=user)
        assert exc.value.status_code == 403


class TestCheDoKepConLaiSurveyProgressVaSurveyReport:
    def test_survey_progress_summary_hai_che_do(self, db, seed, cap_quyen):
        from app.modules.survey_request.model import SurveyRequest, SurveyRequestLine

        cap_quyen(seed.u_nstm_id, "survey_request", scope="all", read=True)
        sr = SurveyRequest(code="YCKS-OLD-01", company_id=seed.company_id, department="Phòng Test",
                           request_date="2026-09-05", status="submitted")
        db.add(sr)
        db.flush()
        db.add(SurveyRequestLine(survey_request_id=sr.id, item_group="Nhãn", assignee=seed.emp_nstm_code))
        db.commit()
        user = SimpleNamespace(id=seed.u_nstm_id)

        old = unwrap(sp_ctrl.progress_summary(make_request(), year="2026", db=db, user=user))
        assert "by_state" in old and old["total"]["lines"] == 1

        new = unwrap(sp_ctrl.progress_summary(
            make_request(preset="custom", date_from="2026-09-01", date_to="2026-09-30",
                        compare="none"), year="", db=db, user=user))
        assert new["totals"]["current"]["lines"] == 1
        assert set(new["meta"].keys()) == {"metrics", "dimensions", "group_by", "rank_by"}

    def test_survey_report_summary_hai_che_do(self, db, seed, cap_quyen):
        cap_quyen(seed.u_nstm_id, "survey", scope="all", read=True)
        user = SimpleNamespace(id=seed.u_nstm_id)
        # `seed` đã dựng 2 phiếu khảo sát "Nhãn"/"Thùng" với dòng SP có `date` rỗng — dựng thêm
        # 1 dòng NCC có ngày rõ ràng để kiểm lọc kỳ (dòng SP của seed không có ngày liên hệ).
        from app.modules.survey.model import SurveySupplierLine

        db.add(SurveySupplierLine(survey_id=seed.sv_nhan_id, supplier_code="NX", supplier_name="Nhà Xuất NX",
                                  line_approve="Đã duyệt", contact_date="2026-09-10"))
        db.commit()

        #  Gọi hàm route TRỰC TIẾP (bỏ qua FastAPI DI) nên PHẢI tự truyền `None` cho mọi tham
        #  số `Query(...)` không dùng tới — để mặc định là để lại đối tượng `Query(None)` sống
        #  (không phải `None` thật), khiến `if preset:` đọc nhầm thành "có truyền".
        _unset = dict(kind=None, item_group=None, supplier=None, q=None, nspt=None,
                      date_from=None, date_to=None, preset=None, compare=None, group_by=None)
        old = unwrap(report_ctrl_survey.report_summary_(**{**_unset}, db=db, user=user))
        assert "by_approve" in old

        new = unwrap(report_ctrl_survey.report_summary_(
            **{**_unset, "preset": "custom", "date_from": "2026-09-01", "date_to": "2026-09-30",
               "compare": "none"}, db=db, user=user))
        assert new["totals"]["current"]["lines_supplier"] == 1
        dim_keys = {d["key"] for d in new["meta"]["dimensions"]}
        assert {"nspt", "item_group", "line_approve", "kind"} == dim_keys


class TestQuyenXuatExcel:
    """Mỗi `/summary/export` PHẢI đòi đúng quyền export của bảng nguồn nó tách ra."""

    def test_report_export_doi_hoi_report_export(self, db, seed):
        deny = SimpleNamespace(id=seed.u_req_id)   # chưa cấp quyền nào
        with pytest.raises(HTTPException) as exc:
            require("report", "export")(user=deny, db=db)
        assert exc.value.status_code == 403

    def test_report_export_co_quyen_thi_qua(self, db, seed, cap_quyen):
        cap_quyen(seed.u_nstm_id, "report", scope="all", export=True)
        user = SimpleNamespace(id=seed.u_nstm_id)
        assert require("report", "export")(user=user, db=db) is user

    def test_purchase_progress_export_or_gate(self, db, seed, cap_quyen):
        deny = SimpleNamespace(id=seed.u_req_id)
        with pytest.raises(HTTPException):
            pp_ctrl._require_progress_export(user=deny, db=db)
        cap_quyen(seed.u_nstm_id, "purchase_request", scope="all", export=True)
        allow = SimpleNamespace(id=seed.u_nstm_id)
        assert pp_ctrl._require_progress_export(user=allow, db=db) is allow

    def test_survey_progress_export_doi_hoi_survey_request_export(self, db, seed, cap_quyen):
        deny = SimpleNamespace(id=seed.u_req_id)
        with pytest.raises(HTTPException):
            sp_ctrl._require_progress_export(user=deny, db=db)
        cap_quyen(seed.u_nstm_id, "survey_request", scope="all", export=True)
        allow = SimpleNamespace(id=seed.u_nstm_id)
        assert sp_ctrl._require_progress_export(user=allow, db=db) is allow

    def test_survey_report_export_doi_hoi_survey_export(self, db, seed, cap_quyen):
        deny = SimpleNamespace(id=seed.u_req_id)
        with pytest.raises(HTTPException):
            require("survey", "export")(user=deny, db=db)
        cap_quyen(seed.u_nstm_id, "survey", scope="all", export=True)
        allow = SimpleNamespace(id=seed.u_nstm_id)
        assert require("survey", "export")(user=allow, db=db) is allow
