"""bao-CR-470 — HQ2 · HQ4 · HQ6 phân hệ Tra cứu giá hải quan: phần ĐỌC.

Canh các luật mà màn hình và trợ lý AI dựa vào — sai là ra con số sai mà không lỗi:

```
biểu đồ chỉ chạy khi đã lọc                 -> toàn bộ dữ liệu là trộn hàng nghìn mặt hàng
không cộng kg với lít                       -> tách theo đơn vị, vẽ từng đơn vị
kỳ trống vẫn có mặt                         -> biểu đồ đứt đoạn, không vẽ 0
kỳ < 5 dòng không được là "kỳ tốt nhất"     -> tháng 8 ATRAZINE "rẻ nhất" chỉ nhờ 3 dòng
bình quân GIA QUYỀN theo lượng              -> lô 5 kg không nặng bằng lô 20 tấn
```
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from fastapi import HTTPException
from pydantic import ValidationError

from app.modules.customs import service as S
from app.modules.customs.constants import PartyType, RegulationList
from app.modules.customs.ingredient import IngredientTagger, derive_aliases, extract_formulation, retag_all
from app.modules.customs.model import (CustomsIngredientAlias, CustomsLine, CustomsParty,
                                       CustomsRegulation, CustomsTariff)
from app.modules.customs.reader import clean_party_name
from app.modules.customs.schema import RegulationCreate


def _line(db, d, name="ATRAZINE 97% TECH", unit="KGM", price=3.0, adj=None, qty=100, importer_id=0, **kw):
    ln = CustomsLine(reg_date=d, product_name=name, unit_code=unit, price_usd=Decimal(str(price)),
                     adj_price_usd=Decimal(str(adj)) if adj is not None else None,
                     quantity=Decimal(str(qty)), importer_id=importer_id, hs_code=kw.pop("hs", "38089319"), **kw)
    db.add(ln)
    return ln


# ── Gắn hoạt chất + hàm lượng (HQ4) ─────────────────────────────────────────
def test_gan_hoat_chat_theo_tu_khoa_va_hon_hop():
    t = IngredientTagger([("EMAMECTIN", "EMAMECTIN BENZOATE"), ("LUFENURON", "LUFENURON")], [])
    assert t.tag("Emamectin benzoate 40g/l + Lufenuron 20EC")[0] == "EMAMECTIN BENZOATE + LUFENURON"


def test_khong_khop_tu_khoa_thi_lui_ve_ten_thuong_mai():
    t = IngredientTagger([], [("BIPYRHONE", "Chlorantraniliprole 200g/l")])
    assert t.tag("Thuốc trừ sâu BIPYRHONE 20EC")[0] == "Chlorantraniliprole 200g/l"
    assert t.tag("Dung dịch khử trùng máy thận")[0] == ""       # không có gì để nối — để trống


def test_tach_ten_hoat_chat_tu_danh_muc_bvtv():
    got = dict(derive_aliases([
        "Bifenazate 277g/l + Etoxazole 166g/l", "Kanamycin sulfate (min 98%)", "Glufosinate-ammonium",
        "Pretilachlor 300g/l + chất an toàn Fenclorim 100g/l", "Dầu tỏi", "SHANDONG JOPHNE BIOTECHNOLOGY CO LTD",
        "Potassium salts of fatty acids",
    ]))
    assert got["BIFENAZATE"] == "BIFENAZATE" and got["ETOXAZOLE"] == "ETOXAZOLE"
    assert got["KANAMYCIN SULFATE"] == "KANAMYCIN SULFATE"
    assert got["KANAMYCIN"] == "KANAMYCIN SULFATE"              # từ đầu đủ dài → trỏ về tên đầy đủ
    assert got["GLUFOSINATE AMMONIUM"] == "GLUFOSINATE AMMONIUM"
    assert "FENCLORIM" in got                                   # bỏ tiền tố «chất an toàn»
    assert not any("D U" in k or "SHANDONG" in k for k in got)   # tên tiếng Việt / tên công ty bị bỏ
    assert "POTASSIUM" not in got                               # từ đầu chung chung không đứng một mình


def test_hoat_chat_suy_ra_khop_nguyen_tu():
    t = IngredientTagger([], [], derive_aliases(["Sulfur", "Cyhalofop-butyl", "Atrazine"]))
    assert t.tag("SULFURIC ACID 98%")[0] == ""                  # không khớp nửa chữ
    assert t.tag("Thuốc trừ cỏ CYHALOFOP BUTYL 10EC")[0] == "CYHALOFOP BUTYL"
    assert t.tag("Thuốc kỹ thuật Atrazine Min 97%TC")[0] == "ATRAZINE"


def test_tu_khoa_tay_thang_hoat_chat_suy_ra():
    t = IngredientTagger([("EMAMECTIN", "EMAMECTIN BENZOATE")], [], derive_aliases(["Emamectin"]))
    assert t.tag("EMAMECTIN 5WG")[0] == "EMAMECTIN BENZOATE"


@pytest.mark.parametrize("name, expected", [
    ("Thuốc kỹ thuật ATRAZINE 97% TECH", "97%"), ("ATRAZINE 80WP", "80WP"),
    ("Pescon 250 SC", "250SC"), ("Glufosinate 150 g/l", "150G/L"), ("Chất tẩy rửa", ""),
])
def test_tach_ham_luong_dang_bao_che(name, expected):
    assert extract_formulation(name) == expected


def test_gan_lai_toan_bo_sau_khi_doi_danh_muc(db):
    _line(db, date(2026, 1, 5), name="ATRAZINE 97% TECH")
    db.commit()
    db.add(CustomsIngredientAlias(keyword="ATRAZINE", canonical="ATRAZINE"))
    db.commit()
    res = retag_all(db)

    ln = db.query(CustomsLine).one()
    assert res == {"total": 1, "tagged": 1}
    assert ln.active_ingredient == "ATRAZINE" and ln.formulation == "97%"


# ── Thống kê theo kỳ (HQ2) ──────────────────────────────────────────────────
def test_bieu_do_chi_chay_khi_da_loc(db):
    with pytest.raises(HTTPException) as e:
        S.compute_stats(db, {}, "month", "adjusted")
    assert e.value.status_code == 400 and "mã HS" in e.value.detail


def test_tach_theo_don_vi_khong_cong_kg_voi_lit(db):
    for _ in range(3):
        _line(db, date(2026, 1, 5), unit="KGM", price=3)
    _line(db, date(2026, 1, 6), unit="LTR", price=9)
    db.commit()
    st = S.compute_stats(db, {"q": "atrazine"}, "month", "adjusted")

    assert st["unit"] == "KGM" and st["units"][0] == {"unit": "KGM", "count": 3}
    assert st["kpi"]["max"] == 3.0                    # giá lít không lọt vào đường kg


def test_ky_trong_van_co_mat_va_danh_dau(db):
    _line(db, date(2026, 5, 1))
    _line(db, date(2026, 7, 1))
    db.commit()
    st = S.compute_stats(db, {"q": "atrazine"}, "month", "adjusted")

    assert [s["label"] for s in st["series"]] == ["05/2026", "06/2026", "07/2026"]
    assert st["series"][1]["count"] == 0 and st["series"][1]["wavg"] is None
    assert st["coverage"]["empty_periods"] == ["06/2026"]


def test_ky_it_dong_khong_duoc_la_ky_tot_nhat(db):
    """Đúng ca ATRAZINE: tháng rẻ hơn nhưng chỉ 3 dòng thì không được gắn nhãn."""
    for _ in range(3):
        _line(db, date(2026, 8, 3), price=2.8)
    for _ in range(6):
        _line(db, date(2026, 5, 3), price=3.05)
    db.commit()
    st = S.compute_stats(db, {"q": "atrazine"}, "month", "adjusted")

    assert st["best_period"] == "2026-05"
    aug = next(s for s in st["series"] if s["period"] == "2026-08")
    assert aug["low_data"] is True and aug["wavg"] == 2.8


def test_binh_quan_gia_quyen_theo_luong(db):
    _line(db, date(2026, 1, 5), price=10, qty=5)
    _line(db, date(2026, 1, 6), price=2, qty=20000)
    db.commit()
    kpi = S.compute_stats(db, {"q": "atrazine"}, "month", "adjusted")["kpi"]
    assert kpi["wavg"] == round((10 * 5 + 2 * 20000) / 20005, 4)       # không phải (10+2)/2


def test_dai_phan_vi_khong_bi_mot_dong_gia_la_keo_lech(db):
    for price in (2.0, 2.5, 3.0, 3.5):
        _line(db, date(2026, 1, 5), price=price)
    _line(db, date(2026, 1, 6), price=250)                            # gói nhỏ / khai sai đơn vị
    db.commit()
    st = S.compute_stats(db, {"q": "atrazine"}, "month", "adjusted")
    s = st["series"][0]
    assert (s["min"], s["max"]) == (2.0, 250.0)                         # thấp / cao vẫn nói thật
    assert (s["p25"], s["median"], s["p75"]) == (2.5, 3.0, 3.5)        # dải vẽ không bị 250 kéo lên
    assert st["kpi"]["outliers"] == 1


def test_uu_tien_gia_dieu_chinh_hay_chi_gia_khai_bao(db):
    _line(db, date(2026, 1, 5), price=3, adj=4)
    db.commit()
    assert S.compute_stats(db, {"q": "atrazine"}, "month", "adjusted")["kpi"]["wavg"] == 4.0
    assert S.compute_stats(db, {"q": "atrazine"}, "month", "declared")["kpi"]["wavg"] == 3.0


def test_ky_thang_quy_nam_cung_tong_luong(db):
    for m in (1, 2, 4, 5):
        _line(db, date(2026, m, 5), qty=100 * m)
    db.commit()
    totals = {p: sum(s["qty"] for s in S.compute_stats(db, {"q": "atrazine"}, p, "adjusted")["series"])
              for p in ("month", "quarter", "year")}
    assert totals["month"] == totals["quarter"] == totals["year"] == 1200


def test_tu_khoa_khop_ca_hoat_chat_da_gan(db):
    _line(db, date(2026, 1, 5), name="Thuốc trừ cỏ XYZ 80WP", active_ingredient="ATRAZINE")
    db.commit()
    total, _ = S.list_lines(db, {"q": "atrazine"}, 0, 50)
    assert total == 1


# ── So sánh, xếp hạng, độ phủ, xuất Excel ───────────────────────────────────
def test_so_sanh_can_it_nhat_hai_tu_khoa(db):
    with pytest.raises(HTTPException):
        S.compare_terms(db, ["atrazine"], {}, "month", "adjusted")


def test_xep_hang_nha_nhap_khau_theo_luong(db):
    a = CustomsParty(party_type=PartyType.DOMESTIC, dedupe_key="0101", tax_code="0101", name="A")
    b = CustomsParty(party_type=PartyType.DOMESTIC, dedupe_key="0202", tax_code="0202", name="B")
    db.add_all([a, b])
    db.commit()
    _line(db, date(2026, 1, 5), qty=100, importer_id=a.id)
    _line(db, date(2026, 1, 6), qty=300, importer_id=b.id)
    db.commit()
    res = S.rank_importers(db, {"q": "atrazine"})

    assert [i["name"] for i in res["items"]] == ["B", "A"]
    assert res["items"][0]["share"] == 0.75


def test_do_phu_theo_thang(db):
    _line(db, date(2026, 3, 1))
    _line(db, date(2026, 3, 2))
    db.commit()
    cv = S.get_coverage(db)
    assert cv["total"] == 2 and cv["years"][0]["months"][2] == 2 and cv["years"][0]["months"][0] == 0


def test_xuat_excel_du_cot(db):
    _line(db, date(2026, 1, 5))
    db.commit()
    ws = openpyxl.load_workbook(io.BytesIO(S.export_lines_xlsx(db, {}))).active
    assert ws.max_column == 36 and ws.max_row == 2   # 32 nguồn + 2 suy ra + 2 VND (bao-CR-493)
    assert ws.cell(1, 1).value == "Ngày đăng ký"


# ── Pháp lý + biểu thuế (HQ6) ───────────────────────────────────────────────
def _reg(db, code, name, cas, **kw):
    db.add(CustomsRegulation(list_code=int(code), name=name, cas_no=cas, is_active=True, **kw))


def test_tra_nghia_vu_theo_cong_thuc_hoa_hoc(db):
    _reg(db, RegulationList.ND24_PL4, "Sulfuric acid", "7664-93-9", threshold_kg=Decimal("1000"),
         legal_basis="NĐ 24/2026/NĐ-CP Phụ lục IV")
    db.commit()
    res = S.lookup_regulations(db, "H2SO4")
    assert res["formula_cas"] == "7664-93-9"
    assert "1000 kg" in res["items"][0]["obligation"]


def test_canh_bao_nguong_va_hoat_chat_cam_theo_tu_khoa(db):
    _reg(db, RegulationList.ND24_PL4, "Ammonia khan (Amoniac)", "7664-41-7", threshold_kg=Decimal("500"),
         legal_basis="NĐ 24/2026/NĐ-CP Phụ lục IV")
    _reg(db, RegulationList.BANNED_TT75, "Aldrin", "309-00-2", banned_year=1994, legal_basis="TT 75/2025/TT-BNNMT")
    _reg(db, RegulationList.ND24_PL2, "Ammonia solution", "1336-21-6")      # PL II không phải cảnh báo
    db.commit()

    assert [a["threshold_kg"] for a in S.match_alerts(db, {"q": "ammonia"})] == [500.0]
    assert "CẤM từ năm 1994" in S.match_alerts(db, {"q": "aldrin"})[0]["obligation"]
    assert S.match_alerts(db, {"q": "ab"}) == []


def test_bieu_thue_kem_ma_cha(db):
    for hs, mfn in (("3808", ""), ("380891", ""), ("38089199", "3"), ("38089290", "0")):
        db.add(CustomsTariff(hs_code=hs, name_vn=hs, rate_mfn=mfn, fta_json="{}"))
    db.commit()
    assert [r["hs_code"] for r in S.lookup_tariff(db, "38089199")] == ["3808", "380891", "38089199"]


def test_schema_danh_muc_chan_chuoi_qua_dai_va_danh_sach_la():
    """SQLite không ép độ dài VARCHAR — chốt phải ở tầng schema (duoc-CR-316)."""
    with pytest.raises(ValidationError):
        RegulationCreate(list_code=4, name="x", cas_no="1" * 41)
    with pytest.raises(ValidationError):
        RegulationCreate(list_code=99, name="x")
    RegulationCreate(list_code=4, name="Chlorine", cas_no="7782-50-5", threshold_kg=Decimal("25"))


# ── Chuẩn hóa tên + bộ đọc danh mục ─────────────────────────────────────────
@pytest.mark.parametrize("raw, expected", [
    ("CôNG TY TNHH BAYER VIệT NAM", "CÔNG TY TNHH BAYER VIỆT NAM"),
    ("Công Ty TNHH Thú Y Toàn Cầu", "Công Ty TNHH Thú Y Toàn Cầu"),
    ("  NOVADAN   APS ", "NOVADAN APS"),
])
def test_chuan_hoa_hoa_thuong_ten_doi_tuong(raw, expected):
    assert clean_party_name(raw) == expected


def test_doc_mang_js_co_ngoac_vuong_trong_ten():
    """Tên hóa chất có sẵn ngoặc vuông (`Benzo[b]…`) — tách theo `[...]` là cắt nhầm."""
    from scripts.load_customs_catalogs import _js_array
    src = "const X = [\n // chu thich [khong phai mang]\n ['Benzo[b]fluoranthene','205-99-2'],\n ['A\\'B','1-2-3', 500],\n];"
    assert _js_array(src, "X") == [["Benzo[b]fluoranthene", "205-99-2"], ["A'B", "1-2-3", 500]]


# ── Trợ lý AI (HQ5) ─────────────────────────────────────────────────────────
def _ctx(db, monkeypatch, allowed: bool):
    from types import SimpleNamespace
    from app.modules.assistant.tools.base import ToolContext
    monkeypatch.setattr(ToolContext, "can", lambda self, entity, action="read": allowed)
    return ToolContext(db=db, user=SimpleNamespace(id=1))


def test_tool_gia_hai_quan_gac_quyen(db, monkeypatch):
    from app.modules.assistant.tools.customs_tool import CUSTOMS_PRICE_STATS_SPEC
    res = CUSTOMS_PRICE_STATS_SPEC.handler(_ctx(db, monkeypatch, False), {"keyword": "atrazine"})
    assert res["denied"] is True


def test_tool_thoi_diem_mua_khong_chon_thang_it_du_lieu_va_noi_do_tin_cay(db, monkeypatch):
    """A-03: tháng rẻ hơn mà ít dòng thì KHÔNG được đề xuất; một năm dữ liệu thì phải nói ra."""
    from app.modules.assistant.tools.customs_tool import CUSTOMS_BUY_TIMING_SPEC
    for _ in range(3):
        _line(db, date(2026, 8, 3), price=2.8)
    for _ in range(6):
        _line(db, date(2026, 5, 3), price=3.05)
    db.commit()
    res = CUSTOMS_BUY_TIMING_SPEC.handler(_ctx(db, monkeypatch, True), {"keyword": "atrazine"})

    assert res["recommended_month"]["label"] == "05/2026"
    assert res["cheapest_month_any"]["label"] == "08/2026" and res["cheapest_month_any"]["reliable"] is False
    assert res["confidence"] == "thấp" and "MỘT năm" in res["caveat"]


def test_tool_thieu_bo_loc_thi_bao_loi_ro(db, monkeypatch):
    from app.modules.assistant.tools.customs_tool import CUSTOMS_PRICE_STATS_SPEC
    res = CUSTOMS_PRICE_STATS_SPEC.handler(_ctx(db, monkeypatch, True), {})
    assert "mã HS" in res["error"]
