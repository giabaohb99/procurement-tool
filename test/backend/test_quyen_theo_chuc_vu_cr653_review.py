"""duoc-CR-653 — HỒI QUY cho các phát hiện của đợt review (plans/reports/code-reviewer-261020-cr653-review.md).

Mỗi bài ghi tên phát hiện (C1, H1…) và đều ĐỎ trên bản trước khi sửa.
"""
import uuid
from datetime import date

import openpyxl
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import event
from sqlalchemy.exc import IntegrityError

from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.hr_work_history_codes import WorkEventType
from app.main import app
from app.modules.audit.model import AuditLog
from app.modules.employee import position_role_service as svc
from app.modules.employee import service as emp_service
from app.modules.employee import work_history_apply_service as applysvc
from app.modules.employee.model import Employee
from app.modules.employee.position_model import JobPosition, JobPositionRole
from app.modules.employee.schema import EmployeeUpdate
from app.modules.employee.work_history_model import EmployeeWorkHistory
from app.modules.import_tool import catalog_import, service as import_service
from app.modules.import_tool.model import ImportBatch, ImportMode, ImportModule, ImportStatus
from app.modules.role.model import Permission, Role
from app.modules.user import service as user_service
from app.modules.user.model import User, UserRole
from app.modules.user.schema import RoleAssign, UserProvision
from cr653_tro_giup import mk_admin


def mk_role(db, code, **perm):
    """perm = {entity: {action, ...}} → ma trận của vai trò."""
    r = Role(code=code, name=f"Vai trò {code}")
    db.add(r)
    db.flush()
    for entity, acts in perm.items():
        db.add(Permission(role_id=r.id, entity=entity, scope="own", **{f"can_{a}": True for a in acts}))
    db.commit()
    return r


def mk_pos(db, name, roles=()):
    p = JobPosition(code=f"cv-{name}".lower(), name=name, is_active=True)
    db.add(p)
    db.flush()
    for r in roles:
        db.add(JobPositionRole(job_position_id=p.id, role_id=r.id))
    db.commit()
    return p


def rows_of(db, uid):
    db.expire_all()
    return {r.role_id: r.granted_by_position_id for r in db.query(UserRole).filter(UserRole.user_id == uid)}


def user_of(db, eid):
    return db.query(User).filter(User.employee_id == eid).one()


def set_pos(db, eid, pid, actor_id):
    emp_service.update_employee(db, eid, EmployeeUpdate(position_id=pid), actor_id)


def savepoint_that(db):
    db.connection().connection.dbapi_connection.isolation_level = None
    event.listen(db.get_bind(), "begin", lambda conn: conn.exec_driver_sql("BEGIN"))


def run_import(db, rows, uid):
    headers = []
    for r in rows:
        headers += [k for k in r if k not in headers]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nhan su"
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    for ri, r in enumerate(rows, 2):
        for i, h in enumerate(headers, 1):
            if h in r:
                ws.cell(row=ri, column=i, value=r[h])
    b = ImportBatch(module=ImportModule.EMPLOYEE, mode=ImportMode.APPLY, filename="t.xlsx", file_id=0,
                    status=ImportStatus.QUEUED, created_by=uid, updated_by=uid)
    db.add(b)
    db.commit()
    db.refresh(b)
    catalog_import.run(db, b, wb, apply=True)
    db.refresh(b)
    return b


@pytest.fixture
def client_as(db):
    def make(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


# ══ C1 — không tự cấp cho chính mình, mỗi đường một bài ════════════════════════════════════════

def test_c1_patch_ho_so_tu_doi_chuc_vu_cua_minh_khong_tu_cap_goi(db, seed):
    ra = mk_role(db, "ZA")
    pos = mk_pos(db, "X", [ra])
    me = user_of(db, seed.emp_req_id)
    set_pos(db, seed.emp_req_id, pos.id, me.id)         # chính tôi đổi chức vụ của tôi
    assert ra.id not in rows_of(db, me.id)


def test_c1_quan_tri_van_tu_cap_duoc_cho_minh(db, seed):
    ra = mk_role(db, "ZA")
    pos = mk_pos(db, "X", [ra])
    admin = mk_admin(db)
    emp = Employee(code="ADM1", full_name="QT", is_active=True)
    db.add(emp)
    db.flush()
    db.get(User, admin).employee_id = emp.id
    db.commit()
    set_pos(db, emp.id, pos.id, admin)
    assert rows_of(db, admin)[ra.id] == pos.id


def test_c1_nhap_excel_nhanh_danh_muc_tu_nhap_cho_minh(db, seed, cap_quyen):
    savepoint_that(db)
    me = seed.u_req_id
    cap_quyen(me, "employee", scope="all", read=True, create=True, write=True)
    ra = mk_role(db, "ZA")
    mk_pos(db, "Chuyên viên", [ra])
    run_import(db, [{"Mã NV *": seed.emp_req_code, "Chức vụ (mã/tên)": "Chuyên viên"}], me)
    assert db.get(Employee, seed.emp_req_id).position == "Chuyên viên"      # đã đổi chức vụ thật
    assert ra.id not in rows_of(db, me)


def test_c1_nhap_excel_nhanh_chu_tu_do_chi_thu_hoi_khong_cap(db, seed, cap_quyen):
    savepoint_that(db)
    me = seed.u_req_id
    cap_quyen(me, "employee", scope="all", read=True, create=True, write=True)
    ra = mk_role(db, "ZA")
    pos = mk_pos(db, "Chuyên viên", [ra])
    set_pos(db, seed.emp_req_id, pos.id, mk_admin(db))               # quản trị cấp trước
    assert rows_of(db, me)[ra.id] == pos.id
    run_import(db, [{"Mã NV *": seed.emp_req_code, "Chức vụ (mã/tên)": "Chức danh lạ hoắc"}], me)
    assert ra.id not in rows_of(db, me)                               # thu hồi vẫn được


def test_c1_hoan_tac_lo_nhap_khong_cap_lai_cho_nguoi_hoan_tac(db, seed, cap_quyen):
    savepoint_that(db)
    me = seed.u_req_id
    cap_quyen(me, "employee", scope="all", read=True, create=True, write=True)
    ra, rb = mk_role(db, "ZA"), mk_role(db, "ZB")
    pa, pb = mk_pos(db, "A", [ra]), mk_pos(db, "B", [rb])
    set_pos(db, seed.emp_req_id, pa.id, mk_admin(db))
    b = run_import(db, [{"Mã NV *": seed.emp_req_code, "Chức vụ (mã/tên)": "B"}], me)
    res = import_service.revert_batch(db, b, me)
    assert res["ok"], res
    assert db.get(Employee, seed.emp_req_id).position_id == pa.id    # đã về chức vụ cũ
    got = rows_of(db, me)
    assert ra.id not in got and rb.id not in got                      # không đường nào cấp cho chính mình


def test_c1_ap_qua_trinh_cong_tac_cua_chinh_minh_khong_cap(world):
    db = world.db
    ra = mk_role(db, "ZA")
    px = mk_pos(db, "Nhân viên KD", [ra])
    a1 = world.grant("a1", "employee", scope="company", actions=("read", "write"))
    row = EmployeeWorkHistory(employee_id=world.emp["a1"], event_type=int(WorkEventType.APPOINT),
                              from_date=date.today(), position_id=px.id, position_label=px.name,
                              created_by=1, updated_by=1)
    db.add(row)
    db.commit()
    from app.core.auth import get_perm_profile
    applysvc.apply(db, row, a1.user, get_perm_profile(db, a1.user))
    db.commit()
    assert db.get(Employee, world.emp["a1"]).position_id == px.id
    assert ra.id not in rows_of(db, a1.user.id)


def test_c1_mo_khoa_chinh_minh_khong_cap(db, seed):
    ra = mk_role(db, "ZA")
    pos = mk_pos(db, "X", [ra])
    me = user_of(db, seed.emp_req_id)
    emp = db.get(Employee, seed.emp_req_id)
    emp.position_id = pos.id
    me.is_active = False
    db.commit()
    user_service.set_active(db, me.id, True, me.id)
    assert ra.id not in rows_of(db, me.id)


# ══ H1 + escalation — chỉ cấp vai trò mà người thao tác đủ quyền ═══════════════════════════════

def test_h1_nguoi_thao_tac_khong_du_quyen_thi_bo_qua_va_ghi_nhat_ky(db, seed, cap_quyen):
    weak = mk_role(db, "ZWEAK", document={"read"})
    strong = mk_role(db, "ZSTRONG", payable={"read", "write"})
    pos = mk_pos(db, "X", [weak, strong])
    hr = user_of(db, seed.emp_tp_id)
    cap_quyen(hr.id, "document", scope="all", read=True)       # HR chỉ có document.read
    set_pos(db, seed.emp_req_id, pos.id, hr.id)
    got = rows_of(db, user_of(db, seed.emp_req_id).id)
    assert weak.id in got and strong.id not in got
    assert db.query(AuditLog).filter(AuditLog.action == "auto_grant_skipped_position").count() == 1


def test_h1_cap_tai_khoan_bang_dat_mat_khau_chi_cap_phan_duoc_phep(db, seed, cap_quyen):
    strong = mk_role(db, "ZSTRONG", payable={"read", "write"})
    weak = mk_role(db, "ZWEAK", document={"read"})
    pos = mk_pos(db, "X", [strong, weak])
    hr = user_of(db, seed.emp_tp_id)
    cap_quyen(hr.id, "document", scope="all", read=True)
    emp = Employee(code="NV777", full_name="Mới", position_id=pos.id, position="X", is_active=True)
    db.add(emp)
    db.commit()
    u = user_service.provision_user(db, UserProvision(employee_id=emp.id, email="m@x.vn", password="Abcdef12!x"), hr.id)
    got = rows_of(db, u.id)
    assert weak.id in got and strong.id not in got


def test_h1_ap_lai_bao_so_vai_tro_bi_bo_qua(db, seed, cap_quyen, client_as):
    strong = mk_role(db, "ZSTRONG", payable={"read", "write"})
    pos = mk_pos(db, "X")
    holder = user_of(db, seed.emp_req_id)
    db.get(Employee, seed.emp_req_id).position_id = pos.id
    db.add(JobPositionRole(job_position_id=pos.id, role_id=strong.id))
    db.commit()
    actor = user_of(db, seed.emp_tp_id)
    cap_quyen(actor.id, "role", scope="all", read=True, write=True)
    s = client_as(actor).post(f"/api/job-positions/{pos.id}/roles/apply").json()["data"]
    assert s["skipped_no_permission"] == 1 and s["granted"] == 0
    assert strong.id not in rows_of(db, holder.id)


# ══ H2 — không tự thêm vai trò vào gói của chính chức vụ mình giữ ══════════════════════════════

def test_h2_giu_chuc_vu_thi_khong_tu_them_vao_goi_nhung_duoc_bo_bot(db, seed, cap_quyen, client_as):
    ra = mk_role(db, "ZA", document={"read"})
    pos = mk_pos(db, "X")
    me = user_of(db, seed.emp_tp_id)
    db.get(Employee, seed.emp_tp_id).position_id = pos.id
    db.commit()
    cap_quyen(me.id, "role", scope="all", read=True, write=True)
    cap_quyen(me.id, "document", scope="all", read=True)
    c = client_as(me)
    assert c.put(f"/api/job-positions/{pos.id}/roles", json={"role_ids": [ra.id]}).status_code == 403
    db.add(JobPositionRole(job_position_id=pos.id, role_id=ra.id))
    db.commit()
    assert c.put(f"/api/job-positions/{pos.id}/roles", json={"role_ids": []}).status_code == 200


# ══ H3 — vai trò đã xóa / chức vụ bị xóa / hồ sơ bị xóa ════════════════════════════════════════

def test_h3_khong_xoa_duoc_vai_tro_dang_nam_trong_goi(db, seed):
    from app.modules.role import service as role_service
    ra = mk_role(db, "ZA")
    mk_pos(db, "X", [ra])
    with pytest.raises(HTTPException) as e:
        role_service.delete_role(db, ra.id, 1)
    assert e.value.status_code in (400, 409) and "mặc định" in str(e.value.detail)


def test_h3_dong_goi_tro_vao_vai_tro_da_xoa_bi_bo_qua_khong_de_dong_mo_coi(db, seed):
    ra, rb = mk_role(db, "ZA"), mk_role(db, "ZB")
    pos = mk_pos(db, "X", [ra, rb])
    db.delete(ra)                      # xóa cứng thẳng (đi vòng qua chốt) để mô phỏng dữ liệu cũ
    db.commit()
    assert svc.bundle_role_ids(db, pos.id) == [rb.id]
    set_pos(db, seed.emp_req_id, pos.id, mk_admin(db))
    assert set(rows_of(db, user_of(db, seed.emp_req_id).id)) == {rb.id}


def test_h3_xoa_chuc_vu_don_goi_quyen_mac_dinh(db, seed, cap_quyen, client_as):
    ra = mk_role(db, "ZA")
    pos = mk_pos(db, "Bỏ")
    db.add(JobPositionRole(job_position_id=pos.id, role_id=ra.id))
    db.commit()
    me = user_of(db, seed.emp_tp_id)
    cap_quyen(me.id, "job_position", scope="all", read=True, delete=True)
    assert client_as(me).delete(f"/api/job-positions/{pos.id}").status_code == 200
    assert db.query(JobPositionRole).filter(JobPositionRole.job_position_id == pos.id).count() == 0


def test_h3_xoa_cung_ho_so_go_dong_he_thong_giu_dong_tay(db, seed):
    ra, rh = mk_role(db, "ZA"), mk_role(db, "ZH")
    pos = mk_pos(db, "X", [ra])
    u = user_of(db, seed.emp_req_id)
    db.add(UserRole(user_id=u.id, role_id=rh.id))
    db.commit()
    set_pos(db, seed.emp_req_id, pos.id, mk_admin(db))
    assert ra.id in rows_of(db, u.id)
    emp_service.detach_users(db, seed.emp_req_id, 1)
    db.commit()
    assert rows_of(db, u.id) == {rh.id: 0}


# ══ M5 — vai trò cấm ═══════════════════════════════════════════════════════════════════════════

@pytest.mark.parametrize("perm", [{"role": {"write"}}, {"user": {"create"}}, {"user": {"delete"}}])
def test_m5_put_tu_choi_vai_tro_quan_ly_phan_quyen(db, seed, cap_quyen, client_as, perm):
    bad = mk_role(db, "ZMGR", **perm)
    pos = mk_pos(db, "X")
    me = user_of(db, seed.emp_tp_id)
    cap_quyen(me.id, "role", scope="all", read=True, write=True)
    assert client_as(me).put(f"/api/job-positions/{pos.id}/roles", json={"role_ids": [bad.id]}).status_code == 400


def test_m5_administrator_cu_bi_tu_choi_va_luc_cap_bi_bo_qua_ca_voi_quan_tri(db, seed):
    legacy = mk_role(db, "ADMINISTRATOR")
    mgr = mk_role(db, "ZMGR", role={"write"})
    ok = mk_role(db, "ZOK", document={"read"})
    pos = mk_pos(db, "X", [legacy, mgr, ok])           # lọt vào gói bằng đường vòng (dữ liệu cũ / script)
    set_pos(db, seed.emp_req_id, pos.id, mk_admin(db))
    assert set(rows_of(db, user_of(db, seed.emp_req_id).id)) == {ok.id}


# ══ M1 + M2 — trùng dòng, unique, ghi đồng thời ════════════════════════════════════════════════

def test_m2_unique_user_role_va_chen_trung_duoc_xu_ly_em(db, seed):
    ra = mk_role(db, "ZA")
    u = user_of(db, seed.emp_req_id)
    db.add(UserRole(user_id=u.id, role_id=ra.id))
    db.commit()
    db.add(UserRole(user_id=u.id, role_id=ra.id))
    with pytest.raises(IntegrityError):
        db.commit()
    db.rollback()
    # đường đồng bộ gặp trùng thì coi như đã có (không nổ, không làm hỏng phiên)
    assert svc._insert_role(db, u.id, ra.id, 5, 1) is False
    db.commit()
    assert db.query(UserRole).filter(UserRole.user_id == u.id, UserRole.role_id == ra.id).count() == 1


def test_m1_admin_luu_lai_danh_sach_co_vai_tro_he_thong_thi_giu_nguyen_nhan(db, seed):
    ra, rb = mk_role(db, "ZA"), mk_role(db, "ZB")
    pos = mk_pos(db, "X", [ra])
    u = user_of(db, seed.emp_req_id)
    set_pos(db, seed.emp_req_id, pos.id, mk_admin(db))
    user_service.assign_roles(db, u.id, RoleAssign(role_ids=[ra.id, rb.id]), 1)
    assert rows_of(db, u.id) == {ra.id: pos.id, rb.id: 0}


# ══ M4 — script đề xuất ════════════════════════════════════════════════════════════════════════

def test_m4_mac_dinh_can_it_nhat_3_nguoi_va_bo_vai_tro_quan_ly_phan_quyen(db, seed):
    from scripts import de_xuat_quyen_theo_chuc_vu as script
    ok, mgr = mk_role(db, "ZOK"), mk_role(db, "ZMGR", user={"write"})
    pos = mk_pos(db, "P")
    big = mk_pos(db, "Q")
    users = [user_of(db, e) for e in (seed.emp_req_id, seed.emp_tp_id, seed.emp_nstm_id)]
    for i, u in enumerate(users):
        for r in (ok, mgr):
            db.add(UserRole(user_id=u.id, role_id=r.id))
        db.get(Employee, u.employee_id).position_id = pos.id if i == 0 else big.id
    db.commit()
    prop = {p["position_id"]: p for p in script.compute_proposal(db)}
    assert not any(r["proposed"] for r in prop[pos.id]["roles"])                 # 1 người < 3
    q_roles = {r["code"]: r["proposed"] for r in prop[big.id]["roles"]}
    assert q_roles == {"ZOK": False}                  # 2 người < 3, và vai trò cấm không xuất hiện
    prop2 = {p["position_id"]: p for p in script.compute_proposal(db, min_holders=2)}
    assert {r["code"]: r["proposed"] for r in prop2[big.id]["roles"]} == {"ZOK": True}


def test_m4_apply_ghi_nhat_ky(db, seed):
    from scripts import de_xuat_quyen_theo_chuc_vu as script
    ok = mk_role(db, "ZOK")
    pos = mk_pos(db, "P")
    for e in (seed.emp_req_id, seed.emp_tp_id, seed.emp_nstm_id):
        db.add(UserRole(user_id=user_of(db, e).id, role_id=ok.id))
        db.get(Employee, e).position_id = pos.id
    db.commit()
    assert script.apply_proposal(db, script.compute_proposal(db)) == 1
    assert db.query(AuditLog).filter(AuditLog.action == "set_position_roles", AuditLog.entity_id == pos.id).count() == 1


# ══ Lows — gói + nhật ký cùng giao dịch, ghi đồng thời → 409 ═══════════════════════════════════

def test_l1_loi_ghi_nhat_ky_thi_goi_khong_duoc_luu(db, seed, monkeypatch):
    ra = mk_role(db, "ZA")
    pos = mk_pos(db, "X")

    def boom(*a, **k):
        raise RuntimeError("audit hỏng")

    monkeypatch.setattr(svc, "record", boom)
    with pytest.raises(RuntimeError):
        svc.set_bundle(db, pos, [ra.id], 1)
    db.rollback()
    assert db.query(JobPositionRole).filter(JobPositionRole.job_position_id == pos.id).count() == 0


def test_l1_ghi_dong_thoi_trung_unique_tra_409(db, seed, monkeypatch):
    ra = mk_role(db, "ZA")
    pos = mk_pos(db, "X")

    def clash(*a, **k):
        raise IntegrityError("x", {}, Exception("dup"))

    monkeypatch.setattr(svc, "record", clash)
    with pytest.raises(HTTPException) as e:
        svc.set_bundle(db, pos, [ra.id], 1)
    assert e.value.status_code == 409
