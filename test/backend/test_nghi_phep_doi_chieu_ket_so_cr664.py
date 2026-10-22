"""duoc-CR-664 (kh-HR2-3) — BẢNG ĐỐI CHIẾU TRƯỚC / SAU KẾT SỔ phép cuối năm.

Bảng dựng lại từ chính sổ quỹ (không có snapshot): số cũ = còn lại + đã chuyển đi,
số mới = còn lại. Bộ bài chốt: con số khớp, quy đổi nhân tỷ lệ, chạy lại không đổi,
dòng chưa kết sổ không lẫn vào, phạm vi dữ liệu có hiệu lực, số truy vấn cố định,
và tệp Excel mở được.
"""
import json
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.core.auth import get_perm_profile
from app.modules.employee.model import Employee
from app.modules.leave import balance_controller as bc
from app.modules.leave import carryover_review_service as review
from app.modules.leave import carryover_service
from app.modules.leave.balance_model import LeaveBalance
from app.modules.leave.catalog_model import LeaveType
from app.modules.leave.constants import YEAR_END_CARRY, YEAR_END_CONVERT, YEAR_END_DROP

from lich_lam_viec_factory import count_queries


def _type(db, code, name, **kw):
    kw.setdefault("counts_balance", True)
    kw.setdefault("annual_quota_days", 12.0)
    kw.setdefault("year_end_mode", YEAR_END_DROP)
    obj = LeaveType(code=code, name=name, **kw)
    db.add(obj)
    db.flush()
    return obj


def _emp(db, code, name, company=1):
    obj = Employee(code=code, full_name=name, company_id=company,
                   hire_date=date(2020, 1, 1))
    db.add(obj)
    db.flush()
    return obj


def _bal(db, emp, lt, year=2026, **kw):
    kw.setdefault("allocated_days", 12.0)
    obj = LeaveBalance(employee_id=emp.id, year=year, leave_type_id=lt.id,
                       company_id=emp.company_id, **kw)
    db.add(obj)
    db.flush()
    return obj


def _close(db, year=2026):
    rows = db.query(LeaveBalance).filter(LeaveBalance.year == year).all()
    return carryover_service.close_year(db, year, rows, actor=1)


@pytest.fixture
def world(db, cap_quyen):
    comp_off = _type(db, "comp_off", "Nghỉ bù", annual_quota_days=0.0)
    annual = _type(db, "annual", "Phép năm", year_end_mode=YEAR_END_CARRY,
                   carry_over_max_days=5.0)
    swap = _type(db, "swap", "Phép đổi", year_end_mode=YEAR_END_CONVERT,
                 convert_to_type_id=comp_off.id, convert_ratio=0.5)
    drop = _type(db, "unpaid", "Không lương")
    a = _emp(db, "NV001", "An")
    b = _emp(db, "NV002", "Bình", company=2)
    cap_quyen(1, "leave_balance", scope="all", read=True, export=True)
    cap_quyen(2, "leave_balance", scope="company", read=True)
    db.commit()
    return SimpleNamespace(db=db, annual=annual, swap=swap, drop=drop, comp_off=comp_off,
                           a=a, b=b, hr=SimpleNamespace(id=1, employee_id=0))


def _review(w, user=None, year=2026):
    user = user or w.hr
    return review.build_review(w.db, user, get_perm_profile(w.db, user), year)


def test_so_cu_so_moi_chenh_khop_voi_so_quy(world):
    w = world
    _bal(w.db, w.a, w.annual, used_days=4.0)      # còn 8, trần mang 5
    _close(w.db)
    row = _review(w)["items"][0]
    #  Còn 8 -> chuyển 5 -> còn 3. Số cũ phải dựng lại ra 8, không phải 3.
    assert (row["before_days"], row["after_days"], row["diff_days"]) == (8.0, 3.0, -5.0)
    assert row["credited_days"] == 5.0 and row["dest_carried_days"] == 5.0
    assert row["dest_type_name"] == "Phép năm"


def test_quy_doi_nhan_ty_le_va_nhan_dung_loai(world):
    w = world
    _bal(w.db, w.a, w.swap, used_days=4.0)        # còn 8 -> nhận 4 nghỉ bù
    _close(w.db)
    row = _review(w)["items"][0]
    assert (row["before_days"], row["after_days"]) == (8.0, 0.0)
    assert (row["ratio"], row["credited_days"], row["dest_carried_days"]) == (0.5, 4.0, 4.0)
    assert row["dest_type_name"] == "Nghỉ bù"


def test_chay_lai_ket_so_khong_doi_bang(world):
    w = world
    _bal(w.db, w.a, w.annual, used_days=4.0)
    _close(w.db)
    first = _review(w)
    _close(w.db)                                   # lượt hai không mang gì thêm
    assert _review(w) == first


def test_dong_chua_ket_so_va_loai_het_nam_la_mat_khong_vao_bang(world):
    w = world
    _bal(w.db, w.a, w.drop)                        # loại "mất" -> không chuyển gì
    _bal(w.db, w.b, w.annual)
    assert _review(w)["items"] == []               # chưa kết sổ: chưa có gì để đối chiếu
    _close(w.db)
    names = [i["source_type_name"] for i in _review(w)["items"]]
    assert names == ["Phép năm"]


def test_nam_khong_co_du_lieu_va_nam_cuc_doan_tra_rong_khong_loi(world):
    for y in (1, 2026, 9999):
        data = _review(world, year=y)
        assert data["items"] == [] and data["totals"]["row_count"] == 0


def test_pham_vi_du_lieu_co_hieu_luc(world):
    w = world
    _bal(w.db, w.a, w.annual)
    _bal(w.db, w.b, w.annual)
    _close(w.db)
    assert len(_review(w)["items"]) == 2           # phạm vi tất cả
    company_user = SimpleNamespace(id=2, employee_id=0, company_id=1)
    scoped = _review(w, company_user)["items"]
    assert [i["employee_code"] for i in scoped] in (["NV001"], [])  # không bao giờ thấy NV002
    assert "NV002" not in [i["employee_code"] for i in scoped]


def test_so_truy_van_co_dinh_khi_tang_so_nguoi(world):
    w = world
    _bal(w.db, w.a, w.annual)
    _close(w.db)
    _review(w)                                     # làm nóng bộ nhớ đệm quyền (60s)
    _, few = count_queries(w.db, lambda: _review(w))
    for i in range(20):
        e = _emp(w.db, f"X{i:03d}", f"Người {i}")
        _bal(w.db, e, w.annual)
        _bal(w.db, e, w.swap)
    _close(w.db)
    result, many = count_queries(w.db, lambda: _review(w))
    assert len(result["items"]) == 41
    assert len(many) == len(few) <= 4


def test_totals_va_endpoint_xuat_excel(world):
    w = world
    _bal(w.db, w.a, w.annual, used_days=4.0)
    _bal(w.db, w.b, w.swap, used_days=4.0)
    _close(w.db)
    body = json.loads(bc.close_year_review(year=2026, db=w.db, user=w.hr).body)
    assert body["data"]["totals"]["before_days"] == 16.0
    assert body["data"]["totals"]["diff_days"] == -13.0

    res = bc.close_year_review_export(year=2026, db=w.db, user=w.hr)
    sheet = load_workbook(BytesIO(res.body)).active
    assert sheet.max_row == 3                      # tiêu đề + 2 dòng
    assert sheet.cell(1, 4).value.startswith("Số cũ")
    assert "doi-chieu-ket-so-phep-2026" in res.headers["content-disposition"]


def test_xuat_excel_loc_theo_pham_vi_cua_quyen_xuat_khong_phai_quyen_xem(world, cap_quyen):
    """Xem «tất cả» nhưng chỉ được XUẤT phần của mình → tệp Excel chỉ có dòng của mình.

    Bản đầu lọc đường xuất bằng grant `read` nên người có phạm vi xuất hẹp vẫn tải
    về được sổ quỹ toàn công ty (lệch với mọi đường xuất khác, xem duoc-CR-660).
    """
    w = world
    _bal(w.db, w.a, w.annual)
    _bal(w.db, w.b, w.annual)
    _close(w.db)
    cap_quyen(3, "leave_balance", scope="all", read=True)
    cap_quyen(3, "leave_balance", scope="own", export=True)
    w.db.commit()
    viewer = SimpleNamespace(id=3, employee_id=w.a.id)
    assert len(_review(w, viewer)["items"]) == 2
    res = bc.close_year_review_export(year=2026, db=w.db, user=viewer)
    sheet = load_workbook(BytesIO(res.body)).active
    assert sheet.max_row == 2                      # tiêu đề + đúng dòng của NV001
    assert sheet.cell(2, 1).value == "NV001"
