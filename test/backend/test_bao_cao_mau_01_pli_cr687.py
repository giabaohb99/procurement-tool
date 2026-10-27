"""duoc-CR-687 — Báo cáo sử dụng lao động mẫu 01/PLI (khóa 20): `/api/labor-usage-report/summary` (+ export, rows).

Canh: chỉ số nhóm vị trí việc làm / loại HĐ / giới tính, suy nhóm vị trí ở danh sách, 403 thiếu quyền ở cả
ba đường, và bố cục tệp Excel đúng mẫu (ô tiêu đề gộp 3 tầng, dòng số 1…27, dòng Tổng, «x», cột HĐ, cột
trống, mỗi pháp nhân một sheet). Kỳ cố định 2025 -> `as_of` = 31/12/2025.
"""
import uuid
from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.subject_match import SUBJECT_EMPLOYEE
from app.main import app
from app.modules.employee.model import Employee
from app.modules.labor_contract.model import LaborContract
from app.modules.labor_register import form_01_pli_layout as lay
from app.modules.report_access.model import ReportAccess

SIGNED = 2
PROBATION, FIXED, INDEFINITE, SERVICE, COLLABORATOR = 1, 2, 3, 4, 5
SUM = "/api/labor-usage-report/summary"
QS = {"preset": "custom", "date_from": "2025-01-01", "date_to": "2025-12-31", "compare": "none"}
_seq = {"n": 0}


@pytest.fixture
def client_as(db):
    def make_client(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make_client
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def _emp(db, world, tag, dept="B.hc", co="B", **kw) -> Employee:
    base = dict(code=f"P{tag}", full_name=f"Người {tag}", company_id=world.co[co], department_id=world.dept[dept],
                is_active=True, hire_date=date(2020, 1, 1), gender=1, date_of_birth=date(1991, 2, 3),
                id_number="079091000111", social_insurance_no="7911111111", position="Chuyên viên")
    emp = Employee(**{**base, **kw})
    db.add(emp)
    db.flush()
    return emp


def _hd(db, world, emp, ctype, start, end=None, salary=12_000_000, allowance=0, co="B", dept="B.hc"):
    _seq["n"] += 1
    db.add(LaborContract(code=f"HP{_seq['n']:05d}", employee_id=emp.id, company_id=world.co[co],
                         department_id=world.dept[dept], contract_type=ctype, status=SIGNED, start_date=start,
                         end_date=end, base_salary=salary, allowance=allowance))
    db.flush()


def _dataset(db, world):
    boss = _emp(db, world, "boss", job_level=4, education_level=4, gender=2, position="Trưởng phòng")
    _hd(db, world, boss, INDEFINITE, date(2021, 1, 1), allowance=1_500_000)
    uni = _emp(db, world, "uni", education_level=5)
    _hd(db, world, uni, FIXED, date(2025, 1, 1), date(2025, 12, 31))
    col = _emp(db, world, "col", education_level=3, gender=2)
    _hd(db, world, col, PROBATION, date(2025, 10, 1), date(2025, 12, 1))
    _emp(db, world, "plain", education_level=1)                                  # chưa có HĐ
    other = _emp(db, world, "coa", dept="A.kt", co="A", education_level=2)
    _hd(db, world, other, INDEFINITE, date(2022, 1, 1), co="A", dept="A.kt")
    db.commit()


def _actor(db, world, client_as, sensitive=True, contract=True, report=True):
    world.grant("a1", "employee", "all", actions=("read", "export"))
    if sensitive:
        world.grant("a1", "employee_sensitive", "all", actions=("read",))
    if contract:
        world.grant("a1", "labor_contract", "all", actions=("read", "export"))
    actor = world.actor("a1")
    if report:
        db.add(ReportAccess(report_key=int(ReportKey.LABOR_USAGE_REPORT), subject_kind=SUBJECT_EMPLOYEE,
                            subject_id=actor.employee.id, effect=1, reason="t", created_by=0, updated_by=0))
        db.commit()
    return client_as(actor.user)


def _get(client, path=SUM, **extra):
    return client.get(path, params={**QS, **extra})


def test_chi_so_nhom_vi_tri_loai_hd_gioi_tinh(db, world, client_as):
    _dataset(db, world)
    client = _actor(db, world, client_as)
    resp = _get(client, department_id=world.dept["B.hc"])
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    assert data["totals"]["current"] == {
        "people": 4, "female": 2, "job_manager": 1, "job_high": 1, "job_mid": 1, "job_other": 1,
        "ct_indefinite": 1, "ct_fixed": 1, "ct_other": 2}                      # thử việc + chưa có HĐ
    assert data["meta"]["drilldown_metrics"] == list(data["totals"]["current"])
    by_co = _get(client, group_by="company").json()["data"]["groups"]
    assert {g["label"]: g["current"]["people"] for g in by_co}["Công ty A"] == 1


def test_rows_suy_nhom_vi_tri(db, world, client_as):
    _dataset(db, world)
    client = _actor(db, world, client_as)
    dept = world.dept["B.hc"]
    for metric, code in (("job_manager", "Pboss"), ("job_high", "Puni"), ("job_mid", "Pcol"), ("job_other", "Pplain")):
        body = _get(client, SUM + "/rows", department_id=dept, metric=metric).json()["data"]
        assert [r["employee_code"] for r in body["rows"]] == [code], metric
    row = _get(client, SUM + "/rows", department_id=dept, metric="job_manager").json()["data"]["rows"][0]
    assert row["job_category"] == "Nhà quản lý" and row["contract_type"] == "Không xác định thời hạn"
    assert "079091000111" not in str(row)
    grouped = _get(client, SUM + "/rows", department_id=dept, metric="people", group_by="job_category",
                   group_key="3").json()["data"]
    assert [r["employee_code"] for r in grouped["rows"]] == ["Pcol"]


@pytest.mark.parametrize("suffix", ["", "/export", "/rows"])
@pytest.mark.parametrize("kw,message", [({"sensitive": False}, "quyền xem thông tin nhạy cảm của nhân sự"),
                                        ({"contract": False}, "quyền xem hợp đồng lao động")])
def test_thieu_quyen_403(db, world, client_as, suffix, kw, message):
    client = _actor(db, world, client_as, **kw)
    resp = _get(client, SUM + suffix, metric="people")
    assert resp.status_code == 403 and message in resp.text


@pytest.mark.parametrize("suffix", ["", "/export", "/rows"])
def test_chua_giao_bao_cao_403(db, world, client_as, suffix):
    client = _actor(db, world, client_as, report=False)
    assert _get(client, SUM + suffix, metric="people").status_code == 403


def _merged(ws) -> set[str]:
    return {str(r) for r in ws.merged_cells.ranges}


def test_excel_dung_bo_cuc_mau_01_pli(db, world, client_as):
    _dataset(db, world)
    client = _actor(db, world, client_as)
    resp = _get(client, SUM + "/export")
    assert resp.status_code == 200
    assert "bao-cao-su-dung-lao-dong-01-pli-2025-01-01-2025-12-31" in resp.headers["content-disposition"]
    wb = load_workbook(BytesIO(resp.content))
    assert set(wb.sheetnames) == {"Công ty A", "Công ty B"}                 # mỗi pháp nhân một sheet
    ws = wb["Công ty B"]
    assert ws["A1"].value == "Tên doanh nghiệp: Công ty B"
    assert ws["P1"].value == "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" and ws["P2"].value == "Độc lập - Tự do - Hạnh phúc"
    assert ws["A5"].value == "BÁO CÁO TÌNH HÌNH SỬ DỤNG LAO ĐỘNG"
    assert ws["A7"].value.startswith("Kính gửi: Sở Nội vụ")
    assert ws["A8"].value == "1. Thông tin chung" and "Mã số thuế" in ws["A11"].value
    head = 15
    merged = _merged(ws)
    #  Gộp ô đối chiếu bản PDF Mẫu số 01/PLI: «Phụ cấp» là tầng giữa M16:Q16, hai cột BHXH (25, 26) đứng riêng.
    for rng in ("A15:A17", "B15:B17", "H15:K15", "L15:Q15", "M16:Q16", "R15:S15", "T15:X15", "U16:V16",
                "W16:X16", "Y15:Y17", "Z15:Z17", "AA15:AA17", "H16:H17", "L16:L17"):
        assert rng in merged, rng
    assert [ws.cell(row=head + 3, column=c).value for c in range(1, 28)] == list(range(1, 28))
    assert ws.cell(row=head, column=8).value.startswith("Vị trí việc làm")
    assert ws.cell(row=head + 1, column=21).value == "Hiệu lực HĐLĐ xác định thời hạn"
    assert ws.cell(row=head + 2, column=22).value == "Ngày kết thúc"
    assert ws.cell(row=head + 1, column=13).value == "Phụ cấp" and ws.cell(row=head + 2, column=14).value == "Thâm niên VK (%)"
    assert ws.cell(row=head, column=7).value == "Cấp bậc, chức vụ, chức danh nghề, nơi làm việc"
    rows = {ws.cell(row=r, column=2).value: [ws.cell(row=r, column=c).value for c in range(1, 28)]
            for r in range(19, 23)}
    boss, uni, col, plain = (rows[f"Người {t}"] for t in ("boss", "uni", "col", "plain"))
    assert boss[7] == "x" and boss[8:11] == [None, None, None] and boss[11] == 12_000_000 and boss[15] == 1_500_000
    assert boss[19].date() == date(2021, 1, 1) and boss[20:24] == [None] * 4
    assert uni[8] == "x" and uni[20].date() == date(2025, 1, 1) and uni[21].date() == date(2025, 12, 31)
    assert col[9] == "x" and col[22].date() == date(2025, 10, 1) and col[23].date() == date(2025, 12, 1)
    assert plain[10] == "x" and plain[11] is None and plain[19:24] == [None] * 5   # chưa có HĐ
    for blank in (12, 13, 14, 16, 17, 18, 24, 25):                                 # cột 13–15, 17–19, 25–26
        assert all(r[blank] is None for r in rows.values()), blank + 1
    assert boss[2] == "7911111111" and boss[5] == "079091000111" and boss[4] == "Nữ"
    assert boss[3].date() == date(1991, 2, 3)
    total = [ws.cell(row=23, column=c).value for c in range(1, 12)]
    assert total[0] == "Tổng (4 lao động)" and total[7:11] == [1, 1, 1, 1]
    texts = [ws.cell(row=r, column=c).value for r in range(24, 36) for c in (1, 20)]
    assert any(t and "ĐẠI DIỆN DOANH NGHIỆP, CƠ QUAN, TỔ CHỨC" in t for t in texts)
    assert any(t and t.startswith("(2) Vị trí việc làm") for t in texts)


def test_excel_khong_co_ai_van_mo_duoc(db, world, client_as):
    client = _actor(db, world, client_as)
    resp = _get(client, SUM + "/export", date_from="2010-01-01", date_to="2010-12-31")
    assert resp.status_code == 200 and load_workbook(BytesIO(resp.content)).sheetnames == ["Không có dữ liệu"]


def test_bo_cuc_khai_du_27_cot():
    leaf = {c for _t1, c1, _t2, c2, _ in lay.HEADER_CELLS for c in range(c1, c2 + 1)}
    assert leaf == set(range(1, 28))


def test_hop_dong_khoan_viec_cong_tac_vien_khong_vao_mau(db, world, client_as):
    from app.core.report_period import parse_period
    from app.modules.labor_register import report_service as svc
    _dataset(db, world)
    for tag, ctype in (("svc", SERVICE), ("ctv", COLLABORATOR)):
        _hd(db, world, _emp(db, world, tag, education_level=2), ctype, date(2025, 2, 1))
    db.commit()
    client = _actor(db, world, client_as)
    dept = world.dept["B.hc"]
    data = _get(client, department_id=dept).json()["data"]
    assert data["totals"]["current"]["people"] == 4 and data["totals"]["current"]["ct_other"] == 2
    assert any(n.startswith("2 người ký hợp đồng khoán việc / cộng tác viên không đưa vào mẫu 01/PLI")
               for n in data["notes"])
    codes = {r["employee_code"] for r in _get(client, SUM + "/rows", department_id=dept,
                                              metric="people").json()["data"]["rows"]}
    assert codes == {"Pboss", "Puni", "Pcol", "Pplain"}
    ws = load_workbook(BytesIO(_get(client, SUM + "/export").content))["Công ty B"]
    assert ws["A4"].value.startswith("2 người ký hợp đồng khoán việc") and ws.cell(row=23, column=1).value == "Tổng (4 lao động)"
    #  Sổ quản lý lao động (B2) vẫn giữ họ: cùng tập người, không loại.
    eq, cq = db.query(Employee).filter(Employee.department_id == dept), db.query(LaborContract)
    book = svc.build_summary(db, svc.BOOK, eq, cq, parse_period(QS), None)
    assert book["totals"]["current"]["people"] == 6
