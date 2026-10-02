"""Báo cáo Hành chính (phase 05, 5.1 + 5.2) — Đặt xe & Duyệt đóng dấu.

Test matrix của `plans/260928-0841-bao-cao-kieu-haravan-da-phan-he/phase-05-admin-reports.md`:

- Đặt xe: kỳ lọc theo NGÀY ĐI (`start_time`), bỏ `BK_DRAFT`, scope phòng ban (`apply_scope`
  thật — không mock) + nhánh TÀI XẾ chỉ thấy chuyến được phân.
- Đóng dấu: kỳ lọc theo `created_at`, bỏ `SEAL_DRAFT`, scope công ty (Giám đốc) chỉ thấy
  phiếu ĐÃ DUYỆT/HOÀN THÀNH, một phiếu NHIỀU công ty → mỗi công ty +1 nhưng Tổng vẫn +1
  (P01 — Tổng cộng trực tiếp trên hàng, không cộng lại từ `groups`).
- Hai bài cuối gọi qua TestClient thật: canh thứ tự router (`/summary` không bị `/{id}`
  nuốt, xem docstring của `report_controller.py`) và đường `/summary/export` trả xlsx.

Gọi thẳng hàm controller (không qua HTTP) cho phần lớn ca — cùng lối với
`test_bao_cao_thu_mua_theo_ky.py`: các route ở đây chỉ đụng `request.query_params`.
"""
import json
import uuid
from datetime import date, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event
from starlette.datastructures import QueryParams

from app.core.auth import get_current_user, get_perm_profile
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.subject_match import SUBJECT_ROLE
from app.main import app
from app.modules.approval.instance_model import INSTANCE_APPROVED, ApprovalInstance
from app.modules.employee.model import Employee
from app.modules.seal_request import model as sm
from app.modules.seal_request import report_controller as seal_rc
from app.modules.seal_request.schema import SealRequestCreate
from app.modules.seal_request.service import create_seal_request
from app.modules.vehicle_booking import model as vm
from app.modules.vehicle_booking import report_controller as veh_rc
from app.modules.vehicle_booking.schema import VehicleBookingCreate
from app.modules.vehicle_booking.service import create_booking


def _req(**params) -> SimpleNamespace:
    """`Request` giả — route chỉ đụng `request.query_params` (cùng khuôn `make_request`
    của `test_bao_cao_thu_mua_theo_ky.py`)."""
    return SimpleNamespace(query_params=QueryParams([(k, str(v)) for k, v in params.items()]))


def _unwrap(resp) -> dict:
    return json.loads(resp.body)["data"]


TODAY = date.today()
WIDE_FROM, WIDE_TO = TODAY - timedelta(days=2), TODAY + timedelta(days=2)


def _period(**extra) -> dict:
    p = {"preset": "custom", "date_from": WIDE_FROM.isoformat(), "date_to": WIDE_TO.isoformat(),
         "compare": "none"}
    p.update(extra)
    return p


def _actor(db, *, uid: int, code: str, dept: int, company: int) -> SimpleNamespace:
    emp = Employee(code=code, full_name=code, email=f"{code}@dego.vn",
                   department_id=dept, company_id=company)
    db.add(emp)
    db.flush()
    return SimpleNamespace(id=uid, employee_id=emp.id, email=f"{code}@dego.vn")


def _count_queries(db, fn):
    """Chạy `fn` và đếm số câu lệnh xuống cơ sở dữ liệu — cùng khuôn
    `test_duyet_dau_gom_cong_ty.py::_count_queries` (review hiệu năng P05, 01/10/2026:
    số truy vấn của `/summary` phải CỐ ĐỊNH, không tăng theo số dòng)."""
    counted: list[str] = []

    def _on_exec(conn, cursor, statement, *args):
        counted.append(statement)

    event.listen(db.get_bind(), "before_cursor_execute", _on_exec)
    try:
        result = fn()
    finally:
        event.remove(db.get_bind(), "before_cursor_execute", _on_exec)
    return result, counted


def _approval(db, entity: str, entity_id: int, *, hours: float) -> None:
    """Một phiên duyệt ĐÃ KẾT THÚC (duyệt) cách nhau đúng `hours` giờ — dữ liệu cho
    `report_turnaround.turnaround_hours`."""
    started = datetime(2026, 1, 1, 8, 0)
    db.add(ApprovalInstance(entity=entity, entity_id=entity_id, flow_id=1,
                            status=INSTANCE_APPROVED, started_at=started,
                            finished_at=started + timedelta(hours=hours)))
    db.flush()


# ══════════════════════════════════════════════════════════════════════════════
#  5.1 Đặt xe
# ══════════════════════════════════════════════════════════════════════════════

DEPT_A, DEPT_B, DEPT_DRV, CTY = 10, 20, 30, 1


def _booking(db, actor, *, start_time: str, submit: bool = True):
    data = VehicleBookingCreate(request_type=1, purpose="Đi khảo sát", start_location="VP",
                                end_location="Kho", start_time=start_time, end_time=start_time,
                                passenger_count=2)
    return create_booking(db, data, actor, submit=submit)


def test_dat_xe_scope_phong_ban_chi_thay_phong_minh(db, cap_quyen):
    creator_a = _actor(db, uid=201, code="NVA", dept=DEPT_A, company=CTY)
    creator_b = _actor(db, uid=202, code="NVB", dept=DEPT_B, company=CTY)
    _booking(db, creator_a, start_time=f"{TODAY.isoformat()}T08:00")
    _booking(db, creator_b, start_time=f"{TODAY.isoformat()}T09:00")

    tp = _actor(db, uid=203, code="TPA", dept=DEPT_A, company=CTY)
    cap_quyen(203, "vehicle_booking", scope="dept", read=True)

    data = _unwrap(veh_rc.booking_summary(_req(**_period()), db, tp))
    assert data["totals"]["current"]["requests"] == 1   # chỉ phiếu của phòng A (creator_a)


def test_dat_xe_tai_xe_chi_thay_chuyen_duoc_giao(db, cap_quyen):
    creator_b = _actor(db, uid=211, code="NVC", dept=DEPT_B, company=CTY)
    _booking(db, creator_b, start_time=f"{TODAY.isoformat()}T08:00")   # không gán tài xế

    drv_user = _actor(db, uid=212, code="TX1", dept=DEPT_DRV, company=CTY)
    driver = vm.Driver(name="Tài xế Một", user_id=212)
    db.add(driver)
    db.flush()
    mine = _booking(db, creator_b, start_time=f"{TODAY.isoformat()}T10:00")
    mine.assigned_driver_id = driver.id
    db.flush()

    cap_quyen(212, "vehicle_booking", scope="assigned", read=True)
    data = _unwrap(veh_rc.booking_summary(_req(**_period()), db, drv_user))
    #  Chỉ chuyến ĐƯỢC PHÂN (mine), không phải chuyến kia của cùng người tạo.
    assert data["totals"]["current"]["requests"] == 1


def test_dat_xe_bo_draft_khong_tinh_vao_so_lieu(db, cap_quyen):
    actor = _actor(db, uid=221, code="NVD", dept=DEPT_A, company=CTY)
    draft = _booking(db, actor, start_time=f"{TODAY.isoformat()}T08:00", submit=False)
    assert draft.status == vm.BK_DRAFT
    _booking(db, actor, start_time=f"{TODAY.isoformat()}T09:00")   # submit=True -> PENDING, tính

    cap_quyen(221, "vehicle_booking", scope="own", read=True)
    data = _unwrap(veh_rc.booking_summary(_req(**_period()), db, actor))
    assert data["totals"]["current"]["requests"] == 1


def test_dat_xe_ty_le_tu_choi_huy_va_tong_km_dung(db, cap_quyen):
    actor = _actor(db, uid=222, code="NVF", dept=DEPT_A, company=CTY)
    ok = _booking(db, actor, start_time=f"{TODAY.isoformat()}T08:00")
    ok.distance_km = 30
    bad = _booking(db, actor, start_time=f"{TODAY.isoformat()}T09:00")
    bad.status = vm.BK_REJECTED
    bad.distance_km = 10
    db.flush()

    cap_quyen(222, "vehicle_booking", scope="own", read=True)
    data = _unwrap(veh_rc.booking_summary(_req(**_period()), db, actor))
    cur = data["totals"]["current"]
    assert cur["requests"] == 2 and cur["distance_km"] == 40
    assert cur["reject_cancel_rate"] == 50.0   # 1/2 từ chối+hủy


def test_dat_xe_thoi_gian_duyet_trung_binh_dung(db, cap_quyen):
    actor = _actor(db, uid=231, code="NVE", dept=DEPT_A, company=CTY)
    b1 = _booking(db, actor, start_time=f"{TODAY.isoformat()}T08:00")
    b2 = _booking(db, actor, start_time=f"{TODAY.isoformat()}T09:00")
    _approval(db, "vehicle_booking", b1.id, hours=2)
    _approval(db, "vehicle_booking", b2.id, hours=6)

    cap_quyen(231, "vehicle_booking", scope="own", read=True)
    data = _unwrap(veh_rc.booking_summary(_req(**_period()), db, actor))
    assert data["totals"]["current"]["avg_approval_hours"] == 4.0   # (2+6)/2


def test_dat_xe_so_luot_hoi_co_dinh_khong_tang_theo_so_dong(db, cap_quyen):
    """`group_by=company` cố tình bật nhánh tra nhãn công ty (nhánh TỐN truy vấn nhất) —
    5 dòng và 45 dòng thêm vào CÙNG một công ty/xe phải ra ĐÚNG cùng số lượt hỏi."""
    actor = _actor(db, uid=241, code="NVG", dept=DEPT_A, company=CTY)
    veh = vm.Vehicle(license_plate="51A-111.11", model="Innova")
    db.add(veh)
    db.flush()
    cap_quyen(241, "vehicle_booking", scope="own", read=True)
    get_perm_profile(db, actor)   # khởi ấm cache quyền (60s) — tránh lệch do cold-cache

    for _ in range(5):
        b = _booking(db, actor, start_time=f"{TODAY.isoformat()}T08:00")
        b.assigned_vehicle_id = veh.id
    db.flush()
    _, five_q = _count_queries(db, lambda: veh_rc.booking_summary(
        _req(**_period(group_by="company")), db, actor))

    for _ in range(45):
        b = _booking(db, actor, start_time=f"{TODAY.isoformat()}T09:00")
        b.assigned_vehicle_id = veh.id
    db.flush()
    _, fifty_q = _count_queries(db, lambda: veh_rc.booking_summary(
        _req(**_period(group_by="company")), db, actor))

    assert len(five_q) == len(fifty_q), (
        f"5 dòng hết {len(five_q)} truy vấn, 50 dòng hết {len(fifty_q)} — "
        "số lượt hỏi đang tăng theo số dòng")


def test_dat_xe_loc_cong_ty_chi_thu_hep_khong_lan_sang_cong_ty_khac(db, cap_quyen):
    """[H1] review 01/10/2026 — `company_id` lọc SAU `apply_scope`, chỉ THU HẸP."""
    other_cty = CTY + 1
    in_cty = _actor(db, uid=251, code="NVH", dept=DEPT_A, company=CTY)
    in_other = _actor(db, uid=252, code="NVI", dept=DEPT_B, company=other_cty)
    _booking(db, in_cty, start_time=f"{TODAY.isoformat()}T08:00")
    _booking(db, in_other, start_time=f"{TODAY.isoformat()}T08:00")

    admin = _actor(db, uid=253, code="ADM253", dept=DEPT_A, company=CTY)
    cap_quyen(253, "vehicle_booking", scope="all", read=True)

    data = _unwrap(veh_rc.booking_summary(_req(**_period(company_id=CTY)), db, admin))
    assert data["totals"]["current"]["requests"] == 1   # không lẫn phiếu của other_cty


def test_dat_xe_company_id_rac_tra_422(db, cap_quyen):
    actor = _actor(db, uid=254, code="ADM254", dept=DEPT_A, company=CTY)
    cap_quyen(254, "vehicle_booking", scope="all", read=True)
    with pytest.raises(HTTPException) as exc:
        veh_rc.booking_summary(_req(**_period(company_id="abc")), db, actor)
    assert exc.value.status_code == 422


# ══════════════════════════════════════════════════════════════════════════════
#  5.2 Duyệt đóng dấu
# ══════════════════════════════════════════════════════════════════════════════

CTY_A, CTY_B, SEAL_DEPT = 51, 52, 60


def _seal(db, actor, company_ids: list[int]):
    return create_seal_request(
        db, SealRequestCreate(purpose="Duyệt dấu hợp đồng", company_ids=company_ids,
                              first_approver_id=900),
        actor, submit=False)


def test_dong_dau_scope_cong_ty_khong_dem_draft_pending(db, cap_quyen):
    creator = _actor(db, uid=301, code="NS301", dept=SEAL_DEPT, company=CTY_A)
    draft = _seal(db, creator, [CTY_A])
    assert draft.status == sm.SEAL_DRAFT
    pending = _seal(db, creator, [CTY_A])
    pending.status = sm.SEAL_PENDING
    approved = _seal(db, creator, [CTY_A])
    approved.status, approved.approved_at = sm.SEAL_APPROVED, "2026-01-01T09:00"
    completed = _seal(db, creator, [CTY_A])
    completed.status = sm.SEAL_COMPLETED
    completed.approved_at, completed.completed_at = "2026-01-01T09:00", "2026-01-01T11:00"
    db.flush()

    #  "Giám đốc" — chỉ `read` phạm vi công ty (không `write` -> không phải Văn thư).
    director = _actor(db, uid=302, code="GD302", dept=SEAL_DEPT, company=CTY_A)
    cap_quyen(302, "seal_request", scope="company", read=True)

    data = _unwrap(seal_rc.seal_summary(_req(**_period()), db, director))
    assert data["totals"]["current"]["requests"] == 2   # chỉ approved + completed


def test_dong_dau_hai_cong_ty_moi_cong_ty_cong_1_tong_1(db, cap_quyen):
    creator = _actor(db, uid=311, code="NS311", dept=SEAL_DEPT, company=CTY_A)
    req = _seal(db, creator, [CTY_A, CTY_B])
    req.status = sm.SEAL_APPROVED
    db.flush()

    admin = _actor(db, uid=312, code="ADM312", dept=SEAL_DEPT, company=CTY_A)
    cap_quyen(312, "seal_request", scope="all", read=True)

    data = _unwrap(seal_rc.seal_summary(_req(**_period(group_by="company")), db, admin))
    assert data["totals"]["current"]["requests"] == 1   # Tổng tính TRỰC TIẾP trên 1 hàng
    groups = {g["key"]: g["current"]["requests"] for g in data["groups"]}
    assert groups[str(CTY_A)] == 1 and groups[str(CTY_B)] == 1   # mỗi công ty +1


def test_dong_dau_thoi_gian_duyet_den_dong_dau_trung_binh(db, cap_quyen):
    creator = _actor(db, uid=321, code="NS321", dept=SEAL_DEPT, company=CTY_A)
    a = _seal(db, creator, [CTY_A])
    a.status, a.approved_at, a.completed_at = sm.SEAL_COMPLETED, "2026-01-01T08:00", "2026-01-01T10:00"
    b = _seal(db, creator, [CTY_A])
    b.status, b.approved_at, b.completed_at = sm.SEAL_COMPLETED, "2026-01-01T08:00", "2026-01-01T12:00"
    db.flush()

    cap_quyen(321, "seal_request", scope="own", read=True)
    data = _unwrap(seal_rc.seal_summary(_req(**_period()), db, creator))
    assert data["totals"]["current"]["avg_approve_complete_hours"] == 3.0   # (2+4)/2


def test_dong_dau_thoi_gian_duyet_trung_binh_dung(db, cap_quyen):
    creator = _actor(db, uid=331, code="NS331", dept=SEAL_DEPT, company=CTY_A)
    a = _seal(db, creator, [CTY_A])
    a.status = sm.SEAL_APPROVED
    b = _seal(db, creator, [CTY_A])
    b.status = sm.SEAL_APPROVED
    db.flush()
    _approval(db, "seal_request", a.id, hours=1)
    _approval(db, "seal_request", b.id, hours=3)

    cap_quyen(331, "seal_request", scope="own", read=True)
    data = _unwrap(seal_rc.seal_summary(_req(**_period()), db, creator))
    assert data["totals"]["current"]["avg_approval_hours"] == 2.0   # (1+3)/2


def test_dong_dau_so_luot_hoi_co_dinh_khong_tang_theo_so_dong(db, cap_quyen):
    """`group_by=company` bật nhánh tra bảng nối NHIỀU công ty (nhánh TỐN truy vấn nhất,
    `get_company_ids_map`) — 5 dòng và 45 dòng thêm vào phải ra ĐÚNG cùng số lượt hỏi."""
    creator = _actor(db, uid=341, code="NS341", dept=SEAL_DEPT, company=CTY_A)
    cap_quyen(341, "seal_request", scope="own", read=True)
    get_perm_profile(db, creator)   # khởi ấm cache quyền (60s) — tránh lệch do cold-cache

    for _ in range(5):
        req = _seal(db, creator, [CTY_A, CTY_B])
        req.status = sm.SEAL_APPROVED
    db.flush()
    _, five_q = _count_queries(db, lambda: seal_rc.seal_summary(
        _req(**_period(group_by="company")), db, creator))

    for _ in range(45):
        req = _seal(db, creator, [CTY_A, CTY_B])
        req.status = sm.SEAL_APPROVED
    db.flush()
    _, fifty_q = _count_queries(db, lambda: seal_rc.seal_summary(
        _req(**_period(group_by="company")), db, creator))

    assert len(five_q) == len(fifty_q), (
        f"5 dòng hết {len(five_q)} truy vấn, 50 dòng hết {len(fifty_q)} — "
        "số lượt hỏi đang tăng theo số dòng")


def test_dong_dau_loc_cong_ty_chi_thu_hep_khong_lan_sang_cong_ty_khac(db, cap_quyen):
    """[H1] review 01/10/2026 — `company_id` lọc SAU `apply_scope`, chỉ THU HẸP, qua BẢNG
    NỐI (không so cột `SealRequest.company_id` đơn)."""
    creator = _actor(db, uid=351, code="NS351", dept=SEAL_DEPT, company=CTY_A)
    a = _seal(db, creator, [CTY_A])
    a.status = sm.SEAL_APPROVED
    b = _seal(db, creator, [CTY_B])
    b.status = sm.SEAL_APPROVED
    db.flush()

    cap_quyen(351, "seal_request", scope="all", read=True)
    data = _unwrap(seal_rc.seal_summary(_req(**_period(company_id=CTY_A)), db, creator))
    assert data["totals"]["current"]["requests"] == 1   # không lẫn phiếu của CTY_B


def test_dong_dau_loc_cong_ty_phieu_2_cong_ty_van_hien_khi_loc_mot_trong_hai(db, cap_quyen):
    creator = _actor(db, uid=352, code="NS352", dept=SEAL_DEPT, company=CTY_A)
    both = _seal(db, creator, [CTY_A, CTY_B])
    both.status = sm.SEAL_APPROVED
    db.flush()

    cap_quyen(352, "seal_request", scope="all", read=True)
    data_a = _unwrap(seal_rc.seal_summary(_req(**_period(company_id=CTY_A)), db, creator))
    data_b = _unwrap(seal_rc.seal_summary(_req(**_period(company_id=CTY_B)), db, creator))
    assert data_a["totals"]["current"]["requests"] == 1
    assert data_b["totals"]["current"]["requests"] == 1


def test_dong_dau_company_id_rac_tra_422(db, cap_quyen):
    actor = _actor(db, uid=353, code="ADM353", dept=SEAL_DEPT, company=CTY_A)
    cap_quyen(353, "seal_request", scope="all", read=True)
    with pytest.raises(HTTPException) as exc:
        seal_rc.seal_summary(_req(**_period(company_id="xyz")), db, actor)
    assert exc.value.status_code == 422


# ══════════════════════════════════════════════════════════════════════════════
#  Cổng HTTP thật: thứ tự router + xuất Excel
# ══════════════════════════════════════════════════════════════════════════════

def _client_as(db, user):
    """Token Bearer GIẢ nhưng DUY NHẤT mỗi lần build — `ReportSummaryCacheMiddleware`
    (gói A2) cache GET `/summary` qua Redis thật theo `(path, query, token)`;
    không có header riêng thì mọi test (và mọi tệp khác) chia cùng khóa (token
    rỗng), một bài trả 200 làm bài kế ăn lại đúng response cũ."""
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})


def test_summary_tra_200_qua_http_khong_bi_id_nuot(db, cap_quyen, gan_bao_cao):
    """`/summary` phải khớp router RIÊNG (đăng ký trước `/{id}`), không bị FastAPI thử ép
    "summary" thành id số rồi trả 422/404."""
    actor = _actor(db, uid=401, code="NS401", dept=DEPT_A, company=CTY)
    role_veh = cap_quyen(401, "vehicle_booking", scope="all", read=True)
    role_seal = cap_quyen(401, "seal_request", scope="all", read=True)
    gan_bao_cao(SUBJECT_ROLE, role_veh.id, ReportKey.VEHICLE_BOOKING)
    gan_bao_cao(SUBJECT_ROLE, role_seal.id, ReportKey.SEAL_REQUEST)
    client = _client_as(db, actor)
    try:
        r1 = client.get("/api/vehicle-bookings/summary",
                        params={"preset": "custom", "date_from": WIDE_FROM.isoformat(),
                                "date_to": WIDE_TO.isoformat(), "compare": "none"})
        r2 = client.get("/api/seal-requests/summary",
                        params={"preset": "custom", "date_from": WIDE_FROM.isoformat(),
                                "date_to": WIDE_TO.isoformat(), "compare": "none"})
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
    assert r1.status_code == 200 and r1.json()["success"] is True
    assert r2.status_code == 200 and r2.json()["success"] is True


def test_summary_export_tra_ve_xlsx(db, cap_quyen, gan_bao_cao):
    actor = _actor(db, uid=402, code="NS402", dept=DEPT_A, company=CTY)
    booking = _booking(db, actor, start_time=f"{TODAY.isoformat()}T08:00")
    booking.status = vm.BK_APPROVED
    seal = _seal(db, actor, [CTY_A])
    seal.status = sm.SEAL_APPROVED
    db.flush()
    role_veh = cap_quyen(402, "vehicle_booking", scope="all", export=True)
    role_seal = cap_quyen(402, "seal_request", scope="all", export=True)
    gan_bao_cao(SUBJECT_ROLE, role_veh.id, ReportKey.VEHICLE_BOOKING)
    gan_bao_cao(SUBJECT_ROLE, role_seal.id, ReportKey.SEAL_REQUEST)
    client = _client_as(db, actor)
    try:
        #  `group_by=status` để có cột chiều + nhãn "Tổng" (không `group_by` thì
        #  `report_export.report_xlsx` KHÔNG dựng cột chiều, xem docstring của nó).
        r1 = client.get("/api/vehicle-bookings/summary/export",
                        params={"preset": "custom", "date_from": WIDE_FROM.isoformat(),
                                "date_to": WIDE_TO.isoformat(), "compare": "none",
                                "group_by": "status"})
        r2 = client.get("/api/seal-requests/summary/export",
                        params={"preset": "custom", "date_from": WIDE_FROM.isoformat(),
                                "date_to": WIDE_TO.isoformat(), "compare": "none",
                                "group_by": "status"})
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)

    for resp in (r1, r2):
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        wb = load_workbook(BytesIO(resp.content))
        assert wb.active.cell(row=2, column=1).value == "Tổng"
