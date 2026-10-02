"""Báo cáo Công việc / Dự án (phase 06) — `backend/app/modules/work/report_service.py`.

Năm ca đại ca chốt trong kế hoạch (`plans/260928-0841-.../phase-06-work-reports.md`)
+ một bài đếm truy vấn (review hiệu năng 01/10/2026): số truy vấn của
`compute_work_summary` phải CỐ ĐỊNH, không lớn theo số việc khớp kỳ.

Gọi thẳng `report_service.compute_work_summary` cho các ca nghiệp vụ (nhanh, dễ
khẳng định số liệu); hai ca "chốt HTTP" (`employee_id=0`, 200 qua TestClient,
xuất Excel) dựng `TestClient` kiểu `client_as` của `test_tu_sua_lien_he_ca_nhan.py`.
"""
import uuid
from datetime import datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event

from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.report_period import parse_period
from app.core.subject_match import SUBJECT_ROLE
from app.main import app
from app.modules.employee.model import Employee
from app.modules.user.model import User
from app.modules.work import list_service, schema, task_service
from app.modules.work import report_service
from app.modules.work.membership_service import Actor
from app.modules.work.model import WorkTaskKind, WorkTaskStatus
from app.modules.work.task_model import WorkTask


# ── Dựng dữ liệu ──────────────────────────────────────────────────────────────

def _emp(db, code: str) -> Employee:
    e = Employee(code=code, full_name=f"Người {code}", is_active=True)
    db.add(e)
    db.flush()
    return e


def _owner(db, code: str) -> Actor:
    emp = _emp(db, code)
    return Actor(user_id=emp.id + 10_000, employee_id=emp.id, company_id=0)


def _project(db, owner: Actor) -> int:
    return list_service.create_list(db, owner, schema.ListCreate(name="Dự án test"))["id"]


def _task(db, actor: Actor, list_id: int, **kw) -> int:
    kw.setdefault("title", "Việc")
    return task_service.create_task(db, actor, schema.TaskCreate(list_id=list_id, **kw))["id"]


def _set_times(db, task_id: int, created_at=None, completed_at=None) -> None:
    """Ghi đè thẳng ngày giờ — `created_at` là `server_default`, không gõ tay
    qua service được, mà bài kiểm cần mốc CỐ ĐỊNH chứ không phải "bây giờ"."""
    t = db.get(WorkTask, task_id)
    if created_at is not None:
        t.created_at = created_at
    if completed_at is not None:
        t.completed_at = completed_at
    db.commit()


def _period(d_from: str, d_to: str, compare: str = "none"):
    return parse_period({"preset": "custom", "date_from": d_from, "date_to": d_to, "compare": compare})


MAY = _period("2026-05-01", "2026-05-31")


# ── Ca 1: không là thành viên thì không thấy việc của dự án đó ────────────────

def test_khong_la_thanh_vien_khong_thay_viec_du_an_do(db):
    owner = _owner(db, "OW1")
    lid = _project(db, owner)
    tid = _task(db, owner, lid)
    _set_times(db, tid, created_at=datetime(2026, 5, 10, 3, 0))

    outsider = _owner(db, "OUT1")   # có hồ sơ nhân sự, KHÔNG phải thành viên dự án trên

    mine = report_service.compute_work_summary(db, owner.employee_id, MAY, {})
    theirs = report_service.compute_work_summary(db, outsider.employee_id, MAY, {})

    assert mine["totals"]["current"]["created"] == 1
    assert theirs["totals"]["current"]["created"] == 0


# ── Ca 2: employee_id=0 bị chặn (như /api/work/overview) ──────────────────────

@pytest.fixture
def client_as(db):
    """Token Bearer GIẢ nhưng DUY NHẤT mỗi lần build — `ReportSummaryCacheMiddleware`
    (gói A2) cache GET `/summary` qua Redis thật theo `(path, query, token)`;
    không có header riêng thì mọi test (và mọi tệp khác) chia cùng khóa (token
    rỗng), một bài trả 200 làm bài kế ăn lại đúng response cũ dù đã mất quyền."""
    def build(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield build
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def test_employee_id_0_bi_chan(db, client_as, cap_quyen, gan_bao_cao):
    user = User(email="khongnhansu@dego.vn", employee_id=0, is_active=True)
    db.add(user)
    db.commit()
    role = cap_quyen(user.id, "work_task", scope="all", read=True)
    gan_bao_cao(SUBJECT_ROLE, role.id, ReportKey.WORK)

    res = client_as(user).get("/api/work/summary")
    assert res.status_code == 400
    assert "chưa gắn" in res.text


# ── Ca 3: việc 2 PIC — mỗi PIC +1, Tổng +1 ─────────────────────────────────────

def test_viec_hai_pic_moi_pic_cong_mot_tong_cong_mot(db):
    owner = _owner(db, "OW3")
    lid = _project(db, owner)
    pic1, pic2 = _emp(db, "P1"), _emp(db, "P2")
    tid = _task(db, owner, lid)
    task_service.set_assignees(db, owner, tid, [pic1.id, pic2.id], [])
    _set_times(db, tid, created_at=datetime(2026, 5, 10, 3, 0))

    data = report_service.compute_work_summary(db, owner.employee_id, MAY, {"group_by": "pic"})

    assert data["totals"]["current"]["created"] == 1          # Tổng = 1 việc, không phải 2
    created_per_group = sorted(g["current"]["created"] for g in data["groups"])
    assert created_per_group == [1, 1]                         # mỗi PIC +1


# ── Ca 4: việc con / cột mốc / xóa mềm không được đếm ──────────────────────────

def test_viec_con_cot_moc_xoa_mem_khong_duoc_dem(db):
    owner = _owner(db, "OW4")
    lid = _project(db, owner)
    when = datetime(2026, 5, 10, 3, 0)

    normal = _task(db, owner, lid, title="Việc thường")
    _set_times(db, normal, created_at=when)

    milestone = _task(db, owner, lid, title="Mốc", kind=int(WorkTaskKind.MILESTONE), due_date="2026-05-10")
    _set_times(db, milestone, created_at=when)

    deleted = _task(db, owner, lid, title="Việc xóa")
    _set_times(db, deleted, created_at=when)
    task_service.delete_task(db, owner, deleted)

    sub = task_service.create_task(db, owner, schema.TaskCreate(title="Việc con", parent_id=normal))["id"]
    _set_times(db, sub, created_at=when)

    data = report_service.compute_work_summary(db, owner.employee_id, MAY, {})

    assert data["totals"]["current"]["created"] == 1   # chỉ "Việc thường"


# ── Ca 5: hoàn thành 23:30 UTC ngày hạn-1 → ngày VN = ngày hạn → đúng hạn ─────

def test_hoan_thanh_23h30_utc_ngay_han_tru_1_tinh_dung_han(db):
    owner = _owner(db, "OW5")
    lid = _project(db, owner)
    tid = _task(db, owner, lid, due_date="2026-05-10")
    task_service.update_task(db, owner, tid, schema.TaskUpdate(status=int(WorkTaskStatus.DONE)))
    #  23:30 UTC ngày 09/05 + 7h = 06:30 giờ VN ngày 10/05 — đúng NGÀY HẠN.
    _set_times(db, tid, created_at=datetime(2025, 1, 1), completed_at=datetime(2026, 5, 9, 23, 30))

    totals = report_service.compute_work_summary(db, owner.employee_id, MAY, {})["totals"]["current"]

    assert totals["completed"] == 1
    assert totals["due_completed_count"] == 1
    assert totals["on_time_count"] == 1
    assert totals["on_time_rate"] == 100.0


# ── M1 (review 01/10/2026): snapshot "Đang mở"/"Quá hạn" đúng TẠI MỐC quá khứ ──

def test_dang_mo_qua_han_tai_moc_qua_khu(db):
    """Mốc = cuối kỳ 31/05. Việc xong TRƯỚC mốc thì không còn "đang mở" ở mốc đó
    (dù hạn đã qua); việc xong SAU mốc vẫn phải tính "đang mở tại mốc" dù hiện
    tại đã xong; việc HỦY không có dấu "lúc hủy" nên loại khỏi cả hai chỉ số."""
    owner = _owner(db, "OW6")
    lid = _project(db, owner)
    when = datetime(2026, 4, 1)

    still_open = _task(db, owner, lid, due_date="2026-05-10")   # chưa xong, quá hạn tại mốc
    _set_times(db, still_open, created_at=when)

    done_before = _task(db, owner, lid, due_date="2026-05-10")   # xong TRƯỚC mốc
    task_service.update_task(db, owner, done_before, schema.TaskUpdate(status=int(WorkTaskStatus.DONE)))
    _set_times(db, done_before, created_at=when, completed_at=datetime(2026, 5, 20, 3, 0))

    done_after = _task(db, owner, lid, due_date="2026-05-15")    # xong SAU mốc — vẫn mở TẠI mốc
    task_service.update_task(db, owner, done_after, schema.TaskUpdate(status=int(WorkTaskStatus.DONE)))
    _set_times(db, done_after, created_at=when, completed_at=datetime(2026, 6, 5, 3, 0))

    cancelled = _task(db, owner, lid, due_date="2026-01-01")     # hủy — không có mốc "lúc hủy"
    task_service.update_task(db, owner, cancelled, schema.TaskUpdate(status=int(WorkTaskStatus.CANCELLED)))
    _set_times(db, cancelled, created_at=when)

    totals = report_service.compute_work_summary(db, owner.employee_id, MAY, {})["totals"]["current"]

    assert totals["open_tasks"] == 2       # still_open + done_after
    assert totals["overdue_tasks"] == 2    # cả hai đều quá hạn 31/05


# ── Chốt HTTP: 200 qua TestClient + xuất Excel ────────────────────────────────

def test_api_summary_tra_200(db, client_as, cap_quyen, gan_bao_cao):
    emp = _emp(db, "API1")
    user = User(email="api1@dego.vn", employee_id=emp.id, is_active=True)
    db.add(user)
    db.commit()
    role = cap_quyen(user.id, "work_task", scope="all", read=True, export=True)
    gan_bao_cao(SUBJECT_ROLE, role.id, ReportKey.WORK)

    owner = Actor(user_id=user.id, employee_id=emp.id, company_id=0)
    lid = _project(db, owner)
    tid = _task(db, owner, lid)
    _set_times(db, tid, created_at=datetime(2026, 5, 10, 3, 0))

    res = client_as(user).get("/api/work/summary", params={
        "preset": "custom", "date_from": "2026-05-01", "date_to": "2026-05-31", "compare": "none"})

    assert res.status_code == 200, res.text
    body = res.json()
    assert body["success"] is True
    assert body["data"]["totals"]["current"]["created"] == 1


def test_xuat_excel_200(db, client_as, cap_quyen, gan_bao_cao):
    emp = _emp(db, "EXP1")
    user = User(email="exp1@dego.vn", employee_id=emp.id, is_active=True)
    db.add(user)
    db.commit()
    role = cap_quyen(user.id, "work_task", scope="all", read=True, export=True)
    gan_bao_cao(SUBJECT_ROLE, role.id, ReportKey.WORK)

    owner = Actor(user_id=user.id, employee_id=emp.id, company_id=0)
    lid = _project(db, owner)
    tid = _task(db, owner, lid)
    _set_times(db, tid, created_at=datetime(2026, 5, 10, 3, 0))

    res = client_as(user).get("/api/work/summary/export", params={
        "preset": "custom", "date_from": "2026-05-01", "date_to": "2026-05-31",
        "compare": "none", "group_by": "list"})

    assert res.status_code == 200, res.text
    wb = load_workbook(BytesIO(res.content))
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "Tổng"


# ── Review hiệu năng 01/10/2026: số truy vấn CỐ ĐỊNH, không lớn theo số việc ──

def _count_queries(db, fn):
    counted: list[str] = []

    def _on_exec(conn, cursor, statement, *args):
        counted.append(statement)

    event.listen(db.get_bind(), "before_cursor_execute", _on_exec)
    try:
        result = fn()
    finally:
        event.remove(db.get_bind(), "before_cursor_execute", _on_exec)
    return result, counted


def test_so_truy_van_co_dinh_khong_tang_theo_so_viec(db):
    """5 việc và 50 việc phải ra ĐÚNG CÙNG một số truy vấn — PIC/ưu tiên tra theo
    LÔ (`report_rows.fetch_rows`), không hỏi lại CSDL trong vòng lặp từng việc."""
    owner = _owner(db, "OW9")
    lid = _project(db, owner)

    for _ in range(5):
        tid = _task(db, owner, lid)
        _set_times(db, tid, created_at=datetime(2026, 5, 10, 3, 0))
    _, counted_5 = _count_queries(
        db, lambda: report_service.compute_work_summary(db, owner.employee_id, MAY, {}))

    for _ in range(45):   # cộng dồn đủ 50 việc
        tid = _task(db, owner, lid)
        _set_times(db, tid, created_at=datetime(2026, 5, 10, 3, 0))
    _, counted_50 = _count_queries(
        db, lambda: report_service.compute_work_summary(db, owner.employee_id, MAY, {}))

    assert len(counted_50) == len(counted_5), (
        f"5 việc hết {len(counted_5)} truy vấn, 50 việc hết {len(counted_50)} — "
        "số truy vấn đang lớn theo số việc (có thể đã quay lại N+1)")


def test_group_by_pic_van_co_dinh_giua_5_va_50_viec(db):
    """`group_by=pic` tra PIC THEO LÔ — vẫn cố định dù mỗi việc một PIC riêng."""
    owner = _owner(db, "OW10")
    lid = _project(db, owner)

    for i in range(5):
        tid = _task(db, owner, lid)
        pic = _emp(db, f"PICA{i}")
        task_service.set_assignees(db, owner, tid, [pic.id], [])
        _set_times(db, tid, created_at=datetime(2026, 5, 10, 3, 0))
    _, counted_5 = _count_queries(
        db, lambda: report_service.compute_work_summary(db, owner.employee_id, MAY, {"group_by": "pic"}))

    for i in range(45):
        tid = _task(db, owner, lid)
        pic = _emp(db, f"PICB{i}")
        task_service.set_assignees(db, owner, tid, [pic.id], [])
        _set_times(db, tid, created_at=datetime(2026, 5, 10, 3, 0))
    _, counted_50 = _count_queries(
        db, lambda: report_service.compute_work_summary(db, owner.employee_id, MAY, {"group_by": "pic"}))

    assert len(counted_50) == len(counted_5), (
        f"5 việc hết {len(counted_5)} truy vấn, 50 việc hết {len(counted_50)} (group_by=pic)")
