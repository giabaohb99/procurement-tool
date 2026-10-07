"""Soạn mẫu hợp đồng NGAY TRÊN WEB (duoc-CR-606, 07/10/2026) — `GET/PUT /{id}/content`.

Góc nhìn kẻ phá: lưu nội dung có biến lạ / cú pháp Jinja lén qua đường mới (đường tải tệp đã chặn,
đường HTML phải chặn y hệt), pháp nhân khác gõ id, người chỉ có quyền đọc bấm Lưu, gửi HTML khổng lồ,
lưu hỏng không được làm mất tệp cũ, và lề trang của bản gốc không được âm thầm đổi.
"""
import zipfile
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Mm

from app.modules.labor_contract.template_web_editor import (DEFAULT_LEFT_MM, DEFAULT_RIGHT_MM,
                                                            page_margins_mm)
from hdld_factory import client_as, docx_text, grant_all, make_docx, storage, upload_template  # noqa: F401

URL = "/api/labor-contract-templates"


@pytest.fixture
def setup(db, world, storage, client_as):  # noqa: F811
    grant_all(world, "a1")
    grant_all(world, "b1")
    a1 = client_as(world.actor("a1"))
    tid = upload_template(a1, company_id=world.co["A"]).json()["data"]["id"]
    return world, a1, client_as, storage, tid


def _docx_with_margins(left_mm: int, right_mm: int, *paragraphs: str) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.left_margin, section.right_margin = Mm(left_mm), Mm(right_mm)
    for text in paragraphs:
        doc.add_paragraph(text)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _download(client, tid) -> bytes:
    res = client.get(f"{URL}/{tid}/file")
    assert res.status_code == 200, res.text
    return res.content


# ── Mở mẫu ra để soạn ──────────────────────────────────────────────────────────────────
def test_mo_mau_ra_html_giu_nguyen_bien_va_le_trang(db, world, storage, client_as):  # noqa: F811
    grant_all(world, "a1")
    a1 = client_as(world.actor("a1"))
    data = _docx_with_margins(25, 15, "Họ tên: {{ ho_ten }}")
    tid = upload_template(a1, company_id=world.co["A"], data=data).json()["data"]["id"]
    res = a1.get(f"{URL}/{tid}/content")
    assert res.status_code == 200, res.text
    body = res.json()["data"]
    assert "{{ ho_ten }}" in body["html"]
    assert (body["margin_left_mm"], body["margin_right_mm"]) == (25, 15)


# ── Lưu nội dung soạn trên web ───────────────────────────────────────────────────────────
def test_luu_noi_dung_thanh_docx_moi_cap_nhat_bien_va_xoa_tep_cu(setup):
    _, a1, _, storage, tid = setup
    old_keys = set(storage)
    html = "<p>Hợp đồng số {{ so_hop_dong }}</p><p>Bên B: <strong>{{ ho_ten }}</strong></p>"
    res = a1.put(f"{URL}/{tid}/content", json={"html": html, "margin_left_mm": 28, "margin_right_mm": 18})
    assert res.status_code == 200, res.text
    assert res.json()["data"]["placeholders"] == ["ho_ten", "so_hop_dong"]
    assert not (old_keys & set(storage)) and len(storage) == 1, "tệp cũ phải bị xóa sau khi lưu"
    saved = _download(a1, tid)
    text = docx_text(saved)
    assert "{{ so_hop_dong }}" in text and "{{ ho_ten }}" in text
    assert page_margins_mm(saved) == (28, 18)
    assert a1.get(f"{URL}/{tid}").json()["data"]["original_filename"].endswith(".docx")


@pytest.mark.parametrize("html, unknown", [
    ("<p>Tên: {{ ten_la_lam }}</p>", ["ten_la_lam"]),
    ("<p>{{ ho_ten }} và {{ HO_TEN }}</p>", ["HO_TEN"]),
])
def test_bien_la_bi_chan_y_nhu_tai_tep_va_tep_cu_con_nguyen(setup, html, unknown):
    _, a1, _, storage, tid = setup
    before = _download(a1, tid)
    res = a1.put(f"{URL}/{tid}/content", json={"html": html})
    assert res.status_code == 422, res.text
    assert res.json()["error"]["details"]["unknown"] == unknown
    assert _download(a1, tid) == before and len(storage) == 1


@pytest.mark.parametrize("html", [
    "<p>{% for x in range(9999999) %}x{% endfor %}</p>",
    "<p>{{ ho_ten.__class__ }}</p>",
    "<p>{{ ho_ten | upper }}</p>",
])
def test_cu_phap_jinja_ngoai_bien_don_bi_chan_qua_duong_html(setup, html):
    _, a1, _, _, tid = setup
    before = _download(a1, tid)
    assert a1.put(f"{URL}/{tid}/content", json={"html": html}).status_code == 422
    assert _download(a1, tid) == before


@pytest.mark.parametrize("body", [
    {"html": ""},
    {"html": "x" * 2_000_001},
    {"html": "<p>a</p>", "margin_left_mm": 0},
    {"html": "<p>a</p>", "margin_right_mm": 999},
    {"html": "<p>a</p>", "company_id": 99},
])
def test_dau_vao_rac_bi_422(setup, body):
    _, a1, _, _, tid = setup
    assert a1.put(f"{URL}/{tid}/content", json=body).status_code == 422


# ── Phạm vi & quyền ─────────────────────────────────────────────────────────────────────
def test_phap_nhan_khac_khong_mo_khong_luu_duoc_noi_dung(setup):
    world, _, make_client, _, tid = setup
    b1 = make_client(world.actor("b1"))
    assert b1.get(f"{URL}/{tid}/content").status_code == 404
    assert b1.put(f"{URL}/{tid}/content", json={"html": "<p>{{ ho_ten }}</p>"}).status_code == 404


def test_chi_quyen_doc_thi_mo_duoc_nhung_khong_luu_duoc(setup):
    world, _, make_client, _, tid = setup
    world.actor("a2").grant("labor_contract_template", "company", actions=("read",))
    a2 = make_client(world.actor("a2"))
    assert a2.get(f"{URL}/{tid}/content").status_code == 200
    assert a2.put(f"{URL}/{tid}/content", json={"html": "<p>x</p>"}).status_code == 403


# ── Đọc lề trang ─────────────────────────────────────────────────────────────────────────
def test_doc_le_trang_tep_hong_hoac_thieu_le_thi_ve_mac_dinh():
    assert page_margins_mm(b"khong phai zip") == (DEFAULT_LEFT_MM, DEFAULT_RIGHT_MM)
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<w:document><w:body><w:p/></w:body></w:document>")
    assert page_margins_mm(buf.getvalue()) == (DEFAULT_LEFT_MM, DEFAULT_RIGHT_MM)


def test_doc_le_trang_keo_ve_khoang_thuoc_cho_phep():
    assert page_margins_mm(_docx_with_margins(2, 90, "x")) == (5, 60)
    assert page_margins_mm(make_docx("x"))[0] > 0


# ── Đích cuối: mẫu lưu từ web phải SINH được hợp đồng thật ────────────────────────────────
def test_mau_soan_tren_web_sinh_hop_dong_dien_dung_du_lieu(setup, db):
    """Biến chèn bằng nút «Chèn biến» rồi tô đậm / nghiêng vẫn phải thay được — biến bị Word cắt
    thành nhiều mảnh chữ là kiểu lỗi âm thầm: tệp lưu được, sinh ra hợp đồng còn nguyên `{{ … }}`."""
    from app.modules.employee.model import Employee
    from hdld_factory import make_contract

    world, a1, _, _, tid = setup
    html = ("<p>Hợp đồng số <em>{{ so_hop_dong }}</em></p>"
            "<p>Bên B: <strong>{{ ho_ten }}</strong></p>"
            "<table><tr><td>Lương</td><td>{{ luong_co_ban }}</td></tr></table>")
    assert a1.put(f"{URL}/{tid}/content", json={"html": html}).status_code == 200
    cid = make_contract(a1, world.emp["a2"], contract_no="HD/2026-99", base_salary=12_000_000)["id"]
    res = a1.post(f"/api/labor-contracts/{cid}/generate", json={"template_id": tid})
    assert res.status_code == 200, res.text
    out = a1.get(f"/api/labor-contracts/{cid}/document").content
    with zipfile.ZipFile(BytesIO(out)) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    name = db.get(Employee, world.emp["a2"]).full_name
    assert "{{" not in xml and "}}" not in xml, "còn biến chưa thay trong hợp đồng sinh ra"
    assert "HD/2026-99" in xml and name in xml and "12.000.000" in xml
