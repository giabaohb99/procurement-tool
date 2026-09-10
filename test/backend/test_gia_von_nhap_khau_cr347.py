"""bao-CR-347 — báo cáo giá vốn lô hàng nhập khẩu (hai cách đọc) + ba mã chi phí mới.

Luật phải giữ:
1. Ba mã chi phí mới 12/13/14 dùng được và hiện đúng tên trên báo cáo.
2. Bảng chỉ bày loại chi phí CÓ PHÁT SINH, đánh số 2.1, 2.2… theo thứ tự danh mục.
3. Tab THEO LÔ: đủ chỉ tiêu theo mẫu, Tổng giá vốn = tiền hàng + chi phí, Giá/Kg chia theo
   tổng kg, cột TỔNG tính lại giá/kg chứ không cộng dồn.
4. Tab THEO DÒNG HÀNG: chi phí của lô chia hết về các dòng (không rơi rớt đồng nào), chia
   bằng đúng `allocate_import_costs` của ĐMH, và có giá vốn mỗi đơn vị.
5. `cost_status` còn trong DB nhưng giao diện đã bỏ hẳn ô Dự kiến (đại ca chốt 10/09/2026) —
   giữ vài test cho tấm lưới an toàn: dòng dự kiến sót lại không được lọt vào công nợ, dòng
   cũ không khai trạng thái vẫn đọc thành Thực tế.
"""
from app.modules.payable.model import Payable
from app.modules.report import import_landed_cost as ilc
from app.modules.purchase_order import service
from app.modules.purchase_order.controller import _out
from app.modules.purchase_order.model import (IMPORT_COST_TYPE_LABELS, ImportCostStatus,
                                              ImportCostType, OrderType, POImportCost, POItem,
                                              PurchaseOrder)
from app.modules.purchase_order.schema import POImportCostIn, POItemIn

RATE = 25_000.0
SRC = service.IMPORT_COST_SOURCE


def _item(code, qty, price, weight):
    return POItemIn(product_code=code, product_name=f"Hàng {code}", item_group="Nhãn", unit="cái",
                    qty_request=qty, qty_order=qty, price=price, vat=0, weight_kg=weight,
                    required_date="2026-09-30", warehouse_code="KHO01")


def _make_po(db, seed, code="PO-NK-323", *, qty=4.0, price=100.0, weight=50.0,
             items=None, **kw):
    """Đơn nhập khẩu một dòng hàng; truyền `items` để dựng đơn nhiều mã.

    Luôn đi qua `service._save_items` chứ không tự dựng POItem: hàm đó chép tỷ giá của đơn
    xuống từng dòng, mà thiếu tỷ giá thì tiền hàng của dòng lệch 25.000 lần và mọi phép chia
    chi phí theo giá trị đều sai.
    """
    kw.setdefault("status", "draft")
    kw.setdefault("currency", "USD")
    po = PurchaseOrder(code=code, company_id=seed.company_id, supplier_code="NX",
                       supplier_name=seed.sup_name, order_date="2026-09-10",
                       order_type=int(OrderType.IMPORT), exchange_rate=RATE, **kw)
    db.add(po)
    db.flush()
    service._save_items(db, po, items or [_item("SP01", qty, price, weight)], user_id=1)
    db.flush()
    return po


# Đơn hai dòng hàng 400$ / 600$ — dùng cho mọi test của tab theo dòng hàng
TWO_LINES = [("SP01", 4.0, 100.0, 50.0), ("SP02", 6.0, 100.0, 150.0)]


def _cost_in(**kw):
    base = dict(cost_type=int(ImportCostType.OCEAN_FREIGHT), description="Cước biển",
                supplier_code="HANGTAU", supplier_name="Hãng tàu ABC",
                currency="VND", amount=1_000_000, vat=0)
    base.update(kw)
    return POImportCostIn(**base)


def _save_costs(db, po, costs):
    service._save_import_costs(db, po, costs, user_id=1)
    db.flush()


def _cost_payables(db, po):
    return (db.query(Payable).filter(Payable.po_id == po.id, Payable.ref_type == SRC)
            .order_by(Payable.ref_id).all())


# ── 1. Ba mã chi phí mới ─────────────────────────────────────────────────────────
def test_ma_chi_phi_moi_dung_duoc_va_dung_ten(db, seed):
    po = _make_po(db, seed, status="approved")
    _save_costs(db, po, [
        _cost_in(cost_type=int(ImportCostType.IMPORT_SERVICE), amount=100_000),
        _cost_in(cost_type=int(ImportCostType.CONTAINER_DAMAGE), amount=200_000),
        _cost_in(cost_type=int(ImportCostType.LATE_INTEREST), amount=300_000),
    ])
    rows = _out(db, po)["import_costs"]
    assert [r["cost_type"] for r in rows] == [12, 13, 14]
    assert rows[1]["cost_type_label"] == "Chi tiền hư container"
    by_type = ilc.compute(db, po_ids=[po.id])["orders"][0]["by_type"]
    assert by_type["12"] == 100_000.0
    assert by_type["13"] == 200_000.0
    assert by_type["14"] == 300_000.0


def test_moi_ma_chi_phi_deu_co_ten_hien_thi():
    """Mã chi phí thiếu tên là một dòng trắng trên báo cáo — không ai đọc được nó là gì."""
    assert all(t in IMPORT_COST_TYPE_LABELS and IMPORT_COST_TYPE_LABELS[t] for t in ImportCostType)


# ── 2. Chỉ bày loại chi phí có phát sinh ─────────────────────────────────────────
def test_chi_bay_loai_chi_phi_co_phat_sinh(db, seed):
    po = _make_po(db, seed, status="approved")
    _save_costs(db, po, [
        _cost_in(cost_type=int(ImportCostType.OCEAN_FREIGHT), amount=1_500_000),
        _cost_in(cost_type=int(ImportCostType.IMPORT_VAT), amount=500_000),
    ])
    cost_types = ilc.compute(db, po_ids=[po.id])["cost_types"]
    # Đúng hai loại, xếp theo thứ tự DANH MỤC chứ không theo thứ tự gõ vào
    assert [g["code"] for g in cost_types] == [int(ImportCostType.OCEAN_FREIGHT),
                                               int(ImportCostType.IMPORT_VAT)]
    assert [g["no"] for g in cost_types] == ["2.1", "2.2"]
    assert cost_types[0]["label"] == IMPORT_COST_TYPE_LABELS[ImportCostType.OCEAN_FREIGHT]


def test_khong_khai_chi_phi_thi_khong_co_dong_loai_nao(db, seed):
    po = _make_po(db, seed, status="approved")
    data = ilc.compute(db, po_ids=[po.id])
    assert data["cost_types"] == []
    assert data["orders"][0]["cost_total"] == 0.0


# ── 3. Tab theo lô ───────────────────────────────────────────────────────────────
def test_bao_cao_theo_lo_du_chi_tieu_va_dung_tong(db, seed):
    po = _make_po(db, seed, qty=4, price=100, weight=50, status="approved",
                  etd_date="2026-09-12", note="Lô thử")
    _save_costs(db, po, [
        _cost_in(cost_type=int(ImportCostType.IMPORT_VAT), amount=500_000),
        _cost_in(cost_type=int(ImportCostType.OCEAN_FREIGHT), amount=1_500_000),
    ])
    o = ilc.compute(db, po_ids=[po.id])["orders"][0]
    assert o["etd_date"] == "2026-09-12"              # "Ngày hàng rời cảng (ETD)"
    assert o["status_label"] == "Đã duyệt"
    assert o["note"] == "Lô thử"
    assert o["qty_total"] == 4.0 and o["weight_total"] == 50.0
    assert o["currency"] == "USD"                     # chỉ tiêu "Đồng tiền"
    assert o["goods_amount"] == 400.0                 # "Tiền hàng (nguyên tệ)"
    assert o["exchange_rate"] == RATE
    goods_base = 400.0 * RATE                          # "1 Tiền hàng (vnd)"
    assert o["goods_base"] == goods_base
    assert o["by_type"][str(int(ImportCostType.IMPORT_VAT))] == 500_000.0
    assert o["by_type"][str(int(ImportCostType.OCEAN_FREIGHT))] == 1_500_000.0
    assert o["cost_total"] == 2_000_000.0
    assert o["landed_total"] == goods_base + 2_000_000.0
    assert o["price_per_kg"] == round((goods_base + 2_000_000.0) / 50.0, 2)


def test_gia_tren_kg_bang_0_khi_chua_khai_can_nang(db, seed):
    po = _make_po(db, seed, weight=0, status="approved")
    _save_costs(db, po, [_cost_in()])
    assert ilc.compute(db, po_ids=[po.id])["orders"][0]["price_per_kg"] == 0.0


def test_cot_tong_tinh_lai_gia_tren_kg(db, seed):
    a = _make_po(db, seed, code="PO-NK-A", qty=4, price=100, weight=50, status="approved")
    b = _make_po(db, seed, code="PO-NK-B", qty=2, price=100, weight=150, status="approved")
    _save_costs(db, a, [_cost_in(amount=1_000_000)])
    _save_costs(db, b, [_cost_in(amount=2_000_000)])
    t = ilc.compute(db, po_ids=[a.id, b.id])["totals"]
    assert t["order_count"] == 2 and t["weight_total"] == 200.0
    goods = (400.0 + 200.0) * RATE
    assert t["goods_base"] == goods
    assert t["cost_total"] == 3_000_000.0
    assert t["landed_total"] == goods + 3_000_000.0
    # Cộng dồn giá/kg của hai đơn sẽ ra số khác — phải chia lại từ tổng
    assert t["price_per_kg"] == round((goods + 3_000_000.0) / 200.0, 2)


def test_cot_tong_bo_trong_tien_nguyen_te_khi_nhieu_dong_tien(db, seed):
    """Lô mua bằng yên đứng chung lô mua bằng đô: cộng hai con số nguyên tệ đó lại là ra
    một số không thuộc đồng tiền nào, nên cột TỔNG phải bỏ trống ô nguyên tệ."""
    a = _make_po(db, seed, code="PO-NK-USD", status="approved", currency="USD")
    b = _make_po(db, seed, code="PO-NK-JPY", status="approved", currency="JPY")
    _save_costs(db, a, [_cost_in(amount=1_000_000)])
    data = ilc.compute(db, po_ids=[a.id, b.id])
    assert [o["currency"] for o in data["orders"]] == ["USD", "JPY"]
    t = data["totals"]
    assert t["currency_mixed"] is True and t["currency"] == ""
    assert t["goods_amount"] == 0.0
    # Cột VNĐ vẫn cộng bình thường vì mỗi đơn đã quy đổi qua tỷ giá của chính nó
    assert t["goods_base"] == 800.0 * RATE


def test_don_lan_dong_tien_ngay_trong_cac_dong_hang_thi_bo_trong_nguyen_te(db, seed):
    """Từng dòng hàng khai được đồng tiền riêng: một đơn USD kèm một dòng VND thì ô nguyên
    tệ của CHÍNH đơn đó cũng là số lẫn, phải bỏ trống chứ không bày ra."""
    po = _make_po(db, seed, status="approved", currency="USD",
                  items=[_item(*x) for x in TWO_LINES])
    for it in db.query(POItem).filter(POItem.po_id == po.id).all():
        if it.product_code == "SP02":
            it.currency, it.exchange_rate = "VND", 1
    db.flush()
    o = ilc.compute(db, po_ids=[po.id])["orders"][0]
    assert o["currency"] == "USD" and o["currency_mixed"] is True
    assert ilc.compute(db, po_ids=[po.id])["totals"]["currency_mixed"] is True


def test_cot_tong_van_cong_tien_nguyen_te_khi_chung_mot_dong_tien(db, seed):
    a = _make_po(db, seed, code="PO-NK-U1", status="approved", currency="USD")
    b = _make_po(db, seed, code="PO-NK-U2", status="approved", currency="USD")
    t = ilc.compute(db, po_ids=[a.id, b.id])["totals"]
    assert t["currency_mixed"] is False and t["currency"] == "USD"
    assert t["goods_amount"] == 800.0


def test_don_khong_khai_dong_tien_thi_hieu_la_vnd(db, seed):
    po = _make_po(db, seed, status="approved", currency="")
    assert ilc.compute(db, po_ids=[po.id])["orders"][0]["currency"] == "VND"


def test_bao_cao_bo_qua_don_trong_nuoc(db, seed):
    dom = PurchaseOrder(code="PO-TN-1", company_id=seed.company_id, supplier_code="NX",
                        supplier_name=seed.sup_name, order_date="2026-09-10",
                        order_type=int(OrderType.DOMESTIC), status="approved")
    db.add(dom)
    db.flush()
    _make_po(db, seed, code="PO-NK-C", status="approved")
    codes = [o["code"] for o in ilc.compute(db, date_from="2026-01-01", date_to="2026-12-31")["orders"]]
    assert codes == ["PO-NK-C"]


# ── 4. Tab theo dòng hàng ────────────────────────────────────────────────────────
def test_theo_dong_hang_chia_het_chi_phi_ve_cac_dong(db, seed):
    """Mỗi đồng chi phí của lô phải nằm ở đúng một dòng hàng — không rơi rớt, không nhân đôi."""
    po = _make_po(db, seed, status="approved",
                  items=[_item(*x) for x in TWO_LINES])
    _save_costs(db, po, [
        _cost_in(cost_type=int(ImportCostType.OCEAN_FREIGHT), amount=1_000_000),
        _cost_in(cost_type=int(ImportCostType.IMPORT_VAT), amount=500_000),
    ])
    data = ilc.compute(db, po_ids=[po.id])
    lines = data["items"]
    assert [r["product_code"] for r in lines] == ["SP01", "SP02"]
    assert round(sum(r["cost_base"] for r in lines), 2) == 1_500_000.0
    for g in data["cost_types"]:
        key = str(g["code"])
        assert round(sum(r["by_type"][key] for r in lines), 2) == data["orders"][0]["by_type"][key]
    # Giá vốn dòng = tiền hàng dòng + chi phí chia về dòng
    for r in lines:
        assert r["landed_base"] == round(r["goods_base"] + r["cost_base"], 2)
        assert r["price_per_unit"] == round(r["landed_base"] / r["qty_order"], 2)


def test_theo_dong_hang_chia_theo_gia_tri(db, seed):
    """Cách mặc định là chia theo giá trị: dòng 400$ / dòng 600$ ăn 40% / 60% chi phí."""
    po = _make_po(db, seed, status="approved",
                  items=[_item(*x) for x in TWO_LINES])
    _save_costs(db, po, [_cost_in(amount=1_000_000)])
    lines = ilc.compute(db, po_ids=[po.id])["items"]
    assert lines[0]["cost_base"] == 400_000.0
    assert lines[1]["cost_base"] == 600_000.0


def test_dong_tong_theo_dong_hang_khong_co_gia_von_don_vi(db, seed):
    """Mỗi dòng một đơn vị tính khác nhau — cộng lại rồi chia ra là con số vô nghĩa."""
    po = _make_po(db, seed, status="approved",
                  items=[_item(*x) for x in TWO_LINES])
    _save_costs(db, po, [_cost_in(amount=1_000_000)])
    data = ilc.compute(db, po_ids=[po.id])
    t = data["item_totals"]
    assert "price_per_unit" not in t
    assert t["line_count"] == 2 and t["qty_order"] == 10.0 and t["weight_kg"] == 200.0
    assert t["cost_base"] == 1_000_000.0
    assert t["landed_base"] == round(sum(r["landed_base"] for r in data["items"]), 2)
    assert t["price_per_kg"] == round(t["landed_base"] / 200.0, 2)


def test_canh_bao_chia_chi_phi_co_ghi_ma_don(db, seed):
    """Chia theo khối lượng mà chưa khai kg thì phải nói rõ đơn nào, không im lặng đổi cách."""
    po = _make_po(db, seed, code="PO-NK-W", qty=4, price=100, weight=0, status="approved")
    _save_costs(db, po, [_cost_in(amount=1_000_000, allocation_method=2)])
    warnings = ilc.compute(db, po_ids=[po.id])["warnings"]
    assert warnings and all(w.startswith("PO-NK-W:") for w in warnings)


# ── 5. Tấm lưới an toàn của `cost_status` ────────────────────────────────────────
def test_dong_du_kien_sot_lai_khong_lot_vao_cong_no(db, seed):
    po = _make_po(db, seed, status="approved")
    _save_costs(db, po, [_cost_in(cost_status=int(ImportCostStatus.ESTIMATED))])
    service.sync_import_cost_payables(db, po, user_id=1)
    assert _cost_payables(db, po) == []
    assert _out(db, po)["import_cost_summary"]["cost_total"] == 0.0


def test_dong_cu_khong_khai_trang_thai_van_la_thuc_te(db, seed):
    """Dữ liệu trước CR-347: `cost_status` để trống / 0 phải đọc thành Thực tế."""
    po = _make_po(db, seed, status="approved")
    _save_costs(db, po, [_cost_in()])
    row = db.query(POImportCost).filter(POImportCost.po_id == po.id).one()
    row.cost_status = 0
    db.flush()
    assert service.is_actual_cost(row) is True
    service.sync_import_cost_payables(db, po, user_id=1)
    assert len(_cost_payables(db, po)) == 1


# ── 6. File Excel ────────────────────────────────────────────────────────────────
def _workbook(data):
    from openpyxl import load_workbook

    from app.modules.report.excel import build_import_landed_cost_workbook
    return load_workbook(build_import_landed_cost_workbook(data))


def _values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_excel_co_hai_sheet_hai_cach_doc(db, seed):
    po = _make_po(db, seed, status="approved",
                  items=[_item(*x) for x in TWO_LINES])
    _save_costs(db, po, [_cost_in(amount=1_000_000)])
    wb = _workbook(ilc.compute(db, po_ids=[po.id]))
    assert wb.sheetnames == ["GIA VON THEO LO", "THEO DONG HANG"]

    lo = _values(wb["GIA VON THEO LO"])
    head = next(r for r in lo if r[0] == "STT")
    assert [v for v in head if v] == ["STT", "CHỈ TIÊU", po.code, "TỔNG"]
    assert any(r[1] == IMPORT_COST_TYPE_LABELS[ImportCostType.OCEAN_FREIGHT] for r in lo)
    assert any(r[1] == "Giá vốn/Kg (vnd)" for r in lo)
    assert any(r[1] == "Người lập biểu" for r in lo)
    # Chỉ tiêu gọi theo trường của đơn, không chép mẫu giấy ("Giá đô", "Công nợ NCC")
    nhan = [r[1] for r in lo]
    assert "Tiền hàng (nguyên tệ)" in nhan and "Đồng tiền" in nhan
    assert "Tiền hàng (vnd)" in nhan and "Ngày hàng rời cảng (ETD)" in nhan
    assert "Giá đô" not in nhan and "Công nợ NCC" not in nhan and "Lần nhận" not in nhan

    hang = _values(wb["THEO DONG HANG"])
    head = next(r for r in hang if r[0] == "STT")
    assert head[1:6] == ["Mã đơn", "Ngày ETD", "Mã hàng", "Tên hàng", "ĐVT"]
    codes = [r[1] for r in hang[hang.index(head) + 1:] if r[0] in (1, 2)]
    assert codes == [po.code, po.code]


def test_excel_dong_tong_theo_dong_hang_khong_lech_cot(db, seed):
    """Dòng TỔNG phải nằm đúng cột với thân bảng, kể cả khi số loại chi phí đổi."""
    po = _make_po(db, seed, status="approved")
    _save_costs(db, po, [
        _cost_in(cost_type=int(ImportCostType.OCEAN_FREIGHT), amount=1_000_000),
        _cost_in(cost_type=int(ImportCostType.IMPORT_VAT), amount=500_000),
    ])
    ws = _workbook(ilc.compute(db, po_ids=[po.id]))["THEO DONG HANG"]
    rows = _values(ws)
    head = next(r for r in rows if r[0] == "STT")
    total = rows[-1] if rows[-1][1] == "TỔNG" else next(r for r in rows if r[1] == "TỔNG")
    i = head.index("Tổng chi phí (vnd)")
    assert total[i] == 1_500_000.0
    # Cột giá vốn/ĐVT của dòng TỔNG bỏ trống, cột giá vốn/Kg thì có số
    assert not total[head.index("Giá vốn/ĐVT (vnd)")]
    assert total[head.index("Giá vốn/Kg (vnd)")] > 0
