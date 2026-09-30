"""bao-CR-525 — xuất Excel Yêu cầu thanh toán THEO DÒNG CHI TIẾT.

Kiểm:
- mỗi DÒNG phiếu một hàng, thông tin phiếu lặp lại; phiếu không có dòng thì không ra hàng nào;
  nhãn trạng thái / loại / hình thức như màn hình; tiền giữ kiểu số;
- endpoint đi qua đúng bộ lọc + phạm vi dữ liệu của màn danh sách (phiếu công ty khác không lọt);
- phạm vi tính theo quyền XUẤT, không theo quyền xem: chỉ có `read` thì tệp rỗng.

Không đụng DB thật — fixture SQLite in-memory ở conftest.
"""
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest
from starlette.datastructures import QueryParams

from app.core.auth import perm_cache_clear
from app.modules.payment_request import controller as C
from app.modules.payment_request.export import COLS, build_rows, load_lines
from app.modules.payment_request.model import PaymentRequest, PaymentRequestLine

KEYS = [c.key for c in COLS]


@pytest.fixture(autouse=True)
def _clear_perm_cache():
    perm_cache_clear()
    yield
    perm_cache_clear()


def _grant(db, user_id: int, scope: str, *, export: bool = True):
    from app.modules.role.model import Permission, Role
    from app.modules.user.model import UserRole
    role = Role(code=f"R{user_id}{scope}{int(export)}", name="Vai trò test")
    db.add(role)
    db.flush()
    db.add(Permission(role_id=role.id, entity="payment_request", scope=scope,
                      can_read=True, can_export=export))
    db.add(UserRole(user_id=user_id, role_id=role.id))
    db.flush()
    perm_cache_clear()


def _request(db, company_id: int, code: str, lines: list[tuple[str, float]], **kw) -> PaymentRequest:
    vals = dict(code=code, company_id=company_id, supplier_code="NX", supplier_name="NCC Xanh",
                source_type="goods", request_date="2026-09-15", payment_method="transfer",
                prepay=0, total=sum(a for _, a in lines), note="", status="submitted")
    vals.update(kw)
    req = PaymentRequest(**vals)
    db.add(req)
    db.flush()
    for po, amount in lines:
        db.add(PaymentRequestLine(request_id=req.id, po_code=po, invoice_no=f"HD-{po}",
                                  invoice_date="2026-09-10", amount=amount))
    db.flush()
    return req


def _req(qs: str = ""):
    return SimpleNamespace(query_params=QueryParams(qs))


def _user(db, seed):
    from app.modules.user.model import User
    return db.get(User, seed.u_req_id)


def _sheet(resp):
    return openpyxl.load_workbook(BytesIO(resp.body)).active


def test_one_row_per_line_with_request_info_repeated(db, seed):
    a = _request(db, seed.company_id, "YCTT001", [("PO-1", 1000), ("PO-2", 2500.5)],
                 payment_method="cash", source_type="shipping", status="paid", prepay=1)
    _request(db, seed.company_id, "YCTT002", [])          # phiếu trắng: không ra hàng nào
    rows = build_rows(db, load_lines(db, [a, db.query(PaymentRequest).filter_by(code="YCTT002").one()]))

    assert [r["po_code"] for r in rows] == ["PO-1", "PO-2"]
    assert {r["code"] for r in rows} == {"YCTT001"}
    first = rows[0]
    assert first["status"] == "Đã chi"
    assert first["source_type"] == "Vận chuyển"
    assert first["payment_method"] == "Tiền mặt"
    assert first["prepay"] == "Có"
    assert float(rows[1]["amount"]) == 2500.5
    assert float(first["request_total"]) == 3500.5
    assert set(first) == set(KEYS), "mọi cột khai báo đều có giá trị, không cột nào rơi"


def test_export_follows_list_filter_and_company_scope(db, seed):
    from app.modules.company.model import Company
    other = Company(name="Cty Khac", code="CT02", is_active=True)
    db.add(other)
    db.flush()
    _request(db, seed.company_id, "YCTT010", [("PO-A", 100)])
    _request(db, seed.company_id, "YCTT011", [("PO-B", 200), ("PO-C", 300)], supplier_code="NY")
    _request(db, other.id, "YCTT012", [("PO-NGOAI", 999)])       # ngoài phạm vi
    user = _user(db, seed)
    _grant(db, user.id, "company")

    ws = _sheet(C.export_xlsx(request=_req(), db=db, user=user))
    po_col = KEYS.index("po_code")
    assert sorted(r[po_col].value for r in ws.iter_rows(min_row=2)) == ["PO-A", "PO-B", "PO-C"]

    ws = _sheet(C.export_xlsx(request=_req("supplier_code=NY"), db=db, user=user))
    assert sorted(r[po_col].value for r in ws.iter_rows(min_row=2)) == ["PO-B", "PO-C"]


def test_ticked_requests_only_but_never_outside_scope(db, seed):
    from app.modules.company.model import Company
    other = Company(name="Cty Khac", code="CT03", is_active=True)
    db.add(other)
    db.flush()
    a = _request(db, seed.company_id, "YCTT030", [("PO-T1", 100)])
    _request(db, seed.company_id, "YCTT031", [("PO-T2", 100)])
    outside = _request(db, other.id, "YCTT032", [("PO-T3", 100)])
    user = _user(db, seed)
    _grant(db, user.id, "company")

    # Tick cả phiếu ngoài phạm vi (gõ tay id lên URL) — vẫn không lọt ra file.
    ws = _sheet(C.export_xlsx(request=_req(f"ids={a.id},{outside.id}"), db=db, user=user))
    po_col = KEYS.index("po_code")
    assert [r[po_col].value for r in ws.iter_rows(min_row=2)] == ["PO-T1"]


def test_scope_is_taken_from_the_export_grant_not_the_read_grant(db, seed):
    # Chỉ có quyền XEM: gọi thẳng hàm (bỏ qua cổng require) vẫn không ra phiếu nào —
    # phạm vi dữ liệu tính theo quyền xuất, không mượn phạm vi của quyền xem.
    _request(db, seed.company_id, "YCTT020", [("PO-X", 100)])
    user = _user(db, seed)
    _grant(db, user.id, "company", export=False)

    ws = _sheet(C.export_xlsx(request=_req(), db=db, user=user))
    assert ws.max_row == 1, "chỉ còn dòng tiêu đề"
