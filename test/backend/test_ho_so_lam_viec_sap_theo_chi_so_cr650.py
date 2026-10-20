"""duoc-CR-650 — /summary và /summary/export nhận `metric`: trả `metric_value`, `sort=±metric` sắp trong SQL, chưa đo luôn cuối."""
from sqlalchemy.orm.attributes import flag_modified

from app.core.report_keys import ReportKey
from app.core.subject_match import SUBJECT_ROLE
from app.modules.export_log.model import ExportLog
from app.modules.usage_profile.model import UserUsageSnapshot

from test_ho_so_lam_viec_api_cr641 import PERIOD, ROOT, admin, client_as, seeded  # noqa: F401  (fixture dùng chung)
from ho_so_lam_viec_factory import make_person
from lich_lam_viec_factory import count_queries

URL = f"{ROOT}/summary"


def _set_values(db, key, by_user: dict):
    """Ghi `metrics[key]` theo user_id (None = không đo: bỏ khóa; 'null' = JSON null)."""
    for snap in db.query(UserUsageSnapshot).all():
        m = snap.metrics
        m["metrics"].pop(key, None)
        v = by_user.get(snap.user_id)
        if v is not None:
            m["metrics"][key] = {"v": None if v == "null" else v, "med": None, "p75": None, "trend": []}
        flag_modified(snap, "metrics")
    db.commit()


def _world(db, values, tag="S1"):
    people, _ = seeded(db, n=len(values), tag=tag)
    uids = [u.id for u, _ in people]
    _set_values(db, "logins", dict(zip(uids, values)))
    return uids


def _get(client, **params):
    return client.get(URL, params={**PERIOD, "page_size": 100, **params})


def test_khoa_chi_so_la_422(db, admin, client_as):
    seeded(db, n=2, tag="V1")
    c = client_as(admin)
    assert _get(c, metric="khong_co_khoa_nay").status_code == 422
    assert _get(c, metric="logins", sort="-metric").status_code == 200
    assert _get(c, sort="metric").status_code == 422          # sắp theo chỉ số mà thiếu metric
    assert c.get(f"{URL}/export", params={**PERIOD, "metric": "xyz"}).status_code == 422
    assert _get(c, metric="x" * 61).status_code == 422        # trần độ dài tham số


def test_khong_metric_thi_khong_co_metric_value(db, admin, client_as):
    seeded(db, n=2, tag="V2")
    assert all("metric_value" not in r for r in _get(client_as(admin)).json()["data"]["items"])


def test_metric_value_dung_gia_tri_tuan_nhan_va_none_khi_chua_do(db, admin, client_as):
    uids = _world(db, [5, 12.5, None, "null"], tag="V3")
    items = _get(client_as(admin), metric="logins").json()["data"]["items"]
    got = {r["user_id"]: r["metric_value"] for r in items}
    assert got == {uids[0]: 5, uids[1]: 12.5, uids[2]: None, uids[3]: None}     # JSON null / thiếu khóa → None, không 0


def test_sap_giam_dan_va_tang_dan_chua_do_luon_cuoi(db, admin, client_as):
    uids = _world(db, [3, None, 10, "null", 7, 0], tag="O1")
    c = client_as(admin)
    desc = [r["user_id"] for r in _get(c, metric="logins", sort="-metric").json()["data"]["items"]]
    asc = [r["user_id"] for r in _get(c, metric="logins", sort="metric").json()["data"]["items"]]
    assert desc[:4] == [uids[2], uids[4], uids[0], uids[5]] and set(desc[4:]) == {uids[1], uids[3]}
    assert asc[:4] == [uids[5], uids[0], uids[4], uids[2]] and set(asc[4:]) == {uids[1], uids[3]}


def test_phan_trang_hai_trang_khong_trung_khong_sot(db, admin, client_as):
    uids = _world(db, [4, None, 9, 1, None, 6], tag="P1")
    c = client_as(admin)
    for sort, expect in (("-metric", [uids[2], uids[5], uids[0], uids[3]]), ("metric", [uids[3], uids[0], uids[5], uids[2]])):
        p1 = _get(c, metric="logins", sort=sort, page_size=4, page=1).json()["data"]
        p2 = _get(c, metric="logins", sort=sort, page_size=4, page=2).json()["data"]
        assert p1["total"] == 6 and [r["user_id"] for r in p1["items"]] == expect
        assert {r["user_id"] for r in p2["items"]} == {uids[1], uids[4]}            # trang 2: hai người chưa đo
        assert all(r["metric_value"] is None for r in p2["items"])


def test_gac_quyen_khong_doi_voi_tham_so_metric(db, client_as):
    user, _ = make_person(db, "NV645")
    c = client_as(user)
    assert c.get(URL, params={**PERIOD, "metric": "logins", "sort": "-metric"}).status_code == 403
    assert c.get(f"{URL}/export", params={**PERIOD, "metric": "logins"}).status_code == 403


def test_xuat_excel_co_cot_chi_so_khi_co_metric(db, admin, client_as):
    from io import BytesIO
    from openpyxl import load_workbook
    _world(db, [3, None, 10], tag="X1")
    c = client_as(admin)
    resp = c.get(f"{URL}/export", params={**PERIOD, "metric": "logins", "sort": "-metric"})
    assert resp.status_code == 200 and resp.content[:2] == b"PK"
    rows = list(load_workbook(BytesIO(resp.content)).active.iter_rows(values_only=True))
    assert rows[0][-1] == "Chỉ số logins" and [r[-1] for r in rows[1:]] == ["10", "3", None]
    plain = c.get(f"{URL}/export", params=PERIOD)
    assert "Chỉ số" not in str(list(load_workbook(BytesIO(plain.content)).active.iter_rows(values_only=True))[0])
    assert db.query(ExportLog).filter_by(entity="usage_profile").count() == 2


def test_so_truy_van_hang_so_theo_so_nguoi_khi_sap_theo_chi_so(db, admin, client_as):
    c = client_as(admin)
    _world(db, [1, 2], tag="Q1")
    params = {**PERIOD, "metric": "logins", "sort": "-metric", "page_size": 50}
    base = count_queries(db, lambda: c.get(URL, params=PERIOD))[1]
    _, few = count_queries(db, lambda: c.get(URL, params=params))
    _world(db, [5, None, 3, 8, 2, 9, 4, 7, 6, 1, 0, 11], tag="Q2")
    resp = None

    def call():
        nonlocal resp
        resp = c.get(URL, params=params)

    _, many = count_queries(db, call)
    assert resp.json()["data"]["total"] == 14
    pick = lambda sql: [s for s in sql if "tab_usage_daily" in s or "tab_user_usage" in s]
    assert len(pick(few)) == len(pick(many)) == len(pick(base)) <= 5


#  Thử tay 20/10: tệp Excel ghi «Chỉ số first_seen_p50_min» và giờ vào là «1025» (số phút) thay vì «17:05».
def test_o_excel_dinh_dang_theo_don_vi_chua_do_de_trong():
    from app.modules.usage_profile.list_export import format_metric_cell
    assert format_metric_cell(1025, "minute_of_day") == "17:05"
    assert format_metric_cell(0, "minute_of_day") == "00:00"
    assert format_metric_cell(0.25, "ratio") == 25.0
    assert format_metric_cell(17.5, "min") == "17.5 phút"
    assert format_metric_cell(5400, "ms") == "5.4 giây"
    assert format_metric_cell(3.14, "h") == "3.1 giờ"
    assert format_metric_cell(1, "flag") == "Có" and format_metric_cell(0, "flag") == "Không"
    assert format_metric_cell(0, "count") == "0"           # 0 là giá trị THẬT
    for unit in ("ratio", "minute_of_day", "count", "flag"):
        assert format_metric_cell(None, unit) is None        # chưa đo → ô trống, không thành 0
