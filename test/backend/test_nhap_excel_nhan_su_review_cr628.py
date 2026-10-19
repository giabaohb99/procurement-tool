"""duoc-CR-628 lượt 2 (review) — hoàn tác an toàn, mã trùng, rò PII, ghi thật một lần…"""
from datetime import timedelta

import pytest
from fastapi import HTTPException
from sqlalchemy.exc import IntegrityError

from app.modules.audit.model import AuditLog
from app.modules.employee import service as employee_service
from app.modules.employee.contact_model import EmployeeContact
from app.modules.employee.model import Employee
from app.modules.employee.position_model import JobPosition
from app.modules.import_tool import controller as imp_ctl
from app.modules.import_tool import service, tasks
from app.modules.import_tool.model import (ImportBatch, ImportMode, ImportModule, ImportStatus,
                                           LogLevel)
from app.modules.attachment.model import StoredFile
from app.modules.user.model import User

from test_nhap_excel_nhan_su_cr628 import emp, errors, hr, logs, make_wb, run_import  # noqa: F401
from test_nhap_excel_nhan_su_cr628 import savepoint_that

CCCD = "079555666777"


def revert(db, b, uid):
    return service.revert_batch(db, b, uid)


# ── H1: hoàn tác không đè sửa đổi về sau ───────────────────────────────────────
def test_h1_hoan_tac_giu_sua_doi_cot_khac_va_tu_choi_khi_cot_da_doi_lai(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "DEMONV", "Điện thoại": "0911111111"}], hr)
    e = emp(db, "DEMONV")
    e.email = "sau-import@x.vn"           # cột KHÔNG thuộc lô → hoàn tác phải giữ
    db.commit()
    res = revert(db, b, hr)
    assert res["ok"], res
    e = emp(db, "DEMONV")
    assert e.email == "sau-import@x.vn" and e.phone == ""

    b2 = run_import(db, [{"Mã NV *": "DEMOTP", "Điện thoại": "0922222222", "Số CCCD": CCCD}], hr)
    e = emp(db, "DEMOTP")
    e.phone = "0999999999"                # cột thuộc lô bị sửa lại → từ chối cả lượt
    db.commit()
    res2 = revert(db, b2, hr)
    assert not res2["ok"] and "DEMOTP" in res2["message"] and "Điện thoại" in res2["message"]
    assert "0999999999" not in res2["message"] and "0922222222" not in res2["message"]
    assert emp(db, "DEMOTP").id_number == CCCD and b2.status == ImportStatus.DONE


def test_h1_nguoi_hoan_tac_thieu_quyen_nhay_cam_bi_tu_choi(db, seed, hr, cap_quyen):
    b = run_import(db, [{"Mã NV *": "DEMONV", "Số CCCD": CCCD}], hr)
    other = seed.u_nstm_id
    cap_quyen(other, "employee", scope="all", read=True, write=True, delete=True)
    res = revert(db, b, other)
    assert not res["ok"] and "nhạy cảm" in res["message"] and CCCD not in res["message"]
    assert emp(db, "DEMONV").id_number == CCCD


# ── H2: mã trùng trong tệp ─────────────────────────────────────────────────────
def test_h2_ma_trung_trong_tep_dong_sau_la_loi(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "DEMONV", "Điện thoại": "0911"},
                        {"Mã NV *": "DEMONV", "Điện thoại": "0922"}], hr)
    assert b.error_count == 1 and "trùng dòng 2" in errors(db, b.id)[0].message
    assert emp(db, "DEMONV").phone == "0911"


def test_l5_ma_tu_sinh_roi_dong_sau_ghi_dung_ma_do_bi_chan(db, seed, hr):
    # seed chưa có mã NSU nào → dòng 1 sinh NSU001; dòng 2 ghi tường minh đúng NSU001 phải bị chặn
    b = run_import(db, [{"Mã NV *": None, "Họ tên *": "Không mã"},
                        {"Mã NV *": "NSU001", "Họ tên *": "Đè lên mã vừa sinh"}], hr)
    assert b.created_count == 1 and b.error_count == 1
    assert errors(db, b.id)[0].category == "duplicate_code"
    assert emp(db, "NSU001").full_name == "Không mã"


def test_l5_ma_nv_so_giu_nguyen_khong_them_dau_cham(db, seed, hr):
    b = run_import(db, [{"Mã NV *": 12345, "Họ tên *": "Mã số"}, {"Mã NV *": "007", "Họ tên *": "Số không"}], hr)
    assert b.created_count == 2
    assert emp(db, "12345") is not None and emp(db, "007") is not None


# ── H3: quyền + phạm vi khi hoàn tác ───────────────────────────────────────────
def test_h3_hoan_tac_thieu_write_hoac_thieu_delete_bi_tu_choi(db, seed, hr, cap_quyen):
    b = run_import(db, [{"Mã NV *": "MOI01", "Họ tên *": "Mới"}, {"Mã NV *": "DEMONV", "Điện thoại": "0933"}], hr)
    other = seed.u_nstm_id
    res = revert(db, b, other)                                   # không có quyền Nhân sự nào
    assert not res["ok"]
    cap_quyen(other, "employee", scope="all", read=True, write=True)      # có write, thiếu delete
    res2 = revert(db, b, other)
    assert not res2["ok"] and "MOI01" in res2["message"] and "xóa" in res2["message"]
    assert emp(db, "MOI01") is not None and emp(db, "DEMONV").phone == "0933"


def test_h3_hoan_tac_ngoai_pham_vi_bi_tu_choi(db, seed, hr, cap_quyen):
    from app.modules.company.model import Company
    co2 = Company(name="Cty Hai", code="CT02", is_active=True)
    db.add(co2)
    db.flush()
    b = run_import(db, [{"Mã NV *": "MOI01", "Họ tên *": "Mới", "Công ty (mã)": "CT02"}], hr)
    other = seed.u_nstm_id
    cap_quyen(other, "employee", scope="company", read=True, write=True, delete=True)
    res = revert(db, b, other)
    assert not res["ok"] and "MOI01" in res["message"]
    assert emp(db, "MOI01") is not None


# ── H4: không rò PII qua lỗi DB ────────────────────────────────────────────────
def test_h4_loi_sqlalchemy_thanh_loi_dong_khong_chua_tham_so(db, seed, hr, monkeypatch):
    def boom(*a, **k):
        raise IntegrityError("INSERT ...", {"id_number": CCCD}, Exception("dup"))
    monkeypatch.setattr(employee_service, "update_employee_core", boom)
    b = run_import(db, [{"Mã NV *": "DEMONV", "Số CCCD": CCCD}], hr)
    assert b.status == ImportStatus.DONE and b.error_count == 1
    for x in logs(db, b.id):
        assert CCCD not in x.message
    assert "IntegrityError" in errors(db, b.id)[0].message


def test_h4_error_summary_lo_nhan_su_khong_chua_gia_tri(db, seed, hr, monkeypatch):
    from app.modules.import_tool import catalog_import
    sf = StoredFile(filename="a.xlsx", file_key="k", content_type="x")
    db.add(sf)
    db.flush()
    b = ImportBatch(module=ImportModule.EMPLOYEE, mode=ImportMode.APPLY, filename="a.xlsx", file_id=sf.id,
                    status=ImportStatus.QUEUED, created_by=hr, updated_by=hr)
    db.add(b)
    db.commit()

    def explode(*a, **k):
        raise RuntimeError(f"(pymysql.err) INSERT ... [parameters: ('{CCCD}',)]")
    monkeypatch.setattr(catalog_import, "run", explode)
    monkeypatch.setattr(tasks, "SessionLocal", lambda: db)
    monkeypatch.setattr(tasks, "_load_workbook", lambda f: make_wb([{"Mã NV *": "A"}]))
    out = tasks.run_import(b.id)
    db.expire_all()
    b = db.get(ImportBatch, b.id)
    assert b.status == ImportStatus.FAILED
    assert CCCD not in b.error_summary and "RuntimeError" in b.error_summary
    assert CCCD not in str(out)


# ── M1: quản lý + cấp dưới tạo cùng lô hoàn tác được ───────────────────────────
def test_m1_quan_ly_va_cap_duoi_tao_cung_lo_hoan_tac_duoc(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "BOSS1", "Họ tên *": "Sếp"},
                        {"Mã NV *": "STAFF1", "Họ tên *": "Lính", "Quản lý trực tiếp (mã NV)": "BOSS1"}], hr)
    assert b.created_count == 2
    res = revert(db, b, hr)
    assert res["ok"] and res["deleted"] == 2
    assert emp(db, "BOSS1") is None and emp(db, "STAFF1") is None


# ── M3: chức vụ không khớp → position_id = 0 + giữ chữ ─────────────────────────
def test_m3_chuc_vu_khong_khop_bo_gan_danh_muc(db, seed, hr):
    jp = JobPosition(code="TP", name="Trưởng phòng", is_active=True)
    db.add(jp)
    db.commit()
    e = emp(db, "DEMONV")
    e.position_id, e.position = jp.id, "Trưởng phòng"
    db.commit()
    b = run_import(db, [{"Mã NV *": "DEMONV", "Chức vụ (mã/tên)": "Chuyên gia lạ"}], hr)
    e = emp(db, "DEMONV")
    assert e.position == "Chuyên gia lạ" and (e.position_id or 0) == 0
    assert logs(db, b.id, LogLevel.REVIEW)
    res = revert(db, b, hr)
    assert res["ok"]
    e = emp(db, "DEMONV")
    assert e.position == "Trưởng phòng" and e.position_id == jp.id


# ── M4: dữ liệu con thêm sau khi nhập → không xóa ──────────────────────────────
def test_m4_co_du_lieu_con_hoac_da_sua_sau_import_thi_khong_hoan_tac(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "MOI01", "Họ tên *": "Mới"}], hr)
    n = emp(db, "MOI01")
    db.add(EmployeeContact(employee_id=n.id, full_name="Người báo tin"))
    db.commit()
    res = revert(db, b, hr)
    assert not res["ok"] and "MOI01" in res["message"] and "liên hệ" in res["message"]
    assert emp(db, "MOI01") is not None

    b2 = run_import(db, [{"Mã NV *": "MOI02", "Họ tên *": "Mới hai"}], hr)
    n2 = emp(db, "MOI02")
    n2.updated_at = n2.created_at + timedelta(minutes=5)
    db.commit()
    res2 = revert(db, b2, hr)
    assert not res2["ok"] and "chỉnh sửa" in res2["message"]


# ── M5: nhật ký từng nhân sự sau hoàn tác ──────────────────────────────────────
def test_m5_hoan_tac_ghi_nhat_ky_tung_nhan_su(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "MOI01", "Họ tên *": "Mới"}, {"Mã NV *": "DEMONV", "Điện thoại": "0944"}], hr)
    nid = emp(db, "MOI01").id
    assert revert(db, b, hr)["ok"]
    q = db.query(AuditLog).filter(AuditLog.message.like("%Hoàn tác lô nhập Excel%"))
    assert {(a.entity_id, a.action) for a in q} >= {(nid, "delete"), (seed.emp_nstm_id, "update")}


# ── M6: phòng ban cấm xóa + gác department.write ───────────────────────────────
def test_m6_phong_ban_cam_xoa_qua_excel_va_gac_department_write(db, seed, cap_quyen):
    import openpyxl

    from app.modules.department.model import Department
    from app.modules.import_tool import catalog_import
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Phong ban"
    ws.append(["Mã *", "Tên phòng ban *", "Hành động"])
    ws.append(["DEPT01", "Phòng Test", "__/delete/__"])
    b = ImportBatch(module=ImportModule.DEPARTMENT, mode=ImportMode.APPLY, filename="a", file_id=0,
                    status=ImportStatus.QUEUED, created_by=1, updated_by=1)
    db.add(b)
    db.commit()
    catalog_import.run(db, b, wb, apply=True)
    assert b.error_count == 1 and db.get(Department, seed.dept_id) is not None

    cap_quyen(seed.u_req_id, "import", scope="all", read=True, create=True, delete=True)
    user = db.get(User, seed.u_req_id)

    class _F:
        filename = "x.xlsx"
        file = None
    with pytest.raises(HTTPException) as e:
        imp_ctl.upload_import(module=ImportModule.DEPARTMENT, mode=0, file=_F(), db=db, user=user)
    assert e.value.status_code == 403


# ── M7: ghi thật một lần ───────────────────────────────────────────────────────
def test_m7_ghi_that_tu_ban_chay_thu_chi_mot_lan(db, seed, cap_quyen, monkeypatch):
    cap_quyen(seed.u_req_id, "import", scope="all", read=True, create=True, delete=True)
    cap_quyen(seed.u_req_id, "employee", scope="all", read=True, create=True, write=True)
    monkeypatch.setattr(imp_ctl.run_import, "delay", lambda bid: None)
    user = db.get(User, seed.u_req_id)
    sf = StoredFile(filename="a.xlsx", file_key="k", content_type="x")
    db.add(sf)
    db.flush()
    dry = ImportBatch(module=ImportModule.EMPLOYEE, mode=ImportMode.DRY_RUN, filename="a", file_id=sf.id,
                      status=ImportStatus.DONE, created_by=user.id, updated_by=user.id)
    db.add(dry)
    db.commit()
    imp_ctl.commit_import(dry.id, db=db, user=user)
    with pytest.raises(HTTPException) as e:
        imp_ctl.commit_import(dry.id, db=db, user=user)
    assert e.value.status_code == 409
    # lô ghi thật bị hoàn tác/lỗi thì cho ghi lại
    applied = db.query(ImportBatch).filter(ImportBatch.mode == ImportMode.APPLY).first()
    applied.status = ImportStatus.FAILED
    db.commit()
    imp_ctl.commit_import(dry.id, db=db, user=user)


# ── M2: chạy thử không gọi generate_code (FOR UPDATE) ──────────────────────────
def test_m2_dry_run_xem_truoc_ma_khong_goi_generate_code(db, seed, hr, monkeypatch):
    def no_lock(*a, **k):
        raise AssertionError("dry-run không được gọi generate_code (khóa hàng)")
    monkeypatch.setattr(employee_service, "generate_code", no_lock)
    b = run_import(db, [{"Mã NV *": None, "Họ tên *": "Xem trước"}], hr, apply=False)
    assert b.created_count == 1 and b.error_count == 0
    assert any("NSU" in x.message for x in logs(db, b.id, LogLevel.INFO))


# ── L4 / L6 ────────────────────────────────────────────────────────────────────
def test_l4_chi_co_quyen_ghi_nhay_cam_thieu_doc_van_bo_cot(db, seed, cap_quyen):
    savepoint_that(db)
    cap_quyen(seed.u_req_id, "employee", scope="all", read=True, create=True, write=True)
    cap_quyen(seed.u_req_id, "employee_sensitive", scope="all", write=True)
    b = run_import(db, [{"Mã NV *": "DEMONV", "Số CCCD": CCCD, "Điện thoại": "0955"}], seed.u_req_id)
    assert (emp(db, "DEMONV").id_number or "") == "" and emp(db, "DEMONV").phone == "0955"
    assert any(x.category == "sensitive_ignored" for x in logs(db, b.id))


def test_l6_so_con_inf_hoac_qua_lon_la_loi_dong_khong_sap(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "DEMONV", "Số con": "1e999"}, {"Mã NV *": "DEMOTP", "Số con": "inf"}], hr)
    assert b.status == ImportStatus.DONE and b.error_count == 2


# ═══════════════════════════ Lượt 3 ═══════════════════════════
def _mk_depts(db, seed):
    from app.modules.department.model import Department
    d3 = Department(code="DEPT03", name="Phòng Ba", company_id=seed.company_id, is_active=True)
    d4 = Department(code="DEPT04", name="Phòng Tư", company_id=seed.company_id, is_active=True)
    db.add_all([d3, d4])
    db.commit()
    return d3.id, d4.id


def test_ha_truy_van_do_lo_apply_la_doc_co_khoa_chia_se_tren_mysql(db):
    """SQLite KHÔNG tái hiện được REPEATABLE READ; ở đây chỉ chứng minh câu dò biên dịch sang MySQL có khóa đọc
    (LOCK IN SHARE MODE) — thiếu là hai cú bấm gần nhau cùng ghi thật (phiên sau đọc ảnh chụp cũ)."""
    from sqlalchemy.dialects import mysql
    b = ImportBatch(id=1, module=ImportModule.EMPLOYEE, file_id=5)
    sql = str(service.live_apply_query(db, b).statement.compile(dialect=mysql.dialect()))
    assert "SHARE MODE" in sql.upper() or "FOR SHARE" in sql.upper()


def test_hb_hoan_tac_khoi_phuc_dung_tap_phong_kiem_nhiem(db, seed, hr):
    from app.modules.employee.department_model import EmployeeDepartment
    from app.modules.employee.department_service import departments_of
    d3, d4 = _mk_depts(db, seed)
    db.add_all([EmployeeDepartment(employee_id=seed.emp_nstm_id, department_id=seed.dept_id, is_primary=True),
                EmployeeDepartment(employee_id=seed.emp_nstm_id, department_id=d4, is_primary=False)])
    db.commit()
    b = run_import(db, [{"Mã NV *": "DEMONV", "Phòng ban (mã)": "DEPT03"}], hr)
    assert b.error_count == 0
    assert set(departments_of(db, seed.emp_nstm_id)) == {d3, seed.dept_id, d4}     # phòng cũ bị hạ xuống kiêm nhiệm
    assert revert(db, b, hr)["ok"]
    db.expire_all()
    rows = {(r.department_id, bool(r.is_primary)) for r in
            db.query(EmployeeDepartment).filter(EmployeeDepartment.employee_id == seed.emp_nstm_id)}
    assert rows == {(seed.dept_id, True), (d4, False)}
    assert emp(db, "DEMONV").department_id == seed.dept_id


def test_hb_tap_phong_ban_da_doi_sau_import_thi_tu_choi(db, seed, hr):
    from app.modules.employee.department_model import EmployeeDepartment
    d3, d4 = _mk_depts(db, seed)
    b = run_import(db, [{"Mã NV *": "DEMONV", "Phòng ban (mã)": "DEPT03"}], hr)
    db.add(EmployeeDepartment(employee_id=seed.emp_nstm_id, department_id=d4, is_primary=False))
    db.commit()
    res = revert(db, b, hr)
    assert not res["ok"] and "DEMONV" in res["message"] and "Phòng ban" in res["message"]


def test_ma_hoan_tac_doi_phong_cua_chinh_minh_la_l1(db, seed, hr, cap_quyen):
    d3, _ = _mk_depts(db, seed)
    e = emp(db, "DEMONV")
    e.department_id = d3
    db.commit()
    b = run_import(db, [{"Mã NV *": "DEMONV", "Phòng ban (mã)": "DEPT01"}], hr)
    cap_quyen(seed.u_nstm_id, "employee", scope="all", read=True, write=True, delete=True)   # u_nstm = DEMONV
    res = revert(db, b, seed.u_nstm_id)
    assert not res["ok"] and "DEMONV" in res["message"] and "phòng ban" in res["message"].lower()


def test_ma_hoan_tac_ve_phong_ngoai_tam_la_l2(db, seed, hr, cap_quyen):
    d3, _ = _mk_depts(db, seed)
    e = emp(db, "DEMOTP")
    e.department_id = d3
    db.commit()
    b = run_import(db, [{"Mã NV *": "DEMOTP", "Phòng ban (mã)": "DEPT01"}], hr)
    backup = db.query(User).filter(User.email == "BACKUP01").first().id
    cap_quyen(backup, "employee", scope="dept", read=True, write=True, delete=True)    # chỉ tầm DEPT01
    res = revert(db, b, backup)
    assert not res["ok"] and "DEMOTP" in res["message"]
    assert emp(db, "DEMOTP").department_id == seed.dept_id


def test_mb_dau_thoi_gian_khong_co_do_tre(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "MOI01", "Họ tên *": "Mới"}], hr)
    n = emp(db, "MOI01")
    n.updated_at = n.updated_at + timedelta(seconds=1)       # 1 giây vẫn là «bị sửa sau» (không còn 60s lỏng)
    db.commit()
    res = revert(db, b, hr)
    assert not res["ok"] and "chỉnh sửa" in res["message"]


def test_mc_hoan_tac_ve_nghi_viec_khoa_lai_tai_khoan_va_ghi_nhat_ky(db, seed, hr):
    e = emp(db, "DEMONV")
    e.status = "resigned"
    db.commit()
    b = run_import(db, [{"Mã NV *": "DEMONV", "Trạng thái": "official"}], hr)
    assert db.get(User, seed.u_nstm_id).is_active is True
    assert revert(db, b, hr)["ok"]
    db.expire_all()
    assert emp(db, "DEMONV").status == "resigned" and db.get(User, seed.u_nstm_id).is_active is False
    assert db.query(AuditLog).filter(AuditLog.entity == "user", AuditLog.entity_id == seed.u_nstm_id,
                                     AuditLog.action == "deactivate").count() == 1


def test_md_hau_ky_hoan_tac_loi_khong_lam_hong_ket_qua(db, seed, hr, monkeypatch):
    from app.modules.import_tool import catalog_import
    b = run_import(db, [{"Mã NV *": "DEMONV", "Điện thoại": "0966"}], hr)

    def boom(*a, **k):
        raise RuntimeError("audit down")
    monkeypatch.setitem(catalog_import.ADAPTERS[ImportModule.EMPLOYEE], "after_revert", boom)
    res = revert(db, b, hr)
    assert res["ok"] and emp(db, "DEMONV").phone == ""


def test_la_chi_bo_qua_cap_duoi_do_lo_tao_hoac_lo_gan_quan_ly(db, seed, hr):
    # lô gán quản lý cho hồ sơ cũ DEMOTP = BOSS1 (mới) → hoàn tác được
    b = run_import(db, [{"Mã NV *": "BOSS1", "Họ tên *": "Sếp"},
                        {"Mã NV *": "DEMOTP", "Quản lý trực tiếp (mã NV)": "BOSS1"}], hr)
    assert revert(db, b, hr)["ok"] and emp(db, "BOSS1") is None and emp(db, "DEMOTP").manager_id == 0
    # nhưng nếu người khác gán BOSS làm quản lý SAU lô → chặn
    b2 = run_import(db, [{"Mã NV *": "BOSS2", "Họ tên *": "Sếp hai"}], hr)
    emp(db, "DEMONV").manager_id = emp(db, "BOSS2").id
    db.commit()
    res = revert(db, b2, hr)
    assert not res["ok"] and "BOSS2" in res["message"]


def test_lb_chuan_hoa_ngay_chi_cho_cot_ngay(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "DEMONV", "Nơi sinh": "2024-05-06 Hà Nội"}], hr)
    e = emp(db, "DEMONV")
    e.place_of_birth = "2024-05-06 Huế"          # khác hẳn nhưng 10 ký tự đầu giống dạng ngày
    db.commit()
    res = revert(db, b, hr)
    assert not res["ok"] and "Nơi sinh" in res["message"]


def test_lc_ma_trung_khong_phan_biet_hoa_thuong(db, seed, hr):
    b = run_import(db, [{"Mã NV *": "DEMONV", "Điện thoại": "0911"}, {"Mã NV *": "demonv", "Điện thoại": "0922"}], hr)
    assert b.error_count == 1 and errors(db, b.id)[0].category == "duplicate_code"
    assert emp(db, "DEMONV").phone == "0911"


@pytest.mark.parametrize("kind", ["deadlock", "pending"])
def test_le_loi_deadlock_hoac_phien_hong_lam_hong_ca_lo(db, seed, hr, monkeypatch, kind):
    from sqlalchemy.exc import OperationalError, PendingRollbackError

    class _Orig(Exception):
        pass

    def boom(*a, **k):
        if kind == "deadlock":
            raise OperationalError("UPDATE ...", {"p": CCCD}, _Orig(1213, "Deadlock found"))
        raise PendingRollbackError("Can't reconnect until invalid transaction is rolled back")
    monkeypatch.setattr(employee_service, "update_employee_core", boom)
    with pytest.raises((OperationalError, PendingRollbackError)):
        run_import(db, [{"Mã NV *": "DEMONV", "Điện thoại": "0911"}], hr)


def test_lf_mau_phong_ban_khong_con_cot_hanh_dong():
    from io import BytesIO

    import openpyxl

    from app.modules.import_tool import catalog_import
    wb = openpyxl.load_workbook(BytesIO(catalog_import.build_template(ImportModule.DEPARTMENT)))
    assert "Hành động" not in [c.value for c in wb.worksheets[0][1]]
