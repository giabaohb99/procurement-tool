"""Mục «Thuốc BVTV» của Tra cứu thị trường (29/09/2026).

Danh mục nạp từ bản cào danhmuc.thuocbvtv.com — tệp `thuoc-bvtv.json` HOẶC `thuoc-bvtv.xlsx`.
Canh ba thứ dễ lủng nhất:
  · hai đường đọc (JSON lồng / Excel phẳng) phải ra CÙNG dữ liệu — lần chạy đầu trên tệp thật
    đã lệch ở thuốc id 3670 (tên hoạt chất dính URL `https://…`, tách ở «:» là cắt giữa URL);
  · nạp lại là THAY TOÀN BỘ, kể cả bảng con phạm vi sử dụng, rồi gắn lại hoạt chất dòng hàng;
  · chuỗi dài hơn cột → MySQL trả 500. SQLite của pytest KHÔNG ép độ dài, nên canh bằng cách
    so trần ở tầng đọc với `String(n)` của model thay vì ghi xuống rồi đọc lên (xanh giả).
"""
import io
import json
from datetime import date
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app.modules.customs import pesticide_reader as R
from app.modules.customs import pesticide_service as S
from app.modules.customs.constants import PesticideStatus
from app.modules.customs.model import CustomsLine, CustomsPesticide, CustomsPesticideUse
from app.modules.customs.pesticide_controller import router


def _raw(**over):
    raw = {
        "id": 2815, "ten_thuoc": "Bipyrhone 20EC", "phan_nhom": "Thuốc trừ bệnh",
        "linh_vuc": "THUỐC SỬ DỤNG TRONG NÔNG NGHIỆP", "tinh_trang": "Còn hiệu lực",
        "cong_ty_dang_ky": "Công ty TNHH Ngân Anh", "hoat_chat": "Chitosan 2% + Polyoxin B 10%",
        "ham_luong": "12% w/w", "so_dang_ky": "514/CNĐKT-BVTV", "ngay_cap": "2023-09-25",
        "ngay_het_han": "2028-09-25", "url": "https://danhmuc.thuocbvtv.com/thuoc/detail/2815/x",
        "nhom_doc": [{"he": "GHS", "nhom": "5", "mo_ta": "GHS - Nhóm 5: Rất ít độc"},
                     {"he": "WHO", "nhom": "4", "mo_ta": "WHO - Nhóm 4: Ít độc"}],
        "quan_ly_tinh_khang": {"hoat_chat": [
            {"ten": "Chitosan", "ma": "", "nhom": "", "phuong_thuc": ""},
            {"ten": "Polyoxin B (Nereistoxihttps://113.190.254.147/x.png)", "ma": "FRAC 19",
             "nhom": "Peptidyl", "phuong_thuc": "Tổng hợp chitin"}]},
        "pham_vi_su_dung": [
            {"cay_trong": "cà rốt", "dich_hai": "đốm vòng", "lieu_luong": "0.4 lít/ha",
             "thoi_gian_cach_ly": "7 ngày", "cach_dung": "Phun khi bệnh chớm"},
            {"cay_trong": "lúa", "dich_hai": "đạo ôn", "lieu_luong": "0.5 lít/ha",
             "thoi_gian_cach_ly": "14 ngày", "cach_dung": "Phun"}],
    }
    raw.update(over)
    return raw


def _json(*raws) -> bytes:
    return json.dumps(list(raws), ensure_ascii=False).encode()


def _xlsx(*raws) -> bytes:
    """Dựng tệp Excel ĐÚNG cách bản cào xuất (hai sheet, ô lồng đã nối chuỗi)."""
    wb = Workbook()
    ws = wb.active
    ws.title = R.LIST_SHEET
    ws.append(["id", "ten_thuoc", "phan_nhom", "linh_vuc", "tinh_trang", "hoat_chat", "ham_luong",
               "cong_ty_dang_ky", "so_dang_ky", "ngay_cap", "ngay_het_han", "nhom_doc",
               "quan_ly_tinh_khang", "url"])
    uses = wb.create_sheet(R.USE_SHEET)
    uses.append(["id", "ten_thuoc", "stt_pham_vi", "cay_trong", "dich_hai", "lieu_luong",
                 "thoi_gian_cach_ly", "cach_dung"])
    for raw in raws:
        doc = "; ".join(f"{t['he']} {t['nhom']} ({t['mo_ta']})" for t in raw["nhom_doc"])
        qltk = "; ".join(f"{a['ten']}: {a['ma']} | {a['nhom']} | {a['phuong_thuc']}"
                         for a in raw["quan_ly_tinh_khang"]["hoat_chat"])
        ws.append([raw["id"], raw["ten_thuoc"], raw["phan_nhom"], raw["linh_vuc"], raw["tinh_trang"],
                   raw["hoat_chat"], raw["ham_luong"], raw["cong_ty_dang_ky"], raw["so_dang_ky"],
                   raw["ngay_cap"], raw["ngay_het_han"], doc, qltk, raw["url"]])
        #  Ghi NGƯỢC thứ tự: sheet phẳng không hứa xếp sẵn, bộ đọc phải tự xếp theo stt.
        for i, u in reversed(list(enumerate(raw["pham_vi_su_dung"], start=1))):
            uses.append([raw["id"], raw["ten_thuoc"], i, u["cay_trong"], u["dich_hai"],
                         u["lieu_luong"], u["thoi_gian_cach_ly"], u["cach_dung"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Đọc tệp ──────────────────────────────────────────────────────────────────────────────

def test_json_record_is_normalized():
    [rec] = R.read_json(_json(_raw()))
    assert rec["status"] == PesticideStatus.ACTIVE
    assert rec["trade_key"] == "BIPYRHONE"
    assert rec["registered_on"] == date(2023, 9, 25) and rec["expires_on"] == date(2028, 9, 25)
    assert rec["toxicity"] == "GHS 5 (GHS - Nhóm 5: Rất ít độc); WHO 4 (WHO - Nhóm 4: Ít độc)"
    #  Hoạt chất chưa xếp nhóm kháng (Chitosan) bỏ hẳn — không bày «Chitosan: » trống trơn.
    assert rec["resistance"] == ("Polyoxin B (Nereistoxihttps://113.190.254.147/x.png): FRAC 19 | "
                                 "Peptidyl | Tổng hợp chitin")
    assert [u["crop"] for u in rec["uses"]] == ["cà rốt", "lúa"]


def test_json_and_xlsx_read_the_same_catalog():
    """Lỗi thật khi nạp tệp thật: id 3670 có URL trong tên hoạt chất, Excel lệch JSON."""
    raws = [_raw(), _raw(id=9, ten_thuoc="Abc 5WP", tinh_trang="Hết hiệu lực", pham_vi_su_dung=[])]
    assert R.read_xlsx(_xlsx(*raws)) == R.read_json(_json(*raws))


@pytest.mark.parametrize("label, code", [
    ("Còn hiệu lực", PesticideStatus.ACTIVE), ("HẾT HIỆU LỰC", PesticideStatus.EXPIRED),
    ("Đang sử dụng", PesticideStatus.IN_USE), ("Tạm đình chỉ", PesticideStatus.UNKNOWN),
    ("", PesticideStatus.UNKNOWN), (None, PesticideStatus.UNKNOWN)])
def test_unknown_status_is_not_guessed(label, code):
    """Tình trạng lạ phải ra UNKNOWN — đoán thành «Còn hiệu lực» là lọt vào bộ lọc mặc định."""
    assert R.read_json(_json(_raw(tinh_trang=label)))[0]["status"] == code


@pytest.mark.parametrize("value", ["", None, "25/9/2028", "0001-01-01", "9999-12-31", "2028-02-30"])
def test_bad_or_absurd_dates_become_empty(value):
    assert R.read_json(_json(_raw(ngay_het_han=value)))[0]["expires_on"] is None


def test_short_trade_names_get_no_trade_key():
    """Tên < 5 ký tự dò trong tên hàng hải quan sẽ khớp bừa — không sinh khóa."""
    assert R.read_json(_json(_raw(ten_thuoc="Ace 5EC")))[0]["trade_key"] == ""


def test_file_without_any_usage_scope_is_refused():
    """Nạp là THAY TOÀN BỘ: tệp không có dòng phạm vi nào mà lọt qua là xóa sạch 15 nghìn dòng."""
    with pytest.raises(ValueError, match="phạm vi"):
        R.read_json(_json(_raw(pham_vi_su_dung=[]), _raw(id=2, pham_vi_su_dung=[])))


def test_xlsx_without_usage_sheet_is_refused():
    wb = Workbook()
    wb.active.title = R.LIST_SHEET
    wb.active.append(["id", "ten_thuoc"])
    wb.active.append([1, "Abc 5WP"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match=R.USE_SHEET):
        R.read_file("thuoc-bvtv.xlsx", buf.getvalue())


def test_xlsx_row_cap_stops_while_reading(monkeypatch):
    """Trần kiểm NGAY LÚC ĐỌC, không đợi đọc hết vào bộ nhớ."""
    monkeypatch.setattr(R, "MAX_USE_ROWS", 1)
    with pytest.raises(ValueError, match="quá trần"):
        R.read_xlsx(_xlsx(_raw()))


@pytest.mark.parametrize("url, kept", [
    ("https://danhmuc.thuocbvtv.com/x", True), ("HTTP://a.vn", True),
    ("javascript:alert(1)", False), ("  javascript:alert(1)", False), ("data:text/html,x", False),
    ("//evil.example", False), ("", False)])
def test_source_url_only_accepts_http(url, kept):
    """Giá trị này thành `<a href>` trên màn chi tiết."""
    assert bool(R.read_json(_json(_raw(url=url)))[0]["source_url"]) is kept


def test_records_without_name_are_dropped_and_empty_file_is_refused():
    assert len(R.read_json(_json(_raw(), _raw(id=2, ten_thuoc="  ")))) == 1
    with pytest.raises(ValueError):
        R.read_json(_json(_raw(ten_thuoc="")))
    with pytest.raises(ValueError):
        R.read_json(b"[]")


@pytest.mark.parametrize("filename, content", [
    ("thuoc.json", b"{\"a\": 1}"), ("thuoc.json", b"\xff\xfe rac"), ("thuoc.json", b"[1, 2"),
    ("thuoc.xlsx", b"khong phai excel"), ("thuoc.csv", b"id,ten"), ("", b"[]")])
def test_malformed_files_raise_value_error(filename, content):
    with pytest.raises(ValueError):
        R.read_file(filename, content)


def test_xlsx_without_list_sheet_is_refused():
    wb = Workbook()
    wb.active.title = "Sheet1"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="thiếu sheet"):
        R.read_file("thuoc-bvtv.xlsx", buf.getvalue())


def test_record_cap_blocks_a_wrong_file(monkeypatch):
    monkeypatch.setattr(R, "MAX_RECORDS", 2)
    with pytest.raises(ValueError, match="quá trần"):
        R.read_json(_json(_raw(id=1), _raw(id=2), _raw(id=3)))


def test_reader_limits_match_model_column_lengths():
    """Trần ở tầng đọc lệch `String(n)` của model = MySQL trả 500 lúc nạp tệp thật."""
    for model, limits in ((CustomsPesticide, R._LIMITS), (CustomsPesticideUse, R._USE_LIMITS)):
        for name, size in limits.items():
            assert model.__table__.c[name].type.length == size, name


def test_overlong_values_are_cut_to_column_size():
    [rec] = R.read_json(_json(_raw(ten_thuoc="X" * 900, so_dang_ky="9" * 900,
                                   pham_vi_su_dung=[{"cay_trong": "c" * 900}])))
    assert len(rec["trade_name"]) == 255 and len(rec["registration_no"]) == 60
    assert len(rec["uses"][0]["crop"]) == 255


# ── Nạp + tra cứu ─────────────────────────────────────────────────────────────────────────

def test_replace_catalog_replaces_everything_including_uses(db):
    S.replace_catalog(db, R.read_json(_json(_raw(), _raw(id=7, ten_thuoc="Cũ 5WP"))), 1, "a.json")
    assert db.query(CustomsPesticide).count() == 2 and db.query(CustomsPesticideUse).count() == 4
    out = S.replace_catalog(db, R.read_json(_json(_raw(id=8, ten_thuoc="Mới 10SC"))), 1, "b.json")
    assert out["pesticides"] == 1 and out["uses"] == 2
    [new] = db.query(CustomsPesticide).all()
    assert new.trade_name == "Mới 10SC"
    #  Id tự cấp 1..N lại từ đầu mỗi lần nạp — phạm vi phải trỏ đúng thuốc MỚI, không mồ côi.
    uses = db.query(CustomsPesticideUse).all()
    assert len(uses) == 2 and {u.pesticide_id for u in uses} == {new.id}


def test_replace_catalog_retags_customs_lines(db):
    """Danh mục đổi thì dòng hàng hải quan đã nạp phải mang hoạt chất theo danh mục MỚI."""
    line = CustomsLine(reg_date=date(2026, 3, 1), product_name="Thuốc BIPYRHONE 20EC nhập khẩu",
                       unit_code="KGM", price_usd=Decimal("1"), quantity=Decimal("1"),
                       hs_code="38089290")
    db.add(line)
    db.commit()
    out = S.replace_catalog(db, R.read_json(_json(_raw(hoat_chat="Bifenazate 277g/l"))), 1, "a.json")
    db.refresh(line)
    assert out["retag"]["tagged"] == 1
    assert "BIFENAZATE" in line.active_ingredient.upper()


def test_list_filters_count_and_order(db):
    S.replace_catalog(db, R.read_json(_json(
        _raw(id=1, ten_thuoc="Zeta 5EC", so_dang_ky="111/CNĐKT-BVTV"),
        _raw(id=2, ten_thuoc="Alpha 5EC", tinh_trang="Hết hiệu lực", pham_vi_su_dung=[]),
        _raw(id=3, ten_thuoc="Beta 5EC", phan_nhom="Thuốc trừ cỏ"))), 1, "a.json")
    total, items = S.list_pesticides(db, "", None, "", "", 0, 50)
    assert total == 3 and [i["trade_name"] for i in items] == ["Alpha 5EC", "Beta 5EC", "Zeta 5EC"]
    assert [i["use_count"] for i in items] == [0, 2, 2]
    total, items = S.list_pesticides(db, "", int(PesticideStatus.ACTIVE), "", "", 0, 50)
    assert total == 2 and all(i["status_label"] == "Còn hiệu lực" for i in items)
    assert S.list_pesticides(db, "111/CNĐKT", None, "", "", 0, 50)[0] == 1, "tìm theo số ĐK"
    assert S.list_pesticides(db, "", None, "Thuốc trừ cỏ", "", 0, 50)[0] == 1
    #  Trang vượt quá tổng: rỗng chứ không lỗi, tổng vẫn đúng.
    assert S.list_pesticides(db, "", None, "", "", 100, 50) == (3, [])


def test_options_and_detail(db):
    S.replace_catalog(db, R.read_json(_json(_raw(id=1), _raw(id=2, tinh_trang="Đang sử dụng"))),
                      1, "a.json")
    opt = S.options(db)
    assert opt["total"] == 2
    assert {s["label"]: s["count"] for s in opt["statuses"]} == {
        "Còn hiệu lực": 1, "Hết hiệu lực": 0, "Đang sử dụng": 1}
    pid = db.query(CustomsPesticide.id).first()[0]
    detail = S.get_pesticide(db, pid)
    assert [u["crop"] for u in detail["uses"]] == ["cà rốt", "lúa"]
    with pytest.raises(HTTPException) as err:
        S.get_pesticide(db, 999_999)
    assert err.value.status_code == 404


def test_empty_catalog_options_do_not_crash(db):
    opt = S.options(db)
    assert opt["total"] == 0 and opt["last_loaded_at"] is None and opt["pest_groups"] == []


# ── Quyền ───────────────────────────────────────────────────────────────────────────────

def _guard(route) -> tuple[str, str]:
    for dep in route.dependant.dependencies:
        code = getattr(dep.call, "__code__", None)
        if code and set(code.co_freevars) >= {"entity", "action"}:
            cells = dict(zip(code.co_freevars, dep.call.__closure__))
            return cells["entity"].cell_contents, cells["action"].cell_contents
    return ("", "")


def test_every_route_is_guarded():
    """Đọc = `customs_price.read` (một mục của màn tra cứu). Sửa danh mục = khóa riêng
    `customs_pesticide` (duoc-CR-490) — nạp tệp (thay danh mục + gắn lại MỌI dòng hàng) = `write`."""
    guards = {(r.path, next(iter(r.methods))): _guard(r) for r in router.routes}
    assert guards == {
        ("/api/customs/pesticides", "GET"): ("customs_price", "read"),
        ("/api/customs/pesticides/options", "GET"): ("customs_price", "read"),
        #  02/10/2026 — xuất Excel (round-trip với «Nạp danh mục»): khóa `export`, không phải
        #  `read` (QĐ-I4, `doc/erp/16-...md`) — ai chỉ xem được không có nút xuất.
        ("/api/customs/pesticides/export", "GET"): ("customs_price", "export"),
        ("/api/customs/pesticides/{pesticide_id}", "GET"): ("customs_price", "read"),
        #  01/10/2026 — «cùng công ty» / «cùng hoạt chất» của trang chi tiết: cùng quyền đọc.
        ("/api/customs/pesticides/{pesticide_id}/related", "GET"): ("customs_price", "read"),
        ("/api/customs/pesticides", "POST"): ("customs_pesticide", "create"),
        ("/api/customs/pesticides/{pesticide_id}", "PATCH"): ("customs_pesticide", "write"),
        ("/api/customs/pesticides/{pesticide_id}", "DELETE"): ("customs_pesticide", "delete"),
        ("/api/customs/pesticides/import", "POST"): ("customs_pesticide", "write"),
    }
