"""duoc-CR-688 — báo cáo D1 «Điều chuyển & bổ nhiệm» (khóa 24): `/api/work-moves/summary`.

Canh: đếm theo loại sự kiện (chỉ dòng có `from_date` trong kỳ, đã áp lẫn chưa áp), số người, số tháng
tới lần bổ nhiệm, kiêm nhiệm tại mốc, chiều `flow` (phòng đi lấy từ dòng liền trước / hồ sơ chưa áp /
«(Chưa rõ)»), phạm vi áp lên NGƯỜI của dòng, `department_id`/`company_id` chỉ thu hẹp, gác khóa báo
cáo + quyền xuất, danh sách KHÔNG trả tệp quyết định, số truy vấn CỐ ĐỊNH.
"""
import uuid
from datetime import datetime, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event

from app.core.auth import get_current_user, get_perm_profile
from app.core.database import get_db
from app.core.hr_work_history_codes import WorkEventType as T
from app.core.report_keys import ReportKey
from app.core.report_period import parse_period
from app.core.subject_match import SUBJECT_EMPLOYEE
from app.core.vn_time import vn_today
from app.main import app
from app.modules.employee import work_move_report_service as svc
from app.modules.employee.model import Employee
from app.modules.employee.work_history_model import EmployeeWorkHistory
from app.modules.report_access.model import ReportAccess

TODAY = vn_today()
SUM = "/api/work-moves/summary"
QS = {"preset": "custom", "date_from": (TODAY - timedelta(days=20)).isoformat(),
      "date_to": TODAY.isoformat(), "compare": "none"}


def _d(days):
    return None if days is None else TODAY + timedelta(days=days)


@pytest.fixture
def client_as(db):
    def make_client(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make_client
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def _emp(db, world, code, hire=-900, **kw) -> Employee:
    base = dict(code=code, full_name=f"Người {code}", company_id=world.co["B"],
                department_id=world.dept["B.hc"], is_active=True, hire_date=_d(hire))
    base.update(kw)
    e = Employee(**base)
    db.add(e)
    db.flush()
    return e


def _line(db, emp, etype, frm, to=None, dept=None, applied=True, decision="", company=None, world=None):
    row = EmployeeWorkHistory(
        employee_id=emp.id, event_type=int(etype), from_date=_d(frm), to_date=_d(to),
        company_id=world.co[company] if company else 0, department_id=world.dept[dept] if dept else 0,
        position_label="Trưởng nhóm", decision_no=decision,
        applied_at=datetime.utcnow() if applied else None)
    db.add(row)
    db.flush()
    return row


def _dataset(db, world):
    w = lambda e, *a, **k: _line(db, e, *a, world=world, **k)   # noqa: E731
    p1 = _emp(db, world, "p1"); w(p1, T.HIRE, -15, dept="B.hc")
    p2 = _emp(db, world, "p2"); w(p2, T.TRANSFER, -300, dept="A.mua"); w(p2, T.TRANSFER, -10, dept="B.hc",
                                                                      decision="QĐ-01")
    p3 = _emp(db, world, "p3"); w(p3, T.TRANSFER, -8, dept="B.kt", applied=False)
    p4 = _emp(db, world, "p4", hire=-365); w(p4, T.APPOINT, -5, dept="B.hc")
    p5 = _emp(db, world, "p5", hire=-730); w(p5, T.APPOINT, -4, dept="B.hc")
    p6 = _emp(db, world, "p6", hire=None); w(p6, T.APPOINT, -3, dept="B.hc")      # thiếu ngày vào làm
    p7 = _emp(db, world, "p7"); w(p7, T.CONCURRENT, -40, to=10, dept="B.kt")       # hiệu lực, hết trong 30 ngày
    p8 = _emp(db, world, "p8"); w(p8, T.CONCURRENT, -3, dept="B.kt")                # hiệu lực, không hạn
    p9 = _emp(db, world, "p9"); w(p9, T.CONCURRENT, -100, to=60, dept="B.kt")       # hiệu lực, hết sau 60 ngày
    p10 = _emp(db, world, "p10"); w(p10, T.CONCURRENT, -50, to=-5, dept="B.kt")     # đã hết
    p11 = _emp(db, world, "p11", status="resigned", resign_date=_d(-30)); w(p11, T.CONCURRENT, -60, dept="B.kt")
    p12 = _emp(db, world, "p12"); w(p12, T.DISMISS, -4)
    p13 = _emp(db, world, "p13", status="resigned", resign_date=_d(-3)); w(p13, T.RESIGN, -3)
    p14 = _emp(db, world, "p14"); w(p14, T.OTHER, -2)
    p15 = _emp(db, world, "p15"); w(p15, T.TRANSFER, -100, dept="B.kt")             # ngoài kỳ
    kt = _emp(db, world, "kt", department_id=world.dept["A.kt"], company_id=world.co["A"])
    w(kt, T.TRANSFER, -6, dept="A.kt", company="A")
    db.commit()


def _actor(db, world, client_as, *, scope="all", report=True, actions=("read", "export"), key="a1"):
    world.grant(key, "employee", scope, actions=actions)
    actor = world.actor(key)
    if report:
        db.add(ReportAccess(report_key=int(ReportKey.WORK_MOVES), subject_kind=SUBJECT_EMPLOYEE,
                            subject_id=actor.employee.id, effect=1, reason="t", created_by=0, updated_by=0))
        db.commit()
    return actor, client_as(actor.user)


def _get(client, path=SUM, **extra):
    dept = extra.pop("department_id", "")
    return client.get(path, params={**QS, "department_id": dept, **extra})


def test_dem_theo_loai_su_kien_va_so_nguoi(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    resp = _get(client, department_id=world.dept["B.hc"])
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    cur = data["totals"]["current"]
    assert (cur["hires"], cur["transfers"], cur["appoints"], cur["concurrents"]) == (1, 2, 3, 1)
    assert (cur["dismisses"], cur["resigns"], cur["others"]) == (1, 1, 1)
    assert cur["people_moved"] == 10                  # p1..p6, p8, p12, p13, p14
    #  p4 = 360 ngày ≈ 11,83 tháng, p5 = 726 ngày ≈ 23,85 tháng; p6 thiếu ngày vào làm nên bị bỏ
    assert cur["avg_months_to_appoint"] == pytest.approx(17.84, abs=0.06)
    assert cur["concurrent_active"] == 3 and cur["concurrent_expiring_30"] == 1
    assert data["meta"]["drilldown_metrics"] == list(svc.DRILLDOWN_METRICS)
    assert [d["key"] for d in data["meta"]["dimensions"]] == ["event_type", "department", "company", "flow"]


def test_xu_huong_theo_from_date(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    data = _get(client, department_id=world.dept["B.hc"]).json()["data"]
    assert len(data["trend"]) > 0
    total = sum((t.get("current") or t.get("values") or t)["appoints"] for t in data["trend"])
    assert total == 3


def test_chieu_loai_su_kien_phong_den_cong_ty(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    dept = world.dept["B.hc"]
    g = {x["key"]: x for x in _get(client, department_id=dept, group_by="event_type").json()["data"]["groups"]}
    assert g[str(int(T.TRANSFER))]["label"] == "Điều chuyển"
    assert g[str(int(T.TRANSFER))]["current"]["transfers"] == 2
    g = {x["key"]: x for x in _get(client, department_id=dept, group_by="department").json()["data"]["groups"]}
    assert g[str(world.dept["B.hc"])]["current"]["appoints"] == 3
    assert g[str(world.dept["B.kt"])]["current"]["transfers"] == 1            # p3 đến B.kt
    g = {x["key"]: x for x in _get(client, department_id=dept, group_by="company").json()["data"]["groups"]}
    assert "" in g                                                              # dòng không khai công ty


def test_chieu_phong_di_phong_den(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    data = _get(client, department_id=world.dept["B.hc"], group_by="flow").json()["data"]
    g = {x["label"]: x["current"] for x in data["groups"]}
    assert g["Phòng Thu mua → Phòng Hành chính"]["transfers"] == 1            # p2: dòng liền trước
    assert g["Phòng Hành chính → Phòng Kế toán"]["transfers"] == 1             # p3: chưa áp → phòng hồ sơ
    assert g["(Chưa rõ) → Phòng Hành chính"]["hires"] == 1                     # p1: tuyển dụng, không có phòng đi
    assert g["(Chưa rõ) → Phòng Hành chính"]["appoints"] == 3                  # bổ nhiệm đã áp, không dòng trước


def test_phong_di_chi_lay_dong_nhom_chinh_truoc_do(db, world, client_as):
    p = _emp(db, world, "f1")
    _line(db, p, T.TRANSFER, -50, dept="A.mua", world=world)
    _line(db, p, T.CONCURRENT, -40, dept="B.kt", world=world)                  # kiêm nhiệm KHÔNG phải phòng đi
    _line(db, p, T.TRANSFER, -5, dept="B.hc", world=world)
    db.commit()
    _, client = _actor(db, world, client_as)
    data = _get(client, department_id=world.dept["B.hc"], group_by="flow").json()["data"]
    assert [g["label"] for g in data["groups"]] == ["Phòng Thu mua → Phòng Hành chính"]


def test_phong_di_dong_chua_ap_khong_lay_phong_hien_tai_khi_co_dong_ap_muon_hon(db, world, client_as):
    p = _emp(db, world, "g1")                                                   # phòng hiện tại B.hc
    _line(db, p, T.TRANSFER, -9, dept="B.kt", applied=False, world=world)       # chưa áp, không dòng trước
    _line(db, p, T.TRANSFER, -3, dept="B.hc", applied=True, world=world)        # đã áp, muộn hơn: đổi phòng hiện tại
    q = _emp(db, world, "g2")
    _line(db, q, T.TRANSFER, -9, dept="B.kt", applied=False, world=world)       # chưa áp, không dòng nào sau
    db.commit()
    _, client = _actor(db, world, client_as)
    data = _get(client, department_id=world.dept["B.hc"], group_by="flow").json()["data"]
    g = {x["label"]: x["current"] for x in data["groups"]}
    assert g["(Chưa rõ) → Phòng Kế toán"]["transfers"] == 1                    # g1
    assert g["Phòng Hành chính → Phòng Kế toán"]["transfers"] == 1             # g2
    assert g["Phòng Kế toán → Phòng Hành chính"]["transfers"] == 1             # g1 dòng thứ hai: dòng liền trước


def test_department_va_company_chi_thu_hep(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    assert _get(client).json()["data"]["totals"]["current"]["transfers"] == 3          # + dòng của "kt"
    assert _get(client, department_id=world.dept["A.kt"]).json()["data"]["totals"]["current"]["transfers"] == 1
    r = _get(client, department_id=world.dept["B.hc"], company_id=world.co["A"])
    assert r.json()["data"]["totals"]["current"]["transfers"] == 0
    assert _get(client, department_id="abc").status_code == 422
    assert _get(client, group_by="khong_co").status_code == 422


def test_pham_vi_ap_len_nguoi_cua_dong(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as, scope="dept")        # a1 chỉ thấy phòng A.kt
    data = _get(client).json()["data"]["totals"]["current"]
    assert data["transfers"] == 1 and data["hires"] == 0 and data["people_moved"] == 1
    assert data["concurrent_active"] == 0
    body = _get(client, SUM + "/rows", metric="people_moved").json()["data"]
    assert [r["employee_code"] for r in body["rows"]] == ["kt"]


def test_rows_dieu_chuyen_va_khong_tra_tep(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    dept = world.dept["B.hc"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="transfers").json()["data"]
    assert body["total"] == 2
    assert [c["key"] for c in body["columns"]] == ["employee_code", "employee_name", "event", "from_date",
                                                  "to_date", "from_department", "department", "position",
                                                  "decision_no"]
    assert [r["employee_code"] for r in body["rows"]] == ["p3", "p2"]          # mới nhất trước
    p2 = body["rows"][1]
    assert p2["from_department"] == "Phòng Thu mua" and p2["department"] == "Phòng Hành chính"
    assert p2["decision_no"] == "QĐ-01" and p2["event"] == "Điều chuyển" and p2["link"].startswith("/hr/employees/")
    assert not {"file", "files", "file_url", "attachments"} & set(p2)
    body = _get(client, SUM + "/rows", department_id=dept, metric="people_moved").json()["data"]
    assert body["total"] == 10                                                  # mỗi người một dòng
    body = _get(client, SUM + "/rows", department_id=dept, metric="avg_months_to_appoint").json()["data"]
    assert sorted(r["employee_code"] for r in body["rows"]) == ["p4", "p5"]
    assert body["columns"][-1]["key"] == "months"
    assert all(isinstance(r["months"], float) for r in body["rows"])           # SỐ, không phải chuỗi
    assert {r["employee_code"]: r["months"] for r in body["rows"]}["p4"] == pytest.approx(11.8, abs=0.1)
    body = _get(client, SUM + "/rows", department_id=dept, metric="appoints").json()["data"]
    assert "months" not in {c["key"] for c in body["columns"]}                 # cột tháng chỉ ở danh sách tháng
    assert all("months" not in r for r in body["rows"])
    body = _get(client, SUM + "/rows", department_id=dept, metric="concurrent_active").json()["data"]
    assert sorted(r["employee_code"] for r in body["rows"]) == ["p7", "p8", "p9"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="concurrent_expiring_30").json()["data"]
    assert [r["employee_code"] for r in body["rows"]] == ["p7"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="others").json()["data"]
    assert [r["employee_code"] for r in body["rows"]] == ["p14"]


def test_rows_loc_nhom_flow(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    key = f"{world.dept['A.mua']}>{world.dept['B.hc']}"
    body = _get(client, SUM + "/rows", department_id=world.dept["B.hc"], metric="transfers",
                group_by="flow", group_key=key).json()["data"]
    assert [r["employee_code"] for r in body["rows"]] == ["p2"]


@pytest.mark.parametrize("metric", ["people_moved_x", "appoint_months_sum", "concurrent_expiring_60", ""])
def test_rows_chi_so_la_422(db, world, client_as, metric):
    _, client = _actor(db, world, client_as)
    assert _get(client, SUM + "/rows", metric=metric).status_code == 422


def test_rows_nhom_gop_422(db, world, client_as):
    _, client = _actor(db, world, client_as)
    assert _get(client, SUM + "/rows", metric="transfers", group_by="flow",
                group_key="__other__").status_code == 422


def test_gac_khoa_bao_cao_o_ca_ba_duong(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as, report=False)
    for path, extra in ((SUM, {}), (SUM + "/export", {}), (SUM + "/rows", {"metric": "transfers"})):
        assert _get(client, path, **extra).status_code == 403, path


def test_xuat_can_quyen_export(db, world, client_as):
    _, client = _actor(db, world, client_as, actions=("read",))
    assert _get(client, SUM + "/export").status_code == 403


def test_xuat_xlsx(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    resp = _get(client, SUM + "/export", department_id=world.dept["B.hc"], group_by="event_type")
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]
    ws = load_workbook(BytesIO(resp.content)).active
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Điều chuyển" in text and "Bổ nhiệm" in text


def test_so_truy_van_co_dinh(db, world):
    world.grant("a1", "employee", scope="all", actions=("read",))
    actor = world.actor("a1")
    prof = get_perm_profile(db, actor.user)
    period = parse_period({**QS, "compare": "previous"})

    def add(n, tag):
        for i in range(n):
            e = _emp(db, world, f"q{tag}-{i}")
            _line(db, e, T.TRANSFER, -60, dept="A.mua", world=world)
            _line(db, e, T.TRANSFER, -3, dept="B.hc", world=world)
            _line(db, e, T.CONCURRENT, -2, to=5, dept="B.kt", world=world)
        db.commit()

    def count():
        stmts = []
        eng = db.get_bind()

        def on(conn, cursor, statement, *a):
            stmts.append(statement)
        event.listen(eng, "before_cursor_execute", on)
        try:
            svc.build_summary(db, actor.user, prof, period, "flow")
        finally:
            event.remove(eng, "before_cursor_execute", on)
        return len(stmts)

    add(5, "a")
    c5 = count()
    add(50, "b")
    assert c5 > 0 and c5 == count()


def test_nguoi_co_bien_dong_danh_sach_mot_dong_moi_nguoi_bang_the(db, world, client_as):
    """N2: một người có 3 sự kiện trong kỳ -> thẻ đếm 1 người, danh sách đúng 1 dòng (mới nhất) + «Số lượt» = 3."""
    p = _emp(db, world, "multi")
    _line(db, p, T.APPOINT, -18, dept="B.hc", world=world)
    _line(db, p, T.CONCURRENT, -12, dept="B.hc", world=world)
    _line(db, p, T.TRANSFER, -5, dept="B.hc", world=world)
    db.commit()
    _, client = _actor(db, world, client_as)
    dept = world.dept["B.hc"]
    kpi = _get(client, department_id=dept).json()["data"]["totals"]["current"]["people_moved"]
    body = _get(client, SUM + "/rows", department_id=dept, metric="people_moved").json()["data"]
    assert body["total"] == kpi == len(body["rows"])
    assert [c["key"] for c in body["columns"]][-1] == "move_count"
    row = next(r for r in body["rows"] if r["employee_code"] == "multi")
    assert row["move_count"] == 3 and row["event"] == "Điều chuyển"
    assert sum(r["employee_code"] == "multi" for r in body["rows"]) == 1
