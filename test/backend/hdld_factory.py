"""Khung dùng chung cho test API HĐLĐ (mẫu hợp đồng + hợp đồng + sinh tệp).

- `storage`: kho tệp GIẢ trong bộ nhớ thay `core.storage` (R2/đĩa) — nghiệp vụ, phạm vi và luật tệp
  vẫn chạy THẬT (guard_upload, inspect_template, render docxtpl); chỉ phần đẩy byte lên R2 là giả.
- `client_as`: TestClient đăng nhập giả bằng một `User` của `world` (ghi đè `get_current_user`).
- `grant_all` / `upload_template` / `make_contract`: dựng dữ liệu theo hình dạng API thật.
"""
import uuid
from io import BytesIO

import pytest
from docx import Document
from fastapi import HTTPException, Request
from fastapi.testclient import TestClient

from app.core.auth import get_current_user
from app.core.database import get_db
from app.main import app

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ALL_ACTIONS = ("read", "create", "write", "delete", "print")
PDF_BYTES = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"
PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def make_docx(*paragraphs: str) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)


@pytest.fixture
def storage(monkeypatch):
    """Kho tệp giả: `key -> bytes`. Test soi `storage` để biết tệp còn / đã bị xóa."""
    files: dict[str, bytes] = {}

    def upload(fileobj, key, content_type=""):
        files[key] = fileobj.read()
        return f"/fake/{key}"

    def download(key):
        if key not in files:
            raise HTTPException(404, "File không tồn tại trên storage")
        return files[key]

    def delete(key):
        files.pop(key, None)

    monkeypatch.setattr("app.modules.attachment.service.upload_fileobj", upload)
    monkeypatch.setattr("app.modules.attachment.service.delete_key", delete)
    monkeypatch.setattr("app.modules.labor_contract.file_response.download_bytes", download)
    monkeypatch.setattr("app.modules.labor_contract.file_response.delete_key", delete)
    return files


@pytest.fixture
def client_as(db):
    """Nhiều client cùng lúc, MỖI client một người: `dependency_overrides` là của cả app nên phải
    tra người theo token Bearer của từng request (ghi đè một lần, ai gọi thì ra người đó)."""
    users: dict[str, object] = {}

    def current_user(request: Request):
        return users[request.headers["Authorization"].removeprefix("Bearer ")]

    def make_client(actor_or_user):
        token = f"test-{uuid.uuid4().hex}"
        users[token] = getattr(actor_or_user, "user", actor_or_user)
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = current_user
        return TestClient(app, headers={"Authorization": f"Bearer {token}"})

    yield make_client
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def grant_all(world, key: str, *, scope: str = "company", template: bool = True, contract: bool = True,
              actions=ALL_ACTIONS, **scope_kw):
    """Cấp cho actor `key` quyền trên 2 khóa mới (hai vai trò = hai grant độc lập)."""
    actor = world.actor(key)
    if contract:
        actor.grant("labor_contract", scope, actions=actions, **scope_kw)
    if template:
        actor.grant("labor_contract_template", scope, actions=actions, **scope_kw)
    return actor


def upload_template(client, *, company_id: int, name: str = "Mẫu thử việc", contract_type: int = 2,
                    data: bytes | None = None, filename: str = "mau.docx", note: str = ""):
    if data is None:
        data = make_docx("HĐLĐ số {{ so_hop_dong }}", "Họ tên: {{ ho_ten }}",
                         "Lương: {{ luong_co_ban }} ({{ luong_co_ban_bang_chu }})")
    return client.post("/api/labor-contract-templates",
                       files={"file": (filename, data, DOCX_MIME)},
                       data={"name": name, "company_id": str(company_id),
                             "contract_type": str(contract_type), "note": note})


def contract_payload(**over) -> dict:
    body = {"contract_type": 2, "start_date": "2026-01-01", "end_date": "2026-12-31",
            "base_salary": 15_000_000, "insurance_salary": 10_000_000, "allowance": 500_000,
            "allowance_note": "Xăng xe", "contract_no": ""}
    body.update(over)
    return body


def make_contract(client, employee_id: int, **over) -> dict:
    res = client.post(f"/api/employees/{employee_id}/labor-contracts", json=contract_payload(**over))
    assert res.status_code == 201, res.text
    return res.json()["data"]["item"]
