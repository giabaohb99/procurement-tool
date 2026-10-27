"""duoc-CR-687 — Báo cáo Biến động nhân sự mở rộng «Cơ cấu» (A2): chiều Nhóm tuổi · Trình độ ·
Địa điểm làm việc · Nhóm vị trí việc làm, che nhóm tuổi nhỏ, chặn danh sách theo nhóm tuổi."""
import uuid
from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event

from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.report_period import parse_period
from app.core.subject_match import SUBJECT_ROLE
from app.main import app
from app.modules.employee import report_service
from app.modules.employee.constants import (EDUCATION_COLLEGE, EDUCATION_POSTGRADUATE,
                                            EDUCATION_UNIVERSITY, JOB_CATEGORY_HIGH_SKILL,
                                            JOB_CATEGORY_MANAGER, JOB_CATEGORY_MID_SKILL,
                                            JOB_LEVEL_DEPUTY_MANAGER)
from app.modules.employee.model import Employee
from app.modules.employee.report_headcount_dims import age_band

SEP = {"preset": "custom", "date_from": "2026-09-01", "date_to": "2026-09-30", "compare": "none"}
ROWS = "/api/employees/summary/rows"


def _emp(db, world, code, dept="A.kt", **kw):
    kw.setdefault("hire_date", date(2020, 1, 1))
    e = Employee(code=code, full_name=f"NV {code}", company_id=world.co["A"],
                 department_id=world.dept[dept], **kw)
    db.add(e)
    db.flush()
    return e


def _groups(db, actor, group_by, **kw):
    data = report_service.build_summary(db, actor.user, actor.profile(), parse_period(SEP), group_by, **kw)
    return data, {g["key"]: g for g in data["groups"]}


@pytest.fixture
def http(db):
    def call(user, path, **params):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        try:
            client = TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})
            return client.get(path, params=params)
        finally:
            app.dependency_overrides.pop(get_db, None)
            app.dependency_overrides.pop(get_current_user, None)
    return call


@pytest.fixture
def hr(db, world):
    """Người xem có quyền nhạy cảm."""
    world.grant("a1", "employee", scope="all", actions=("read",))
    world.grant("a1", "employee_sensitive", scope="all", actions=("read",))
    return world.actor("a1")


@pytest.fixture
def plain(db, world):
    """Người xem KHÔNG có quyền nhạy cảm."""
    world.grant("a1", "employee", scope="all", actions=("read",))
    return world.actor("a1")


class TestNhomTuoi:
    @pytest.mark.parametrize("dob,expected", [
        (date(2001, 9, 30), "25-34"),    # đúng 25 tuổi ngày 30/09/2026
        (date(2001, 10, 1), "lt25"),     # mai mới đủ 25
        (date(1991, 9, 30), "35-44"),    # đúng 35
        (date(1991, 10, 1), "25-34"),    # 34
        (date(1971, 9, 30), "55+"),      # đúng 55
        (date(1971, 10, 1), "45-54"),    # 54
        (None, ""),                      # thiếu ngày sinh
        (date(2030, 1, 1), ""),          # ngày sinh sau mốc: dữ liệu sai → chưa gắn
    ])
    def test_ranh_gioi(self, dob, expected):
        assert age_band(dob, date(2026, 9, 30))[0] == expected

    def test_tinh_tai_ngay_mot_khong_phai_hom_nay(self):
        dob = date(2001, 1, 1)
        assert age_band(dob, date(2026, 1, 1))[0] == "25-34"
        assert age_band(dob, date(2025, 12, 31))[0] == "lt25"

    def test_gom_nhom_dinh_bien_va_su_kien(self, db, world, hr):
        _emp(db, world, "T1", date_of_birth=date(2003, 1, 1))
        _emp(db, world, "T2", date_of_birth=date(2002, 5, 5))
        _emp(db, world, "T3", date_of_birth=date(2001, 1, 1))
        _emp(db, world, "T4", date_of_birth=date(1990, 1, 1))
        _emp(db, world, "T5", date_of_birth=date(1990, 6, 1), hire_date=date(2026, 9, 10))
        _emp(db, world, "T6")                                  # thiếu ngày sinh
        db.commit()
        data, g = _groups(db, hr, "age_band")
        assert g["lt25"]["current"]["headcount_end"] == 2
        assert g["25-34"]["current"]["headcount_end"] == 1
        assert g["35-44"]["current"]["headcount_end"] == 2 and g["35-44"]["current"]["new_hires"] == 1
        assert g[""]["label"] == "(Chưa gắn)" and g[""]["current"]["headcount_end"] == 1
        assert data["totals"]["current"]["headcount_end"] == 6

    def test_tuoi_tai_ngay_su_kien_nghi_viec(self, db, world, hr):
        #  Sinh 01/10/2001: lúc nghỉ 20/09/2026 mới 24 tuổi → «Dưới 25», dù nay đã 25.
        _emp(db, world, "R1", date_of_birth=date(2001, 10, 1), status="resigned",
             resign_date=date(2026, 9, 20))
        db.commit()
        _, g = _groups(db, hr, "age_band")
        assert g["lt25"]["current"]["departures"] == 1


class TestChieuKhac:
    def test_trinh_do_dia_diem_nhom_vi_tri(self, db, world, hr):
        _emp(db, world, "E1", education_level=EDUCATION_UNIVERSITY, work_location="Hà Nội")
        _emp(db, world, "E2", education_level=EDUCATION_POSTGRADUATE, work_location="  hà nội ")
        _emp(db, world, "E3", education_level=EDUCATION_COLLEGE, work_location="HÀ  NỘI")
        _emp(db, world, "E4", education_level=EDUCATION_UNIVERSITY, job_level=JOB_LEVEL_DEPUTY_MANAGER,
             work_location="Đà Nẵng")
        _emp(db, world, "E5")
        db.commit()
        _, edu = _groups(db, hr, "education")
        assert edu[str(EDUCATION_UNIVERSITY)]["current"]["headcount_end"] == 2
        assert edu[str(EDUCATION_UNIVERSITY)]["label"] == "Đại học"
        assert edu[""]["current"]["headcount_end"] == 1

        _, loc = _groups(db, hr, "work_location")
        assert loc["hà nội"]["current"]["headcount_end"] == 3
        assert loc["hà nội"]["label"] in ("Hà Nội", "hà nội", "HÀ NỘI")
        assert loc["đà nẵng"]["current"]["headcount_end"] == 1
        assert loc[""]["current"]["headcount_end"] == 1

        _, cat = _groups(db, hr, "job_category")
        assert cat[str(JOB_CATEGORY_MANAGER)]["current"]["headcount_end"] == 1
        assert cat[str(JOB_CATEGORY_HIGH_SKILL)]["current"]["headcount_end"] == 2  # E1, E2
        assert cat[str(JOB_CATEGORY_MID_SKILL)]["current"]["headcount_end"] == 1   # E3

    def test_su_kien_vao_moi_theo_trinh_do(self, db, world, hr):
        _emp(db, world, "N1", education_level=EDUCATION_UNIVERSITY, hire_date=date(2026, 9, 3))
        db.commit()
        _, g = _groups(db, hr, "education")
        assert g[str(EDUCATION_UNIVERSITY)]["current"]["new_hires"] == 1


def _band(db, world, prefix, n, dob, **kw):
    for i in range(n):
        _emp(db, world, f"{prefix}{i}", date_of_birth=dob, **kw)


def _visible(g) -> bool:
    return not g.get("suppressed")


class TestCheNhomNho:
    def _seed(self, db, world):
        _band(db, world, "BIG", 5, date(1990, 1, 1))     # 35–44: 5 người
        _band(db, world, "MID", 4, date(1996, 1, 1))     # 25–34: 4 người
        _band(db, world, "TRI", 3, date(1978, 1, 1))     # 45–54: 3 người (nhỏ nhất còn hiện)
        _emp(db, world, "S1", date_of_birth=date(2003, 1, 1))     # dưới 25: 1 người
        _emp(db, world, "S2", date_of_birth=date(2002, 1, 1))     # dưới 25: 2 người
        db.commit()

    def test_thieu_quyen_che_nhom_duoi_3_va_nhom_ke_tiep(self, db, world, plain):
        self._seed(db, world)
        data, g = _groups(db, plain, "age_band")
        assert all(v is None for v in g["lt25"]["current"].values())
        assert g["lt25"]["current"].keys() >= {"headcount_end", "new_hires"}
        #  Chỉ MỘT nhóm bị che thì Tổng − Σ nhóm hiện lộ đúng nhóm đó → che thêm nhóm hiện nhỏ nhất (45–54).
        assert all(v is None for v in g["45-54"]["current"].values())
        assert g["35-44"]["current"]["headcount_end"] == 5 and g["25-34"]["current"]["headcount_end"] == 4
        assert "Nhóm tuổi dưới 3 người được che số để không đoán ra tuổi từng người" in data["notes"]
        assert data["totals"]["current"]["headcount_end"] == 14   # tổng giữ nguyên

    def test_che_ca_ky_so_sanh(self, db, world, plain):
        self._seed(db, world)
        period = parse_period({**SEP, "compare": "previous"})
        data = report_service.build_summary(db, plain.user, plain.profile(), period, "age_band")
        g = {x["key"]: x for x in data["groups"]}
        assert g["lt25"]["compare"] is not None and all(v is None for v in g["lt25"]["compare"].values())

    def test_co_quyen_thay_so_that(self, db, world, hr):
        self._seed(db, world)
        data, g = _groups(db, hr, "age_band")
        assert g["lt25"]["current"]["headcount_end"] == 2 and g["45-54"]["current"]["headcount_end"] == 3
        assert not any("được che số" in n for n in data["notes"]) and all(_visible(x) for x in g.values())

    def test_chieu_khac_khong_che(self, db, world, plain):
        _emp(db, world, "X1", education_level=EDUCATION_UNIVERSITY)
        db.commit()
        data, g = _groups(db, plain, "education")
        assert g[str(EDUCATION_UNIVERSITY)]["current"]["headcount_end"] == 1
        assert not any("được che số" in n for n in data["notes"])

    def test_tong_tru_nhom_hien_khong_lo_nhom_61_tuoi(self, db, world, plain):
        _band(db, world, "Y", 4, date(1996, 1, 1))       # 30 tuổi
        _emp(db, world, "OLD", date_of_birth=date(1965, 1, 1))   # 61 tuổi
        db.commit()
        dept = world.dept["A.kt"]
        data, g = _groups(db, plain, "age_band", department_id=dept)
        total = data["totals"]["current"]["headcount_end"]
        shown = sum(x["current"]["headcount_end"] or 0 for x in g.values())
        assert total == 5 and total - shown >= 3 and total - shown != 1     # không tách riêng được người 61 tuổi
        assert g["55+"]["suppressed"] is True and g["25-34"]["suppressed"] is True

    def test_co_quyen_nhay_cam_thay_het_ca_nhom_mot_nguoi(self, db, world, hr):
        _band(db, world, "Y", 4, date(1996, 1, 1))
        _emp(db, world, "OLD", date_of_birth=date(1965, 1, 1))
        db.commit()
        _, full = _groups(db, hr, "age_band", department_id=world.dept["A.kt"])
        assert full["55+"]["current"]["headcount_end"] == 1 and full["25-34"]["current"]["headcount_end"] == 4
        assert not any(x.get("suppressed") for x in full.values())

    def test_dau_ky_hoac_vao_moi_nho_cung_che(self, db, world, plain):
        #  25–34: cuối kỳ 3 nhưng đầu kỳ 1 và vào mới 2 → mỗi con số đó đã đủ chỉ ra từng người.
        _emp(db, world, "O1", date_of_birth=date(1996, 1, 1))
        _band(db, world, "N", 2, date(1996, 2, 2), hire_date=date(2026, 9, 10))
        _band(db, world, "BIG", 6, date(1990, 1, 1))
        db.commit()
        _, g = _groups(db, plain, "age_band")
        cur = g["25-34"]["current"]
        assert cur["headcount_end"] is None and cur["headcount_start"] is None and cur["new_hires"] is None

    def test_ky_so_sanh_co_so_dem_1_cung_che(self, db, world, plain):
        #  35–44: kỳ này đủ lớn (5/5/0/0) nhưng kỳ trước (tháng 8) có đúng 1 người vào mới.
        _band(db, world, "OLD", 4, date(1990, 1, 1))
        _emp(db, world, "AUG", date_of_birth=date(1990, 3, 3), hire_date=date(2026, 8, 15))
        _band(db, world, "MID", 6, date(1996, 1, 1))
        db.commit()
        period = parse_period({**SEP, "compare": "previous"})
        data = report_service.build_summary(db, plain.user, plain.profile(), period, "age_band")
        g = {x["key"]: x for x in data["groups"]}
        assert g["35-44"]["current"]["headcount_end"] is None and g["35-44"]["suppressed"] is True

    def test_xuat_xlsx_nhom_bi_che_de_trong_khong_phai_0(self, db, world, plain, http, gan_bao_cao):
        gan_bao_cao(SUBJECT_ROLE, plain.roles[-1].id, ReportKey.HR_HEADCOUNT)
        world.grant("a1", "employee", scope="all", actions=("export",))
        _emp(db, world, "L1", date_of_birth=date(2003, 1, 1))     # dưới 25: 1 người
        _band(db, world, "TRI", 3, date(1990, 1, 1))              # 35–44: 3 người, bị che thêm
        _band(db, world, "MID", 4, date(1996, 1, 1))              # 25–34: 4 người, hiện
        db.commit()
        r = http(plain.user, "/api/employees/summary/export", group_by="age_band", **SEP)
        assert r.status_code == 200, r.text[:200]
        ws = load_workbook(BytesIO(r.content)).active
        rows = {row[0].value: [c.value for c in row[1:]] for row in ws.iter_rows(min_row=2)}
        assert all(v is None for v in rows["Dưới 25"]) and all(v is None for v in rows["35–44"])
        assert "4" in rows["25–34"] and "8" in rows["Tổng"]     # định biên ghi dạng chữ (kind text)


class TestDanhSachTheoNhomTuoi:
    @pytest.fixture(autouse=True)
    def _bao_cao(self, gan_bao_cao):
        self.gan = gan_bao_cao

    def test_thieu_quyen_nhay_cam_403(self, db, world, plain, http):
        self.gan(SUBJECT_ROLE, plain.roles[-1].id, ReportKey.HR_HEADCOUNT)
        _emp(db, world, "Z1", date_of_birth=date(2003, 1, 1))
        db.commit()
        r = http(plain.user, ROWS, metric="headcount_end", group_by="age_band", group_key="lt25", **SEP)
        assert r.status_code == 403
        r = http(plain.user, ROWS, metric="headcount_end", group_by="age_band", group_key="", **SEP)
        assert r.status_code == 403

    def test_khong_nhom_hoac_chieu_khac_van_xem_duoc(self, db, world, plain, http):
        self.gan(SUBJECT_ROLE, plain.roles[-1].id, ReportKey.HR_HEADCOUNT)
        _emp(db, world, "Z1", date_of_birth=date(2003, 1, 1), education_level=EDUCATION_UNIVERSITY)
        db.commit()
        assert http(plain.user, ROWS, metric="headcount_end", **SEP).status_code == 200
        r = http(plain.user, ROWS, metric="headcount_end", group_by="education",
                 group_key=str(EDUCATION_UNIVERSITY), **SEP)
        assert r.status_code == 200 and r.json()["data"]["total"] == 1

    def test_co_quyen_loc_dung_nhom_va_khong_lo_ngay_sinh(self, db, world, hr, http):
        self.gan(SUBJECT_ROLE, hr.roles[0].id, ReportKey.HR_HEADCOUNT)
        _emp(db, world, "Z1", date_of_birth=date(2003, 1, 1))
        _emp(db, world, "Z2", date_of_birth=date(1980, 1, 1))
        db.commit()
        r = http(hr.user, ROWS, metric="headcount_end", group_by="age_band", group_key="lt25", **SEP)
        d = r.json()["data"]
        assert r.status_code == 200 and d["total"] == 1 and d["rows"][0]["code"] == "Z1"
        assert "date_of_birth" not in str(d) and "2003" not in str(d["rows"])


class TestSoTruyVan:
    @pytest.mark.parametrize("group_by", ["age_band", "education", "work_location", "job_category"])
    def test_khong_tang_theo_so_nhan_su(self, db, world, plain, group_by):
        prof = plain.profile()
        period = parse_period(SEP)

        def add(n, start):
            for i in range(n):
                _emp(db, world, f"BULK{start + i}", date_of_birth=date(1990, 1, 1),
                     hire_date=date(2026, 9, 5), work_location=f"Kho {i % 4}")
            db.commit()

        def count():
            stmts = []
            eng = db.get_bind()

            def on(conn, cursor, statement, *a):
                stmts.append(statement)
            event.listen(eng, "before_cursor_execute", on)
            try:
                report_service.build_summary(db, plain.user, prof, period, group_by)
            finally:
                event.remove(eng, "before_cursor_execute", on)
            return len(stmts)

        add(5, 0)
        c5 = count()
        add(45, 5)
        assert c5 > 0 and c5 == count()


class TestXuatTep:
    def test_xuat_xlsx_nhom_tuoi_da_che_khong_loi(self, db, world, plain, http, gan_bao_cao):
        gan_bao_cao(SUBJECT_ROLE, plain.roles[-1].id, ReportKey.HR_HEADCOUNT)
        world.grant("a1", "employee", scope="all", actions=("export",))
        _emp(db, world, "XL1", date_of_birth=date(2003, 1, 1))
        db.commit()
        r = http(plain.user, "/api/employees/summary/export", group_by="age_band", **SEP)
        assert r.status_code == 200, r.text[:200]
