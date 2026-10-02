"""Xuất Excel mục «Thuốc BVTV» — round-trip với nút «Nạp danh mục» (02/10/2026).

Ba điều canh chặt nhất:
  · `scope=all` xuất RA SAO thì `pesticide_reader` phải NẠP LẠI ĐƯỢC đúng như vậy — so trực tiếp
    với `read_json` trên cùng dữ liệu gốc, không chỉ "nạp không lỗi";
  · thuốc TỰ THÊM (`is_manual`) phải vẫn thấy trong tệp (đối chiếu/lưu) nhưng KHÔNG được bộ đọc
    coi là thuốc nguồn — nạp lại không nhân đôi, dù hai thuốc tự thêm cùng `source_id = 0`;
  · `scope=page` phải khớp Y HỆT thứ tự + nội dung của `list_pesticides` với cùng bộ lọc/trang.

Thêm (review 02/10/2026, sửa C1/C2/C3/H1/H2/M1/M2/M3/L2/L6): `_xuat` sheet ẩn + chế độ nạp
CẬP NHẬT theo trang, id ÂM cho thuốc tự thêm/trùng `source_id`, ô chuỗi an toàn (không phải
công thức/ký tự cấm XML), khóa xuất liên tiến trình, dọn tệp tạm khi lỗi giữa chừng, trần kích
thước tệp xuất, nhãn nhật ký đọc được, giờ VN dùng chung.
"""
import io
import json
import os
import uuid
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.auth import get_current_user
from app.core.database import get_db
from app.main import app
from app.modules.attachment.model import FileLink, StoredFile
from app.modules.customs import pesticide_edit_service as E
from app.modules.customs import pesticide_export_columns as C
from app.modules.customs import pesticide_export_lock as lock
from app.modules.customs import pesticide_export_marker as marker
from app.modules.customs import pesticide_export_service as X
from app.modules.customs import pesticide_merge_service as M
from app.modules.customs import pesticide_reader as R
from app.modules.customs import pesticide_service as S
from app.modules.customs.constants import PesticideStatus
from app.modules.customs.model import CustomsPesticide, CustomsPesticideUse
from app.modules.customs.pesticide_schema import PesticideIn
from app.modules.export_log.model import ExportLog
from app.modules.user.model import User

USER = SimpleNamespace(id=1)


def _raw(id_: int, ten_thuoc: str, **over) -> dict:
    raw = {
        "id": id_, "ten_thuoc": ten_thuoc, "phan_nhom": "Thuốc trừ bệnh",
        "linh_vuc": "THUỐC SỬ DỤNG TRONG NÔNG NGHIỆP", "tinh_trang": "Còn hiệu lực",
        "hoat_chat": "Chitosan 2% + Polyoxin B 10%", "ham_luong": "12% w/w",
        "cong_ty_dang_ky": "Công ty TNHH Ngân Anh", "so_dang_ky": f"{id_}/CNĐKT-BVTV",
        "ngay_cap": "2023-09-25", "ngay_het_han": "2028-09-25",
        "url": "https://danhmuc.thuocbvtv.com/thuoc/detail/x",
        "tom_tat_su_dung": "Thuốc trừ bệnh dùng trên lúa, phòng đạo ôn.",
        "nhom_doc": [{"he": "GHS", "nhom": "5", "mo_ta": "Rất ít độc"}],
        #  Nhiều hoạt chất CÓ đủ mã/nhóm/phương thức — round-trip nguyên vẹn (xem
        #  `pesticide_export_columns`/`pesticide_reader._resistance`, bản rỗng mới bị bộ đọc bỏ).
        "quan_ly_tinh_khang": {"hoat_chat": [
            {"ten": "Chitosan", "ma": "FRAC 19", "nhom": "Peptidyl", "phuong_thuc": "Tổng hợp chitin"},
            {"ten": "Polyoxin B", "ma": "FRAC 19", "nhom": "Peptidyl", "phuong_thuc": "Tổng hợp chitin"},
        ]},
        "pham_vi_su_dung": [
            {"cay_trong": "lúa", "dich_hai": "đạo ôn", "lieu_luong": "0.4 lít/ha",
             "thoi_gian_cach_ly": "7 ngày", "cach_dung": "Phun khi bệnh chớm"},
            {"cay_trong": "ngô", "dich_hai": "sâu đục thân", "lieu_luong": "0.5 lít/ha",
             "thoi_gian_cach_ly": "10 ngày", "cach_dung": "Phun định kỳ"},
        ],
    }
    raw.update(over)
    return raw


def _load_json(db, *raws, filename="seed.json") -> list[dict]:
    content = json.dumps(list(raws), ensure_ascii=False).encode()
    S.replace_catalog(db, R.read_json(content), user_id=1, filename=filename)
    return list(raws)


def _manual_body(name: str, **over) -> PesticideIn:
    body = {"trade_name": name, "active_ingredient": "Abamectin 18g/l",
            "status": int(PesticideStatus.ACTIVE),
            "uses": [{"crop": "lúa", "pest": "sâu cuốn lá", "dosage": "0.3 lít/ha"}]}
    body.update(over)
    return PesticideIn(**body)


def _export(db, scope="all", q="", status=None, pest_group="", sector="", banned_only=False,
           banned_regulation_id=None, page=1, offset=0, limit=50) -> tuple[bytes, str]:
    """Gọi service, đọc tệp tạm ra bytes rồi TỰ XÓA — không để lại tệp rác giữa các bài test."""
    path, filename = X.export_to_tempfile(db, USER, scope, q, status, pest_group, sector,
                                          banned_only, banned_regulation_id, page, offset, limit)
    try:
        with open(path, "rb") as f:
            return f.read(), filename
    finally:
        os.remove(path)


# ── Round-trip toàn bộ danh mục ──────────────────────────────────────────────────────────

def test_export_all_then_reimport_is_a_pure_round_trip(db):
    #  Ba thuốc TỪ NGUỒN: một bình thường, một hết hiệu lực, một có ngày rỗng + tên có HTML
    #  entity (bản cào hay trả `&amp;` cho dấu &).
    raws = _load_json(
        db,
        _raw(101, "Alpha 10SC"),
        _raw(102, "Beta 5WP", tinh_trang="Hết hiệu lực",
             pham_vi_su_dung=[{"cay_trong": "cà phê", "dich_hai": "rệp sáp"}]),
        _raw(103, "Gamma &amp; Delta 20EC", ngay_cap=None, ngay_het_han=None,
             pham_vi_su_dung=[{"cay_trong": "tiêu", "dich_hai": "tuyến trùng"}]),
    )
    #  Hai thuốc TỰ THÊM — CÙNG `source_id = 0` (luôn vậy, `pesticide_edit_service`): đây đúng là
    #  ca "nhiều thuốc tự thêm cùng id 0" mà bài này phải canh không bị gộp nhầm phạm vi.
    m1 = E.create_pesticide(db, _manual_body("Thủ Công Một"), user_id=1)
    m2 = E.create_pesticide(db, _manual_body("Thủ Công Hai",
                            uses=[{"crop": "cà phê", "pest": "rệp"}]), user_id=1)

    content, filename = _export(db, scope="all")
    assert filename.startswith("thuoc-bvtv-toan-bo-") and filename.endswith(".xlsx")

    records = R.read_file("x.xlsx", content)
    #  Thuốc tự thêm xuất hiện trong TỆP (đối chiếu được) nhưng bộ đọc phải BỎ QUA lúc nạp lại.
    assert {"Thủ Công Một", "Thủ Công Hai"}.isdisjoint({r["trade_name"] for r in records})

    expected = R.read_json(json.dumps(raws, ensure_ascii=False).encode())
    assert (sorted(records, key=lambda r: r["source_id"])
            == sorted(expected, key=lambda r: r["source_id"]))

    #  Nạp lại: danh mục không đổi — 3 thuốc nguồn giữ nguyên, 2 thuốc tự thêm GIỮ, KHÔNG nhân đôi.
    result = S.replace_catalog(db, records, user_id=1, filename=filename)
    assert result["kept_manual"] == 2 and result["dropped"] == 0
    assert db.query(CustomsPesticide).count() == 5
    manual = {p.trade_name: len(p.trade_name) for p in
             db.query(CustomsPesticide).filter_by(is_manual=True)}
    assert set(manual) == {"Thủ Công Một", "Thủ Công Hai"}
    #  id của hai thuốc tự thêm không bị đụng/cấp lại qua lần nạp.
    assert {p.id for p in db.query(CustomsPesticide).filter_by(is_manual=True)} == {m1["id"], m2["id"]}

    #  Nạp lần NỮA (cùng tệp) — vẫn không nhân đôi gì thêm.
    S.replace_catalog(db, records, user_id=1, filename=filename)
    assert db.query(CustomsPesticide).count() == 5


def test_export_logs_one_row_per_call(db):
    _load_json(db, _raw(201, "Zeta 1SC"))
    _export(db, scope="all")
    log = db.query(ExportLog).filter_by(entity="customs_price").one()
    assert log.fmt == "xlsx" and log.row_count == 1 and log.file_id == 0
    assert log.filename.startswith("thuoc-bvtv-toan-bo-")


# ── scope=page khớp đúng list cùng tham số + trang ──────────────────────────────────────

def test_export_page_matches_list_filter_sort_and_page(db):
    #  `replace_catalog` thay TOÀN BỘ thuốc từ nguồn — phải nạp chung MỘT lượt, nạp hai lượt sẽ
    #  xóa mất lượt trước (đúng như hành vi thật của «Nạp danh mục»).
    omega = [_raw(300 + i, name) for i, name in
            enumerate(["Omega 1", "Omega 2", "Omega 3", "Omega 4", "Omega 5"])]
    _load_json(db, *omega, _raw(999, "Khác hẳn"))   # "Khác hẳn" không khớp bộ lọc "Omega"

    total, expected_items = S.list_pesticides(db, "Omega", None, "", "", offset=2, limit=2)
    assert total == 5
    expected_names = [it["trade_name"] for it in expected_items]
    assert expected_names == ["Omega 3", "Omega 4"]   # trang 2, cỡ trang 2, sắp theo tên

    content, filename = _export(db, scope="page", q="Omega", page=2, offset=2, limit=2)
    assert filename.startswith("thuoc-bvtv-trang-2-")
    got_names = [r["trade_name"] for r in R.read_file("x.xlsx", content)]
    assert got_names == expected_names


# ── Quyền + trần + chặn bấm dồn ──────────────────────────────────────────────────────────

@pytest.fixture
def client_as(db):
    """Token Bearer GIẢ nhưng DUY NHẤT mỗi lần build — tránh đụng khóa cache của middleware
    báo cáo khác (cùng mẹo với `test_bao_cao_van_ban.py`)."""
    def build(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})
    yield build
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def test_khong_co_quyen_export_thi_403(db, seed, cap_quyen, client_as):
    v = db.query(User).filter(User.employee_id == seed.emp_tp_id).one()
    cap_quyen(v.id, "customs_price", scope="all", read=True, export=False)
    assert client_as(v).get("/api/customs/pesticides/export").status_code == 403


def test_co_quyen_export_tra_ve_file_xlsx_qua_http(db, seed, cap_quyen, client_as):
    _load_json(db, _raw(401, "Http 1SC"))
    v = db.query(User).filter(User.employee_id == seed.emp_tp_id).one()
    cap_quyen(v.id, "customs_price", scope="all", read=True, export=True)
    resp = client_as(v).get("/api/customs/pesticides/export?scope=all")
    assert resp.status_code == 200
    assert "spreadsheet" in resp.headers["content-type"]
    assert 'filename="thuoc-bvtv-toan-bo-' in resp.headers["content-disposition"]
    assert db.query(ExportLog).filter_by(entity="customs_price").count() == 1


def test_vuot_tran_so_thuoc_la_422(db, monkeypatch):
    _load_json(db, _raw(501, "Một"), _raw(502, "Hai"), _raw(503, "Ba"))
    monkeypatch.setattr(R, "MAX_RECORDS", 2)
    with pytest.raises(HTTPException) as exc:
        X.export_to_tempfile(db, USER, "all", "", None, "", "", False, None, 1, 0, 20)
    assert exc.value.status_code == 422 and "vượt trần" in exc.value.detail


def test_vuot_tran_so_dong_pham_vi_la_422(db, monkeypatch):
    _load_json(db, _raw(601, "Một"))
    monkeypatch.setattr(R, "MAX_USE_ROWS", 1)
    with pytest.raises(HTTPException) as exc:
        X.export_to_tempfile(db, USER, "all", "", None, "", "", False, None, 1, 0, 20)
    assert exc.value.status_code == 422


def test_bam_don_xuat_toan_bo_la_429(db):
    """Lượt thứ hai CỦA CHÍNH người đó trong lúc lượt đầu còn chạy phải bị chặn."""
    lock.EXPORTING_ALL.add(USER.id)
    try:
        with pytest.raises(HTTPException) as exc:
            X.export_to_tempfile(db, USER, "all", "", None, "", "", False, None, 1, 0, 20)
        assert exc.value.status_code == 429
    finally:
        lock.EXPORTING_ALL.discard(USER.id)
    #  Khóa đã được dọn — lượt kế tiếp (người khác, hoặc chính người đó sau khi xong) chạy bình thường.
    assert USER.id not in lock.EXPORTING_ALL


# ── M1 — khóa chặn bấm dồn đứng vững qua NHIỀU tiến trình uvicorn ──────────────────────────

class _FakeMySQLServer:
    """Mô phỏng GET_LOCK/RELEASE_LOCK CÓ TRẠNG THÁI THẬT (không như MagicMock vô tri) — để hai
    `_FakeDb` CÙNG `server` (hai "tiến trình" khác nhau) thấy đúng khóa của nhau."""

    def __init__(self):
        self.locks: dict[str, bool] = {}

    def get_lock(self, name: str) -> int:
        if self.locks.get(name):
            return 0
        self.locks[name] = True
        return 1

    def release_lock(self, name: str) -> int:
        self.locks.pop(name, None)
        return 1


class _FakeConn:
    def __init__(self, server: _FakeMySQLServer):
        self.server = server

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def execute(self, stmt, params):
        sql, name = str(stmt), params["n"]
        value = (self.server.get_lock(name) if "GET_LOCK" in sql
                 else self.server.release_lock(name))
        return SimpleNamespace(scalar=lambda: value)


class _FakeDb:
    """`db` giả CHỈ đủ cho `_export_lock` (`.get_bind().dialect.name` + `.connect()`) — không
    phải ORM Session thật, vì bài này kiểm LOGIC khóa, không kiểm truy vấn."""

    def __init__(self, server: _FakeMySQLServer):
        self._bind = SimpleNamespace(dialect=SimpleNamespace(name="mysql"),
                                     connect=lambda: _FakeConn(server))

    def get_bind(self):
        return self._bind


def test_khoa_xuat_dung_vung_qua_nhieu_tien_trinh_uvicorn():
    """Tái hiện lỗ hổng M1: set `_EXPORTING_ALL` CHỈ sống trong MỘT tiến trình — hai yêu cầu của
    CÙNG một người rơi vào hai tiến trình uvicorn khác nhau (`--workers 2`) sẽ KHÔNG thấy nhau
    nếu chỉ có lớp đó. Giả lập "tiến trình khác" bằng cách XÓA người đó khỏi `_EXPORTING_ALL`
    ngay trong lúc khóa MySQL (dùng chung `server`) vẫn đang giữ — chỉ khóa MySQL còn chặn được."""
    server = _FakeMySQLServer()
    db1, db2 = _FakeDb(server), _FakeDb(server)
    uid = 8888
    with lock.export_lock(db1, uid):
        lock.EXPORTING_ALL.discard(uid)   # mô phỏng "tiến trình 2" — set trong-tiến-trình rỗng
        with pytest.raises(HTTPException) as exc:
            with lock.export_lock(db2, uid):
                pass
        assert exc.value.status_code == 429
    #  Lượt 1 đã thoát — khóa MySQL được nhả — lượt MỚI của chính người đó chạy được.
    with lock.export_lock(db1, uid):
        pass
    assert uid not in lock.EXPORTING_ALL


def test_khoa_mysql_duoc_nha_trong_finally_du_than_lenh_nem_loi():
    server = _FakeMySQLServer()
    db = _FakeDb(server)
    uid = 7777
    with pytest.raises(RuntimeError):
        with lock.export_lock(db, uid):
            raise RuntimeError("lỗi giữa lúc đang xuất")
    #  Khóa MySQL đã được nhả (không set) — lượt sau chạy được ngay, không kẹt vĩnh viễn.
    assert not server.locks
    with lock.export_lock(db, uid):
        pass


# ── C1 — tệp xuất THEO TRANG nạp lại là CẬP NHẬT, không xóa phần còn lại ───────────────────

def test_xuat_theo_trang_ghi_sheet_an_dung_scope_va_so_trang(db):
    _load_json(db, _raw(1001, "Trang Một"), _raw(1002, "Trang Hai"))
    content, _ = _export(db, scope="page", page=3, offset=0, limit=20)
    wb = load_workbook(io.BytesIO(content))
    assert marker.MARKER_SHEET in wb.sheetnames
    ws = wb[marker.MARKER_SHEET]
    assert ws.sheet_state == "hidden"
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("scope", "page", "exported_at")
    assert rows[1][0] == "page" and rows[1][1] == 3


def test_che_do_nap_theo_tung_loai_tep(db):
    """`scope=page` → merge; `scope=all` → replace; JSON (bản cào gốc) → replace."""
    _load_json(db, _raw(1010, "Chế độ một"), _raw(1011, "Chế độ hai"))
    all_content, all_name = _export(db, scope="all")
    page_content, page_name = _export(db, scope="page", page=1, offset=0, limit=20)
    json_content = json.dumps([_raw(1012, "JSON gốc")], ensure_ascii=False).encode()

    _, mode_all = R.read_file_with_mode(all_name, all_content)
    _, mode_page = R.read_file_with_mode(page_name, page_content)
    _, mode_json = R.read_file_with_mode("thuoc-bvtv.json", json_content)
    assert mode_all == "replace" and mode_page == "merge" and mode_json == "replace"


def test_tep_xlsx_khong_co_sheet_xuat_la_ban_cao_goc_thi_thay_toan_bo():
    """Bản cào THẬT (danhmuc.thuocbvtv.com) không biết khái niệm `_xuat` — phải vẫn THAY
    TOÀN BỘ như hành vi gốc, không bị coi nhầm là tệp theo trang."""
    wb = Workbook()
    ws = wb.active
    ws.title = R.LIST_SHEET
    ws.append(["id", "ten_thuoc"])
    ws.append([5000, "Bản cào gốc"])
    uses = wb.create_sheet(R.USE_SHEET)
    uses.append(["id", "stt_pham_vi", "cay_trong"])
    uses.append([5000, 1, "lúa"])
    buf = io.BytesIO()
    wb.save(buf)
    _, mode = R.read_file_with_mode("thuoc-bvtv.xlsx", buf.getvalue())
    assert mode == "replace"


def test_nap_tep_xuat_theo_trang_da_sua_chi_cap_nhat_dung_cac_thuoc_do(db):
    """Ca chính của C1: nạp tệp theo trang đã sửa vài ô → CHỈ đúng các thuốc trong tệp đổi,
    tổng số thuốc KHÔNG đổi, thuốc KHÁC (ngoài trang) và tệp đính kèm của nó còn nguyên."""
    raws = [_raw(1100 + i, f"Trang {i}") for i in range(5)]
    _load_json(db, *raws)
    manual = E.create_pesticide(db, _manual_body("Thủ công giữ nguyên"), user_id=1)

    #  Tệp đính kèm của một thuốc KHÔNG bị sửa — phải còn nguyên sau khi nạp CẬP NHẬT.
    untouched = db.query(CustomsPesticide).filter_by(source_id=1104).one()
    sf = StoredFile(filename="ho-so.pdf", file_key="k1", content_type="application/pdf",
                    size=1, created_by=1, updated_by=1)
    db.add(sf)
    db.flush()
    db.add(FileLink(file_id=sf.id, entity="customs_price", entity_id=untouched.id))
    db.commit()
    untouched_id = untouched.id

    content, filename = _export(db, scope="page", page=1, offset=0, limit=10)
    #  Giả lập người dùng SỬA một ô trên Excel rồi nạp lại (đổi tên thuốc source_id=1100).
    wb = load_workbook(io.BytesIO(content))
    ws = wb[R.LIST_SHEET]
    header = [c.value for c in ws[1]]
    id_idx, name_idx = header.index("id"), header.index("ten_thuoc")
    for row in ws.iter_rows(min_row=2):
        if row[id_idx].value == 1100:
            row[name_idx].value = "Trang 0 đã sửa"
    buf = io.BytesIO()
    wb.save(buf)

    records, mode = R.read_file_with_mode(filename, buf.getvalue())
    assert mode == "merge"
    result = M.merge_catalog(db, records, user_id=1, filename=filename)
    assert result["mode"] == "merge" and result["updated"] == 5 and result["added"] == 0

    #  Tổng số thuốc KHÔNG đổi (5 nguồn + 1 thủ công) — merge không xóa ai.
    assert db.query(CustomsPesticide).count() == 6
    changed = db.query(CustomsPesticide).filter_by(source_id=1100).one()
    assert changed.trade_name == "Trang 0 đã sửa"
    #  Thuốc thủ công không bị đụng.
    assert db.get(CustomsPesticide, manual["id"]).is_manual is True
    #  id DB của thuốc đổi/không đổi đều GIỮ NGUYÊN (CR-494) — tệp đính kèm còn nguyên.
    assert db.get(CustomsPesticide, untouched_id) is not None
    assert (db.query(FileLink).filter_by(entity="customs_price", entity_id=untouched_id).count() == 1)


def test_nap_tep_xuat_toan_bo_van_la_thay_toan_bo_qua_http(db, seed, cap_quyen, client_as):
    """`mode` trả về đúng cho FE — xuất toàn bộ rồi nạp lại qua HTTP vẫn THAY TOÀN BỘ như cũ."""
    _load_json(db, _raw(1200, "Toàn bộ một"), _raw(1201, "Toàn bộ hai"))
    v = db.query(User).filter(User.employee_id == seed.emp_tp_id).one()
    cap_quyen(v.id, "customs_pesticide", scope="all", read=True, write=True)
    content, filename = _export(db, scope="all")
    resp = client_as(v).post("/api/customs/pesticides/import",
                             files={"file": (filename, content,
                                             "application/vnd.openxmlformats-officedocument"
                                             ".spreadsheetml.sheet")})
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert data["mode"] == "replace" and resp.json()["message"].startswith("Đã thay toàn bộ")


# ── C2 — id của thuốc thủ công / nguồn trùng source_id không bao giờ đụng id dương khác ────

def test_export_id_tra_am_cho_thu_cong_va_cho_nguon_trung_source_id():
    manual = CustomsPesticide(id=50, is_manual=True, source_id=0, trade_name="x")
    sourced = CustomsPesticide(id=7, is_manual=False, source_id=50, trade_name="y")
    assert C.export_id(manual) == -50
    assert C.export_id(sourced) == 50
    #  `source_id=50` này đang TRÙNG với một thuốc nguồn khác (danh sách trùng truyền vào) →
    #  xuất ÂM theo `id` hệ thống của CHÍNH nó, không phải `source_id` nữa.
    assert C.export_id(sourced, dup_source_ids=frozenset({50})) == -7


def test_thuoc_tu_them_co_id_trung_source_id_khac_khong_bi_gop_nham_pham_vi(db):
    """Tái hiện lỗi C2: thuốc tự thêm mang `id` hệ thống TRÙNG `source_id` của một thuốc nguồn
    khác (hai không gian id khác nhau, không gì ngăn trùng số). Bản `export_id()` CŨ
    (`p.id if is_manual else p.source_id`) xuất CẢ HAI dòng list-sheet CÙNG id, và CẢ HAI dòng
    phạm vi CÙNG id ở use-sheet → nạp lại gộp nhầm phạm vi của hai thuốc khác nhau."""
    _load_json(db, _raw(777, "Nguồn Bảy Trăm Bảy Mươi Bảy",
                        pham_vi_su_dung=[{"cay_trong": "lúa", "dich_hai": "sâu",
                                          "lieu_luong": "1L"}]))
    manual = CustomsPesticide(id=777, trade_name="Thủ Công Bảy Trăm Bảy Mươi Bảy", is_manual=True,
                              source_id=0, status=int(PesticideStatus.ACTIVE),
                              created_by=1, updated_by=1)
    db.add(manual)
    db.flush()
    db.add(CustomsPesticideUse(pesticide_id=777, sort_order=1, crop="cà phê", pest="rệp"))
    db.commit()

    content, _ = _export(db, scope="all")
    [rec] = R.read_file("x.xlsx", content)
    assert rec["source_id"] == 777
    #  PHẢI chỉ còn phạm vi của THUỐC NGUỒN — không lẫn phạm vi của thuốc thủ công.
    assert [u["crop"] for u in rec["uses"]] == ["lúa"]


def test_hai_thuoc_nguon_trung_source_id_van_xuat_duoc_tep_nap_lai_duoc(db):
    """DB nhiễm sẵn hai thuốc NGUỒN cùng `source_id` (câu hỏi "nêu rõ" của C2) — xuất vẫn phải
    ra tệp nạp lại được, không tự ném ValueError vì chính tệp mình xuất ra có id trùng. Xử lý:
    CẢ HAI bản trùng xuất id ÂM (không chọn "ai thắng") — bộ đọc bỏ cả hai, admin tự xử lý trùng
    ở DB; phần CÒN LẠI của danh mục (không trùng gì) vẫn xuất/nạp lại bình thường."""
    _load_json(db, _raw(900, "Trùng Một"), _raw(901, "Không trùng ai"))
    dup = CustomsPesticide(trade_name="Trùng Hai", source_id=900, is_manual=False,
                           status=int(PesticideStatus.ACTIVE), created_by=1, updated_by=1)
    db.add(dup)
    db.flush()
    db.add(CustomsPesticideUse(pesticide_id=dup.id, sort_order=1, crop="ngô", pest="sâu"))
    db.commit()

    content, _ = _export(db, scope="all")   # KHÔNG được ném ValueError id trùng
    records = R.read_file("x.xlsx", content)
    #  Cặp trùng (900) CẢ HAI bị bộ đọc bỏ; thuốc không trùng (901) vẫn còn nguyên.
    assert len(records) == 1 and records[0]["source_id"] == 901


def test_doc_tep_co_hai_dong_thuoc_cung_id_duong_bi_tu_choi():
    """Bất kể tệp từ đâu ra — HAI dòng list-sheet cùng id DƯƠNG phải bị TỪ CHỐI rõ ràng, không
    lặng lẽ gộp nhầm phạm vi của hai thuốc khác nhau vào một id."""
    wb = Workbook()
    ws = wb.active
    ws.title = R.LIST_SHEET
    ws.append(["id", "ten_thuoc"])
    ws.append([42, "Thuốc A"])
    ws.append([42, "Thuốc B"])
    uses = wb.create_sheet(R.USE_SHEET)
    uses.append(["id", "stt_pham_vi", "cay_trong"])
    uses.append([42, 1, "lúa"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="trùng nhau"):
        R.read_xlsx(buf.getvalue())


def test_doc_bo_qua_moi_dong_id_am_hoac_bang_khong():
    """id <= 0 ở CẢ HAI sheet đều bị bỏ — không chỉ sheet thuốc."""
    wb = Workbook()
    ws = wb.active
    ws.title = R.LIST_SHEET
    ws.append(["id", "ten_thuoc"])
    ws.append([-5, "Bỏ qua (âm)"])
    ws.append([0, "Bỏ qua (không)"])
    ws.append([9, "Giữ lại"])
    uses = wb.create_sheet(R.USE_SHEET)
    uses.append(["id", "stt_pham_vi", "cay_trong"])
    uses.append([-5, 1, "rác âm"])
    uses.append([0, 1, "rác không"])
    uses.append([9, 1, "lúa"])
    buf = io.BytesIO()
    wb.save(buf)
    [rec] = R.read_xlsx(buf.getvalue())
    assert rec["source_id"] == 9
    assert [u["crop"] for u in rec["uses"]] == ["lúa"]


# ── C3 + H1 — ô Excel an toàn: không phải công thức, không ký tự điều khiển cấm ────────────

def test_chuoi_bat_dau_bang_dau_bang_ghi_dang_chuoi_khong_phai_cong_thuc(db):
    _load_json(db, _raw(1300, "=SUM(A1:A9)", cach_dung="=1+1",
                        tom_tat_su_dung='=HYPERLINK("http://x")',
                        pham_vi_su_dung=[{"cay_trong": "lúa", "dich_hai": "rầy",
                                          "lieu_luong": "=A1", "cach_dung": "=1+1"}]))
    content, _ = _export(db, scope="all")
    wb = load_workbook(io.BytesIO(content))
    ws = wb[R.LIST_SHEET]
    header = [c.value for c in ws[1]]
    row2 = ws[2]
    name_cell = row2[header.index("ten_thuoc")]
    summary_cell = row2[header.index("tom_tat_su_dung")]
    #  Ô CHUỖI ('s'), KHÔNG phải CÔNG THỨC ('f') — openpyxl mặc định suy 'f' cho chuỗi "=...".
    assert name_cell.value == "=SUM(A1:A9)" and name_cell.data_type == "s"
    assert summary_cell.data_type == "s"

    uses_ws = wb[R.USE_SHEET]
    use_header = [c.value for c in uses_ws[1]]
    use_row2 = uses_ws[2]
    dosage_cell = use_row2[use_header.index("lieu_luong")]
    assert dosage_cell.value == "=A1" and dosage_cell.data_type == "s"

    #  Round-trip qua bộ đọc vẫn ra đúng chữ — không bị "tính" mất lúc mở lại bằng Excel.
    [rec] = R.read_file("x.xlsx", content)
    assert rec["trade_name"] == "=SUM(A1:A9)"
    assert rec["uses"][0]["dosage"] == "=A1"


def test_ky_tu_dieu_khien_cam_khong_lam_vo_luot_xuat(db):
    """H1 — openpyxl ném `IllegalCharacterError` ngay khi `.append()` gặp `\\x00`/`\\x0b` chưa
    lọc; export PHẢI thành công, không ném lỗi, không bắt người dùng sửa tệp nguồn bằng tay."""
    _load_json(db, _raw(1301, "Có ký tự lạ",
                        tom_tat_su_dung="Tóm tắt có \x00 và \x0b bên trong"))
    content, filename = _export(db, scope="all")   # KHÔNG được ném IllegalCharacterError
    assert filename
    [rec] = R.read_file("x.xlsx", content)
    assert "\x00" not in rec["summary"]


# ── H2 — cột nguyên văn `quan_ly_tinh_khang_raw` giữ chữ tự do có ';'/':' round-trip đúng ──

def test_resistance_tu_do_co_hai_cham_va_cham_phay_round_trip_nguyen_van(db):
    """Chữ tự do CÓ ';'/':' không phải dấu phân tách — bản tách-rồi-ghép qua `_resistance()`
    (cột `quan_ly_tinh_khang` cũ) hiểu nhầm là dấu phân tách, mất/méo chữ. `RAW_RESISTANCE_
    COLUMN` giữ nguyên văn, bộ đọc ưu tiên cột đó."""
    tricky = "A: 1 | 2; B; C: x:y|z"
    rec = R._record(_raw(1400, "Resistance Test"), toxicity="", resistance=tricky,
                    uses=[{"crop": "lúa", "pest": "rầy", "dosage": "", "pre_harvest_interval": "",
                          "usage": ""}])
    S.replace_catalog(db, [rec], user_id=1, filename="seed.json")

    content, _ = _export(db, scope="all")
    [got] = R.read_file("x.xlsx", content)
    assert got["resistance"] == tricky


# ── M2 + M3 — lỗi SAU khi đã ghi xong tệp phải dọn tệp tạm, tệp quá khổ bị chặn ────────────

def test_loi_sau_khi_ghi_xong_tep_thi_don_tep_tam_roi_nem_lai(db, monkeypatch):
    _load_json(db, _raw(1500, "Dọn tệp tạm"))
    captured: dict = {}
    original_write = X._write_file

    def _spy_write_file(*a, **kw):
        path, n, m = original_write(*a, **kw)
        captured["path"] = path
        return path, n, m

    monkeypatch.setattr(X, "_write_file", _spy_write_file)

    def _boom(*a, **kw):
        raise RuntimeError("lỗi ghi nhật ký")

    monkeypatch.setattr(X, "_record_log", _boom)
    with pytest.raises(RuntimeError):
        X.export_to_tempfile(db, USER, "all", "", None, "", "", False, None, 1, 0, 20)
    assert captured.get("path") and not os.path.exists(captured["path"])


def test_tep_xuat_vuot_tran_kich_thuoc_cua_bo_nap_la_422_va_xoa_tep_tam(db, monkeypatch):
    _load_json(db, _raw(1501, "Quá khổ lớn"))
    monkeypatch.setattr(R, "MAX_UPLOAD_BYTES", 10)
    captured: dict = {}
    original_write = X._write_file

    def _spy_write_file(*a, **kw):
        path, n, m = original_write(*a, **kw)
        captured["path"] = path
        return path, n, m

    monkeypatch.setattr(X, "_write_file", _spy_write_file)
    with pytest.raises(HTTPException) as exc:
        X.export_to_tempfile(db, USER, "all", "", None, "", "", False, None, 1, 0, 20)
    assert exc.value.status_code == 422 and "MB" in exc.value.detail
    assert captured.get("path") and not os.path.exists(captured["path"])


# ── L2 — nhật ký xuất hiện nhãn đọc được + L6 — giờ VN dùng chung ──────────────────────────

def test_nhat_ky_xuat_hien_nhan_doc_duoc_tren_system_exports(db):
    _load_json(db, _raw(1600, "Nhãn đọc được"))
    _export(db, scope="all")
    log = db.query(ExportLog).filter_by(entity="customs_price").one()
    assert log.filter_summary.startswith("Thuốc BVTV — toàn bộ")

    from app.modules.export_log.controller import _log_out
    out = _log_out(db, log)
    assert out["entity_label"] == "Thuốc BVTV"


def test_vn_offset_dung_chung_voi_app_core_export_xlsx():
    from app.core import export_xlsx
    assert X.VN_OFFSET is export_xlsx.VN_OFFSET
