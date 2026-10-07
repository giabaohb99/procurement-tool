"""bao-CR-603 — Tra cứu thị trường: năm cột TÙY CHỌN trên tệp nạp.

```
thiếu cột                 -> nạp như cũ: «Nước nhận hàng» trống, hoạt chất / hàm lượng suy ra, giá VND tính
tệp có cột, ô có chữ      -> lấy của tệp, cờ _from_file = 1, retag_all KHÔNG ghi đè
tệp có cột, ô trống       -> suy ra / tính như cũ (cờ = 0)
mã băm chống trùng        -> KHÔNG đổi theo cột tùy chọn: cùng 32 cột là cùng một dòng
tệp Excel xuất ra         -> nạp lại được, bốn cột cuối được nhận diện
```
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from app.modules.customs import importer, reader
from app.modules.customs import service as S
from app.modules.customs.constants import COLUMNS, OPTIONAL_COLUMNS, OPTIONAL_LABELS
from app.modules.customs.dedupe import HASH_FIELDS
from app.modules.customs.ingredient import retag_all
from app.modules.customs.model import CustomsIngredientAlias, CustomsLine
from app.modules.import_tool.model import ImportBatch, ImportLog, ImportMode, ImportModule, ImportStatus

ROW = {
    "reg_date": "13-01-2026", "office_code": "HQHPKV3", "importer_tax_code": "'0500590269",
    "importer_name": "CÔNG TY TNHH THÚ Y TOÀN CẦU", "partner_name": "NOVADAN APS",
    "hs_code": "'38089990", "line_no": 1, "product_name": "ATRAZINE 97% TECH",
    "price_usd": 3.0, "price_nt": 3.0, "adj_price_usd": None, "adj_price_nt": None,
    "currency": "USD", "fx_rate": 26130, "usd_rate": 26130, "quantity": 880,
    "unit_code": "KGM", "origin_country": "DK", "contract_no": "770710",
    "contract_date": None, "incoterm": "CIF", "transport_mode": "2-Đường biển (container)",
    "rate_import": 5, "rate_excise": 0, "rate_vat": 5, "rate_safeguard": 0, "tax_import": 0,
    "tax_excise": 0, "tax_vat": 0, "tax_environment": 0, "tax_safeguard": 0, "import_country": "VN",
}


def _xlsx(rows: list[dict], headers: list[tuple[str, str]]) -> bytes:
    """`headers` = [(khóa, tiêu đề)] — chọn được cột nào có, cột nào thiếu, tiêu đề ghi sao."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([label for _, label in headers])
    for r in rows:
        ws.append([r.get(key) for key, _ in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _batch(db) -> ImportBatch:
    b = ImportBatch(module=ImportModule.CUSTOMS_DECLARATION, mode=ImportMode.APPLY, filename="gtt02.xlsx",
                    sheet_info="", error_summary="", status=ImportStatus.RUNNING, created_by=1)
    db.add(b)
    db.commit()
    return b


FULL = list(COLUMNS)
WITH_OPTIONAL = FULL + [(k, labels[0]) for k, labels in OPTIONAL_COLUMNS]


@pytest.fixture
def alias(db):
    db.add(CustomsIngredientAlias(keyword="ATRAZINE", canonical="ATRAZINE"))
    db.commit()


# ── Thiếu cột thì nạp như cũ ────────────────────────────────────────────────
def test_missing_import_country_is_accepted_and_left_blank():
    headers = [c for c in FULL if c[0] != "import_country"]
    reader.check_headers(_xlsx([ROW], headers))             # không từ chối lô
    res = reader.parse(_xlsx([ROW], headers), today=date(2026, 9, 23))
    assert res.rows[0]["import_country"] == ""
    assert res.optional_columns == []
    #  Tệp GTT02 chuẩn (có «Nước nhận hàng») cũng không bị kể là «có cột tùy chọn»
    assert reader.parse(_xlsx([ROW], FULL), today=date(2026, 9, 23)).optional_columns == []


def test_other_required_columns_still_rejected():
    headers = [c for c in FULL if c[0] != "product_name"]
    with pytest.raises(reader.CustomsFileError, match="«Tên hàng»"):
        reader.check_headers(_xlsx([ROW], headers))


def test_file_without_optional_columns_derives_and_computes_as_before(db, alias):
    b = _batch(db)
    importer.run(db, b, _xlsx([ROW], FULL), apply=True)
    ln = db.query(CustomsLine).one()
    assert ln.active_ingredient == "ATRAZINE" and ln.formulation == "97%"
    assert ln.active_ingredient_from_file is False and ln.formulation_from_file is False
    assert ln.price_vnd_flat is None and ln.price_vnd_line_tax is None
    row = S.serialize_lines(db, [ln])[0]
    assert row["price_vnd_flat"] == round(3 * 26130 * 1.07)
    assert row["price_vnd_line_tax"] == round(3 * 26130 * 1.05)
    assert row["price_vnd_flat_from_file"] is False and row["active_ingredient_from_file"] is False
    #  Không có dòng nhật ký «cột tùy chọn» khi tệp không có cột nào
    assert not db.query(ImportLog).filter(ImportLog.batch_id == b.id, ImportLog.message.like("Tệp có cột%")).count()


# ── Tệp có cột: ô có chữ thì lấy của tệp, ô trống mới suy ra ────────────────
def test_file_values_win_and_blank_cells_fall_back_to_derivation(db, alias):
    rows = [
        {**ROW, "line_no": 1, "active_ingredient": "Atrazine + Mesotrione", "formulation": "80wp",
         "price_vnd_flat": 90000, "price_vnd_line_tax": 88000.5},
        {**ROW, "line_no": 2, "active_ingredient": "", "formulation": None,
         "price_vnd_flat": None, "price_vnd_line_tax": ""},
    ]
    b = _batch(db)
    importer.run(db, b, _xlsx(rows, WITH_OPTIONAL), apply=True)
    first, second = db.query(CustomsLine).order_by(CustomsLine.line_no).all()

    assert first.active_ingredient == "Atrazine + Mesotrione" and first.active_ingredient_from_file is True
    assert first.formulation == "80WP" and first.formulation_from_file is True      # viết hoa cho khớp ô lọc
    assert first.price_vnd_flat == Decimal("90000") and first.price_vnd_line_tax == Decimal("88000.5")

    assert second.active_ingredient == "ATRAZINE" and second.active_ingredient_from_file is False
    assert second.formulation == "97%" and second.formulation_from_file is False
    assert second.price_vnd_flat is None and second.price_vnd_line_tax is None

    out = {r["line_no"]: r for r in S.serialize_lines(db, [first, second])}
    assert out[1]["price_vnd_flat"] == 90000 and out[1]["price_vnd_flat_from_file"] is True
    assert out[1]["price_vnd_line_tax"] == 88000.5 and out[1]["price_vnd_line_tax_from_file"] is True
    assert out[2]["price_vnd_flat"] == round(3 * 26130 * 1.07) and out[2]["price_vnd_flat_from_file"] is False
    #  Lô ghi lại tệp có cột tùy chọn nào
    log = db.query(ImportLog).filter(ImportLog.batch_id == b.id, ImportLog.message.like("Tệp có cột%")).one()
    assert "«Hoạt chất»" in log.message and "«Đơn giá quy đổi VND (thuế NK 7%)»" in log.message


def test_vnd_columns_can_come_one_at_a_time(db):
    """Chỉ có cột 7% trong tệp → cột 7% lấy của tệp, cột theo thuế suất vẫn tính."""
    headers = FULL + [("price_vnd_flat", "Giá VND (thuế NK 7%)")]
    b = _batch(db)
    importer.run(db, b, _xlsx([{**ROW, "price_vnd_flat": 12345}], headers), apply=True)
    row = S.serialize_lines(db, db.query(CustomsLine).all())[0]
    assert row["price_vnd_flat"] == 12345 and row["price_vnd_flat_from_file"] is True
    assert row["price_vnd_line_tax"] == round(3 * 26130 * 1.05) and row["price_vnd_line_tax_from_file"] is False


def test_bad_number_in_vnd_column_warns_and_falls_back():
    res = reader.parse(_xlsx([{**ROW, "price_vnd_flat": "chín mươi nghìn"}], WITH_OPTIONAL), today=date(2026, 9, 23))
    assert res.rows[0]["price_vnd_flat"] is None
    assert any("Đơn giá quy đổi VND (thuế NK 7%)" in m and "tính như cũ" in m for _, _, m in res.logs)


def test_long_ingredient_text_is_cut_to_column_limit():
    res = reader.parse(_xlsx([{**ROW, "active_ingredient": "A" * 300, "formulation": "B" * 50}], WITH_OPTIONAL),
                       today=date(2026, 9, 23))
    assert len(res.rows[0]["active_ingredient"]) == 255 and len(res.rows[0]["formulation"]) == 40
    assert sum("đã cắt" in m for _, _, m in res.logs) == 2


def test_header_aliases_are_recognized():
    """Tiêu đề cũ «(suy ra)» của Excel xuất trước CR-603 và tiêu đề ngắn của màn hình đều khớp."""
    headers = FULL + [("active_ingredient", "Hoạt chất (suy ra)"), ("formulation", "Hàm lượng / dạng (suy ra)"),
                      ("price_vnd_flat", "Giá VND (thuế NK 7%)"), ("price_vnd_line_tax", "Giá VND (thuế suất dòng)")]
    res = reader.parse(_xlsx([{**ROW, "active_ingredient": "X", "formulation": "20EC",
                               "price_vnd_flat": 1, "price_vnd_line_tax": 2}], headers), today=date(2026, 9, 23))
    assert res.optional_columns == ["active_ingredient", "formulation", "price_vnd_flat", "price_vnd_line_tax"]
    assert res.rows[0]["active_ingredient"] == "X" and res.rows[0]["price_vnd_line_tax"] == Decimal("2")


# ── retag_all không ghi đè giá trị lấy từ tệp ──────────────────────────────
def test_retag_keeps_values_taken_from_file(db, alias):
    rows = [{**ROW, "line_no": 1, "active_ingredient": "HOẠT CHẤT TAY", "formulation": "", },
            {**ROW, "line_no": 2}]
    b = _batch(db)
    importer.run(db, b, _xlsx(rows, WITH_OPTIONAL), apply=True)
    #  Danh mục đổi: từ khóa ATRAZINE nay trỏ sang tên chuẩn khác → dòng suy ra phải đổi theo
    db.query(CustomsIngredientAlias).update({"canonical": "ATRAZINE (NEW)"})
    db.commit()
    retag_all(db)
    first, second = db.query(CustomsLine).order_by(CustomsLine.line_no).all()
    assert first.active_ingredient == "HOẠT CHẤT TAY"           # từ tệp — giữ nguyên
    assert first.formulation == "97%"                           # ô trống trong tệp → suy ra, được gắn lại
    assert second.active_ingredient == "ATRAZINE (NEW)"         # suy ra — đổi theo danh mục


# ── Mã băm chống trùng không đổi ────────────────────────────────────────────
def test_optional_columns_do_not_change_row_hash(db, alias):
    assert not ({"price_vnd_flat", "price_vnd_line_tax", "active_ingredient", "formulation"} & set(HASH_FIELDS))
    first = _batch(db)
    importer.run(db, first, _xlsx([ROW], FULL), apply=True)
    second = _batch(db)
    importer.run(db, second, _xlsx([{**ROW, "active_ingredient": "KHÁC", "price_vnd_flat": 1}], WITH_OPTIONAL),
                 apply=True)
    assert db.query(CustomsLine).count() == 1                   # dòng thứ hai là «Đã có», bỏ qua
    assert second.created_count == 0


# ── Excel xuất ra nạp lại được ──────────────────────────────────────────────
def test_exported_excel_round_trips_through_the_reader(db, alias):
    b = _batch(db)
    importer.run(db, b, _xlsx([{**ROW, "active_ingredient": "ABC", "price_vnd_flat": 5000}], WITH_OPTIONAL), apply=True)
    raw = S.export_lines_xlsx(db, {})
    ws = openpyxl.load_workbook(io.BytesIO(raw)).active
    assert ws.max_column == 38                  # bao-CR-608: + «ID» đầu, «Thao tác» cuối
    assert [ws.cell(1, 34 + i).value for i in range(4)] == [OPTIONAL_LABELS[k] for k, _ in OPTIONAL_COLUMNS]
    res = reader.parse(raw, today=date(2026, 9, 23))
    assert res.optional_columns == ["active_ingredient", "formulation", "price_vnd_flat", "price_vnd_line_tax"]
    #  Excel xuất ghi ngày dạng ISO và phương tiện vận chuyển dạng nhãn — bộ đọc phải nhận, không bỏ dòng.
    assert res.skipped == 0 and len(res.rows) == 1
    assert res.rows[0]["reg_date"] == date(2026, 1, 13) and res.rows[0]["transport_mode"] == 2
    assert res.rows[0]["active_ingredient"] == "ABC" and res.rows[0]["price_vnd_flat"] == Decimal("5000")
    assert res.rows[0]["price_vnd_line_tax"] == Decimal(str(round(3 * 26130 * 1.05)))
    #  Nạp lại tệp xuất ra: dòng đã có → bỏ qua, không nhân đôi (cột tùy chọn không vào mã băm).
    #  bao-CR-608: tệp xuất có cột «ID» — dòng trỏ đúng dòng đã lưu mà dữ liệu y hệt → «Đã có».
    again = _batch(db)
    importer.run(db, again, raw, apply=True)
    assert db.query(CustomsLine).count() == 1 and again.created_count == 0
