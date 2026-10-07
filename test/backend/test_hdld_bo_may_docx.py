"""Bộ máy .docx của HĐLĐ: kiểm mẫu, render, ngữ cảnh, danh mục biến — KHÔNG DB, KHÔNG HTTP.

Mẫu .docx được dựng ngay trong test bằng `python-docx` (không commit tệp nhị phân).
Mục tiêu: tìm chỗ bộ máy LỦNG (SSTI, bom nén, macro, ký tự phá XML, run bị tách), không
phải chứng minh đường đẹp chạy được.
"""
import time
import zipfile
from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from app.modules.labor_contract import docx_engine
from app.modules.labor_contract.context_builder import (
    build_context, describe_duration, fmt_money,
)
from app.modules.labor_contract.docx_engine import TemplateRejected, inspect_template, render
from app.modules.labor_contract.placeholder_catalog import (
    KNOWN_KEYS, PLACEHOLDERS, sample_context,
)


# ── Dựng .docx trong bộ nhớ ─────────────────────────────────────────────────────
def make_docx(*paragraphs: str) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)


def rewrap_zip(data: bytes, extra: dict[str, bytes] | None = None, content_types_patch=None) -> bytes:
    """Chép gói zip, thêm thành phần / vá [Content_Types].xml."""
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "[Content_Types].xml" and content_types_patch:
                body = content_types_patch(body)
            dst.writestr(info.filename, body)
        for name, body in (extra or {}).items():
            dst.writestr(name, body)
    return out.getvalue()


# ── Danh mục biến ──────────────────────────────────────────────────────────────
def test_danh_muc_bien_khong_trung_va_la_ascii():
    keys = [p.key for p in PLACEHOLDERS]
    assert len(keys) == len(set(keys)) == len(KNOWN_KEYS)
    assert all(k.isascii() and k == k.lower() and k.replace("_", "").isalnum() for k in keys)
    assert set(sample_context()) == KNOWN_KEYS


# ── inspect_template: đường đẹp ──────────────────────────────────────────────────
def test_mau_hop_le_tra_ve_danh_sach_bien_dung():
    data = make_docx("Họ tên: {{ ho_ten }}", "Lương: {{ luong_co_ban }} ({{ luong_co_ban_bang_chu }})")
    assert inspect_template(data) == ["ho_ten", "luong_co_ban", "luong_co_ban_bang_chu"]


def test_mau_khong_co_bien_nao_van_hop_le():
    assert inspect_template(make_docx("Văn bản tĩnh")) == []


def test_chu_thich_jinja_vo_hai_duoc_bo_qua():
    assert inspect_template(make_docx("{# ghi chú soạn mẫu #}Họ tên: {{ ho_ten }}")) == ["ho_ten"]


# ── inspect_template: CHỈ nhận `{{ ten_bien }}` (đóng lỗ DoS vòng lặp lồng) ─────────────
NESTED_FOR = "{% for a in range(100000) %}{% for b in range(100000) %}x{% endfor %}{% endfor %}"


@pytest.mark.parametrize("payload", [
    NESTED_FOR,
    "{% for a in range(3) %}x{% endfor %}",
    "{% if phu_cap_ghi_chu %}Có phụ cấp: {{ phu_cap }}{% endif %}",
    "{% set x = 1 %}{{ ho_ten }}",
    "{%p for a in range(100000) %}{%p for b in range(100000) %}x{%p endfor %}{%p endfor %}",
    "{%tr for a in range(100000) %}{%tr endfor %}",
    "{%r for a in range(100000) %}{%r endfor %}",
    "{{ ho_ten|upper }}",
    "{{ ho_ten ~ so_cccd }}",
    "{{ range(5) }}",
    "{{ ho_ten.__class__ }}",
    "{{ 7*7 }}",
    "{{ 'x' }}",
])
def test_mau_chi_duoc_dung_bien_thuan(payload):
    start = time.monotonic()
    with pytest.raises(TemplateRejected):
        inspect_template(make_docx(payload))
    assert time.monotonic() - start < 3   # bị chặn ở bước phân tích, KHÔNG chạy vòng lặp


def test_render_ban_luu_san_co_vong_lap_cung_bi_tu_choi():
    """`render` dùng cùng môi trường: tệp trên storage bị thay bằng bản có vòng lặp vẫn không chạy."""
    with pytest.raises(TemplateRejected):
        render(make_docx(NESTED_FOR), sample_context())


# ── inspect_template: biến lạ ─────────────────────────────────────────────────────
def test_bien_la_bi_tu_choi_va_liet_ke_het():
    data = make_docx("{{ ho_ten }} {{ ten_con_meo }} {{ luong_tang }}")
    with pytest.raises(TemplateRejected) as e:
        inspect_template(data)
    assert e.value.unknown == ["luong_tang", "ten_con_meo"]


def test_bien_viet_hoa_hay_sai_dau_la_bien_la():
    """`{{ Ho_Ten }}` KHÔNG phải `ho_ten` — render ra rỗng im lặng nếu cho qua."""
    with pytest.raises(TemplateRejected) as e:
        inspect_template(make_docx("{{ Ho_Ten }}"))
    assert e.value.unknown == ["Ho_Ten"]


# ── Word tách run ───────────────────────────────────────────────────────────────
def test_run_bi_tach_boi_in_dam_giua_bien_van_render_dung():
    """Word tách `{{ ho_` + `ten }}` thành hai run (định dạng giữa chừng). docxtpl gộp được
    trường hợp phổ biến này; nếu một bản docxtpl sau này hết gộp thì test đỏ ở đây."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Tên: {{ ho_")
    p.add_run("ten }}").bold = True
    buf = BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    assert inspect_template(data) == ["ho_ten"]
    assert "Tên: Lê Văn A" in docx_text(render(data, {**sample_context(), "ho_ten": "Lê Văn A"}))


def test_bien_bi_cat_dut_thanh_tu_khoa_khac_nhau_bi_chan():
    """`{{ ho_ten` thiếu `}}` → lỗi cú pháp có thông báo, không phải 500."""
    with pytest.raises(TemplateRejected) as e:
        inspect_template(make_docx("{{ ho_ten"))
    assert "cú pháp" in e.value.message


# ── SSTI ────────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("payload", [
    "{{ ''.__class__.__mro__[1].__subclasses__() }}",
    "{{ ho_ten.__class__ }}",
    "{{ cycler.__init__.__globals__.os.popen('id').read() }}",
    "{{ lipsum.__globals__ }}",
    "{{ ho_ten.__class__.__base__.__subclasses__() }}",
])
def test_ssti_bi_chan_khi_tai_len(payload):
    with pytest.raises(TemplateRejected):
        inspect_template(make_docx(payload))


def test_ssti_bi_chan_ca_o_render_va_khong_ro_dau_ra():
    with pytest.raises(TemplateRejected):
        render(make_docx("{{ ho_ten.__class__.__mro__ }}"), sample_context())


def test_range_khong_lo_bi_san_box_chan():
    """Vòng lặp khổng lồ không được treo máy chủ."""
    start = time.monotonic()
    with pytest.raises(TemplateRejected):
        inspect_template(make_docx("{% for i in range(100000000000) %}x{% endfor %}"))
    assert time.monotonic() - start < 5


# ── Ký tự phá XML ────────────────────────────────────────────────────────────────
def test_ten_co_ky_tu_dac_biet_van_ra_docx_mo_lai_duoc():
    tpl = make_docx("Bên A: {{ ten_cong_ty }} - {{ dia_chi_cong_ty }}")
    nasty = 'A&B <Co> "Quote" \'x\' </w:t></w:p>'
    out = render(tpl, {**sample_context(), "ten_cong_ty": nasty, "dia_chi_cong_ty": "1 & 2 < 3 > 0"})
    text = docx_text(out)  # mở lại được = XML còn nguyên
    assert nasty in text
    assert "1 & 2 < 3 > 0" in text


def test_gia_tri_chua_cu_phap_jinja_khong_bi_chay_lan_hai():
    """Dữ liệu NGƯỜI DÙNG nhập (họ tên) chứa `{{ 7*7 }}` phải ra nguyên văn, không thành 49."""
    out = render(make_docx("{{ ho_ten }}"), {**sample_context(), "ho_ten": "{{ 7*7 }} {% if 1 %}x{% endif %}"})
    assert docx_text(out).strip() == "{{ 7*7 }} {% if 1 %}x{% endif %}"


def test_gia_tri_rong_va_unicode_dai():
    out = render(make_docx("[{{ ho_ten }}] [{{ ghi_chu }}]"),
                 {**sample_context(), "ho_ten": "Đặng Thị Ngọc Hạnh" * 50, "ghi_chu": ""})
    text = docx_text(out)
    assert "[]" in text and text.count("Đặng Thị Ngọc Hạnh") == 50


# ── Gói zip ───────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("junk", [b"", b"not a zip", b"PK\x03\x04" + b"\x00" * 10, b"%PDF-1.4 ..."])
def test_khong_phai_zip_bi_tu_choi(junk):
    with pytest.raises(TemplateRejected):
        inspect_template(junk)


def test_zip_hop_le_nhung_khong_phai_docx():
    out = BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(TemplateRejected) as e:
        inspect_template(out.getvalue())
    assert "document.xml" in e.value.message


def test_zip_bomb_vuot_tran_giai_nen_bi_chan():
    bomb = rewrap_zip(make_docx("x"), extra={"word/media/bomb.bin": b"\x00" * (70 * 1024 * 1024)})
    assert len(bomb) < 1024 * 1024  # nén cực mạnh: nhỏ để tải lên, to khi bung
    with pytest.raises(TemplateRejected) as e:
        inspect_template(bomb)
    assert "giải nén" in e.value.message


def test_zip_qua_nhieu_muc_bi_chan(monkeypatch):
    many = rewrap_zip(make_docx("x"), extra={f"word/f{i}.txt": b"a" for i in range(docx_engine.MAX_ZIP_ENTRIES + 1)})
    with pytest.raises(TemplateRejected) as e:
        inspect_template(many)
    assert "quá nhiều" in e.value.message


def test_docx_co_macro_bi_chan():
    data = rewrap_zip(make_docx("{{ ho_ten }}"), extra={"word/vbaProject.bin": b"\xd0\xcf\x11\xe0"})
    with pytest.raises(TemplateRejected) as e:
        inspect_template(data)
    assert "macro" in e.value.message


def test_docm_doi_duoi_bi_chan_qua_content_type():
    data = rewrap_zip(
        make_docx("x"),
        content_types_patch=lambda b: b.replace(b"</Types>", b'<Override PartName="/x" ContentType="application/vnd.ms-word.document.macroEnabled.12"/></Types>'),
    )
    with pytest.raises(TemplateRejected):
        inspect_template(data)


def test_docx_thieu_content_types():
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(make_docx("x"))) as src, zipfile.ZipFile(out, "w") as dst:
        for i in src.infolist():
            if i.filename != "[Content_Types].xml":
                dst.writestr(i.filename, src.read(i.filename))
    with pytest.raises(TemplateRejected):
        inspect_template(out.getvalue())


def test_document_xml_hong_khong_ra_500():
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(make_docx("x"))) as src, zipfile.ZipFile(out, "w") as dst:
        for i in src.infolist():
            dst.writestr(i.filename, b"<<<khong phai xml" if i.filename == "word/document.xml" else src.read(i.filename))
    with pytest.raises(TemplateRejected):
        inspect_template(out.getvalue())


# ── Lỗi cú pháp Jinja ────────────────────────────────────────────────────────────
def test_if_thieu_endif_bao_loi_ro():
    with pytest.raises(TemplateRejected) as e:
        inspect_template(make_docx("{% if phu_cap %}co phu cap"))
    assert "cú pháp" in e.value.message and "endif" in e.value.message


def test_endif_mo_coi():
    with pytest.raises(TemplateRejected):
        inspect_template(make_docx("{% endif %}"))


# ── Render đầy đủ + hiệu năng ──────────────────────────────────────────────────────
def test_render_mau_nhieu_trang_duoi_1_giay():
    paras = [f"Điều {i}: {{{{ ho_ten }}}} - {{{{ luong_co_ban }}}} - {{{{ ten_cong_ty }}}} " + "Lorem ipsum " * 40
             for i in range(60)]
    data = make_docx(*paras)
    start = time.monotonic()
    out = render(data, sample_context())
    assert time.monotonic() - start < 1.0
    assert "Nguyễn Văn An" in docx_text(out) and "{{" not in docx_text(out)


def test_render_kiem_lai_goi_khong_tin_tep_da_luu():
    """Tệp lưu bị thay bằng gói có macro thì render cũng từ chối."""
    bad = rewrap_zip(make_docx("{{ ho_ten }}"), extra={"word/vbaProject.bin": b"x"})
    with pytest.raises(TemplateRejected):
        render(bad, sample_context())


# ── Ngữ cảnh ─────────────────────────────────────────────────────────────────────
def _objs(**over):
    contract = SimpleNamespace(
        code="HDLD001", contract_no="", contract_type=2, sign_date=date(2026, 10, 1),
        start_date=date(2026, 10, 1), end_date=date(2027, 9, 30), job_title="Mua hàng",
        work_location="HN", note="", base_salary=15_000_000, insurance_salary=12_000_000,
        allowance=2_000_000, allowance_note="Ăn trưa")
    employee = SimpleNamespace(
        full_name="An", code="NV1", gender=1, date_of_birth=date(1990, 2, 1), place_of_birth="HN",
        ethnicity="Kinh", id_number="1", id_issue_date=None, id_issue_place="", permanent_address="a",
        current_address="b", phone="0", personal_email="", email="w@x", tax_code="", social_insurance_no="",
        bank_account_no="", bank_name="", bank_branch="", education_level=4, major="")
    company = SimpleNamespace(name="Cty", short_name="C", tax_code="1", address="d",
                              legal_rep=SimpleNamespace(full_name="Bình"), legal_rep_title="GĐ")
    department = SimpleNamespace(name="Mua hàng")
    for k, v in over.items():
        setattr({"contract": contract, "employee": employee, "company": company}[k.split("__")[0]], k.split("__")[1], v)
    return contract, employee, company, department


def test_build_context_du_khoa_va_toan_chuoi():
    ctx = build_context(*_objs(), today=date(2026, 10, 5))
    assert set(ctx) == KNOWN_KEYS
    assert all(isinstance(v, str) for v in ctx.values())
    assert ctx["so_hop_dong"] == "HDLD001"  # contract_no rỗng → dùng code
    assert ctx["luong_co_ban"] == "15.000.000" and ctx["tong_thu_nhap"] == "17.000.000"
    assert ctx["tong_thu_nhap_bang_chu"] == "Mười bảy triệu đồng chẵn"
    assert ctx["thoi_han"] == "12 tháng" and ctx["email"] == "w@x" and ctx["gioi_tinh"] == "Nam"
    assert (ctx["ngay_ky_ngay"], ctx["ngay_ky_thang"], ctx["ngay_ky_nam"]) == ("01", "10", "2026")


def test_build_context_thieu_het_khong_ra_none():
    contract, employee, _, _ = _objs()
    contract.sign_date = contract.end_date = None
    contract.base_salary = contract.allowance = contract.insurance_salary = None
    employee.gender = employee.education_level = 0
    employee.date_of_birth = None
    ctx = build_context(contract, employee, None, None, today=date(2026, 10, 5))
    assert set(ctx) == KNOWN_KEYS
    assert not any("None" in v for v in ctx.values())
    assert ctx["luong_co_ban_bang_chu"] == "Không đồng" and ctx["ten_cong_ty"] == "" and ctx["nguoi_dai_dien"] == ""


def test_build_context_loai_khong_ro_hoac_ma_la():
    contract, employee, company, dept = _objs()
    contract.contract_type = 77
    assert build_context(contract, employee, company, dept, date(2026, 1, 1))["loai_hop_dong"] == ""


def test_build_context_so_am_khong_nhan_doi_dau():
    contract, employee, company, dept = _objs()
    contract.base_salary = -5
    ctx = build_context(contract, employee, company, dept, date(2026, 1, 1))
    assert ctx["luong_co_ban_bang_chu"] == "Không đồng"


@pytest.mark.parametrize("start,end,ctype,expected", [
    (date(2026, 10, 1), date(2027, 9, 30), 2, "12 tháng"),
    (date(2026, 1, 15), date(2026, 7, 14), 1, "6 tháng"),
    (date(2026, 1, 31), date(2026, 3, 1), 1, "1 tháng 2 ngày"),
    (date(2026, 10, 1), date(2026, 10, 20), 1, "20 ngày"),
    (date(2026, 10, 1), date(2026, 10, 1), 1, "1 ngày"),
    (date(2026, 10, 1), date(2026, 9, 1), 2, ""),          # kết thúc trước bắt đầu
    (date(2026, 10, 1), None, 2, ""),
    (date(2026, 10, 1), None, 3, "Không xác định thời hạn"),
    (date(2026, 10, 1), date(2030, 1, 1), 3, "Không xác định thời hạn"),  # loại 3 bỏ qua end_date
    (date(2024, 2, 29), date(2025, 2, 28), 2, "12 tháng"),  # năm nhuận
])
def test_thoi_han(start, end, ctype, expected):
    assert describe_duration(start, end, ctype) == expected


@pytest.mark.parametrize("n,expected", [(0, "0"), (None, "0"), (999, "999"), (1000, "1.000"), (15_000_000, "15.000.000")])
def test_fmt_money(n, expected):
    assert fmt_money(n) == expected
