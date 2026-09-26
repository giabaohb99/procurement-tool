"""bao-CR-493 — đợt 1 yêu cầu phòng Thu mua (chị Mi, FR-PROC-2026-001) cho Tra cứu giá hải quan.

```
hai cột VND: 7% tạm tính · theo thuế suất XNK dòng   -> gốc là GIÁ HIỆU LỰC × tỷ giá USD, tròn đồng
thiếu tỷ giá -> cả hai trống; thiếu thuế suất -> chỉ cột theo dòng trống (không lén điền 7%)
doanh nghiệp / đối tác / lô chọn NHIỀU               -> lặp trên URL hoặc «1,2,3», nhiều là HOẶC
khoảng đơn giá lọc trên giá HIỆU LỰC                  -> dòng có giá điều chỉnh thì so giá điều chỉnh
xuất Excel                                            -> 32 cột nguồn + 2 suy ra + 2 VND = 36
```
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl

from app.modules.customs import service as S
from app.modules.customs.constants import FLAT_IMPORT_TAX_RATE
from app.modules.customs.model import CustomsLine, CustomsParty
from app.modules.customs.constants import PartyType
from app.modules.import_tool.model import ImportBatch, ImportMode, ImportModule, ImportStatus


def _line(db, d=date(2026, 3, 1), name="ATRAZINE 97% TECH", price=3.0, adj=None, usd_rate=26130,
          rate_import=None, currency="USD", fx_rate=None, qty=100, **kw):
    ln = CustomsLine(reg_date=d, product_name=name, unit_code=kw.pop("unit", "KGM"),
                     price_usd=Decimal(str(price)), adj_price_usd=Decimal(str(adj)) if adj is not None else None,
                     usd_rate=Decimal(str(usd_rate)) if usd_rate is not None else None,
                     fx_rate=Decimal(str(fx_rate)) if fx_rate is not None else None,
                     rate_import=Decimal(str(rate_import)) if rate_import is not None else None,
                     currency=currency, quantity=Decimal(str(qty)), hs_code="38089319", **kw)
    db.add(ln)
    db.flush()
    return ln


# ── Hai cột VND ─────────────────────────────────────────────────────────────
def test_vnd_uses_effective_price_and_rounds_to_dong():
    ln = CustomsLine(price_usd=Decimal("3.0"), adj_price_usd=Decimal("2.8"), usd_rate=Decimal("26130"),
                     rate_import=Decimal("3"), currency="USD")
    flat, line_tax = S.compute_vnd_prices(ln)
    assert flat == round(2.8 * 26130 * (1 + FLAT_IMPORT_TAX_RATE))
    assert line_tax == round(2.8 * 26130 * 1.03)
    assert flat != line_tax


def test_vnd_missing_rate_blanks_both_but_missing_line_tax_blanks_only_second():
    assert S.compute_vnd_prices(CustomsLine(price_usd=Decimal("3"), usd_rate=None, currency="EUR")) == (None, None)
    flat, line_tax = S.compute_vnd_prices(CustomsLine(price_usd=Decimal("3"), usd_rate=Decimal("26000"),
                                                      rate_import=None))
    assert flat == round(3 * 26000 * 1.07) and line_tax is None


def test_vnd_falls_back_to_fx_rate_only_for_usd_lines():
    usd = CustomsLine(price_usd=Decimal("2"), fx_rate=Decimal("25000"), currency="USD")
    eur = CustomsLine(price_usd=Decimal("2"), fx_rate=Decimal("28000"), currency="EUR")
    assert S.compute_vnd_prices(usd)[0] == round(2 * 25000 * 1.07)
    assert S.compute_vnd_prices(eur) == (None, None)     # tỷ giá EUR không nhân được với giá USD
    assert S.compute_vnd_prices(CustomsLine(price_usd=Decimal("2"), usd_rate=Decimal("0"))) == (None, None)


def test_serialize_and_export_carry_the_two_vnd_columns(db):
    _line(db, price=3.0, rate_import=5)
    db.commit()
    row = S.serialize_lines(db, db.query(CustomsLine).all())[0]
    assert row["price_vnd_flat"] == round(3 * 26130 * 1.07)
    assert row["price_vnd_line_tax"] == round(3 * 26130 * 1.05)
    ws = openpyxl.load_workbook(io.BytesIO(S.export_lines_xlsx(db, {}))).active
    assert ws.max_column == 36
    assert ws.cell(1, 35).value.startswith("Đơn giá quy đổi VND (thuế NK 7%)")
    assert ws.cell(2, 35).value == row["price_vnd_flat"] and ws.cell(2, 36).value == row["price_vnd_line_tax"]


# ── Bộ lọc mới ──────────────────────────────────────────────────────────────
def test_id_list_accepts_repeat_list_comma_string_and_drops_junk():
    assert S._id_list(["3", "5"]) == [3, 5]
    assert S._id_list("3,5, x,0,-1") == [3, 5]
    assert S._id_list(7) == [7]
    assert S._id_list("") == [] and S._id_list(None) == [] and S._id_list([]) == []


def _parties(db):
    a = CustomsParty(party_type=PartyType.DOMESTIC, dedupe_key="0101", tax_code="0101", name="A")
    b = CustomsParty(party_type=PartyType.DOMESTIC, dedupe_key="0202", tax_code="0202", name="B")
    c = CustomsParty(party_type=PartyType.DOMESTIC, dedupe_key="0303", tax_code="0303", name="C")
    db.add_all([a, b, c])
    db.flush()
    return a, b, c


def _names(db, f):
    return sorted(r.product_name for r in S.apply_line_filters(db.query(CustomsLine), f).all())


def test_many_importers_in_one_box_mean_or(db):
    a, b, c = _parties(db)
    _line(db, name="của A", importer_id=a.id)
    _line(db, name="của B", importer_id=b.id)
    _line(db, name="của C", importer_id=c.id)
    db.commit()
    assert _names(db, {"importer_id": [str(a.id), str(c.id)]}) == ["của A", "của C"]
    assert _names(db, {"importer_id": f"{a.id},{b.id}"}) == ["của A", "của B"]
    assert _names(db, {"importer_id": a.id}) == ["của A"]          # đường gọi cũ: một số
    assert len(_names(db, {"importer_id": []})) == 3               # rỗng = không lọc


def test_price_range_filters_on_effective_price(db):
    _line(db, name="khai 5 chỉnh 2", price=5.0, adj=2.0)
    _line(db, name="khai 3", price=3.0)
    _line(db, name="khai 9", price=9.0)
    db.commit()
    assert _names(db, {"price_min": "2.5", "price_max": "4"}) == ["khai 3"]
    assert _names(db, {"price_max": "2"}) == ["khai 5 chỉnh 2"]     # giá điều chỉnh thắng giá khai
    assert _names(db, {"price_min": "abc"}) and len(_names(db, {"price_min": "abc"})) == 3   # rác → bỏ qua


def test_quantity_rate_currency_incoterm_and_batch_filters(db):
    _line(db, name="lô 1 EUR", qty=50, usd_rate=26000, currency="EUR", incoterm="CIF", batch_id=1)
    _line(db, name="lô 2 USD", qty=500, usd_rate=26500, currency="USD", incoterm="FOB", batch_id=2)
    db.commit()
    assert _names(db, {"qty_min": "100"}) == ["lô 2 USD"]
    assert _names(db, {"rate_max": "26100"}) == ["lô 1 EUR"]
    assert _names(db, {"currency": "eur"}) == ["lô 1 EUR"]
    assert _names(db, {"incoterm": "fob"}) == ["lô 2 USD"]
    assert _names(db, {"batch_id": ["2"]}) == ["lô 2 USD"]


def test_options_expose_currencies_incoterms_and_applied_batches(db):
    done = ImportBatch(module=ImportModule.CUSTOMS_DECLARATION, mode=ImportMode.APPLY,
                       status=ImportStatus.DONE, filename="1.xls", created_count=10)
    dry = ImportBatch(module=ImportModule.CUSTOMS_DECLARATION, mode=ImportMode.DRY_RUN,
                      status=ImportStatus.DONE, filename="thu.xls", created_count=0)
    db.add_all([done, dry])
    db.flush()
    _line(db, currency="USD", incoterm="CIF", batch_id=done.id)
    db.commit()
    opts = S.list_options(db)
    assert [o["value"] for o in opts["currencies"]] == ["USD"]
    assert [o["value"] for o in opts["incoterms"]] == ["CIF"]
    assert [b["value"] for b in opts["batches"]] == [str(done.id)]   # lô chạy thử không vào ô lọc
    assert opts["batches"][0]["label"].endswith("1.xls")
