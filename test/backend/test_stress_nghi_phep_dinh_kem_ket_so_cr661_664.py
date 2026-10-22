"""ÉP TẢI — CR-661 (kh-HR1-2) + CR-664 (kh-HR2-3): ĐÍNH KÈM BẮT BUỘC + ĐỐI CHIẾU KẾT SỔ

Stress test các yêu cầu:
  1. CR-661: Loại nghỉ có require_attachment=True phải có tệp khi gửi duyệt, không chặn khi lưu nháp.
     Ép tải: nhiều loại, đơn nhiều loại (chỉ cần dòng có `require_attachment=True` bị chặn),
     thêm/xóa/thêm lại tệp, gửi duyệt lần 2+.
  2. CR-664: Bảng đối chiếu phép cuối năm có tính đúng số, xử lý edge case (negative, zero, huge),
     không rò bản ngoài phạm vi, query count phải cố định khi N tăng.
"""
from datetime import date, timedelta
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.core.auth import get_perm_profile, perm_cache_clear
from app.modules.attachment.model import FileLink
from app.modules.employee.model import Employee
from app.modules.leave import balance_controller as bc
from app.modules.leave import carryover_review_service as review
from app.modules.leave import carryover_service, request_service
from app.modules.leave.balance_model import LeaveBalance
from app.modules.leave.catalog_model import LeaveType
from app.modules.leave.constants import YEAR_END_CARRY, YEAR_END_CONVERT, YEAR_END_DROP
from app.modules.leave.schema import LeaveLineItem, LeaveRequestCreate, LeaveRequestUpdate

from lich_lam_viec_factory import count_queries

MONDAY = date(2026, 1, 5)


# ── CR-661: ĐÍNH KÈM BẮT BUỘC ──────────────────────────────────────────────────

def _type(db, code, name, **kw):
    kw.setdefault("counts_balance", True)
    kw.setdefault("annual_quota_days", 30.0)
    obj = LeaveType(code=code, name=name, **kw)
    db.add(obj)
    db.flush()
    return obj


def _emp(db, code, name):
    obj = Employee(code=code, full_name=name, company_id=1, department_id=7)
    db.add(obj)
    db.flush()
    return obj


def _attach(db, obj, entity="leave_request"):
    db.add(FileLink(file_id=1, entity=entity, entity_id=obj.id, created_by=1, updated_by=1))
    db.flush()


def _draft(db, user, lines):
    """Lưu nháp đơn với danh sách (loại nghỉ, số ngày)."""
    return request_service.create(db, LeaveRequestCreate(
        lines=[LeaveLineItem(leave_type_id=lt.id, days=d) for lt, d in lines],
        from_date=MONDAY, to_date=MONDAY + timedelta(days=2), reason="Test"), user)


@pytest.fixture
def cr661_world(db):
    """Cảnh CR-661: 10 loại nghỉ (5 bắt buộc, 5 không bắt buộc)."""
    types = {}
    for i in range(5):
        types[f"req{i}"] = _type(db, f"req{i}", f"Loại bắt {i}", require_attachment=True)
    for i in range(5):
        types[f"opt{i}"] = _type(db, f"opt{i}", f"Loại tùy {i}", require_attachment=False)

    user = SimpleNamespace(id=1, employee_id=_emp(db, "EMP1", "Người 1").id)
    return SimpleNamespace(db=db, user=user, types=types)


def test_cr661_mot_loai_bat_co_tep_qua(cr661_world):
    """Một loại bắt buộc → có tệp → qua."""
    w = cr661_world
    obj = _draft(w.db, w.user, [(w.types["req0"], 1)])
    _attach(w.db, obj)
    assert request_service.prepare_submit(w.db, obj, w.user) is not None


def test_cr661_mot_loai_bat_khong_tep_chan(cr661_world):
    """Một loại bắt buộc → không tệp → chặn, ghi đúng tên loại."""
    w = cr661_world
    obj = _draft(w.db, w.user, [(w.types["req0"], 1)])
    with pytest.raises(Exception) as exc:
        request_service.prepare_submit(w.db, obj, w.user)
    assert "Loại bắt 0" in str(exc.value.detail) and "đính kèm" in str(exc.value.detail)


def test_cr661_loai_khong_bat_khong_can_tep(cr661_world):
    """Loại không bắt buộc → không tệp → qua (điểm nguy: bỏ qua cách chỉ lọc loại bắt)."""
    w = cr661_world
    obj = _draft(w.db, w.user, [(w.types["opt0"], 1)])
    assert request_service.prepare_submit(w.db, obj, w.user) is not None


def test_cr661_don_nhieu_loai_chi_loai_bat_can_tep(cr661_world):
    """Đơn 3 loại: 1 bắt + 2 không bắt → chỉ loại bắt yêu tệp, 2 loại không bắt tự qua."""
    w = cr661_world
    obj = _draft(w.db, w.user, [
        (w.types["req0"], 1),
        (w.types["opt0"], 1),
        (w.types["opt1"], 1)
    ])
    with pytest.raises(Exception):
        request_service.prepare_submit(w.db, obj, w.user)
    _attach(w.db, obj)
    assert request_service.prepare_submit(w.db, obj, w.user) is not None


def test_cr661_luu_nhap_khong_bi_chan(cr661_world):
    """Lưu nháp tệp thiếu phải qua (chặn chỉ lúc gửi duyệt)."""
    w = cr661_world
    obj = _draft(w.db, w.user, [(w.types["req0"], 1)])
    assert obj.id > 0


def test_cr661_them_xoa_them_tep_va_gui_lai(cr661_world):
    """Thêm tệp → gửi → OK; tệp bị xóa → lưu → OK; thêm lại → gửi lần 2 → OK.

    Trường hợp: phòng gửi lại đơn sau khi tệp bị xóa nhầm, phải gắn lại.
    """
    w = cr661_world
    obj = _draft(w.db, w.user, [(w.types["req0"], 1)])
    _attach(w.db, obj)
    request_service.prepare_submit(w.db, obj, w.user)  # gửi lần 1 ✓

    # Xóa tệp
    w.db.query(FileLink).filter(FileLink.entity_id == obj.id).delete()
    w.db.flush()

    # Lưu lại nháp → phải qua (refetch object after changes)
    obj = w.db.get(obj.__class__, obj.id)
    request_service.update(w.db, obj, LeaveRequestUpdate(reason="Sửa lý do"), w.user)

    # Gắn tệp lại + gửi lần 2
    _attach(w.db, obj)
    obj = w.db.get(obj.__class__, obj.id)
    assert request_service.prepare_submit(w.db, obj, w.user) is not None


def test_cr661_nhieu_don_cung_loai_bat_buoc(cr661_world):
    """10 đơn, mỗi đơn 1 dòng loại bắt, không tệp → tất cả đều bị chặn."""
    w = cr661_world
    for i in range(10):
        obj = _draft(w.db, w.user, [(w.types["req0"], 1)])
        with pytest.raises(Exception):
            request_service.prepare_submit(w.db, obj, w.user)


def test_cr661_nhieu_don_tep_them_xoa_lien_tuc(cr661_world):
    """5 đơn, mỗi đơn thêm-xóa-thêm tệp, sau đó gửi duyệt."""
    w = cr661_world
    objs = []
    for i in range(5):
        obj = _draft(w.db, w.user, [(w.types["req0"], 1)])
        objs.append(obj)

    for i, obj in enumerate(objs):
        _attach(w.db, obj, entity="leave_request")
        # Xóa lại
        w.db.query(FileLink).filter(FileLink.entity_id == obj.id).delete()
        # Thêm lần 2
        _attach(w.db, obj)
        # Gửi
        request_service.prepare_submit(w.db, obj, w.user)


# ── CR-664: ĐỐI CHIẾU KẾT SỔ ───────────────────────────────────────────────────

def _type_cr664(db, code, name, **kw):
    kw.setdefault("counts_balance", True)
    kw.setdefault("annual_quota_days", 12.0)
    kw.setdefault("year_end_mode", YEAR_END_DROP)
    obj = LeaveType(code=code, name=name, **kw)
    db.add(obj)
    db.flush()
    return obj


def _emp_cr664(db, code, name, company=1):
    obj = Employee(code=code, full_name=name, company_id=company, hire_date=date(2020, 1, 1))
    db.add(obj)
    db.flush()
    return obj


def _bal(db, emp, lt, year=2026, **kw):
    kw.setdefault("allocated_days", 12.0)
    obj = LeaveBalance(employee_id=emp.id, year=year, leave_type_id=lt.id,
                       company_id=emp.company_id, **kw)
    db.add(obj)
    db.flush()
    return obj


def _close(db, year=2026):
    rows = db.query(LeaveBalance).filter(LeaveBalance.year == year).all()
    return carryover_service.close_year(db, year, rows, actor=1)


def _review(db, user, year=2026):
    return review.build_review(db, user, get_perm_profile(db, user), year)


@pytest.fixture
def cr664_world(db, cap_quyen):
    """Cảnh CR-664: 3 loại nghỉ + nhiều nhân sự."""
    comp_off = _type_cr664(db, "comp_off", "Nghỉ bù", annual_quota_days=0.0)
    annual = _type_cr664(db, "annual", "Phép năm", year_end_mode=YEAR_END_CARRY,
                        carry_over_max_days=5.0)
    swap = _type_cr664(db, "swap", "Phép đổi", year_end_mode=YEAR_END_CONVERT,
                      convert_to_type_id=comp_off.id, convert_ratio=0.5)
    cap_quyen(1, "leave_balance", scope="all", read=True, export=True)
    db.commit()
    return SimpleNamespace(
        db=db, annual=annual, swap=swap, comp_off=comp_off,
        hr=SimpleNamespace(id=1, employee_id=0)
    )


def test_cr664_so_cu_so_moi_chenh_khop_voi_so_quy(cr664_world):
    """Số cũ = còn lại + đã chuyển, số mới = còn lại (quy định CR-664)."""
    w = cr664_world
    emp = _emp_cr664(w.db, "E001", "Nhân 1")
    _bal(w.db, emp, w.annual, used_days=4.0)  # 12 - 4 = 8 còn, 5 mang → 3 sau
    _close(w.db)

    rows = _review(w.db, w.hr)["items"]
    assert len(rows) == 1
    row = rows[0]
    assert row["before_days"] == 8.0
    assert row["after_days"] == 3.0
    assert row["diff_days"] == -5.0
    assert row["credited_days"] == 5.0


def test_cr664_quy_doi_nhan_ty_le(cr664_world):
    """Phép đổi: 8 phút → 4 nghỉ bù (× 0.5 tỷ lệ)."""
    w = cr664_world
    emp = _emp_cr664(w.db, "E002", "Nhân 2")
    _bal(w.db, emp, w.swap, used_days=4.0)  # 12 - 4 = 8
    _close(w.db)

    row = _review(w.db, w.hr)["items"][0]
    assert row["before_days"] == 8.0
    assert row["after_days"] == 0.0
    assert row["ratio"] == 0.5
    assert row["credited_days"] == 4.0


def test_cr664_zero_used_toan_so_mang(cr664_world):
    """Used = 0 → còn 12, có 5 mang → 7 sau."""
    w = cr664_world
    emp = _emp_cr664(w.db, "E003", "Nhân 3")
    _bal(w.db, emp, w.annual, used_days=0.0)
    _close(w.db)

    row = _review(w.db, w.hr)["items"][0]
    assert row["before_days"] == 12.0
    assert row["after_days"] == 7.0
    assert row["diff_days"] == -5.0


def test_cr664_negative_used_la_loi_du_lieu(cr664_world):
    """Có nhân sự used < 0 (dữ liệu cũ lỗi) → tính đúng (-5 = còn 17)."""
    w = cr664_world
    emp = _emp_cr664(w.db, "E004", "Nhân 4")
    _bal(w.db, emp, w.annual, used_days=-5.0)  # không bao giờ có, nhưng test edge
    _close(w.db)

    row = _review(w.db, w.hr)["items"][0]
    assert row["before_days"] == 17.0
    assert row["after_days"] == 12.0  # 17 - 5 (trần)


def test_cr664_massive_balance(cr664_world):
    """Nhân sự có 1000 ngày phép (dữ liệu quấy): trần mang 5 → còn 995."""
    w = cr664_world
    emp = _emp_cr664(w.db, "E005", "Nhân 5")
    _bal(w.db, emp, w.annual, allocated_days=1000.0, used_days=0.0)
    _close(w.db)

    row = _review(w.db, w.hr)["items"][0]
    assert row["before_days"] == 1000.0
    assert row["after_days"] == 995.0
    assert row["credited_days"] == 5.0


def test_cr664_loai_drop_khong_vao_bang_doi_chieu(cr664_world):
    """Loại `YEAR_END_DROP` không chuyển, không vào bảng đối chiếu (nó bỏ đi)."""
    w = cr664_world
    emp = _emp_cr664(w.db, "E006", "Nhân 6")
    drop = _type_cr664(w.db, "drop", "Không lương")
    _bal(w.db, emp, drop, used_days=5.0)
    _close(w.db)

    rows = _review(w.db, w.hr)["items"]
    assert len(rows) == 0  # loại drop không lấy xuống


def test_cr664_50_nhan_su_so_truy_van_co_dinh(cr664_world):
    """50 nhân sự → số truy vấn phải cố định (không N+1)."""
    w = cr664_world

    # Tạo 5 nhân sự + kết sổ
    emps = []
    for i in range(5):
        emp = _emp_cr664(w.db, f"E0A{i:02d}", f"Nhân A{i}")
        _bal(w.db, emp, w.annual, used_days=i)
        emps.append(emp)
    _close(w.db)
    _review(w.db, w.hr)  # làm nóng cache quyền

    _, few = count_queries(w.db, lambda: _review(w.db, w.hr))

    # Thêm 50 nhân sự nữa
    for i in range(50):
        emp = _emp_cr664(w.db, f"E0B{i:02d}", f"Nhân B{i}")
        _bal(w.db, emp, w.annual, used_days=i % 10)
    _close(w.db)

    result, many = count_queries(w.db, lambda: _review(w.db, w.hr))

    # 5 + 50 = 55 người, all có annual type
    assert len(result["items"]) == 55
    # Không quá 5 truy vấn (có độc lập với N)
    assert len(many) == len(few) <= 5, f"few={len(few)}, many={len(many)}"


def test_cr664_pham_vi_du_lieu_khong_ro(cr664_world, cap_quyen):
    """Phạm vi dữ liệu: tài khoản X chỉ thấy công ty 1, không thấy công ty 2."""
    w = cr664_world

    # Công ty 1
    emp1 = _emp_cr664(w.db, "E_C1", "Nhân công ty 1", company=1)
    _bal(w.db, emp1, w.annual)

    # Công ty 2
    emp2 = _emp_cr664(w.db, "E_C2", "Nhân công ty 2", company=2)
    _bal(w.db, emp2, w.annual)

    _close(w.db)

    # User toàn bộ → 2 dòng
    all_rows = _review(w.db, w.hr)["items"]
    assert len(all_rows) == 2

    # User phạm vi công ty 1 → chỉ 1 dòng (phải loại ra công ty 2)
    # Tạo employee cho user 2 ở công ty 1
    emp_user2 = _emp_cr664(w.db, "USR2", "User 2", company=1)
    cap_quyen(2, "leave_balance", scope="company", read=True)
    w.db.commit()
    user1 = SimpleNamespace(id=2, employee_id=emp_user2.id, company_id=1)
    perm_cache_clear(2)

    scoped_rows = _review(w.db, user1)["items"]
    assert len(scoped_rows) == 1
    assert scoped_rows[0]["employee_code"] == "E_C1"


def test_cr664_run_lai_khong_doi(cr664_world):
    """Chạy kết sổ lần 2 không tạo dòng mới (idempotent)."""
    w = cr664_world
    emp = _emp_cr664(w.db, "E007", "Nhân 7")
    _bal(w.db, emp, w.annual, used_days=3.0)

    _close(w.db)
    first = _review(w.db, w.hr)

    _close(w.db)
    second = _review(w.db, w.hr)

    assert first == second


def test_cr664_excel_export_mo_duoc_va_co_so(cr664_world):
    """Xuất Excel: file mở được, có tiêu đề + dữ liệu."""
    w = cr664_world
    emp = _emp_cr664(w.db, "E008", "Nhân 8")
    _bal(w.db, emp, w.annual, used_days=2.0)
    _close(w.db)

    res = bc.close_year_review_export(year=2026, db=w.db, user=w.hr)
    sheet = load_workbook(BytesIO(res.body)).active

    assert sheet.max_row == 2  # tiêu đề + 1 dòng
    assert sheet.cell(1, 4).value.startswith("Số cũ")
    assert "doi-chieu" in res.headers["content-disposition"]


def test_cr664_khong_co_du_lieu_nam_cu_thi_khong_loi(cr664_world):
    """Xem năm chưa có dữ liệu (1, 2025, 9999) → không lỗi, trả rỗng."""
    for year in (1, 2025, 9999):
        data = _review(cr664_world.db, cr664_world.hr, year=year)
        assert data["items"] == []
        assert data["totals"]["row_count"] == 0
