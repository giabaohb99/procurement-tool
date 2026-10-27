"""duoc-CR-688 — báo cáo H1 «Phúc lợi & sự kiện» (khóa 22): `/api/benefit-audience/summary`.

Canh: ngày dịp = cuối kỳ, từng bộ lọc, người đang nghỉ, đếm con, ngân sách hai mức, 403 lọc nhạy cảm,
422 đầu vào, phạm vi, gác khóa báo cáo ở cả ba đường, Excel và số truy vấn CỐ ĐỊNH.
Mọi người thử nằm ở phòng B.hc (trống trong thế giới mẫu) và được lọc bằng `department_id`.
"""
import uuid
from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event

from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.report_period import parse_period
from app.core.scoping import apply_scope
from app.core.subject_match import SUBJECT_EMPLOYEE
from app.main import app
from app.modules.employee import benefit_report_service as svc
from app.modules.employee.benefit_report_inputs import Filters
from app.modules.employee.contact_model import EmployeeFamily
from app.modules.employee.model import Employee
from app.modules.labor_contract.model import LaborContract
from app.modules.leave.request_model import LeaveRequest
from app.modules.report_access.model import ReportAccess

D = date(2026, 6, 15)           # ngày dịp
URL = "/api/benefit-audience/summary"
QS = {"preset": "custom", "date_from": "2026-06-15", "date_to": "2026-06-15", "compare": "none"}
_n = {"v": 0}


@pytest.fixture
def client_as(db):
    def make(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def _emp(db, world, tag, **kw) -> Employee:
    base = dict(code=f"P{tag}", full_name=f"Người {tag}", company_id=world.co["B"],
                department_id=world.dept["B.hc"], is_active=True, hire_date=date(2023, 3, 1),
                job_level=1, position_id=1, position="Chuyên viên")
    base.update(kw)
    e = Employee(**base)
    db.add(e)
    db.flush()
    return e


def _contract(db, emp, ctype):
    _n["v"] += 1
    db.add(LaborContract(code=f"HDP{_n['v']}", employee_id=emp.id, company_id=emp.company_id,
                         contract_type=ctype, status=2, start_date=date(2026, 1, 1),
                         end_date=date(2026, 12, 31)))
    db.flush()


def _kid(db, emp, dob, relation=6):
    db.add(EmployeeFamily(employee_id=emp.id, full_name="Con kín", relation=relation, date_of_birth=dob))
    db.flush()


def _dataset(db, world):
    _emp(db, world, "mgr", gender=1, job_level=4, date_of_birth=date(1990, 6, 10), position_id=5,
         position="Trưởng phòng", hire_date=date(2021, 6, 1))                    # 5 năm, tháng 6
    s = _emp(db, world, "staff", gender=2, date_of_birth=date(1985, 1, 1), hire_date=date(2023, 6, 20))
    _kid(db, s, date(2020, 1, 1))                                               # con 6 tuổi
    _kid(db, s, date(2010, 1, 1), relation=7)                                   # con gái 16 tuổi
    _kid(db, s, date(2022, 1, 1), relation=4)                                   # vợ: KHÔNG phải con
    _emp(db, world, "future", hire_date=date(2027, 1, 1))
    _emp(db, world, "gone", status="resigned", resign_date=date(2026, 6, 1))
    _emp(db, world, "leaving", status="resigned", resign_date=date(2026, 7, 1))   # còn làm vào D
    _contract(db, _emp(db, world, "svc"), 4)
    lv = _emp(db, world, "leave")
    db.add(LeaveRequest(code="NP-T-1", employee_id=lv.id, company_id=lv.company_id,
                        department_id=lv.department_id, from_date=date(2026, 6, 14), to_date=date(2026, 6, 16),
                        status=3, is_deleted=False))
    db.commit()


def _actor(db, world, client_as, *, sensitive=True, scope="all", report=True, key="a1", leave=True):
    world.grant(key, "employee", scope, actions=("read", "export"))
    if leave:
        world.grant(key, "leave_request", "all", actions=("read",))
    if sensitive:
        world.grant(key, "employee_sensitive", "all", actions=("read",))
    actor = world.actor(key)
    if report:
        db.add(ReportAccess(report_key=int(ReportKey.BENEFIT_AUDIENCE), subject_kind=SUBJECT_EMPLOYEE,
                            subject_id=actor.employee.id, effect=1, reason="t", created_by=0, updated_by=0))
        db.commit()
    return actor, client_as(actor.user)


def _get(client, world, path=URL, **extra):
    return client.get(path, params={**QS, "department_id": world.dept["B.hc"], **extra})


def _totals(client, world, **extra):
    r = _get(client, world, **extra)
    assert r.status_code == 200, r.text
    return r.json()["data"]["totals"]["current"]


def _codes(client, world, **extra):
    r = _get(client, world, URL + "/rows", metric="people", **extra)
    assert r.status_code == 200, r.text
    return {x["employee_code"] for x in r.json()["data"]["rows"]}


def test_ngay_dip_va_loc_mac_dinh(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    # mgr, staff, leaving, leave còn làm; future/gone ngoài ngày dịp; svc bị loại mặc định
    assert _codes(client, world) == {"Pmgr", "Pstaff", "Pleaving", "Pleave"}
    assert "Psvc" in _codes(client, world, exclude_contract_types="")
    d = _get(client, world).json()["data"]
    assert d["trend"] == [] and d["meta"]["drilldown_metrics"] == ["people"]
    assert {m["key"] for m in d["meta"]["metrics"]} == {"people", "managers", "staff", "budget"}


def test_nguoi_dang_nghi(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    assert "Pleave" in _codes(client, world, include_on_leave="1")
    assert "Pleave" not in _codes(client, world, include_on_leave="0")


def test_bo_nguoi_dang_nghi_can_quyen_xem_don_nghi_phep(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as, leave=False)
    r = _get(client, world, URL + "/rows", metric="people", include_on_leave="0")
    assert r.status_code == 403 and "nghỉ phép" in r.json()["error"]["message"]
    assert "Pleave" in _codes(client, world, include_on_leave="1")      # tính cả người nghỉ thì không cần quyền


def test_thieu_ngay_vao_lam_lay_ngay_tao_ho_so_va_co_ghi_chu(db, world, client_as):
    from datetime import datetime
    _dataset(db, world)
    _emp(db, world, "nohire", hire_date=None, created_at=datetime(2026, 5, 1, 3, 0))     # tạo trước ngày dịp
    _emp(db, world, "nohirelate", hire_date=None, created_at=datetime(2026, 7, 1, 3, 0))  # tạo sau ngày dịp
    _, client = _actor(db, world, client_as)
    assert {"Pnohire"} <= _codes(client, world) and "Pnohirelate" not in _codes(client, world)
    notes = _get(client, world).json()["data"]["notes"]
    assert any("1 người thiếu ngày vào làm" in n for n in notes)


def test_loc_gioi_tuoi_thang_sinh_tham_nien_cap_bac(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    assert _codes(client, world, gender=2) == {"Pstaff"}
    assert _codes(client, world, age_min=35, age_max=40) == {"Pmgr"}          # 36 tuổi; staff 41
    assert _codes(client, world, birth_month=6) == {"Pmgr"}
    assert _codes(client, world, seniority_years="5,10") == {"Pmgr"}          # tháng 6, tròn 5 năm
    assert _codes(client, world, seniority_years="3") == {"Pstaff"}
    assert _codes(client, world, job_level_min=3) == {"Pmgr"}
    assert _codes(client, world, position_id=5) == {"Pmgr"}


def test_dem_con_va_ngan_sach_hai_muc(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    cur = _totals(client, world, child_max_age=10, amount_staff=100000, amount_manager=300000)
    assert cur["people"] == 1 and cur["children"] == 1 and cur["staff"] == 1 and cur["managers"] == 0
    assert cur["budget"] == 100000                                           # chỉ con 6 tuổi, vợ không tính
    cur = _totals(client, world, child_max_age=16, amount_staff=100000)
    assert cur["children"] == 2 and cur["budget"] == 200000
    body = _get(client, world, URL + "/rows", metric="people", child_max_age=16, amount_staff=100000).json()["data"]
    assert [c["key"] for c in body["columns"]][-3:] == ["children", "amount", "signature"]
    assert body["rows"][0]["children"] == 2 and body["rows"][0]["amount"] == 200000
    assert body["rows"][0]["signature"] == ""
    cur = _totals(client, world, amount_staff=100000, amount_manager=300000)  # theo người
    assert cur["managers"] == 1 and cur["staff"] == 3 and cur["budget"] == 600000
    assert "children" not in cur


def test_khong_lo_ngay_sinh_ten_con(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    for path in (URL, URL + "/rows"):
        text = _get(client, world, path, metric="people", child_max_age=16).text
        assert "1985" not in text and "1990" not in text and "Con kín" not in text and "date_of_birth" not in text


@pytest.mark.parametrize("extra", [{"age_min": 30}, {"age_max": 50}, {"birth_month": 6}, {"child_max_age": 5}])
def test_loc_nhay_cam_thieu_quyen_403(db, world, client_as, extra):
    _dataset(db, world)
    _, client = _actor(db, world, client_as, sensitive=False)
    for suffix in ("", "/export", "/rows"):
        r = _get(client, world, URL + suffix, metric="people", **extra)
        assert r.status_code == 403
        assert "employee_sensitive" not in r.text and "quyền xem ngày sinh" in r.text
    assert _get(client, world, gender=2).status_code == 200                   # lọc thường vẫn chạy


@pytest.mark.parametrize("extra", [
    {"amount_staff": -1}, {"amount_manager": 1_000_000_001}, {"amount_staff": "abc"}, {"age_min": 121},
    {"age_min": 50, "age_max": 40}, {"birth_month": 0}, {"birth_month": 13}, {"gender": 9},
    {"job_level_min": 7}, {"seniority_years": "x"}, {"exclude_contract_types": "77"},
    {"include_on_leave": "2"}, {"child_max_age": 121}, {"occasion_label": "x" * 101}, {"group_by": "xyz"}])
def test_dau_vao_sai_422(db, world, client_as, extra):
    _, client = _actor(db, world, client_as)
    assert _get(client, world, **extra).status_code == 422


def test_xem_theo_va_loc_nhom(db, world, client_as):
    _dataset(db, world)
    _, client = _actor(db, world, client_as)
    d = _get(client, world).json()["data"]                                    # mặc định: pháp nhân
    assert [g["label"] for g in d["groups"]] == ["Công ty B"]
    for dim in ("department", "position", "job_level"):
        assert _get(client, world, group_by=dim).json()["data"]["groups"]
    body = _get(client, world, URL + "/rows", metric="people", group_by="job_level", group_key="4").json()["data"]
    assert [r["employee_code"] for r in body["rows"]] == ["Pmgr"]
    assert _get(client, world, URL + "/rows", metric="budget").status_code == 422
    assert _get(client, world, URL + "/rows", metric="people", group_by="company",
                group_key="__other__").status_code == 422


def test_pham_vi_chi_thu_hep(db, world, client_as):
    _emp(db, world, "pvB")
    _emp(db, world, "pvA", company_id=world.co["A"], department_id=world.dept["A.kt"])
    db.commit()
    _, client = _actor(db, world, client_as, scope="dept")                    # a1 thuộc A.kt
    got = {r["employee_code"] for r in client.get(URL + "/rows", params={**QS, "metric": "people"}).json()["data"]["rows"]}
    assert "PpvA" in got and "PpvB" not in got


@pytest.mark.parametrize("suffix", ["", "/export", "/rows"])
def test_chua_giao_bao_cao_403(db, world, client_as, suffix):
    _, client = _actor(db, world, client_as, report=False)
    assert _get(client, world, URL + suffix, metric="people").status_code == 403


def test_excel_ky_nhan_moi_phap_nhan_mot_sheet(db, world, client_as):
    _dataset(db, world)
    _emp(db, world, "coA", company_id=world.co["A"], department_id=world.dept["A.kt"])
    db.commit()
    _, client = _actor(db, world, client_as)
    label = '=HYPERLINK("x")'
    r = client.get(URL + "/export", params={**QS, "occasion_label": label, "amount_staff": 50000,
                                            "amount_manager": 90000})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = load_workbook(BytesIO(r.content))
    assert len(wb.sheetnames) >= 2 and any("Công ty B" in n for n in wb.sheetnames)
    ws = wb[next(n for n in wb.sheetnames if "Công ty B" in n)]
    assert ws["A2"].value == "DANH SÁCH NHẬN QUÀ " + label and ws["A2"].data_type == "s"
    assert [c.value for c in ws[4]][:3] == ["STT", "Mã NV", "Họ tên"] and ws["A5"].value == 1
    vals = [c.value for row in ws.iter_rows() for c in row]
    assert "Tổng tiền" in vals and "Người lập" in vals and "Người duyệt" in vals
    total_row = next(row for row in ws.iter_rows() if row[0].value == "Tổng tiền")
    assert 50000 * 3 + 90000 in [c.value for c in total_row]                  # mgr + staff + leaving + leave
    empty = client.get(URL + "/export", params={**QS, "gender": 3, "department_id": world.dept["B.hc"]})
    assert empty.status_code == 200 and "Không có dữ liệu" in load_workbook(BytesIO(empty.content)).sheetnames


def test_so_truy_van_khong_tang_theo_so_nhan_su(db, world):
    world.grant("a1", "employee", "all", actions=("read",))
    world.grant("a1", "employee_sensitive", "all", actions=("read",))
    actor = world.actor("a1")
    prof = actor.profile()
    period = parse_period(QS)
    flt = Filters(child_max_age=10, include_on_leave=False, age_min=1, amount_staff=1000)

    def count() -> int:
        stmts: list[str] = []

        def listener(conn, cur, stmt, *args):
            stmts.append(stmt)

        engine = db.get_bind()
        event.listen(engine, "before_cursor_execute", listener)
        try:
            q = apply_scope(db.query(Employee), Employee, "employee", actor.user, prof)
            svc.build_summary(db, q, period, "department", flt)
        finally:
            event.remove(engine, "before_cursor_execute", listener)
        return len(stmts)

    def add(n, start):
        for i in range(n):
            e = _emp(db, world, f"q{start + i}", date_of_birth=date(1990, 1, 1))
            _kid(db, e, date(2020, 1, 1))
            _contract(db, e, 2)
            db.add(LeaveRequest(code=f"NP-Q-{start + i}", employee_id=e.id, from_date=D, to_date=D, status=3,
                                is_deleted=False))
        db.commit()

    add(5, 0)
    c5 = count()
    add(45, 5)
    c50 = count()
    assert c5 > 0 and c5 == c50, (c5, c50)
