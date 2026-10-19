"""duoc-CR-641 phase 05 — API đọc hồ sơ: gác kép (setting + ReportKey 14), nội dung, lọc/trang, 404/422, số truy vấn cố định.

Đi qua `TestClient` vì `require_report` gắn ở `dependencies=[...]` — gọi hàm Python trực tiếp sẽ không chạy gác này. Mỗi client
dùng token giả DUY NHẤT (cache `/summary` khóa theo token — xem chú thích fixture `client_as` của test gác đường).
"""
import uuid
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient

from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.subject_match import SUBJECT_ROLE
from app.core.usage_codes import UsageStyle
from app.main import app
from app.modules.export_log.model import ExportLog
from app.modules.report_access.seed_defaults import ensure_report_access_defaults
from app.modules.role.model import Permission, Role
from app.modules.usage_profile.model import UsageDaily, UserUsageLabel, UserUsageSnapshot
from app.modules.usage_profile.snapshot_builder import build_snapshots
from app.modules.user.model import UserRole

from ho_so_lam_viec_factory import make_company, make_dept, make_person
from lich_lam_viec_factory import count_queries
from test_ho_so_lam_viec_snapshot_cr641 import FIRST, WEEK, make_world

ROOT = "/api/usage-profiles"
PERIOD = {"preset": "custom", "date_from": "2026-01-05", "date_to": "2026-02-01"}


@pytest.fixture
def client_as(db):
    def make(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


@pytest.fixture
def admin(db, cap_quyen, gan_bao_cao):
    user, _ = make_person(db, "ADM")
    role = cap_quyen(user.id, "setting", scope="all", read=True, export=True)
    gan_bao_cao(SUBJECT_ROLE, role.id, ReportKey.USAGE_PROFILE)
    db.commit()
    return user


def all_paths(uid=1):
    return [f"{ROOT}/summary", f"{ROOT}/summary/export", f"{ROOT}/cohorts", f"{ROOT}/{uid}"]


def seeded(db, n=6, **kw):
    people, role = make_world(db, n, **kw)
    build_snapshots(db, WEEK)
    return people, role


# ── gác quyền ───────────────────────────────────────────────────────────────

def test_khong_co_setting_read_403_ca_4_duong(db, client_as):
    user, _ = make_person(db, "NV")
    client = client_as(user)
    for path in all_paths(user.id):
        assert client.get(path, params=PERIOD).status_code == 403, path


def test_co_setting_read_nhung_chua_gan_khoa_14_van_403(db, cap_quyen, client_as):
    user, _ = make_person(db, "MGR")
    cap_quyen(user.id, "setting", scope="all", read=True, export=True)
    client = client_as(user)
    for path in all_paths(user.id):
        assert client.get(path, params=PERIOD).status_code == 403, path


def test_vai_tro_admin_sau_ensure_report_access_defaults_duoc_200(db, client_as):
    people, _ = seeded(db)
    admin_user, _ = make_person(db, "BOSS")
    role = Role(code="admin", name="Quản trị hệ thống")
    db.add(role)
    db.flush()
    db.add(UserRole(user_id=admin_user.id, role_id=role.id))
    db.add(Permission(role_id=role.id, entity="setting", scope="all", can_read=True, can_export=True))
    db.commit()
    ensure_report_access_defaults(db)
    client = client_as(admin_user)
    for path in all_paths(people[0][0].id):
        assert client.get(path, params=PERIOD).status_code == 200, path


def test_nhan_vien_thuong_xem_ho_so_cua_chinh_minh_403(db, client_as):
    people, _ = seeded(db)
    me = people[0][0]
    assert client_as(me).get(f"{ROOT}/{me.id}").status_code == 403


# ── danh sách ───────────────────────────────────────────────────────────────

def test_danh_sach_hinh_dang_json_va_nhom_so_sanh(db, admin, client_as):
    people, role = seeded(db)
    resp = client_as(admin).get(f"{ROOT}/summary", params=PERIOD)
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert data["total"] == 6 and data["page"] == 1 and data["page_size"] == 20
    assert data["period"] == {"date_from": "2026-01-05", "date_to": "2026-02-01"}
    assert data["data_since"] == "2026-01-05" and data["label_week"] == "2026-01-26"
    row = next(r for r in data["items"] if r["user_id"] == people[0][0].id)
    assert set(row) == {"user_id", "full_name", "emp_code", "department_name", "company_name", "cohort",
                        "active_days", "workdays_net", "active_ratio", "api_writes", "docs_created", "approvals",
                        "exports", "http_422_per_100_writes", "labels", "last_active_day"}
    assert row["cohort"] == {"kind": 1, "key": role.id, "size": 6, "name": role.name}
    assert row["active_days"] == 24 and row["workdays_net"] == 24.0 and row["active_ratio"] == 1.0
    assert row["api_writes"] == 28 * 4 and row["last_active_day"] == "2026-01-31"     # CN 01/02 không làm việc


def test_loc_phong_ban_q_nhan_trang_2_va_sap_xep(db, admin, client_as):
    people, _ = seeded(db, tag="L1")
    other_co = make_company(db, "OTHER")
    other_dept = make_dept(db, other_co, "OD")
    ou, oe = make_person(db, "ZZ99", other_co, other_dept)
    for d in range(28):
        db.add(UsageDaily(user_id=ou.id, day=FIRST + timedelta(days=d), employee_id=oe.id, is_workday=True,
                          work_capacity=1.0, active=d % 2 == 0, api_writes=d))
    db.add(UserUsageSnapshot(user_id=ou.id, period_start=WEEK, metrics={}))
    db.flush()
    snap_id = db.query(UserUsageSnapshot).filter_by(user_id=ou.id).one().id
    db.add(UserUsageLabel(snapshot_id=snap_id, user_id=ou.id, period_start=WEEK, label=int(UsageStyle.LOW_USAGE),
                          score=0.4, evidence=["active_ratio"]))
    db.commit()
    c = client_as(admin)
    names = lambda params: [r["emp_code"] for r in c.get(f"{ROOT}/summary", params={**PERIOD, **params}).json()["data"]["items"]]
    assert names({"department_id": other_dept.id}) == ["ZZ99"]
    assert names({"company_id": other_co.id}) == ["ZZ99"]
    assert names({"q": "zz9"}) == ["ZZ99"] and names({"q": "L100"}) == ["L100"]
    assert names({"q": "%"}) == []                                  # ký tự LIKE bị escape, không khớp tất cả
    assert names({"label": int(UsageStyle.LOW_USAGE)}) == ["ZZ99"]
    assert names({"label": int(UsageStyle.POWER_USER)}) == []
    page2 = c.get(f"{ROOT}/summary", params={**PERIOD, "page_size": 3, "page": 2, "sort": "emp_code"}).json()["data"]
    assert page2["total"] == 7 and [r["emp_code"] for r in page2["items"]] == ["L103", "L104", "L105"]
    desc = names({"sort": "-api_writes", "page_size": 1})
    assert desc == ["ZZ99"]                                         # 0+1+…+27 = 378 lượt ghi > 112 của mỗi người L1xx
    assert names({"sort": "api_writes", "page_size": 1})[0].startswith("L1")
    assert c.get(f"{ROOT}/summary", params={**PERIOD, "sort": "password"}).status_code == 422


def test_tham_so_ngoai_dai_la_422(db, admin, client_as):
    c = client_as(admin)
    for params in ({"page_size": 101}, {"page": 0}, {"label": 99}, {"q": "x" * 101}, {"preset": "bogus"}):
        assert c.get(f"{ROOT}/summary", params={**PERIOD, **params}).status_code == 422, params
    assert c.get(f"{ROOT}/{admin.id}", params={"weeks": 27}).status_code == 422
    assert c.get(f"{ROOT}/{admin.id}", params={"weeks": 0}).status_code == 422


# ── hồ sơ một người ─────────────────────────────────────────────────────────

def test_ho_so_404_khi_khong_co_nguoi(db, admin, client_as):
    assert client_as(admin).get(f"{ROOT}/999999").status_code == 404


def test_ho_so_hinh_dang_va_khong_co_truong_van_ban_tu_do(db, admin, client_as):
    people, role = seeded(db, tag="P1")
    u = people[0][0]
    db.query(UsageDaily).filter_by(user_id=u.id, day=FIRST).update(
        {"page_views": 3, "screens": {"/a": [3, 900, 1]}})
    db.commit()
    resp = client_as(admin).get(f"{ROOT}/{u.id}", params={"weeks": 26})
    assert resp.status_code == 200
    data = resp.json()["data"]
    assert set(data) == {"subject", "latest", "weekly", "daily", "heatmap", "screens", "screens_source",
                         "metric_meta", "deferred_metrics", "data_since", "tracking_since"}
    assert set(data["subject"]) == {"user_id", "full_name", "emp_code", "position", "department_name", "company_name",
                                    "role_names", "hire_date", "tenure_days"}
    assert data["subject"]["role_names"] == [role.name]
    assert len(data["heatmap"]) == 7 and all(len(r) == 24 for r in data["heatmap"])
    assert data["latest"]["cohort"]["size"] == 6 and data["latest"]["period"]["window_weeks"] == 4
    assert data["latest"]["metrics"]["active_ratio"]["v"] == 1.0
    assert [w["period_start"] for w in data["weekly"]] == ["2026-01-26"]
    assert data["screens"][0]["route"] == "/a" and data["screens_source"] == 1
    assert data["metric_meta"]["active_ratio"]["unit"] == "ratio"
    assert data["tracking_since"] == "2026-01-05"


def test_ho_so_khong_co_snapshot_van_tra_latest_none(db, admin, client_as):
    user, _ = make_person(db, "EMPTY")
    data = client_as(admin).get(f"{ROOT}/{user.id}").json()["data"]
    assert data["latest"] is None and data["weekly"] == [] and data["daily"] == []


def test_nhom_so_sanh_theo_tuan(db, admin, client_as):
    _, role = seeded(db, tag="C1")
    data = client_as(admin).get(f"{ROOT}/cohorts", params={"week": "2026-01-26"}).json()["data"]
    assert data["week"] == "2026-01-26"
    assert data["cohorts"][0]["cohort_kind"] == 1 and data["cohorts"][0]["name"] == role.name
    assert data["cohorts"][0]["size"] == 6 and data["cohorts"][0]["medians"]["active_ratio"] == 1.0
    assert client_as(admin).get(f"{ROOT}/cohorts").json()["data"]["week"] == "2026-01-26"      # mặc định = tuần mới nhất


# ── xuất Excel ──────────────────────────────────────────────────────────────

def test_xuat_excel_ghi_nhat_ky_xuat(db, admin, client_as):
    seeded(db, tag="E1")
    resp = client_as(admin).get(f"{ROOT}/summary/export", params=PERIOD)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"] and resp.content[:2] == b"PK"
    log = db.query(ExportLog).filter_by(entity="usage_profile").one()
    assert log.row_count == 6 and log.fmt == "xlsx" and log.created_by == admin.id


# ── số truy vấn cố định ─────────────────────────────────────────────────────

def _selects_of(sql):
    return [s for s in sql if s.lstrip().upper().startswith("SELECT")]


def test_danh_sach_2_vs_15_nguoi_bang_nhau_va_toi_da_5(db, admin, client_as):
    c = client_as(admin)
    seeded(db, n=2, tag="K1")
    _, few = count_queries(db, lambda: c.get(f"{ROOT}/summary", params=PERIOD))
    make_world(db, 13, tag="K2")
    build_snapshots(db, WEEK)
    resp = None

    def call():
        nonlocal resp
        resp = c.get(f"{ROOT}/summary", params={**PERIOD, "page_size": 50})

    _, many = count_queries(db, call)
    assert resp.json()["data"]["total"] == 15
    sql_few = [s for s in few if "tab_usage_daily" in s or "tab_user_usage" in s]
    sql_many = [s for s in many if "tab_usage_daily" in s or "tab_user_usage" in s]
    assert len(sql_few) == len(sql_many) <= 5


def test_ho_so_4_vs_12_tuan_bang_nhau_va_toi_da_7(db, admin, client_as):
    people, _ = seeded(db, n=6, tag="K3")
    u = people[0][0]
    for w in range(1, 12):                       # thêm snapshot nhiều tuần cũ
        db.add(UserUsageSnapshot(user_id=u.id, period_start=WEEK - timedelta(weeks=w), metrics={"metrics": {}}))
    db.commit()
    c = client_as(admin)
    c.get(f"{ROOT}/{u.id}", params={"weeks": 4})              # khởi động: nạp cache hồ sơ quyền
    _, four = count_queries(db, lambda: c.get(f"{ROOT}/{u.id}", params={"weeks": 4}))
    r12, twelve = count_queries(db, lambda: c.get(f"{ROOT}/{u.id}", params={"weeks": 12}))
    assert len(r12.json()["data"]["weekly"]) == 12
    assert len(_selects_of(four)) == len(_selects_of(twelve))
    usage = [s for s in twelve if "tab_usage_daily" in s or "tab_user_usage" in s or "tab_employee" in s]
    assert 0 < len(usage) <= 7


def test_loc_nhieu_kieu_lam_viec_la_hoac_ca_summary_va_export(db, admin, client_as):
    """Bug H2 (CR-641 review): FE gửi label=2&label=5, BE chỉ nhận số cuối nên lọc sai im lặng. Phải là OR trên mọi nhãn."""
    seeded(db, tag="M1", n=3)
    snaps = []
    for code, style in (("ZA01", UsageStyle.LOW_USAGE), ("ZA02", UsageStyle.POWER_USER), ("ZA03", UsageStyle.WANDERER)):
        u, e = make_person(db, code)
        for d in range(28):
            db.add(UsageDaily(user_id=u.id, day=FIRST + timedelta(days=d), employee_id=e.id, is_workday=True,
                              work_capacity=1.0, active=True))
        db.add(UserUsageSnapshot(user_id=u.id, period_start=WEEK, metrics={}))
        db.flush()
        sid = db.query(UserUsageSnapshot).filter_by(user_id=u.id).one().id
        db.add(UserUsageLabel(snapshot_id=sid, user_id=u.id, period_start=WEEK, label=int(style), score=0.5, evidence=[]))
        snaps.append(u)
    db.commit()
    c = client_as(admin)
    pick = lambda labels: sorted(r["emp_code"] for r in c.get(
        f"{ROOT}/summary", params={**PERIOD, "label": labels}).json()["data"]["items"])
    assert pick([int(UsageStyle.LOW_USAGE)]) == ["ZA01"]
    assert pick([int(UsageStyle.LOW_USAGE), int(UsageStyle.POWER_USER)]) == ["ZA01", "ZA02"]   # trước đây chỉ ra ZA02
    assert c.get(f"{ROOT}/summary", params={**PERIOD, "label": [int(UsageStyle.LOW_USAGE), 99]}).status_code == 422
    resp = c.get(f"{ROOT}/summary/export", params={**PERIOD, "label": [int(UsageStyle.LOW_USAGE), int(UsageStyle.POWER_USER)]})
    assert resp.status_code == 200
    from io import BytesIO
    from openpyxl import load_workbook
    cells = {str(x.value) for row in load_workbook(BytesIO(resp.content)).active.iter_rows() for x in row}
    assert {"ZA01", "ZA02"} <= cells and "ZA03" not in cells
