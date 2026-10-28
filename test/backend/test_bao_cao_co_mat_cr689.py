"""duoc-CR-689 — báo cáo C3 «Có mặt theo ngày» (khóa 26): `/api/attendance-report/summary` (+ export, rows).

Canh: định nghĩa người-ngày (lịch riêng, ngày lễ, ngày vào làm, nghỉ cả ngày / nửa buổi / theo giờ, đơn chờ
duyệt không tính), nhóm theo phòng ban, phạm vi, thiếu quyền xem đơn nghỉ, trần kỳ 92 ngày, danh sách `/rows`,
gác khóa báo cáo ở cả ba đường và số truy vấn CỐ ĐỊNH.
"""
import uuid
from datetime import date, time

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import event

from app.core.auth import get_current_user
from app.core.database import get_db
from app.core.report_keys import ReportKey
from app.core.report_period import parse_period
from app.core.scoping import apply_scope
from app.core.subject_match import SUBJECT_EMPLOYEE
from app.main import app
from app.modules.employee.model import Employee
from app.modules.leave.catalog_model import Holiday, LeaveType
from app.modules.leave.constants import (LR_APPROVED, LR_PENDING, SESSION_FULL, SESSION_HOURLY,
                                         SESSION_MORNING)
from app.modules.leave.request_model import LeaveRequest
from app.modules.report_access.model import ReportAccess
from app.modules.work_schedule import attendance_report_service as service

BASE = "/api/attendance-report/summary"
P = {"preset": "custom", "date_from": "2026-06-01", "date_to": "2026-06-14", "compare": "none"}
_seq = {"n": 0}


@pytest.fixture
def client_as(db):
    def make(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def _emp(db, world, tag, dept="B.hc", co="B", **kw) -> Employee:
    e = Employee(code=f"C{tag}", full_name=f"Người {tag}", company_id=world.co[co],
                 department_id=world.dept[dept], is_active=True, hire_date=date(2026, 1, 1), **kw)
    db.add(e)
    db.flush()
    return e


def _leave_type(db) -> LeaveType:
    lt = db.query(LeaveType).first()
    if lt is None:
        lt = LeaveType(code="PN", name="Phép năm")
        db.add(lt)
        db.flush()
    return lt


def _leave(db, world, emp, start, end=None, *, status=LR_APPROVED, fs=SESSION_FULL, ts=SESSION_FULL,
           ft=None, tt=None, dept="B.hc", co="B") -> LeaveRequest:
    _seq["n"] += 1
    r = LeaveRequest(code=f"CL{_seq['n']:05d}", employee_id=emp.id, company_id=world.co[co],
                     department_id=world.dept[dept], leave_type_id=_leave_type(db).id,
                     from_date=date.fromisoformat(start), to_date=date.fromisoformat(end or start),
                     from_session=fs, to_session=ts, from_time=ft, to_time=tt, status=status)
    db.add(r)
    db.flush()
    return r


def _actor(db, world, client_as, scope="all", leave=True, report=True, key="a1"):
    world.grant(key, "employee", scope, actions=("read", "export"))
    if leave:
        world.grant(key, "leave_request", "all", actions=("read",))
    actor = world.actor(key)
    if report:
        db.add(ReportAccess(report_key=int(ReportKey.ATTENDANCE_DAILY), subject_kind=SUBJECT_EMPLOYEE,
                            subject_id=actor.employee.id, effect=1, reason="t", created_by=0, updated_by=0))
        db.commit()
    return actor, client_as(actor.user)


def _get(client, path=BASE, **extra):
    return client.get(path, params={**P, **extra})


def _setup(db, world):
    """e1 + e2 (vào làm 08/06) ở B.hc; ngày lễ 03/06. Kỳ 01–14/06/2026 (T2…CN, DEGO làm T2–T7)."""
    e1 = _emp(db, world, "1")
    e2 = _emp(db, world, "2")
    e2.hire_date = date(2026, 6, 8)
    db.add(Holiday(company_id=0, date=date(2026, 6, 3), name="Lễ thử", is_recurring=False, is_active=True))
    _leave(db, world, e1, "2026-06-04", "2026-06-05")                                   # 2 ngày
    _leave(db, world, e1, "2026-06-09", fs=SESSION_MORNING, ts=SESSION_MORNING)         # nửa ngày
    _leave(db, world, e1, "2026-06-10", fs=SESSION_HOURLY, ts=SESSION_HOURLY,
           ft=time(10, 0), tt=time(11, 0))                                              # 1/8 ngày
    _leave(db, world, e1, "2026-06-11", status=LR_PENDING)                              # chờ duyệt: bỏ qua
    db.commit()
    return e1, e2


def test_dinh_nghia_nguoi_ngay(db, world, client_as):
    _setup(db, world)
    _, client = _actor(db, world, client_as)
    resp = _get(client)
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    cur = data["totals"]["current"]
    #  e1: 11 ngày làm (T2–T7 trừ lễ 03/06); e2: 6 ngày (từ 08/06)
    assert cur["scheduled"] == 17
    assert cur["on_leave"] == pytest.approx(2 + 0.5 + 0.125)
    assert cur["present"] == pytest.approx(17 - 2.625)
    assert cur["absence_rate"] == pytest.approx(round(2.625 / 17 * 100, 1))
    assert cur["present_rate"] == pytest.approx(round(14.375 / 17 * 100, 1))
    assert data["meta"]["drilldown_metrics"] == ["on_leave"]
    assert [m["key"] for m in data["meta"]["metrics"]][:3] == ["scheduled", "on_leave", "present"]
    trend = data["trend"]
    assert sum(t["current"]["scheduled"] for t in trend) == 17


def test_nhom_theo_phong_ban_va_phap_nhan(db, world, client_as):
    e1, _ = _setup(db, world)
    _emp(db, world, "3", dept="B.kt")
    db.commit()
    _, client = _actor(db, world, client_as)
    groups = _get(client, group_by="department").json()["data"]["groups"]
    by = {g["label"]: g["current"] for g in groups}
    assert by["Phòng Hành chính"]["scheduled"] == 17 and by["Phòng Kế toán"]["scheduled"] == 11
    assert by["Phòng Kế toán"]["on_leave"] == 0
    co = _get(client, group_by="company").json()["data"]["groups"]
    assert len(co) == 1 and co[0]["current"]["scheduled"] == 28
    assert _get(client, group_by="bogus").status_code == 422
    only = _get(client, department_id=world.dept["B.hc"]).json()["data"]["totals"]["current"]
    assert only["scheduled"] == 17


def test_phamvi_nhan_su_chi_thay_phong_minh(db, world, client_as):
    _setup(db, world)
    _emp(db, world, "x", dept="A.kt", co="A")                       # phòng của a1
    db.commit()
    _, client = _actor(db, world, client_as, scope="dept")
    cur = _get(client).json()["data"]["totals"]["current"]
    assert cur["scheduled"] == 11                                   # chỉ người ở A.kt


def test_thieu_quyen_xem_don_nghi(db, world, client_as):
    _setup(db, world)
    _, client = _actor(db, world, client_as, leave=False)
    data = _get(client).json()["data"]
    cur = data["totals"]["current"]
    assert cur["scheduled"] == 17
    assert not {"on_leave", "present", "present_rate", "absence_rate"} & set(cur)
    assert not {"on_leave", "present"} & {m["key"] for m in data["meta"]["metrics"]}
    xlsx = _get(client, BASE + "/export")                           # tệp Excel không có cột nghỉ (không ra 0)
    assert xlsx.status_code == 200
    from io import BytesIO
    from openpyxl import load_workbook
    header = [c.value for c in next(load_workbook(BytesIO(xlsx.content)).active.iter_rows(max_row=1))]
    assert not any(h and "nghỉ" in str(h).lower() for h in header)
    assert any("quyền xem đơn nghỉ" in n for n in data["notes"])
    assert data["meta"]["drilldown_metrics"] == []
    assert _get(client, BASE + "/rows", metric="on_leave").status_code == 403


def test_tran_ky_92_ngay(db, world, client_as):
    _setup(db, world)
    _, client = _actor(db, world, client_as)
    ok = _get(client, date_from="2026-06-01", date_to="2026-09-01")      # đúng 93 ngày? 30+31+31+1=93
    assert ok.status_code == 422 and "92 ngày" in ok.json()["error"]["message"]
    fine = _get(client, date_from="2026-06-01", date_to="2026-08-31")     # 92 ngày
    assert fine.status_code == 200
    for suffix, extra in (("/export", {}), ("/rows", {"metric": "on_leave"})):
        assert _get(client, BASE + suffix, date_from="2026-06-01", date_to="2026-09-01", **extra).status_code == 422


def test_rows_khop_so_nguoi_ngay_nghi(db, world, client_as):
    _setup(db, world)
    _, client = _actor(db, world, client_as)
    rows = _get(client, BASE + "/rows", metric="on_leave").json()["data"]
    assert [c["key"] for c in rows["columns"]] == ["employee_code", "employee_name", "department", "day",
                                                   "leave_type", "session"]
    assert rows["total"] == 4                                         # ngày 04, 05, 09 (sáng), 10 (theo giờ)
    sessions = sorted(r["session"] for r in rows["rows"])
    assert "Buổi sáng" in sessions and any(s.startswith("Theo giờ 10:00") for s in sessions)
    assert all(r["leave_type"] == "Phép năm" for r in rows["rows"])
    grp = _get(client, BASE + "/rows", metric="on_leave", group_by="department",
               group_key=str(world.dept["B.hc"])).json()["data"]
    assert grp["total"] == rows["total"]
    none = _get(client, BASE + "/rows", metric="on_leave", group_by="department",
                group_key=str(world.dept["B.kt"])).json()["data"]
    assert none["total"] == 0
    for bad in ({"metric": "scheduled"}, {}, {"metric": "on_leave", "group_by": "department",
                                              "group_key": "__other__"}):
        assert _get(client, BASE + "/rows", **bad).status_code == 422


@pytest.mark.parametrize("suffix", ["", "/export", "/rows"])
def test_chua_giao_bao_cao_403_ca_ba_duong(db, world, client_as, suffix):
    _, client = _actor(db, world, client_as, report=False)
    assert _get(client, BASE + suffix, metric="on_leave").status_code == 403


def test_export_xlsx(db, world, client_as):
    _setup(db, world)
    _, client = _actor(db, world, client_as)
    r = _get(client, BASE + "/export", group_by="department")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]


def test_so_truy_van_khong_tang_theo_so_nguoi(db, world):
    world.grant("a1", "employee", "all", actions=("read",))
    world.grant("a1", "leave_request", "all", actions=("read",))
    actor = world.actor("a1")
    prof = actor.profile()
    period = parse_period({**P, "compare": "previous"})

    def count() -> int:
        stmts: list[str] = []

        def listener(conn, cur, stmt, *args):
            stmts.append(stmt)

        eng = db.get_bind()
        event.listen(eng, "before_cursor_execute", listener)
        try:
            service.build_summary(
                db, apply_scope(db.query(Employee), Employee, "employee", actor.user, prof),
                apply_scope(db.query(LeaveRequest), LeaveRequest, "leave_request", actor.user, prof),
                True, period, "department")
        finally:
            event.remove(eng, "before_cursor_execute", listener)
        return len(stmts)

    def add(n, start):
        for i in range(n):
            e = _emp(db, world, f"q{start + i}", dept="B.hc" if i % 2 else "B.kt")
            _leave(db, world, e, "2026-06-04", "2026-06-05")
        db.commit()

    add(5, 0)
    c5 = count()
    add(45, 5)
    assert c5 > 0 and c5 == count()
