"""duoc-CR-694 — hồ sơ nhân sự nhập NHIỀU số điện thoại (`extra_phones`; `phone` vẫn là số chính).

Phủ: luật schema (định dạng, trùng, quá 4), lưu/đọc qua API, trùng số chính, tự sửa ở Trang cá nhân,
đặt số phụ làm số chính, tìm kiếm theo số phụ, xuất xlsx + nhập Excel.
"""
import io
import uuid

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app.core.auth import get_current_user
from app.core.database import get_db
from app.main import app
from app.modules.employee import phone_rules
from app.modules.employee.model import Employee
from app.modules.employee.schema import EmployeeCreate, EmployeeUpdate, SelfContactUpdate

from test_nhap_excel_nhan_su_cr628 import emp, errors, hr, run_import, savepoint_that  # noqa: F401


@pytest.fixture
def client_as(db):
    def make(user):
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app, headers={"Authorization": f"Bearer test-{uuid.uuid4().hex}"})

    yield make
    app.dependency_overrides.pop(get_db, None)
    app.dependency_overrides.pop(get_current_user, None)


def _hr(world, key="a1"):
    world.grant(key, "employee", "all", actions=("read", "write", "create", "export"))
    return world.actor(key)


# ---------- schema ----------
def test_so_hop_le_duoc_cat_khoang_trang_va_nhan_tu_do():
    u = EmployeeUpdate(extra_phones=[{"number": " 0901 234 567 ", "label": " Zalo "}, {"number": "+84 (28) 3822-1234"}])
    assert u.extra_phones[0].number == "0901 234 567" and u.extra_phones[0].label == "Zalo"
    assert u.extra_phones[1].label == ""


@pytest.mark.parametrize("bad", ["abc12345678", "0901-234-567 ext", "1234567", "1234567890123456", ""])
def test_sai_dinh_dang_hoac_so_chu_so(bad):
    with pytest.raises(ValidationError) as e:
        EmployeeUpdate(extra_phones=[{"number": bad}])
    assert "Số điện thoại" in str(e.value)


def test_trung_nhau_sau_khi_bo_ky_tu_phan_cach():
    with pytest.raises(ValidationError) as e:
        EmployeeUpdate(extra_phones=[{"number": "0901 234 567"}, {"number": "0901.234.567"}])
    assert "trùng" in str(e.value)


def test_qua_bon_so_phu():
    nums = [{"number": f"090123456{i}"} for i in range(5)]
    with pytest.raises(ValidationError) as e:
        EmployeeUpdate(extra_phones=nums)
    assert "tối đa 4" in str(e.value)
    assert len(EmployeeUpdate(extra_phones=nums[:4]).extra_phones) == 4


def test_do_dai_o_khop_cot():
    with pytest.raises(ValidationError):
        EmployeeUpdate(extra_phones=[{"number": "0" * 21}])
    with pytest.raises(ValidationError):
        EmployeeUpdate(extra_phones=[{"number": "0901234567", "label": "x" * 51}])


def test_chuyen_doi_chu_hai_chieu():
    items = [{"number": "(028) 3822 1234", "label": "Nhà riêng"}, {"number": "0901234567", "label": ""}]
    text = phone_rules.format_phones_text(items)
    assert text == "(028) 3822 1234 (Nhà riêng); 0901234567"
    assert phone_rules.parse_phones_text(text) == items


# ---------- API hồ sơ ----------
def test_luu_doc_giu_nguyen_khi_vang_va_xoa_bang_rong(db, world, client_as):
    actor = _hr(world)
    client = client_as(actor.user)
    eid = world.emp["a2"]
    body = {"phone": "0911111111", "extra_phones": [{"number": "0922222222", "label": "Zalo"}]}
    r = client.patch(f"/api/employees/{eid}", json=body)
    assert r.status_code == 200, r.text
    assert r.json()["data"]["extra_phones"] == [{"number": "0922222222", "label": "Zalo"}]
    # PATCH vắng = giữ nguyên; null = giữ nguyên
    client.patch(f"/api/employees/{eid}", json={"full_name": "Tên mới"})
    client.patch(f"/api/employees/{eid}", json={"extra_phones": None})
    assert client.get(f"/api/employees/{eid}").json()["data"]["extra_phones"][0]["number"] == "0922222222"
    # [] = xóa hết
    client.patch(f"/api/employees/{eid}", json={"extra_phones": []})
    assert client.get(f"/api/employees/{eid}").json()["data"]["extra_phones"] == []


def test_trung_so_chinh_bi_422(db, world, client_as):
    client = client_as(_hr(world).user)
    eid = world.emp["a2"]
    r = client.patch(f"/api/employees/{eid}", json={"phone": "0911 111 111",
                                                    "extra_phones": [{"number": "0911.111.111"}]})
    assert r.status_code == 422 and "trùng với số chính" in r.text
    # chỉ gửi số phụ trùng số chính đã lưu
    client.patch(f"/api/employees/{eid}", json={"phone": "0911111111"})
    r = client.patch(f"/api/employees/{eid}", json={"extra_phones": [{"number": "0911111111"}]})
    assert r.status_code == 422


def test_sai_dinh_dang_qua_api_la_422_tieng_viet(db, world, client_as):
    client = client_as(_hr(world).user)
    r = client.patch(f"/api/employees/{world.emp['a2']}", json={"extra_phones": [{"number": "abc"}]})
    assert r.status_code == 422 and "Số điện thoại chỉ gồm" in r.text


def test_dat_so_phu_lam_so_chinh_hoan_doi(db, world, client_as):
    client = client_as(_hr(world).user)
    eid = world.emp["a2"]
    client.patch(f"/api/employees/{eid}", json={"phone": "0911111111",
                                                "extra_phones": [{"number": "0922222222", "label": "Zalo"}]})
    client.patch(f"/api/employees/{eid}", json={"phone_label": "Cá nhân"})
    assert client.get(f"/api/employees/{eid}").json()["data"]["phone_label"] == "Cá nhân"
    # hoán đổi CẢ số lẫn nhãn
    r = client.patch(f"/api/employees/{eid}", json={
        "phone": "0922222222", "phone_label": "Zalo",
        "extra_phones": [{"number": "0911111111", "label": "Cá nhân"}]})
    assert r.status_code == 200
    d = r.json()["data"]
    assert d["phone"] == "0922222222" and d["phone_label"] == "Zalo"
    assert d["extra_phones"] == [{"number": "0911111111", "label": "Cá nhân"}]
    # nhãn số chính vắng = giữ nguyên; quá 50 ký tự = 422
    client.patch(f"/api/employees/{eid}", json={"full_name": "X"})
    assert client.get(f"/api/employees/{eid}").json()["data"]["phone_label"] == "Zalo"
    assert client.patch(f"/api/employees/{eid}", json={"phone_label": "x" * 51}).status_code == 422


def test_tao_moi_kem_so_phu_va_chan_trung(db, world, client_as):
    client = client_as(_hr(world).user)
    r = client.post("/api/employees", json={"full_name": "Người Mới", "phone": "0933333333",
                                            "extra_phones": [{"number": "0944444444", "label": "Công việc"}]})
    assert r.status_code in (200, 201), r.text
    assert r.json()["data"]["extra_phones"][0]["label"] == "Công việc"
    bad = client.post("/api/employees", json={"full_name": "Trùng", "phone": "0933333333",
                                              "extra_phones": [{"number": "0933 333 333"}]})
    assert bad.status_code == 422


def test_tim_kiem_theo_so_phu(db, world, client_as):
    client = client_as(_hr(world).user)
    eid = world.emp["a2"]
    client.patch(f"/api/employees/{eid}", json={"extra_phones": [{"number": "0977665544", "label": "Zalo"}]})
    def ids_of(kw):
        data = client.get("/api/employees", params={"search": kw}).json()["data"]
        return [x["id"] for x in (data["items"] if isinstance(data, dict) else data)]

    assert eid in ids_of("0977665544")
    # từ khóa chữ không khớp nhầm «number»/«label» trong JSON
    assert eid not in ids_of("Zalo")
    assert eid not in ids_of("number")


# ---------- Tự sửa ở Trang cá nhân ----------
def test_tu_sua_lien_he_nhan_so_phu(db, world, client_as):
    me = world.actor("a1")           # không cấp quyền employee.write
    client = client_as(me.user)
    r = client.patch("/api/employees/me/contact", json={"phone": "0955555555", "phone_label": "Công việc",
                                                        "extra_phones": [{"number": "0966666666", "label": "Nhà riêng"}]})
    assert r.status_code == 200, r.text
    assert r.json()["data"]["extra_phones"] == [{"number": "0966666666", "label": "Nhà riêng"}]
    assert r.json()["data"]["phone_label"] == "Công việc"
    # gửi lại y nguyên → không đổi
    r = client.patch("/api/employees/me/contact", json={"extra_phones": [{"number": "0966666666", "label": "Nhà riêng"}]})
    assert r.json()["message"] == "Không có thay đổi nào"
    # trùng số chính → 422; quá 4 → 422
    assert client.patch("/api/employees/me/contact", json={"extra_phones": [{"number": "0955 555 555"}]}).status_code == 422
    five = [{"number": f"090123456{i}"} for i in range(5)]
    assert client.patch("/api/employees/me/contact", json={"extra_phones": five}).status_code == 422
    # xóa hết
    client.patch("/api/employees/me/contact", json={"extra_phones": []})
    assert db.get(Employee, me.user.employee_id).extra_phones == []


def test_tu_sua_van_cam_khoa_la():
    with pytest.raises(ValidationError):
        SelfContactUpdate(extra_phones=[], bank_name="x")


# ---------- Xuất / nhập Excel ----------
def test_xuat_xlsx_co_cot_so_dien_thoai_khac(db, world, client_as):
    client = client_as(_hr(world).user)
    eid = world.emp["a2"]
    client.patch(f"/api/employees/{eid}", json={"extra_phones": [{"number": "0922222222", "label": "Zalo"},
                                                                  {"number": "0933333333"}]})
    r = client.get("/api/employees/export/xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert "Nhãn số chính" in header
    col = header.index("Số điện thoại khác")
    values = [row[col].value for row in ws.iter_rows(min_row=2)]
    assert "0922222222 (Zalo); 0933333333" in values


def test_nhap_excel_so_dien_thoai_khac(db, hr):
    b = run_import(db, [{"Mã NV *": "PH001", "Họ tên *": "Ba Số", "Điện thoại": "0911111111",
                         "Nhãn số chính": "Cá nhân",
                         "Số điện thoại khác": "0922222222 (Zalo); (028) 3822 1234 (Nhà riêng)"}], hr)
    assert not errors(db, b.id)
    assert emp(db, "PH001").phone_label == "Cá nhân"
    assert emp(db, "PH001").extra_phones == [{"number": "0922222222", "label": "Zalo"},
                                              {"number": "(028) 3822 1234", "label": "Nhà riêng"}]
    # ô trống = giữ nguyên; sentinel = xóa
    run_import(db, [{"Mã NV *": "PH001", "Số điện thoại khác": ""}], hr)
    assert len(emp(db, "PH001").extra_phones) == 2
    run_import(db, [{"Mã NV *": "PH001", "Số điện thoại khác": "__/empty_value/__"}], hr)
    assert emp(db, "PH001").extra_phones == []


@pytest.mark.parametrize("cell", ["abc", "0911111111", "0922222222; 0922 222 222",
                                  "0900000001; 0900000002; 0900000003; 0900000004; 0900000005"])
def test_nhap_excel_sai_la_loi_dong(db, hr, cell):
    b = run_import(db, [{"Mã NV *": "PH002", "Họ tên *": "Lỗi", "Điện thoại": "0911111111",
                         "Số điện thoại khác": cell}], hr)
    assert errors(db, b.id)
    assert emp(db, "PH002") is None
