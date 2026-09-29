#!/usr/bin/env python3
"""Cào danh mục thuốc BVTV từ https://danhmuc.thuocbvtv.com (trang HTML công khai).

Nguồn:
  - Danh sách: trang chủ phân trang `/?trang=N` (50 dòng/trang) — đếm đối chiếu + cột danh sách.
  - Sitemap `/sitemap-thuoc.xml?page=N` — tập URL chi tiết thứ hai để đối chiếu.
  - Chi tiết: `/thuoc/detail/<id>/<slug>` (robots.txt cho phép; KHÔNG gọi /api/, /thuoc/search).

Chạy:  python3 scrape-thuoc-bvtv.py [--rps 3] [--workers 3] [--export-only]
Checkpoint trong ./checkpoint/ — chạy lại sẽ bỏ qua phần đã xong.
"""
import argparse, csv, json, re, sys, threading, time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import requests
from bs4 import BeautifulSoup

BASE = "https://danhmuc.thuocbvtv.com"
OUT = Path(__file__).resolve().parent
CK = OUT / "checkpoint"
LIST_CK, SITEMAP_CK = CK / "list.json", CK / "sitemap.json"
DETAIL_CK, ERR_CK = CK / "details.jsonl", CK / "errors.json"
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 14_0) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/128.0 Safari/537.36")


class Fetcher:
    """requests.Session + giới hạn tốc độ toàn cục + retry/backoff (429/5xx/lỗi mạng)."""

    def __init__(self, rps):
        self.s = requests.Session()
        self.s.headers.update({"User-Agent": UA, "Accept-Language": "vi,en;q=0.8"})
        self.gap, self.lock, self.next_t = 1.0 / rps, threading.Lock(), 0.0

    def _wait(self):
        with self.lock:
            now = time.monotonic()
            t = max(now, self.next_t)
            self.next_t = t + self.gap
        time.sleep(max(0, t - now))

    def get(self, url, tries=6):
        for i in range(tries):
            self._wait()
            try:
                r = self.s.get(url, timeout=30)
                if r.status_code == 200:
                    r.encoding = "utf-8"
                    return r.text
                if r.status_code == 404:
                    raise FileNotFoundError(url)
                if r.status_code in (429, 500, 502, 503, 504):
                    ra = r.headers.get("Retry-After")
                    delay = int(ra) if ra and ra.isdigit() else 2 ** i * 2
                    print(f"  HTTP {r.status_code} {url} -> chờ {delay}s", flush=True)
                    time.sleep(delay)
                    continue
                raise RuntimeError(f"HTTP {r.status_code}")
            except (requests.ConnectionError, requests.Timeout) as e:
                print(f"  lỗi mạng {url}: {e} (lần {i+1})", flush=True)
                time.sleep(2 ** i * 2)
        raise RuntimeError(f"hết lượt thử: {url}")


def txt(el):
    return re.sub(r"\s+", " ", el.get_text(" ", strip=True)).strip() if el else ""


def id_from_url(u):
    m = re.search(r"/thuoc/detail/(\d+)/([^/?#]*)", u)
    return (int(m.group(1)), m.group(2)) if m else (None, None)


def load_json(p, default):
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default


def save_json(p, data):
    tmp = p.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(p)


# ---------------- Danh sách + sitemap ----------------
def scrape_list(f):
    ck = load_json(LIST_CK, {"total": None, "pages": None, "rows": {}})
    if ck["pages"] is None:
        html = f.get(f"{BASE}/?trang=1")
        s = BeautifulSoup(html, "lxml")
        m = re.search(r"([\d.,]+)\s*sản phẩm", txt(s.select_one("#san-pham")) or s.get_text(" "))
        ck["total"] = int(re.sub(r"\D", "", m.group(1))) if m else None
        pages = [int(x) for x in re.findall(r"[?&]trang=(\d+)", html)]
        ck["pages"] = max(pages) if pages else 1
        save_json(LIST_CK, ck)
    for p in range(1, ck["pages"] + 1):
        if str(p) in ck["rows"]:
            continue
        s = BeautifulSoup(f.get(f"{BASE}/?trang={p}"), "lxml")
        rows = []
        for tr in s.select("table tbody tr"):
            a = tr.select_one("a.thuoc-name")
            if not a:
                continue
            tid, slug = id_from_url(a["href"])
            ct = tr.select_one("a.congty-link")
            rows.append({"id": tid, "slug": slug, "url": a["href"],
                         "stt": txt(tr.select_one(".col-num")), "ten_thuoc": txt(a),
                         "hoat_chat": txt(tr.select_one(".thuoc-hoatchat")),
                         "phan_nhom": txt(tr.select_one(".col-nhom")),
                         "cong_ty": txt(ct) if ct else txt(tr.select_one(".col-congty"))})
        ck["rows"][str(p)] = rows
        save_json(LIST_CK, ck)
        print(f"danh sách trang {p}/{ck['pages']}: {len(rows)} dòng", flush=True)
    return ck


def scrape_sitemap(f):
    if SITEMAP_CK.exists():
        return load_json(SITEMAP_CK, [])
    idx = f.get(f"{BASE}/sitemap.xml")
    urls = []
    for sm in re.findall(r"<loc>([^<]*sitemap-thuoc[^<]*)</loc>", idx):
        urls += re.findall(r"<loc>([^<]*/thuoc/detail/[^<]*)</loc>", f.get(sm.replace("&amp;", "&")))
    save_json(SITEMAP_CK, urls)
    return urls


# ---------------- Chi tiết ----------------
def parse_date(d):
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$", d or "")
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else ""


def parse_frac(box):
    out = {"nhom": [], "hoat_chat": []}
    for g in box.select(".frac-group-badge"):
        out["nhom"].append({"ma": [txt(c) for c in g.select(".fgb-code")],
                            "loai": txt(g.select_one(".fgb-type"))})
    for row in box.select(".frac-badge-row"):
        code = row.select_one(".frac-tag-code")
        out["hoat_chat"].append({
            "ten": txt(row.select_one(".frac-ingredient-name")),
            "ma": txt(code).replace("▾", "").strip() if code else "",
            "nhom": txt(row.select_one(".frac-tag-nhom")),
            "phuong_thuc": txt(row.select_one(".frac-tag-phuonthuc")),
            "cung_nhom": [txt(a) for a in row.select("div[id^=frac-tip] a")],
        })
    extra = txt(box) if not out["nhom"] and not out["hoat_chat"] else ""
    if extra:
        out["ghi_chu"] = extra
    return out


PV_KEYS = {"Liều lượng": "lieu_luong", "Thời gian cách ly": "thoi_gian_cach_ly", "Cách dùng": "cach_dung"}
TARGET_KEYS = {"Cây trồng": "cay_trong", "Dịch hại": "dich_hai"}


def parse_detail(html, url):
    s = BeautifulSoup(html, "lxml")
    tid, slug = id_from_url(url)
    r = {"id": tid, "slug": slug, "url": url,
         "ten_thuoc": txt(s.select_one("h1.detail-title")),
         "phan_nhom": txt(s.select_one(".detail-badges .badge-group")),
         "linh_vuc": txt(s.select_one(".detail-badges .badge-sector")),
         "tinh_trang": txt(s.select_one(".detail-badges .badge-status")),
         "cong_ty_dang_ky": "", "cong_ty_url": "", "hoat_chat": "", "ham_luong": "",
         "quan_ly_tinh_khang": {"nhom": [], "hoat_chat": []},
         "so_dang_ky": "", "thoi_han_dang_ky": "", "ngay_cap": "", "ngay_het_han": "",
         "nhom_doc": [], "tom_tat_su_dung": txt(s.select_one(".summary-usage-desc")),
         "pham_vi_su_dung": [], "nguon_ecofarm": "", "thong_tin_khac": {}}
    if not r["ten_thuoc"]:
        raise ValueError("không thấy tiêu đề h1 — trang chi tiết lạ")
    other_badges = [txt(b) for b in s.select(".detail-badges .badge")
                    if not {"badge-group", "badge-sector", "badge-status"} & set(b.get("class", []))]
    if other_badges:
        r["thong_tin_khac"]["badge_khac"] = other_badges
    for it in s.select(".detail-info-grid .info-item"):
        label = re.sub(r"^[^\wÀ-ỹ]+", "", txt(it.select_one("label"))).strip()
        val_el = it.find("div")
        lk = label.lower()
        if lk.startswith("công ty"):
            a = it.select_one("a")
            r["cong_ty_dang_ky"], r["cong_ty_url"] = txt(val_el), a["href"] if a else ""
        elif lk.startswith("hoạt chất"):
            props = {p.select_one("meta")["content"]: txt(p.select_one("[itemprop=value]"))
                     for p in it.select("[itemprop=additionalProperty]")}
            r["hoat_chat"] = props.get("Hoạt chất", "") or txt(val_el)
            r["ham_luong"] = props.get("Hàm lượng", "")
        elif "tính kháng" in lk:
            r["quan_ly_tinh_khang"] = parse_frac(it)
        elif lk.startswith("số đăng ký"):
            r["so_dang_ky"] = txt(val_el)
        elif lk.startswith("thời hạn"):
            r["thoi_han_dang_ky"] = txt(val_el)
            parts = re.split(r"->|→", r["thoi_han_dang_ky"])
            r["ngay_cap"] = parse_date(parts[0]) if parts else ""
            r["ngay_het_han"] = parse_date(parts[1]) if len(parts) > 1 else ""
        elif lk.startswith("nhóm độc"):
            for b in it.select(".badge"):
                r["nhom_doc"].append({"he": txt(b.select_one("strong")),
                                      "nhom": txt(b.select_one(".nhom-doc-so")),
                                      "mo_ta": b.get("title", "")})
        else:
            r["thong_tin_khac"][label] = txt(val_el)
    for card in s.select(".pham-vi-section .pham-vi-card"):
        pv = {"cay_trong": "", "dich_hai": "", "lieu_luong": "", "thoi_gian_cach_ly": "", "cach_dung": ""}
        for t in card.select(".pv-target"):
            lb = txt(t.select_one(".pv-label"))
            pv[TARGET_KEYS.get(lb, lb)] = txt(t.select_one("strong"))
        for row in card.select(".pv-detail-row"):
            lb = txt(row.select_one(".pv-detail-label")).rstrip(":").strip()
            pv[PV_KEYS.get(lb, lb)] = txt(row.select_one(".pv-detail-val"))
        r["pham_vi_su_dung"].append(pv)
    src = s.select_one(".detail-source a[href]")
    r["nguon_ecofarm"] = src["href"] if src else ""
    for sc in s.select('script[type="application/ld+json"]'):
        try:
            j = json.loads(sc.string or "")
        except ValueError:
            continue
        if isinstance(j, dict) and j.get("@type") == "Product":
            for p in j.get("additionalProperty", []):
                if p.get("name") not in ("Hoạt chất", "Hàm lượng", "Tình trạng đăng ký", "Năm đăng ký"):
                    r["thong_tin_khac"]["ld_" + p.get("name", "")] = p.get("value")
            if not r["so_dang_ky"]:
                r["so_dang_ky"] = j.get("identifier", "")
    return r


def scrape_details(f, targets, workers):
    done = set()
    if DETAIL_CK.exists():
        for line in DETAIL_CK.read_text(encoding="utf-8").splitlines():
            if line.strip():
                done.add(json.loads(line)["id"])
    errors = {}
    todo = [(i, u) for i, u in targets.items() if i not in done]
    print(f"chi tiết: đã có {len(done)}, còn {len(todo)}", flush=True)
    wlock, n = threading.Lock(), 0

    def job(item):
        tid, url = item
        try:
            return tid, url, parse_detail(f.get(url), url), None
        except Exception as e:  # ghi lại, không dừng cả mẻ
            return tid, url, None, f"{type(e).__name__}: {e}"

    with ThreadPoolExecutor(workers) as ex, DETAIL_CK.open("a", encoding="utf-8") as fh:
        futs = [ex.submit(job, it) for it in todo]
        for fu in as_completed(futs):
            tid, url, rec, err = fu.result()
            if err:
                errors[str(tid)] = {"url": url, "loi": err}
            else:
                with wlock:
                    fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
                    fh.flush()
            n += 1
            if n % 200 == 0:
                print(f"  {n}/{len(todo)} (lỗi {len(errors)})", flush=True)
    save_json(ERR_CK, errors)
    return errors


# ---------------- Xuất ----------------
def flat_frac(q):
    parts = [f"{h['ten']}: {h['ma']} | {h['nhom']} | {h['phuong_thuc']}" for h in q.get("hoat_chat", [])]
    return "; ".join(parts) or q.get("ghi_chu", "")


def export(records, list_by_id):
    records.sort(key=lambda r: (r["ten_thuoc"].lower(), r["id"]))
    for r in records:
        lr = list_by_id.get(r["id"], {})
        r["hoat_chat_danh_sach"] = lr.get("hoat_chat", "")
    (OUT / "thuoc-bvtv.json").write_text(json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8")

    cols = ["id", "ten_thuoc", "phan_nhom", "linh_vuc", "tinh_trang", "hoat_chat", "ham_luong",
            "hoat_chat_danh_sach", "cong_ty_dang_ky", "so_dang_ky", "thoi_han_dang_ky", "ngay_cap",
            "ngay_het_han", "nhom_doc", "nhom_khang_thuoc", "quan_ly_tinh_khang", "so_pham_vi",
            "cay_trong", "dich_hai", "pham_vi_su_dung", "tom_tat_su_dung", "thong_tin_khac",
            "url", "cong_ty_url", "nguon_ecofarm"]

    def flat(r):
        pv = r["pham_vi_su_dung"]
        uniq = lambda k: "; ".join(dict.fromkeys(p.get(k, "") for p in pv if p.get(k)))
        return {**{k: r.get(k, "") for k in cols},
                "nhom_doc": "; ".join(f"{d['he']} {d['nhom']} ({d['mo_ta']})" for d in r["nhom_doc"]),
                "nhom_khang_thuoc": "; ".join(f"{','.join(g['ma'])} {g['loai']}".strip()
                                              for g in r["quan_ly_tinh_khang"].get("nhom", [])),
                "quan_ly_tinh_khang": flat_frac(r["quan_ly_tinh_khang"]),
                "so_pham_vi": len(pv), "cay_trong": uniq("cay_trong"), "dich_hai": uniq("dich_hai"),
                "pham_vi_su_dung": " || ".join(" | ".join(f"{k}: {v}" for k, v in p.items() if v) for p in pv),
                "thong_tin_khac": json.dumps(r["thong_tin_khac"], ensure_ascii=False) if r["thong_tin_khac"] else ""}

    flats = [flat(r) for r in records]
    with (OUT / "thuoc-bvtv.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(flats)

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sach thuoc"
    ws.append(cols)
    for fr in flats:
        ws.append([str(fr[c])[:32000] if isinstance(fr[c], str) else fr[c] for c in cols])
    pv_cols = ["id", "ten_thuoc", "so_dang_ky", "phan_nhom", "tinh_trang", "stt_pham_vi",
               "cay_trong", "dich_hai", "lieu_luong", "thoi_gian_cach_ly", "cach_dung", "khac"]
    ws2 = wb.create_sheet("Pham vi su dung")
    ws2.append(pv_cols)
    for r in records:
        for i, p in enumerate(r["pham_vi_su_dung"], 1):
            extra = {k: v for k, v in p.items() if k not in pv_cols}
            ws2.append([r["id"], r["ten_thuoc"], r["so_dang_ky"], r["phan_nhom"], r["tinh_trang"], i,
                        p.get("cay_trong", ""), p.get("dich_hai", ""), p.get("lieu_luong", ""),
                        p.get("thoi_gian_cach_ly", ""), p.get("cach_dung", ""),
                        json.dumps(extra, ensure_ascii=False) if extra else ""])
    for sh in (ws, ws2):
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for i in range(1, sh.max_column + 1):
            sh.column_dimensions[get_column_letter(i)].width = 22
    wb.save(OUT / "thuoc-bvtv.xlsx")
    print(f"xuất: {len(records)} thuốc, {ws2.max_row - 1} dòng phạm vi", flush=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rps", type=float, default=3.0)
    ap.add_argument("--workers", type=int, default=3)
    ap.add_argument("--export-only", action="store_true")
    a = ap.parse_args()
    CK.mkdir(exist_ok=True)
    f = Fetcher(a.rps)
    lst = scrape_list(f) if not a.export_only else load_json(LIST_CK, {"rows": {}})
    list_rows = [x for p in sorted(lst["rows"], key=int) for x in lst["rows"][p]]
    list_by_id = {x["id"]: x for x in list_rows}
    sm = scrape_sitemap(f) if not a.export_only else load_json(SITEMAP_CK, [])
    targets = {x["id"]: x["url"] for x in list_rows}
    for u in sm:
        tid, _ = id_from_url(u)
        targets.setdefault(tid, u)
    print(f"trang báo {lst.get('total')} | danh sách {len(list_rows)} dòng, {len(list_by_id)} id "
          f"| sitemap {len(sm)} | hợp {len(targets)}", flush=True)
    if not a.export_only:
        errs = scrape_details(f, targets, a.workers)
        print(f"lỗi chi tiết: {len(errs)}", flush=True)
    recs = {}
    for line in DETAIL_CK.read_text(encoding="utf-8").splitlines():
        if line.strip():
            rec = json.loads(line)
            recs[rec["id"]] = rec
    export(list(recs.values()), list_by_id)


if __name__ == "__main__":
    sys.exit(main())
